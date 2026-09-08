"""The corrections run as a two-sheet Excel workbook: "Applied" and "Not applied".

Every correction the run received lands on exactly one of the two sheets, with
the page the reviewer marked it on and the page InDesign shows it on in the
corrected file. It is a pure function over the payload `report.py` writes to
`corrections.json` — the same dict, not a re-derivation — so the spreadsheet
can never disagree with the report.

The unit is the reviewer comment when a marked proof drove the run (one row per
comment, as the report counts), and the edit when a typed or Word list did.
"""
from __future__ import annotations

import logging
from pathlib import Path

from .model import (APPLIED_EXACTLY, DEVIATES, DISP_APPLIED, DISP_FLAGGED,
                    DISP_NO_OP, DISP_NOT_EXTRACTED, MISSING)
from .report import FLAG_TITLES

log = logging.getLogger("docproof.corrections.spreadsheet")

SHEET_APPLIED = "Applied"
SHEET_NOT_APPLIED = "Not applied"

HEADER = ("#", "Page (as marked)", "Page in IDML", "Page confirmed",
          "Correction (as written)", "Marked text", "Find", "Replace",
          "Status", "Reason / detail", "Edit id(s)", "Comment id")

# Column widths, in Excel character units, in HEADER order.
_WIDTHS = (6, 14, 14, 14, 50, 40, 50, 50, 34, 60, 16, 16)

# The header row sits below the summary block (rows 1-7) and one blank row.
SUMMARY_ROWS = 7
HEADER_ROW = SUMMARY_ROWS + 2

STATUS_APPLIED_EXACTLY = "Applied exactly"
STATUS_APPLIED_DEVIATES = "Applied (differs from what was asked)"
STATUS_APPLIED = "Applied"
STATUS_RESOLVED = "Applied (resolved in review)"
STATUS_NO_OP = "No change needed"
STATUS_NOT_EXTRACTED = "Could not be turned into an edit"
STATUS_FLAGGED = "Flagged for a human"
STATUS_SET_ASIDE = "Set aside in review"
STATUS_MISSING = "Not in the corrected document"


def write_spreadsheet(payload: dict, out_path: Path, *,
                      corrected_name: str = "") -> Path:
    """Write the two-sheet workbook for `payload` (the `corrections.json` dict) to
    `out_path` and return it. `corrected_name` names the corrected file in the
    summary block; the payload's `after` basename stands in when it is empty."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    applied, not_applied = build_rows(payload)
    total = len(applied) + len(not_applied)
    _check_total(payload, total)

    after = payload.get("after") or ""
    corrected = corrected_name or (Path(after).name if after else "")
    pages = payload.get("pages") or {}
    summary = (
        ("Source file", payload.get("source_name") or ""),
        ("Corrected file", corrected),
        ("Generated", (payload.get("generated_at") or "")[:19].replace("T", " ")
         + (" UTC" if payload.get("generated_at") else "")),
        ("Total corrections received", total),
        ("Applied", len(applied)),
        ("Not applied", len(not_applied)),
        ("Pages placed / total",
         f"{pages.get('placed', 0)} / {pages.get('total', 0)}"),
    )
    assert len(summary) == SUMMARY_ROWS

    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    first = wb.active
    first.title = SHEET_APPLIED
    second = wb.create_sheet(SHEET_NOT_APPLIED)
    for ws, rows in ((first, applied), (second, not_applied)):
        for i, (k, v) in enumerate(summary, start=1):
            ws.cell(row=i, column=1, value=k).font = bold
            ws.cell(row=i, column=2, value=v)
        for col, name in enumerate(HEADER, start=1):
            c = ws.cell(row=HEADER_ROW, column=col, value=name)
            c.font = bold
            c.alignment = wrap
        for n, row in enumerate(rows, start=1):
            values = (n,) + tuple(row.get(h, "") for h in HEADER[1:])
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=HEADER_ROW + n, column=col, value=v)
                c.alignment = wrap
        for col, width in enumerate(_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)
        last_row = HEADER_ROW + max(len(rows), 1)
        ws.auto_filter.ref = (f"A{HEADER_ROW}:"
                              f"{get_column_letter(len(HEADER))}{last_row}")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    log.info("Wrote %s (%d applied, %d not applied)", out_path,
             len(applied), len(not_applied))
    return out_path


def build_rows(payload: dict) -> tuple[list[dict], list[dict]]:
    """The two sheets' rows as dicts keyed by HEADER name (minus `#`, numbered
    at write time): `(applied, not_applied)`."""
    ap = payload.get("apply")
    verify = payload.get("verify") or {}
    recon = {r["id"]: r for r in verify.get("reconciliations") or []}
    edits = _edits_by_id(payload)
    items = (payload.get("comments") or {}).get("items") or []
    if items:
        applied, not_applied, covered = _comment_rows(items, edits, recon)
        # A typed or hand-added edit that rode along with the proof's comments has
        # no comment to be counted under; it still has to land on one sheet.
        extra = [e for eid, e in edits.items()
                 if eid not in covered and not _synthetic(eid)
                 and not (e.get("source") and e["source"] in
                          {c.get("id") for c in items})]
        a2, n2 = _edit_rows(extra, recon)
        return applied + a2, not_applied + n2
    if ap is not None:
        return _edit_rows(list(edits.values()), recon)
    # Verify-only mode: no apply outcomes, so the reconciliations are the ledger.
    return _edit_rows([dict(r) for r in recon.values()], recon)


def _synthetic(edit_id: str) -> bool:
    """An outcome the apply added on its own (an emptied paragraph's removal, a
    companion format edit) — accounted for on the edit it rode along with."""
    return edit_id.endswith("-para") or edit_id.endswith("-fmt")


def _edits_by_id(payload: dict) -> dict[str, dict]:
    """Every edit the payload knows about, by id, each row carrying its apply
    status. Applied first so a flagged duplicate id never masks a landing."""
    ap = payload.get("apply") or {}
    out: dict[str, dict] = {}
    for key in ("applied_items", "flagged", "no_op"):
        for o in ap.get(key) or []:
            out.setdefault(o["id"], o)
    return out


def _page_cells(row: dict) -> dict:
    """The three page columns for a row carrying `page` and maybe `page_label`.
    Never invents a page: no page and no label reads as blank / blank / —."""
    page = row.get("page") or 0
    label = row.get("page_label")
    if label:
        return {"Page (as marked)": page or "", "Page in IDML": label,
                "Page confirmed": "Yes"}
    if page:
        return {"Page (as marked)": page, "Page in IDML": page,
                "Page confirmed": "No"}
    return {"Page (as marked)": "", "Page in IDML": "", "Page confirmed": "—"}


def _applied_status(edit_ids, recon: dict[str, dict]) -> str:
    """"Applied exactly" unless verify found any of the edits carried out
    differently than asked; plain "Applied" when verify has no word on it."""
    statuses = [recon[e]["status"] for e in edit_ids if e in recon]
    if any(s == DEVIATES for s in statuses):
        return STATUS_APPLIED_DEVIATES
    if statuses and all(s == APPLIED_EXACTLY for s in statuses):
        return STATUS_APPLIED_EXACTLY
    return STATUS_APPLIED


def _flag_status(o: dict | None, fallback: str = STATUS_FLAGGED) -> str:
    if o is None:
        return fallback
    status = o.get("status") or ""
    return FLAG_TITLES.get(status) or (status.replace("_", " ").capitalize()
                                        if status else fallback)


def _join(values) -> str:
    seen = [v for v in values if v]
    return " | ".join(dict.fromkeys(seen))


def _comment_rows(items: list[dict], edits: dict[str, dict],
                  recon: dict[str, dict]) -> tuple[list[dict], list[dict], set]:
    applied: list[dict] = []
    not_applied: list[dict] = []
    covered: set[str] = set()
    for c in items:
        ids = [e for e in (c.get("edit_ids") or []) if e]
        covered.update(ids)
        its = [edits[e] for e in ids if e in edits]
        row = {**_page_cells(c),
               "Correction (as written)": c.get("instruction") or "",
               "Marked text": c.get("anchor") or "",
               "Find": _join(e.get("find", "") for e in its),
               "Replace": _join(e.get("replace", "") for e in its),
               "Edit id(s)": ", ".join(ids),
               "Comment id": c.get("id") or ""}
        disp = c.get("disposition")
        detail = c.get("detail") or ""
        if c.get("dismissed"):
            row.update({"Status": STATUS_SET_ASIDE, "Reason / detail": detail})
            not_applied.append(row)
        elif c.get("resolved"):
            row.update({"Status": STATUS_RESOLVED, "Reason / detail": detail})
            applied.append(row)
        elif disp == DISP_APPLIED:
            row.update({"Status": _applied_status(ids, recon),
                        "Reason / detail": _join(
                            [detail] + [recon[e].get("detail", "") for e in ids
                                        if e in recon
                                        and recon[e]["status"] != APPLIED_EXACTLY])})
            applied.append(row)
        elif disp == DISP_NO_OP:
            row.update({"Status": STATUS_NO_OP, "Reason / detail": detail})
            not_applied.append(row)
        elif disp == DISP_NOT_EXTRACTED:
            row.update({"Status": STATUS_NOT_EXTRACTED,
                        "Reason / detail": detail})
            not_applied.append(row)
        else:                                        # DISP_FLAGGED or unknown
            flagged = [e for e in its if e.get("status") not in (None, "applied")]
            o = flagged[0] if flagged else None
            row.update({"Status": _flag_status(o),
                        "Reason / detail": _join(
                            [o.get("detail", "") if o else "", detail])})
            not_applied.append(row)
    return applied, not_applied, covered


def _edit_rows(outcomes: list[dict],
               recon: dict[str, dict]) -> tuple[list[dict], list[dict]]:
    """One row per apply outcome (or, in verify mode, per reconciliation)."""
    applied: list[dict] = []
    not_applied: list[dict] = []
    for o in outcomes:
        eid = o.get("id") or ""
        row = {**_page_cells(o),
               "Correction (as written)": o.get("instruction") or "",
               "Marked text": "",
               "Find": o.get("find") or "",
               "Replace": o.get("replace") or "",
               "Edit id(s)": eid,
               "Comment id": o.get("source") or ""}
        if o.get("format"):
            row["Replace"] = (row["Replace"] + f"  [set {o['format']}]").strip()
        status = o.get("status") or ""
        detail = o.get("detail") or ""
        if o.get("dismissed"):
            row.update({"Status": STATUS_SET_ASIDE, "Reason / detail": detail})
            not_applied.append(row)
        elif o.get("resolved"):
            row.update({"Status": STATUS_RESOLVED, "Reason / detail": detail})
            applied.append(row)
        elif status == "applied":
            r = recon.get(eid)
            row.update({"Status": _applied_status([eid], recon),
                        "Reason / detail": _join(
                            [detail, r.get("detail", "")
                             if r and r["status"] != APPLIED_EXACTLY else ""])})
            applied.append(row)
        elif status in (APPLIED_EXACTLY, DEVIATES):    # verify-mode rows
            row.update({"Status": (STATUS_APPLIED_EXACTLY
                                   if status == APPLIED_EXACTLY
                                   else STATUS_APPLIED_DEVIATES),
                        "Reason / detail": detail})
            applied.append(row)
        elif status == MISSING:
            row.update({"Status": STATUS_MISSING, "Reason / detail": detail})
            not_applied.append(row)
        elif status == "no_op":
            row.update({"Status": STATUS_NO_OP, "Reason / detail": detail})
            not_applied.append(row)
        else:
            row.update({"Status": _flag_status(o), "Reason / detail": detail})
            not_applied.append(row)
    return applied, not_applied


def _check_total(payload: dict, total: int) -> None:
    """The reviewer's own count is the one the sheets must add up to. A mismatch
    is a bug worth hearing about, not one worth sinking a finished run over."""
    com = payload.get("comments") or {}
    expected = com.get("total") or 0
    if expected and total < expected:
        log.warning("Corrections spreadsheet lists %d row(s) but the run "
                    "received %d reviewer comment(s) — a comment is missing "
                    "from the ledger.", total, expected)
    elif expected and total > expected:
        log.info("Corrections spreadsheet lists %d row(s) for %d reviewer "
                 "comment(s): %d edit(s) arrived without a comment.",
                 total, expected, total - expected)
