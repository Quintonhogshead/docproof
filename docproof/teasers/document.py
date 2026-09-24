"""One editable author document: the five options, then the two-page dos and don'ts."""
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from . import AUTHOR_WARNING, guide

GUIDE_TITLE = "Dos and don’ts for your teaser"


def _shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _cell(cell, text, *, bold=False, size=9.5):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def _table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Inches(1.25), Inches(2.8), Inches(2.8))
    for cell, heading, fill in zip(table.rows[0].cells, ("", "Do", "Don’t"), ("FFFFFF", "E3F1E6", "F8E4E1")):
        _cell(cell, heading, bold=True, size=10)
        _shade(cell, fill)
    for focus, do, dont in rows:
        cells = table.add_row().cells
        _cell(cells[0], focus, bold=True)
        _cell(cells[1], do)
        _cell(cells[2], dont)
    # Word honours cell widths, Google Docs and LibreOffice the grid: set both.
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    grid = table._tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(width.twips)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    return table


def write_document(path: Path, draft, *, book_label: str):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ("Normal", "Body Text"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 24), ("Subtitle", 12), ("Heading 1", 18), ("Heading 2", 13)):
        doc.styles[name].font.name = "Arial"
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string("000000")
        doc.styles[name].paragraph_format.keep_with_next = True
        for border in doc.styles[name].element.iter(qn("w:pBdr")):
            border.getparent().remove(border)
    doc.add_heading(draft.title or book_label, 0)
    if draft.author:
        doc.add_paragraph(draft.author, "Subtitle")
    doc.add_paragraph("Five back-cover teaser options", "Subtitle")
    doc.add_paragraph(AUTHOR_WARNING)
    doc.add_paragraph("Choose and adapt any of these five options. The two pages of dos and don’ts "
                      "at the end of this document can guide your revision.")
    for teaser in sorted(draft.teasers, key=lambda t: t.number):
        doc.add_page_break()
        doc.add_heading(f"Option {teaser.number}", 1)
        doc.add_paragraph(teaser.angle, "Subtitle")
        for paragraph in teaser.paragraphs:
            p = doc.add_paragraph(paragraph)
            p.paragraph_format.keep_together = True

    doc.add_page_break()
    doc.add_heading(GUIDE_TITLE + ": the story", 1)
    doc.add_paragraph(guide.INTRODUCTION)
    _table(doc, guide.STORY)

    doc.add_page_break()
    doc.add_heading(GUIDE_TITLE + ": the writing", 1)
    _table(doc, guide.CRAFT)
    doc.add_heading("Applying the guidance", 2)
    for label, text in guide.APPLYING:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.add_run(label + ". ").bold = True
        p.add_run(text).font.size = Pt(10)

    doc.core_properties.title = (draft.title or book_label) + " — Author teasers"
    doc.core_properties.author = "DocProof"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
