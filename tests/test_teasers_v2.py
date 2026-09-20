"""Contract, isolation, recovery and delivery tests for the promoted teaser workflow."""
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
import hashlib
import json
import threading

import pytest
from docx import Document

from test_teasers import queued, story, draft, approved
from app.teasers import accept_story, generate_draft, accept_review, Queue
from docproof.providers.base import ProviderResult
from docproof.teasers import AUTHOR_WARNING, DEEPSEEK_MODEL
from docproof.teasers.models import OptionBrief, Draft, Teaser, digest, draft_issues
from docproof.teasers.facts import writer_prompt


@pytest.fixture
def v2(queued, story, draft):
    queue, job, task = queued
    task['version'] = 3
    queue.save(task)
    story.writer_brief.author_copy = Draft(teasers=[], opening_hooks=[], editorial_note='',
        elements=[], best_practices=[], modification_checklist=[])
    story.writer_brief.option_briefs = [OptionBrief(number=n, angle=f'Angle {n}',
        facts=['Mara returns to repair her father’s ferry.', 'Her brother wants to sell it.'],
        direction='Focus on the siblings and their competing plans.') for n in range(1, 6)]
    story.protected_revelations = ['PRIVATE_ENDING_SENTINEL']
    task = accept_story(queue, task, story.model_dump())
    task['feedback'] = ['PRIVATE_REJECTED_COPY_SENTINEL']
    queue.save(task)
    outputs = {}
    for t in draft.teasers:
        words = ' '.join(t.paragraphs).split()
        outputs[t.number] = Teaser(number=t.number, angle=f'Angle {t.number}',
            paragraphs=[' '.join(words[:55]), ' '.join(words[55:110]), ' '.join(words[110:])])
    return queue, job, task, outputs


def test_enqueue_when_formatting_starts_and_dedupe_completion(queued):
    queue, job, old = queued
    fresh = replace(job, id='new-format', state='running')
    task_id = queue.add(fresh)
    assert queue.get(task_id)['version'] == 3
    assert queue.get(task_id)['state'] == 'queued'
    assert queue.add(replace(fresh, state='done')) == task_id
    assert queue.add(replace(fresh, id='unstarted', state='queued')) is None


def test_five_concurrent_public_only_calls_and_hash_bound_approval(v2):
    queue, job, task, outputs = v2
    barrier = threading.Barrier(5)
    seen = []
    class Provider:
        def complete_structured(self, **kw):
            assert kw['model'] == DEEPSEEK_MODEL
            assert 'PRIVATE_' not in kw['system'] + kw['user']
            payload = json.loads(kw['user'])
            assert set(payload) == {'number', 'angle', 'selected_facts', 'direction'}
            seen.append(payload['number'])
            barrier.wait(timeout=5)
            return ProviderResult(parsed=outputs[payload['number']].model_dump())
    task = generate_draft(queue, task, provider=Provider())
    assert task['state'] == 'drafted'
    assert sorted(seen) == [1, 2, 3, 4, 5]
    result = Draft.model_validate(task['drafts'][-1]['content'])
    assert result.version == 2 and not draft_issues(result)
    task = accept_review(queue, task, approved(result).model_dump())
    assert task['state'] == 'approved'


def test_partial_failure_retries_only_missing_option(v2):
    queue, job, task, outputs = v2
    calls = []
    class Provider:
        def complete_structured(self, **kw):
            number = json.loads(kw['user'])['number']
            calls.append(number)
            if number == 3 and calls.count(3) == 1:
                return ProviderResult(parsed=None, stop_reason='error', error='temporary outage')
            return ProviderResult(parsed=outputs[number].model_dump())
    provider = Provider()
    task = generate_draft(queue, task, provider=provider)
    assert task['state'] == 'retry_wait'
    assert len(task['option_results']) == 4
    task['retry_at'] = 0
    queue.save(task)
    queue.recover()
    task = generate_draft(queue, queue.get(task['id']), provider=provider)
    assert task['state'] == 'drafted'
    assert calls.count(3) == 2 and len(calls) == 6


@pytest.mark.parametrize('count,paragraphs,valid', [(149,3,False),(150,3,True),(200,3,True),(201,3,False),(180,2,False),(180,4,False)])
def test_word_and_paragraph_contract(count, paragraphs, valid):
    from docproof.teasers.models import teaser_issues
    words = ['word'] * count
    parts = [' '.join(words[i*count//paragraphs:(i+1)*count//paragraphs]) for i in range(paragraphs)]
    assert (not teaser_issues(Teaser(number=1,angle='Test',paragraphs=parts),version=2)) == valid


def test_v1_hash_stays_compatible(draft):
    raw = draft.model_dump()
    raw.pop('version')
    expected = hashlib.sha256(json.dumps(raw,sort_keys=True,ensure_ascii=False,separators=(',',':')).encode()).hexdigest()
    assert digest(draft) == expected


def test_author_document_contains_exact_warning_and_unranked_options(v2, tmp_path):
    from docproof.teasers.document import write_document
    from docproof.teasers.models import Storysheet
    queue, job, task, outputs = v2
    result = Draft(version=2, teasers=list(outputs.values()), opening_hooks=[], editorial_note='',
                   elements=[], best_practices=[], modification_checklist=[])
    path = write_document(tmp_path/'teasers.docx',Storysheet.model_validate(task['storysheet']),
                          result,approved(result),book_label=task['book_label'])
    text = '\n'.join(p.text for p in Document(path).paragraphs)
    assert AUTHOR_WARNING in text
    assert 'Recommended' not in text
    assert text.index('Option 1') < text.index('Option 2') < text.index('Option 5')
    assert 'PRIVATE_' not in text


def test_delivery_waits_for_both_guides_and_reuses_saved_ids(v2, monkeypatch):
    from app import teaser_delivery as delivery
    queue, job, task, outputs = v2
    task['state'] = 'approved'
    queue.save(task)
    uploaded = []
    records = {}
    def upload(token, folder_id, path, **kwargs):
        file_id = Path(path).suffix[1:]
        uploaded.append(file_id)
        records[file_id] = {'mimeType':kwargs['mime_type'], 'parents':[folder_id],
            'md5Checksum':hashlib.md5(Path(path).read_bytes()).hexdigest(),
            'webViewLink':'https://drive.google.com/file/d/'+file_id}
        return file_id
    monkeypatch.setattr(delivery.drive, 'search_files', lambda *a,**k:[])
    monkeypatch.setattr(delivery.drive, 'upload', upload)
    monkeypatch.setattr(delivery.drive, '_json_call', lambda request,**k: records[request.full_url.split('/files/')[1].split('?')[0]])
    delivery.deliver_guides(queue, task, 'token', 'folder')
    delivery.deliver_guides(queue, task, 'token', 'folder')
    assert uploaded == ['xlsx','pdf']
    assert task['guide_url'].endswith('/pdf')
    records['pdf']['md5Checksum'] = 'wrong'
    with pytest.raises(ValueError, match='complete two-page guide'):
        delivery.deliver_guides(queue, task, 'token', 'folder')


def test_unfinished_production_v2_migrates_without_erasing_audit(queued):
    queue, job, task = queued
    task.update(version=2, storysheet={"private": "original brief"}, drafts=[{"old": "draft"}], reviews=[])
    queue.save(task, "story_ready")
    migrated = queue.claim("worker")
    assert migrated["version"] == 3 and migrated["state"] == "queued"
    assert "storysheet" not in migrated and migrated["drafts"] == []
    assert migrated["prior_workflows"][0]["drafts"] == [{"old":"draft"}]


def test_folder_search_rejects_case_insensitive_name_match(queued, monkeypatch):
    from app import teaser_delivery as delivery
    queue, job, task = queued
    old = SimpleNamespace(id="old", name="Author teasers", is_folder=True)
    new = SimpleNamespace(id="new", name="author teasers", is_folder=True)
    queue.configure(folder_id="old")
    monkeypatch.setattr(delivery.drive,"get_file",lambda token, file_id, **kw:old if file_id=="old" else new)
    monkeypatch.setattr(delivery.drive,"search_files",lambda *a,**k:[old])
    created=[]
    def create(token,parent,name,**kwargs):
        created.append(name)
        return "new"
    monkeypatch.setattr(delivery.drive,"create_folder",create)
    assert delivery.ensure_folder(queue,"token") == "new"
    assert created == ["author teasers"]
    assert delivery.ensure_folder(queue,"token") == "new"
    assert created == ["author teasers"]
