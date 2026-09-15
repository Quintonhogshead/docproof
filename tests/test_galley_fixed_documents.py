"""Native Word round trips and source-bound fixed package recovery, offline."""
from __future__ import annotations

import base64
import copy
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.shared import Pt

from docproof.cleancopy import has_markup
from docproof.providers import NormalizedUsage, ProviderResult
from docproof.utils.xml_helpers import DocxPackage, qn, walk_package
from galley import fixed_documents as fd
from galley.driver import Driver, seed_workspace
from galley.fixed_calls import FixedCalls, FixedCallError
from galley.fixed_policy import configuration
from galley.fixed_workflow import workflow_plan
from galley.manifest import sha256_file


@pytest.fixture
def manuscript(tmp_path):
    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("Ann")
    first.font.name, first.font.size, first.bold = "Arial", Pt(16), True
    paragraph.add_run(" met Ann by teh harbor.")
    document.add_comment(first, text="Keep this author note.", author="Original Author", initials="OA")
    title = document.add_paragraph("She read Aeneid.")
    document.add_paragraph("Already italic").runs[0].italic = True
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Teh table entry."
    document.sections[0].header.paragraphs[0].text = "Teh running head"
    document.add_picture(BytesIO(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGOoqKgA"
        "AALUAWneSzH7AAAAAElFTkSuQmCC")))
    path = tmp_path / "Writer - Book 1.docx"
    document.save(path)
    return path


def corrected(source):
    return {pid: text.replace("Ann met", "Anne met").replace("teh", "the").replace("Teh", "The")
            for pid, text in fd.paragraph_views(source).items()}


def test_native_output_preserves_source_styles_tables_images_and_comments(manuscript, tmp_path):
    original = fd.paragraph_views(manuscript)
    accepted = corrected(manuscript)
    body = next(pid for pid, text in original.items() if text.startswith("Ann met"))
    title = next(pid for pid, text in original.items() if "Aeneid" in text)
    already = next(pid for pid, text in original.items() if text == "Already italic")
    question = {"id": "author-q", "para_id": body, "quote": "Ann", "occurrence": 2, "question": "Which Ann is present here?",
                "missing_knowledge": "Identity of the second Ann"}
    formats = [{"para_id": title, "snapshot": original[title], "start": 9, "end": 15,
                "format": "italic", "reason": "Book title."},
               {"para_id": already, "snapshot": original[already], "start": 0, "end": len(original[already]),
                "format": "italic", "reason": "Already correctly set."}]
    source_bytes = manuscript.read_bytes()
    tracked, clean, findings = fd.write_manuscripts(manuscript, tmp_path / "final", accepted, [question], formats)
    assert manuscript.read_bytes() == source_bytes
    assert fd.paragraph_views(tracked, "reject") == original
    assert fd.paragraph_views(tracked) == fd.paragraph_views(clean) == accepted
    assert has_markup(tracked) and not has_markup(clean)
    assert len(Document(clean).inline_shapes) == 1
    assert Document(clean).tables[0].cell(0, 0).text == "The table entry."
    assert Document(clean).sections[0].header.paragraphs[0].text == "The running head"
    first = Document(clean).paragraphs[0].runs[0]
    assert first.font.name == "Arial" and first.font.size.pt == 16 and first.bold
    assert next(r for r in Document(clean).paragraphs[1].runs if r.text == "Aeneid").italic
    assert next(row for row in findings if row.get("already_set"))["applied"] is False
    source_comments = fd._comments(DocxPackage(manuscript))
    final_comments = fd._comments(DocxPackage(tracked))
    assert all(final_comments[key] == value for key, value in source_comments.items())
    assert len(final_comments) == len(source_comments) + 1
    assert fd._comments(DocxPackage(clean)) == {}
    query = next(row for row in findings if row["finding_id"] == "author-q")
    assert query["anchor"]["start"] == original[body].index("Ann", 1)
    assert query["occurrence"] == 2
    from galley.comment_reconcile import _comment_ranges
    new_comment_id = next(iter(set(final_comments) - set(source_comments)))
    paragraph = next(p.element for p in walk_package(DocxPackage(tracked)) if p.para_id == body)
    assert _comment_ranges(paragraph)[new_comment_id] == (query["anchor"]["start"], query["anchor"]["end"])
    with ZipFile(manuscript) as source, ZipFile(tracked) as final:
        for name in source.namelist():
            if name.startswith("word/media/") or name in {"word/styles.xml", "word/numbering.xml"}:
                assert source.read(name) == final.read(name)


def test_final_question_on_inserted_word_maps_back_without_losing_text(manuscript, tmp_path):
    accepted = fd.paragraph_views(manuscript)
    pid = next(pid for pid, text in accepted.items() if text.startswith("Ann met"))
    accepted[pid] = accepted[pid].replace("met", "first met")
    question = {"id": "first-q", "para_id": pid, "quote": "first", "question": "Was this their first meeting?",
                "missing_knowledge": "Earlier meetings"}
    tracked, clean, findings = fd.write_manuscripts(manuscript, tmp_path / "inserted", accepted, [question])
    assert fd.paragraph_views(tracked) == fd.paragraph_views(clean) == accepted
    assert next(x for x in findings if x["finding_id"] == "first-q")["queried"]


@pytest.mark.parametrize("defect", ["missing_paragraph", "unknown_format", "unanchored_comment"])
def test_unwritable_changes_fail_instead_of_certifying_partial_output(manuscript, tmp_path, defect):
    accepted = fd.paragraph_views(manuscript)
    pid = next(iter(accepted))
    questions, formats = [], []
    if defect == "missing_paragraph":
        accepted.pop(pid)
    elif defect == "unknown_format":
        formats = [{"para_id": pid, "snapshot": accepted[pid], "start": 0, "end": 3,
                    "format": "unbold", "reason": "unsupported"}]
    else:
        questions = [{"id": "bad", "para_id": pid, "quote": "not in this paragraph",
                      "question": "Who?", "missing_knowledge": "Identity"}]
    with pytest.raises(ValueError):
        fd.write_manuscripts(manuscript, tmp_path / "refused", accepted, questions, formats)


class OfflineReader:
    def complete_structured(self, **kwargs):
        return ProviderResult(parsed={"ok": True}, usage=NormalizedUsage(input_tokens=10, output_tokens=2, billed=False))


@pytest.fixture
def completed(manuscript, tmp_path):
    workspace = seed_workspace(manuscript, "writer", workspace_root=tmp_path / "work")
    driver = Driver(manuscript, "writer", workspace_root=tmp_path / "work", execution_mode="fixed", source_id="drive-source")
    directory = workspace / "runs/fixed"
    cfg = configuration(True)
    identity = {"source_sha256": sha256_file(manuscript), "configuration": cfg.model_dump(mode="json"), "recipe": workflow_plan()}
    calls = FixedCalls(directory / "calls", identity, cfg, provider_factory=lambda *a, **kw: OfflineReader())
    schema = {"type": "object", "properties": {"ok": {"type": "boolean"}}, "required": ["ok"], "additionalProperties": False}
    for stage in ("poetry", "spelling"):
        calls.ask(stage, model="claude-sonnet-5", system="Offline proofread fixture", user="Read this fixture.", schema=schema, schema_name="fixture")
    accepted = corrected(manuscript)
    stages = []
    for stage in ("poetry", "typed", "poetry_complete"):
        payload = {"stage": stage, "accepted_sha256": fd._hash(accepted), "questions": [], "evidence": {"fixture": True}}
        path = directory / "stages" / f"{stage}.json"
        path.parent.mkdir(exist_ok=True)
        fd._save(path, payload)
        stages.append({"stage": stage, "path": str(path), "sha256": fd._hash(payload)})
    result = {"identity": identity, "execution_mode": "fixed", "status": "completed", "source": str(manuscript),
              "original": fd.paragraph_views(manuscript), "accepted": accepted, "questions": [], "formats": [],
              "history": [], "stages": stages, "poetry_only": True, "editorial_verdict": "ready", "usage": calls.usage_summary()}
    result["result_sha256"] = fd._hash({k: v for k, v in result.items() if k != "usage"})
    fd._save(directory / "result.json", result)
    fd._save(directory / "workflow.json", {"identity": identity, "execution_mode": "fixed", "status": "completed", "result_sha256": result["result_sha256"]})
    return driver, result


def test_package_has_watch_compatible_names_and_reuses_exact_bytes(completed, monkeypatch):
    driver, result = completed
    package = fd.package_result(driver, result)
    assert package["source_id"] == "drive-source" and package["outcome"] == "done" and package["reason"]
    # A pre-proofread goes to a proofreader, not to the author and not to
    # HubSpot: no clean reading copy, and no outcome.json for DocWatch to
    # commit on.
    assert {row["name"] for row in package["artifacts"]} == {
        "Writer - Book 2 - Pre-Proofread.docx", "Writer - Book 2 - proofreading report.md",
        "Writer - Book 2 - review evidence.json", "Writer - Book 2 - fixed certificate.json"}
    assert fd.validate_delivery_package(package)["delivery_ready"] is True
    original = {row["path"]: Path(row["path"]).read_bytes() for row in package["artifacts"]}
    monkeypatch.setattr(fd, "write_manuscripts", lambda *a, **kw: pytest.fail("rebuilt completed documents"))
    assert fd.package_result(driver, result) == package
    assert all(Path(path).read_bytes() == contents for path, contents in original.items())


@pytest.mark.parametrize("target", ["source", "stage", "workflow", "result", "response", "missing_call", "tracked", "clean", "certificate", "handoff", "missing_artifact", "wrong_outcome"])
def test_package_tampering_prevents_delivery(completed, target):
    driver, result = completed
    package = fd.package_result(driver, result)
    directory = driver.workspace / "runs/fixed"
    if target == "missing_artifact":
        package["artifacts"].pop()
    elif target == "wrong_outcome":
        package["outcome"] = "needs_human"
    elif target == "missing_call":
        next((directory / "calls/calls").glob("*/receipt.json")).unlink()
    else:
        certificate = json.loads(Path(package["certificate"]).read_text())
        path = {"source": driver.book, "stage": Path(result["stages"][0]["path"]),
                "workflow": directory / "workflow.json", "result": directory / "result.json",
                "response": next((directory / "calls/calls").glob("*/attempts/*/response.json")),
                "tracked": Path(certificate["tracked"]["path"]), "clean": Path(certificate["clean"]["path"]),
                "certificate": Path(package["certificate"]), "handoff": Path(package["artifacts"][0]["path"])}[target]
        path.write_bytes(path.read_bytes() + b"changed")
    with pytest.raises((ValueError, FixedCallError, OSError)):
        fd.validate_delivery_package(package)


def test_report_describes_runover_joins_and_not_revisions(manuscript):
    original = fd.paragraph_views(manuscript)
    first = next(iter(original))
    result = {"identity": {"intake": {"version": "fixed-intake-v3"}}, "source": str(manuscript),
              "poetry_only": False, "questions": [], "stages": [], "history": []}
    receipt = {"resolved_revision_elements": {}, "runover_joins": [
        {"para_id": "body-0000", "baseline_para_id": first, "absorbed": ["body-0001"], "seam_offsets": [70]}]}
    report = fd._report(result, [], receipt)
    assert "## Page-runover paragraphs joined at intake" in report
    assert "1 paragraphs that the typeset export had split across page boundaries (1 continuation lines)" in report
    assert "- Paragraph 1: joined 1 continuation line(s) (seam at 70)" in report
    assert "Incoming tracked changes" not in report
    revised = fd._report(result, [], {"resolved_revision_elements": {"word/document.xml": 2}, "runover_joins": []})
    assert "Incoming tracked changes" in revised and "Page-runover" not in revised


def test_a_question_about_a_running_head_is_carried_by_the_first_body_paragraph(tmp_path):
    """Word keeps comments in the body only; a running-head question must not
    block the delivery of 569 good corrections (Wilder, 2026-09-14)."""
    from docx import Document
    from docproof.utils.xml_helpers import DocxPackage, qn
    source = tmp_path / "Writer - Book 1.docx"
    document = Document()
    document.add_paragraph("PROLOGUE")
    document.add_paragraph("The boat arrived at dawn.")
    document.sections[0].header.paragraphs[0].text = "CHAPTER ONE"
    document.save(source)
    accepted = fd.paragraph_views(source)
    header = next(pid for pid, text in accepted.items() if text == "CHAPTER ONE")
    first = next(pid for pid, text in accepted.items() if text == "PROLOGUE")
    assert header.startswith("header")
    question = {"id": "q-head", "para_id": header, "quote": "CHAPTER ONE",
                "question": "Should this be PROLOGUE?", "missing_knowledge": "The intended label"}
    tracked, clean, details = fd.write_manuscripts(source, tmp_path / "out", accepted, [question])
    row = next(x for x in details if x["finding_id"] == "q-head")
    assert row["queried"] and row["relocated_from"] == header and row["para_id"] == first
    comments = DocxPackage(tracked).tree("word/comments.xml")
    texts = ["".join(t.text or "" for t in c.iter(qn("w:t"))) for c in comments if c.tag == qn("w:comment")]
    assert texts and "About the running head \u201cCHAPTER ONE\u201d: Should this be PROLOGUE?" in texts[0]
    assert fd.paragraph_views(clean) == accepted
    result = {"identity": {}, "source": str(source), "poetry_only": False, "questions": [question],
              "stages": [], "history": []}
    report = fd._report(result, details)
    assert "because Word keeps comments in the body" in report
