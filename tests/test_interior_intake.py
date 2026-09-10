from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfWriter

from docproof.interior.intake import build_packet


def _commented_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("A paragraph outside the table.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Old wording"
    table.cell(0, 1).text = "Additional Notes: blank pages should be consistent"
    document.save(path)
    # Add a real comments part and an anchor in the table-cell paragraph.  The
    # intake reader works at package level so table comments are retained even
    # when python-docx cannot author them.
    with ZipFile(path, "r") as source:
        members = {name: source.read(name) for name in source.namelist()}
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    import xml.etree.ElementTree as ET
    xml = ET.fromstring(members["word/document.xml"])
    cell_paragraph = xml.findall(".//w:tbl//w:p", ns)[0]
    start = ET.Element(qn("w:commentRangeStart"), {qn("w:id"): "4"})
    end = ET.Element(qn("w:commentRangeEnd"), {qn("w:id"): "4"})
    cell_paragraph.insert(0, start)
    cell_paragraph.append(end)
    members["word/document.xml"] = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
    comments = ET.Element(qn("w:comments"), {"xmlns:w": ns["w"]})
    comment = ET.SubElement(comments, qn("w:comment"), {qn("w:id"): "4", qn("w:author"): "Terry"})
    paragraph = ET.SubElement(comment, qn("w:p"))
    run = ET.SubElement(paragraph, qn("w:r"))
    text = ET.SubElement(run, qn("w:t"))
    text.text = "Please check this table row."
    members["word/comments.xml"] = ET.tostring(comments, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w") as target:
        for name, data in members.items():
            target.writestr(name, data)


def test_docx_tables_comments_and_text_are_complete(tmp_path):
    docx = tmp_path / "review.docx"
    _commented_docx(docx)
    packet = build_packet([docx], "Additional Notes: keep the source wording", tmp_path / "work")
    source = packet["attachments"][0]
    assert source["tables"]
    assert any("blank pages" in cell for row in source["tables"][0]["rows"] for cell in row)
    assert source["comments"][0]["text"] == "Please check this table row."
    assert source["comments"][0]["anchors"]
    assert any(row["comments"] for row in source["paragraphs"])
    assert packet["text_source_id"] in packet["source_ids"]
    assert packet["text_source_id"] in {row["source_id"] for row in packet["evidence"]}


def test_supplied_text_decodes_entities_without_losing_raw_text(tmp_path):
    packet = build_packet([], "Thanks! <-- Terry Ardell&#x20;", tmp_path / "work")
    assert packet["text"].endswith("Ardell ")
    assert packet["text_raw"].endswith("Ardell&#x20;")
    source = packet["sources"][0]
    assert source["text"] == packet["text"]
    assert source["raw_text"] == packet["text_raw"]


def test_docx_struck_deletion_and_explicit_false_survive_intake(tmp_path):
    path = tmp_path / 'corrections.docx'
    document = Document()
    paragraph = document.add_paragraph('Keep this ')
    paragraph.add_run('delete this').font.strike = True
    paragraph.add_run('keep this too').font.strike = False
    document.save(path)
    packet = build_packet([path], '', tmp_path / 'work')
    runs = next(row['runs'] for row in packet['evidence'] if row['kind'] == 'docx_paragraph')
    assert runs[1]['formatting']['strike'] is True
    assert runs[2]['formatting']['strike'] is False
    assert 'strike' in runs[1]['formatting_xml']


def test_flattened_pdf_keeps_every_page_as_visual_evidence(tmp_path):
    pdf = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    with pdf.open("wb") as stream:
        writer.write(stream)
    packet = build_packet([pdf], "", tmp_path / "work")
    source = packet["attachments"][0]
    assert source["page_count"] == 2
    assert [row["page_index"] for row in source["pages"]] == [0, 1]
    assert len([row for row in packet["evidence"] if row["kind"] == "pdf_page"]) == 2
    assert all("note" in row for row in packet["evidence"] if row["kind"] == "pdf_page")
