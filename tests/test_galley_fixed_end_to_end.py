"""Real orchestration, receipts and Word delivery with offline model endpoints."""
from __future__ import annotations

import json
import re

import pytest
from docx import Document

from docproof.providers import NormalizedUsage, ProviderResult
from galley import codex_runner, driver as gd, fixed_calls as fc
from galley.fixed_documents import paragraph_views, validate_delivery_package
from galley.fixed_workflow import ASTRA, FABLE, LUNA, OPUS, SOL, SONNET


class ScriptedReaders:
    def __init__(self, poetry):
        self.poetry = poetry
        self.requests = []

    def answer(self, model, user, schema):
        self.requests.append((model, user))
        fields = schema["properties"]
        if "classification" in fields:
            return {"classification": "poetry" if self.poetry else "prose", "reason": "Fixture classification."}
        if "narration" in fields:
            return {"narration": "Third person, past tense.", "characters": [], "notes": []}
        if "reviewed_paragraph_ids" in fields:
            paragraphs = re.findall(r'<paragraph id="([^"]+)">\n(.*?)\n</paragraph>', user, re.S)
            item = fields["findings"]["items"]
            if "$ref" in item:
                item = schema["$defs"][item["$ref"].split("/")[-1]]
            props = item["properties"]
            kinds = props["error_type"].get("enum", [props["error_type"].get("const")])
            findings = []
            if "spelling" in kinds:
                for pid, text in paragraphs:
                    if "recieved" in text:
                        row = dict(para_id=pid, error_type="spelling", original_text="recieved",
                                   occurrence=1, corrected_text="received", confidence="high")
                        if "explanation" in props:
                            row["explanation"] = "Correct the misspelling."
                        findings.append(row)
            return {"findings": findings, "reviewed_paragraph_ids": [p[0] for p in paragraphs]}
        payload = json.loads(user)
        if "reviewed_ids" in fields:
            owned = payload.get("paragraphs", [])
            ids = ([x["id"] for x in payload["sites"]] if "sites" in payload else [x["id"] for x in owned])
            return {"reviewed_ids": ids, "findings": [], "comment_decisions": [], "editorial_verdict": "ready"}
        if "decisions" in fields:
            if "changes" in payload:
                return {"decisions": [{"id": x["id"], "verdict": "approve", "reason": "The spelling correction preserves meaning."}
                                       for x in payload["changes"]]}
            return {"decisions": [{"id": x["id"], "action": "drop", "replacement": "", "reason": "No clear error.",
                                    "missing_knowledge": "", "question": ""} for x in payload["sites"]]}
        raise AssertionError(f"Unexpected reader contract: {list(fields)}")

    def complete_structured(self, **request):
        body = self.answer(request["model"], request["user"], request["schema"])
        return ProviderResult(parsed=body, actual_model=request["model"],
                              usage=NormalizedUsage(input_tokens=100, output_tokens=30),
                              resource_usage={"input_tokens": 100, "output_tokens": 30,
                                              "cache_read_input_tokens": 0, "cache_creation_input_tokens": 0})

    def subscription(self, prompt, schema, directory, **options):
        assert options["no_tools"] is True
        assert options["model"] in {SOL, ASTRA}
        return self.answer(options["model"], prompt.rsplit("\n\n", 1)[-1], schema)


@pytest.mark.parametrize("poetry", [False, True])
def test_real_fixed_driver_delivers_and_resumes_without_new_generations(tmp_path, monkeypatch, poetry):
    source = tmp_path / "Writer - Book 1.docx"
    document = Document()
    document.add_paragraph("She recieved two letters.")
    document.save(source)
    readers = ScriptedReaders(poetry)
    monkeypatch.setattr(fc, "_default_provider", lambda *a, **k: readers)
    monkeypatch.setattr(codex_runner, "run_structured", readers.subscription)
    monkeypatch.setattr(gd, "spawn_claude", lambda *a, **k: pytest.fail("A supervising Brain started"))
    worker = gd.Driver(source, "writer", workspace_root=tmp_path / "work", execution_mode="fixed")
    result = worker.run()
    assert result.outcome == "done", result.reason
    package = json.loads((worker.workspace / "runs/driver/package.json").read_text())
    assert validate_delivery_package(package)["delivery_ready"] is True
    documents = [p for p in result.handoff if p.suffix == ".docx"]
    assert len(documents) == 2
    for path in documents:
        assert list(paragraph_views(path).values()) == ["She received two letters."]
    corrected = next(p for p in documents if "Clean" not in p.name)
    assert paragraph_views(corrected, "reject") == paragraph_views(source)
    models = {model for model, _ in readers.requests}
    assert models == ({SONNET} if poetry else {SONNET, LUNA, OPUS, SOL, FABLE, ASTRA})
    count = len(readers.requests)
    hashes = {x["name"]: x["sha256"] for x in package["artifacts"]}
    again = gd.Driver(source, "writer", workspace_root=tmp_path / "work", execution_mode=None).run()
    assert again.outcome == "done", again.reason
    assert len(readers.requests) == count
    reused = json.loads((worker.workspace / "runs/driver/package.json").read_text())
    assert {x["name"]: x["sha256"] for x in reused["artifacts"]} == hashes


def test_interrupted_stage_resumes_from_paid_read_receipts(tmp_path, monkeypatch):
    from galley.fixed_workflow import FixedWorkflow

    source = tmp_path / "Writer - Book 1.docx"
    document = Document()
    document.add_paragraph("She recieved two letters.")
    document.save(source)
    readers = ScriptedReaders(False)
    monkeypatch.setattr(fc, "_default_provider", lambda *a, **k: readers)
    monkeypatch.setattr(codex_runner, "run_structured", readers.subscription)
    stage_start = FixedWorkflow._stage

    def interrupt(self, stage):
        if stage == "fable":
            raise OSError("Simulated local interruption between completed reads")
        stage_start(self, stage)

    monkeypatch.setattr(FixedWorkflow, "_stage", interrupt)
    worker = gd.Driver(source, "writer", workspace_root=tmp_path / "work", execution_mode="fixed")
    stopped = worker.run()
    assert stopped.outcome == "blocked" and "Simulated local interruption" in stopped.reason
    count = len(readers.requests)
    assert count > 0 and FABLE not in {model for model, _ in readers.requests}
    monkeypatch.setattr(FixedWorkflow, "_stage", stage_start)
    resumed = gd.Driver(source, "writer", workspace_root=tmp_path / "work", execution_mode=None).run()
    assert resumed.outcome == "done", resumed.reason
    assert [model for model, _ in readers.requests[count:]] == [FABLE, ASTRA]
    package = json.loads((worker.workspace / "runs/driver/package.json").read_text())
    assert validate_delivery_package(package)["delivery_ready"] is True
