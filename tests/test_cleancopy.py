"""The clean copy: the tracked-changes proofread with every change accepted and
every comment removed, derived at hand-off (docproof/cleancopy.py).

What is held here: the clean text is exactly the accepted view the verifiers
read; a deleted paragraph mark joins paragraphs the way Word does; an inserted
mark (a speaker split) leaves the paragraph standing; comments vanish from the
text AND from the comments part; a file with no markup round-trips unchanged;
and a file that is not a .docx is refused by name.
"""
from __future__ import annotations

import zipfile
from pathlib import Path

import docx
import pytest

from docproof.cleancopy import (CleanCopyError, accept_all, has_markup,
                                strip_comments, write_clean_copy)
from docproof.reassembler import paragraph_view_text
from docproof.utils.xml_helpers import DocxPackage, paragraph_text, walk_package

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
REV = 'w:author="Atmosphere Press Proofreader" w:date="2026-09-08T00:00:00Z"'


def _run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _docx(path: Path, body: str, comments: str | None = None) -> Path:
    """A .docx from raw body XML, built on python-docx's package so styles,
    rels and content types are real; the comments part, if given, is wired in
    the way docproof.reassembler wires it."""
    d = docx.Document()
    d.add_paragraph("placeholder")
    d.save(str(path))
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                xml = data.decode("utf-8")
                start = xml.index("<w:body>") + len("<w:body>")
                end = xml.index("<w:sectPr")
                data = (xml[:start] + body + xml[end:]).encode("utf-8")
            elif comments and item.filename == "[Content_Types].xml":
                data = data.decode("utf-8").replace(
                    "</Types>",
                    '<Override PartName="/word/comments.xml" ContentType='
                    '"application/vnd.openxmlformats-officedocument.'
                    'wordprocessingml.comments+xml"/></Types>').encode()
            elif comments and item.filename == "word/_rels/document.xml.rels":
                data = data.decode("utf-8").replace(
                    "</Relationships>",
                    '<Relationship Id="rId99" Type="http://schemas.openxml'
                    'formats.org/officeDocument/2006/relationships/comments" '
                    'Target="comments.xml"/></Relationships>').encode()
            zout.writestr(item, data)
        if comments:
            zout.writestr("word/comments.xml",
                          f'<?xml version="1.0" encoding="UTF-8" standalone='
                          f'"yes"?><w:comments {W}>{comments}</w:comments>')
    tmp.replace(path)
    return path


def _texts(path: Path) -> list[str]:
    pkg = DocxPackage(str(path))
    return [paragraph_text(wp.element) for wp in walk_package(pkg)]


def _accepted_view(path: Path) -> list[str]:
    pkg = DocxPackage(str(path))
    return [paragraph_view_text(wp.element, "accept")
            for wp in walk_package(pkg)]


TRACKED = (
    "<w:p>" + _run("The cat ") +
    f'<w:del w:id="1" {REV}><w:r><w:delText>sat</w:delText></w:r></w:del>'
    f'<w:ins w:id="2" {REV}>{_run("sits")}</w:ins>' + _run(" on the mat.") +
    "</w:p>"
    "<w:p>" + _run("Plain ") +
    f'<w:r><w:rPr><w:i/><w:rPrChange w:id="3" {REV}><w:rPr/></w:rPrChange>'
    f'</w:rPr><w:t>italic</w:t></w:r>' + _run(" run.") + "</w:p>"
)


def test_the_clean_text_is_the_accepted_view(tmp_path):
    src = _docx(tmp_path / "tracked.docx", TRACKED)
    assert has_markup(src)
    expected = _accepted_view(src)
    out = write_clean_copy(src, tmp_path / "clean" / "clean.docx")
    assert out.is_file() and not has_markup(out)
    assert _texts(out) == expected == ["The cat sits on the mat.",
                                       "Plain italic run."]
    # The formatting the change record wrapped survives; only the record went.
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "<w:i/>" in xml and "rPrChange" not in xml and "delText" not in xml
    # It is a Word file Word will open.
    assert [p.text for p in docx.Document(str(out)).paragraphs] == expected


def test_a_deleted_paragraph_mark_joins_the_next_paragraph(tmp_path):
    body = (
        f'<w:p><w:pPr><w:pStyle w:val="Heading1"/><w:rPr><w:del w:id="1" '
        f'{REV}/></w:rPr></w:pPr>' + _run("First half,") + "</w:p>"
        '<w:p><w:pPr><w:pStyle w:val="Quote"/></w:pPr>' +
        _run(" second half.") + "</w:p>"
        "<w:p>" + _run("Untouched.") + "</w:p>"
    )
    src = _docx(tmp_path / "join.docx", body)
    out = write_clean_copy(src, tmp_path / "clean.docx")
    assert _texts(out) == ["First half, second half.", "Untouched."]
    # Word's rule: the surviving mark is the second paragraph's.
    pkg = DocxPackage(str(out))
    first = next(walk_package(pkg)).element
    style = first.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle")
    assert style.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "Quote"


def test_an_inserted_paragraph_mark_leaves_the_split_standing(tmp_path):
    # What a speaker split writes: the new paragraph's mark tagged inserted.
    body = (
        f'<w:p><w:pPr><w:rPr><w:ins w:id="1" {REV}/></w:rPr></w:pPr>' +
        _run("“Hello,” she said.") + "</w:p>"
        "<w:p>" + _run("“Hi,” he said.") + "</w:p>"
    )
    src = _docx(tmp_path / "split.docx", body)
    out = write_clean_copy(src, tmp_path / "clean.docx")
    assert _texts(out) == ["“Hello,” she said.", "“Hi,” he said."]
    assert not has_markup(out)


def test_comments_leave_the_text_and_the_comments_part(tmp_path):
    body = (
        '<w:p>' + _run("Is this ") +
        '<w:commentRangeStart w:id="0"/>' + _run("Kathryn") +
        '<w:commentRangeEnd w:id="0"/>'
        '<w:r><w:rPr><w:rStyle w:val="CommentReference"/></w:rPr>'
        '<w:commentReference w:id="0"/></w:r>' +
        _run(" or Katherine?") + "</w:p>"
    )
    comments = (f'<w:comment w:id="0" {REV} w:initials="dp"><w:p>' +
                _run("AU: the name is spelled both ways; which is right?") +
                "</w:p></w:comment>")
    src = _docx(tmp_path / "queried.docx", body, comments)
    out = write_clean_copy(src, tmp_path / "clean.docx")
    assert _texts(out) == ["Is this Kathryn or Katherine?"]
    assert not has_markup(out)
    z = zipfile.ZipFile(out)
    # The part stays (its relationship and content type still name it) but
    # carries nothing: the author's open question is not hidden in the file.
    assert "word/comments.xml" in z.namelist()
    assert "Katherine" not in z.read("word/comments.xml").decode("utf-8")
    assert "<w:comment " not in z.read("word/comments.xml").decode("utf-8")
    assert docx.Document(str(out)).paragraphs[0].text == "Is this Kathryn or Katherine?"


def test_a_file_without_markup_round_trips_its_text(tmp_path):
    d = docx.Document()
    d.add_paragraph("Nothing to accept here.")
    d.add_paragraph("Nor here.")
    src = tmp_path / "plain.docx"
    d.save(str(src))
    pkg = DocxPackage(str(src))
    assert accept_all(pkg) == 0 and strip_comments(pkg) == 0
    out = write_clean_copy(src, tmp_path / "clean.docx")
    assert _texts(out) == ["Nothing to accept here.", "Nor here."]


def test_a_non_docx_is_refused_by_name(tmp_path):
    bad = tmp_path / "book.docx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(CleanCopyError, match="book.docx"):
        write_clean_copy(bad, tmp_path / "clean.docx")
