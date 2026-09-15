"""Protect fixed-stage boundaries and broad, exact number/poetry selections."""
from __future__ import annotations

import pytest

from galley.fixed_policy import (DIAGNOSTIC_ONLY_TYPES, LOCAL_CANDIDATE_TYPES,
                                 NUMBER_POLICY, PROOFREADING_POLICY,
                                 configuration, extract_numbers, poetry_samples)


def _flat_types(cfg):
    return [key for entry in cfg.error_types
            for key in ([entry] if isinstance(entry, str) else
                        entry["group"] if isinstance(entry, dict) else entry)]


@pytest.mark.parametrize("poetry", [False, True])
def test_fixed_recipe_cannot_launch_unscheduled_model_stages(poetry):
    cfg = configuration(poetry)
    for name in ("glossary", "storysheet", "continuity", "chapter_continuity",
                 "adjudicate", "rewrite", "sapling", "chapter_sweep",
                 "repair", "smoothing", "factcheck", "toccheck", "meaning_check",
                 "fix_check", "residuals", "recurrence"):
        assert not getattr(cfg, name).enabled, name
    assert not cfg.low_confidence.confirm
    assert not cfg.smoothing.edits
    assert not cfg.ensemble.verifies
    assert cfg.ensemble.verifier_model is None
    assert cfg.ensemble.verify_policy == "none"
    assert cfg.candidate_screening.mode == "off"
    assert not cfg.candidate_screening.judgment_enabled
    assert not cfg.examination_graph.enabled
    assert not cfg.examination_graph.production_verdicts
    assert not cfg.examination_graph.judgment.enabled
    assert cfg.rounds.count == 1
    assert cfg.flights.posture == "strict"
    assert not cfg.normalize.quotes and not cfg.normalize.spaces
    assert not cfg.speaker_split.enabled
    assert not cfg.not_applied_comments and not cfg.excluded_words_comment
    assert cfg.spellcheck.enabled
    assert cfg.api.claude_lane == "subagent"


def test_number_read_replaces_typed_number_group_without_pruning_other_mechanics():
    cfg = configuration()
    assert [(d.model, d.effort) for d in cfg.ensemble.detectors] == [
        ("claude-sonnet-5", "low"), ("gpt-5.6-luna", "low")]
    assert "number_style" not in _flat_types(cfg)
    assert "currency_style" not in _flat_types(cfg)
    assert {"spelling", "subject_verb_agreement", "introductory_comma",
            "missing_word", "unnecessary_comma"} <= set(_flat_types(cfg))
    assert len(cfg.error_types) == 9
    assert "sweep_dash" in cfg.sweeps
    cfg.sweeps.clear()
    assert configuration().sweeps  # no mutable config shared between manuscripts


def test_poetry_is_house_mechanics_without_sentence_level_lanes():
    cfg = configuration(poetry=True)
    flat = [key for group in cfg.error_types for key in ([group] if isinstance(group, str) else group)]
    assert flat == ["spelling", "homophone_confusion", "apostrophe_error", "capitalization",
                    "serial_comma", "missing_word", "ly_adverb_hyphen"]
    assert "number_style" not in flat            # the number stage reads verse too
    assert not cfg.ensemble.enabled
    assert cfg.api.model == "claude-sonnet-5"
    assert "sweep_dash" in cfg.sweeps and "sweep_elision_apostrophe" in cfg.sweeps
    assert "sweep_terminal_period" not in cfg.sweeps and "sweep_doubled_word" not in cfg.sweeps
    assert not cfg.languagetool.enabled and not cfg.consistency.enabled
    assert not cfg.style.heading_title_case
    assert not cfg.style.heading_vocab_queries
    assert not cfg.style.unclosed_quote_queries
    assert not any(getattr(cfg.genre_scans, kind).enabled
                   for kind in ("anachronism", "citation_format", "reading_level"))


def test_local_inventory_includes_grammar_and_comma_floor_but_separates_style():
    cfg = configuration()
    assert cfg.languagetool.enabled and cfg.consistency.enabled
    assert cfg.candidate_screening.candidate_types == LOCAL_CANDIDATE_TYPES
    # The exhaustive comma sweep stays out: 33 applied of 6,961 screened sites
    # on the first production book, at half the run's wall-clock.
    assert "comma_boundary" not in LOCAL_CANDIDATE_TYPES
    assert {"grammar", "compound_sentence_comma", "homophone",
            "heading_sequence", "list_punctuation", "term_consistency",
            "word_echo"} <= set(LOCAL_CANDIDATE_TYPES)
    assert not {"number_style", "currency_style"} & set(LOCAL_CANDIDATE_TYPES)
    assert DIAGNOSTIC_ONLY_TYPES == {"word_echo", "reading_level"}
    assert "internal diagnostics only" in PROOFREADING_POLICY
    assert cfg.genre_scans.anachronism.enabled
    assert cfg.genre_scans.anachronism.era is None


def test_preparation_really_runs_local_consistency_and_citation_checks(tmp_path):
    import docx
    from docproof.pipeline import prepare
    from pathlib import Path

    book = docx.Document()
    texts = [
        "The blood-cursed rider never returns from the northern marches.",
        "A blood-cursed name is a heavy thing for a child to carry.",
        "She had heard of the bloodcursed before, in stories told at night.",
        "The study found strong effects (Smith, 2019).",
        "A later paper confirmed this (Jones, 2021).",
        "Other researchers agree [12].",
        "This was also noted elsewhere [13, 14].",
        'She said, "I  keep my spaces."',
    ]
    for text in texts:
        book.add_paragraph(text)
    source = tmp_path / "local.docx"
    book.save(source)
    original_package = source.read_bytes()
    prepared = prepare(configuration(), source,
                       Path(__file__).resolve().parents[1] / "config/error_types")

    assert [p.text for p in prepared.doc.paragraphs] == texts
    assert source.read_bytes() == original_package
    assert any("bloodcursed" in f.original_text for f in prepared.consistency_findings)
    assert sum(f.error_type == "citation_format" for f in prepared.genre_findings) == 2
    assert all(f.force_query for f in prepared.genre_findings)
    # The fixed adapter, not the old paid adjudication/grammar path, owns these.
    assert prepared.adjudicate_candidates == []
    assert prepared.candidate_screening is None


def test_full_house_number_exceptions_are_carried_to_readers():
    for phrase in ("type 2 diabetes", "stage four liver cancer", "ten thousand",
                   "4:00 AM", "a five-dollar", "two cents", "45 BC", "a quarter"):
        assert phrase in NUMBER_POLICY
    assert "decimal only when" in NUMBER_POLICY
    assert "not an error" in NUMBER_POLICY
    assert "Do not polish" in PROOFREADING_POLICY
    assert "Never invent facts" in PROOFREADING_POLICY


def test_numeric_exceptions_are_collected_without_prescribing_corrections():
    paragraphs = {"body-1": (
        "At 4:30 p.m., around 4, he paid $12.50 and 90 cents for a five-dollar bill. "
        "Type 2 diabetes, stage IV cancer, 9mm, July 14, 1989, room 4, chapter 5, "
        "40 percent, 2/3 and 2½ all stay available for contextual judgment.")}
    rows = extract_numbers(paragraphs)
    values = [row["text"] for row in rows]
    for expected in ("4:30 p.m.", "$12.50", "90 cents", "five-dollar", "2", "IV",
                     "9", "14", "1989", "5", "40 percent", "2/3", "2½"):
        assert expected in values
    assert {row["kind"] for row in rows} >= {"time", "currency", "percentage", "fraction", "roman"}
    assert all("correction" not in row for row in rows)
    for row in rows:
        assert paragraphs[row["para_id"]][row["start"]:row["end"]] == row["text"]
        assert row["text"] in row["context"]


def test_compound_written_numbers_fractions_and_ordinals_keep_exact_spans():
    text = ("One hundred and twenty-five; twenty-first; two-thirds; one-and-a-half; "
            "ten thousand; three dozen; first; .5; 1 1/2; 10–12; 212-555-0123; "
            "USD 1.2 million; 5+; −4; 1e-3; halves; forties.")
    rows = extract_numbers({"p": text})
    assert [row["text"] for row in rows] == [
        "One hundred and twenty-five", "twenty-first", "two-thirds", "one-and-a-half",
        "ten thousand", "three dozen", "first", ".5", "1 1/2", "10–12", "212-555-0123",
        "USD 1.2 million", "5+", "−4", "1e-3", "halves", "forties"]
    assert {row["text"]: row["kind"] for row in rows}["twenty-first"] == "ordinal"


def test_candidates_do_not_merge_unrelated_counts_or_partial_number_words():
    rows = extract_numbers({"p": "Someone won once; alone, a stone. One and two. 4 amazing dogs; 8 PM."})
    assert [row["text"] for row in rows] == ["One", "two", "4", "8 PM."]
    assert rows[2]["kind"] == "numeral"


def test_fractions_are_not_confused_with_compound_ordinals():
    rows = extract_numbers({"p": "forty-third; twenty‑fifth; one fifth; one-third; half a dozen; quarter of a million"})
    assert [(row["text"], row["kind"]) for row in rows] == [
        ("forty-third", "ordinal"), ("twenty‑fifth", "ordinal"),
        ("one fifth", "fraction"), ("one-third", "fraction"),
        ("half a dozen", "fraction"), ("quarter of a million", "fraction")]


def test_number_ids_are_repeatable_and_independent_of_other_paragraphs():
    paragraphs = {"p": "Twenty birds saw 20 birds and 20 nests."}
    rows = extract_numbers(paragraphs)
    assert rows == extract_numbers(paragraphs)
    assert len({row["id"] for row in rows}) == len(rows)
    assert rows == [row for row in extract_numbers({"new": "100 dogs", **paragraphs})
                    if row["para_id"] == "p"]
    changed = extract_numbers({"p": "Twenty birds saw 21 birds and 20 nests."})
    assert changed[1]["id"] != rows[1]["id"]


def test_poetry_sampling_covers_book_and_middle_verse_block():
    paragraphs = {f"p{i}": ("This is an ordinary paragraph of narrative prose. " * 6)
                  for i in range(30)}
    for i in (7, 8, 9, 10, 11):
        paragraphs[f"p{i}"] = f"A short verse line {i}"
    samples = poetry_samples(paragraphs)
    sampled = {sample["para_id"] for sample in samples}
    assert {"p0", "p14", "p29"} <= sampled
    assert len(sampled & {"p7", "p8", "p9", "p10", "p11"}) >= 2
    assert len(samples) <= 6
    assert samples == poetry_samples(paragraphs)
    assert [int(row["para_id"][1:]) for row in samples] == sorted(int(row["para_id"][1:]) for row in samples)


def test_poetry_samples_mark_truncation_and_preserve_short_paragraphs():
    paragraphs = {"start": "A" * 200, "verse": "line one\nline two\nline three", "end": "B" * 200}
    samples = poetry_samples(paragraphs, char_budget=40)
    assert samples[0]["start"] == 0 and samples[-1]["end"] == 200
    assert samples[1]["text"] == paragraphs["verse"]
    assert not samples[1]["truncated"]
    assert samples[0]["truncated"] and samples[-1]["truncated"]
    assert all(len(row["text"]) <= 40 for row in samples)
    assert all(paragraphs[row["para_id"]][row["start"]:row["end"]] == row["text"] for row in samples)
    assert poetry_samples({}) == []
    with pytest.raises(ValueError):
        poetry_samples(paragraphs, char_budget=0)


@pytest.mark.parametrize("before,after,category,reason", [
    ("teh", "the", "spelling", None),
    ("Moon", "moon", "spelling", None),                         # mid-line recase is mechanics
    ("smokey", "smoky", "spelling", None),
    (" -- ", "—", "sweep_dash", None),
    ("waits", "Waits", "spelling", "Verse keeps the case of its line heads"),
    ("quiet", "quiet.", "punctuation", "Verse takes no terminal mark the poet did not set"),
    ("\n  ", " ", "punctuation", "Verse keeps its line breaks"),
    ("waits", "wait", "grammar", "Verse takes house mechanics only; sentence-level judgments stay out"),
    ("waits", "waits", "continuity", "Verse takes house mechanics only; sentence-level judgments stay out"),
])
def test_verse_safe_guards_structure_and_category(before, after, category, reason):
    from galley.fixed_policy import verse_safe
    text = "teh Moon\n  waits, smokey -- quiet"
    lo = text.index(before)
    row = {"start": lo, "end": lo + len(before), "before": before, "replacement": after,
           "category": category, "format": ""}
    assert verse_safe(row, text) == reason


def test_verse_safe_refuses_formatting():
    from galley.fixed_policy import verse_safe
    row = {"start": 0, "end": 3, "before": "The", "replacement": "The", "category": "spelling", "format": "italic"}
    assert verse_safe(row, "The Moon") == "Verse takes no formatting changes"
