"""Content and emphasis must survive the actual files prep hands back."""
from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path
import zipfile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docproof import prep
from docproof.config import load_config
from docproof.prep.ingest import BODY_PART
from docproof.prep.nontext import protected_content
from docproof.prep.verify import VerificationFailed, verify_output
from docproof.prep.writers.indesign_idml import body_style_names, discover_body_story
from docproof.utils.xml_helpers import DocxPackage


CONFIG_DIR = Path(__file__).parents[1] / "config"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/"
    "x8AAwMCAO+jR1kAAAAASUVORK5CYII=")


def _prepared(document, tmp_path, *, verify=True):
    source = tmp_path / "manuscript.docx"
    document.save(source)
    cfg = load_config(CONFIG_DIR / "default.yaml")
    cfg.prep.verify = verify
    prepared = prep.prepare(cfg, source, config_dir=CONFIG_DIR)
    tags, usage = prep.run_mock(prepared)
    return cfg, prepared, tags, usage, source


def _finish(inputs, tmp_path, outputs=("book", "tracked")):
    cfg, prepared, tags, usage, source = inputs
    return prep.finish(prepared, tags, usage, cfg, out_dir=tmp_path / "out",
                       source_path=source, outputs=list(outputs))


def _picture_document(*, google=False):
    document = Document()
    document.add_paragraph("Before the illustration.")
    shape = document.add_picture(BytesIO(PNG))
    run = shape._inline.getparent().getparent()
    paragraph = run.getparent()
    control = OxmlElement("w:sdt")
    props = OxmlElement("w:sdtPr")
    props.append(OxmlElement("w:picture"))
    if google:
        alias = OxmlElement("w:alias")
        alias.set(qn("w:val"), "goog_rdk_picture")
        props.append(alias)
    content = OxmlElement("w:sdtContent")
    content.append(run)
    control.append(props)
    control.append(content)
    paragraph.append(control)
    document.add_paragraph("After the illustration.")
    return document


def _equation_document(*, inline=False):
    document = Document()
    document.add_paragraph("Before the equation.")
    paragraph = document.add_paragraph("The equation is " if inline else "")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "E=mc²"
    run.append(text)
    math.append(run)
    if inline:
        paragraph._p.append(math)
        paragraph.add_run(" in this example.")
    else:
        wrapper = OxmlElement("m:oMathPara")
        wrapper.append(math)
        paragraph._p.append(wrapper)
    document.add_paragraph()
    document.add_paragraph("After the equation.")
    return document


@pytest.mark.parametrize("google", [False, True])
def test_picture_controls_keep_the_image_in_both_word_outputs(tmp_path, google):
    inputs = _prepared(_picture_document(google=google), tmp_path)
    prepared = inputs[1]
    assert prepared.structure.paragraphs[1].preserved
    result = _finish(inputs, tmp_path)
    for path in result.documents.values():
        pkg = DocxPackage(path)
        assert protected_content(pkg.tree(BODY_PART)) == prepared.structure.protected_content
        assert len(pkg.tree(BODY_PART).findall(".//" + qn("w:drawing"))) == 1
        assert pkg.raw("word/media/image1.png") == PNG
    assert {c.view for c in result.verifications} == {"clean", "accept", "reject"}
    assert all(c.ok for c in result.verifications)


@pytest.mark.parametrize("inline", [False, True])
def test_native_equations_survive_both_word_outputs(tmp_path, inline):
    inputs = _prepared(_equation_document(inline=inline), tmp_path)
    prepared = inputs[1]
    equation = prepared.structure.paragraphs[1]
    assert equation.has_math and not equation.is_blank
    assert equation.preserved == (not inline)
    result = _finish(inputs, tmp_path)
    for path in result.documents.values():
        pkg = DocxPackage(path)
        assert protected_content(pkg.tree(BODY_PART)) == prepared.structure.protected_content
        assert pkg.tree(BODY_PART).find(".//" + qn("m:t")).text == "E=mc²"
    assert all(c.ok for c in result.verifications)


@pytest.mark.parametrize("kind", ["picture", "equation"])
def test_verification_refuses_nontext_loss_even_when_words_match(tmp_path, kind):
    document = _picture_document() if kind == "picture" else _equation_document()
    inputs = _prepared(document, tmp_path)
    result = _finish(inputs, tmp_path, ("book",))
    path = result.documents["book"]
    pkg = DocxPackage(path)
    tag = "w:drawing" if kind == "picture" else "m:oMathPara"
    element = pkg.tree(BODY_PART).find(".//" + qn(tag))
    element.getparent().remove(element)
    pkg.mark_modified(BODY_PART)
    pkg.save(path)
    with pytest.raises(VerificationFailed, match="picture.*equation"):
        verify_output(inputs[1].structure, path, result.plan.glyph)


@pytest.mark.parametrize("kind", ["picture", "equation"])
@pytest.mark.parametrize("verify", [False, True])
def test_idml_refuses_content_it_cannot_preserve(tmp_path, kind, verify):
    document = _picture_document() if kind == "picture" else _equation_document(inline=True)
    inputs = _prepared(document, tmp_path, verify=verify)
    with pytest.raises(VerificationFailed, match="Use the Word output"):
        _finish(inputs, tmp_path, ("indesign",))
    assert not list((tmp_path / "out").glob("*.idml"))
    assert "cannot preserve" in (tmp_path / "out/prep_notes.md").read_text()


def _emphasis_document():
    document = Document()
    inherited = document.styles.add_style("InheritedEmphasis", WD_STYLE_TYPE.CHARACTER)
    inherited.base_style = document.styles["Emphasis"]
    inherited.font.bold = True
    double_toggle = document.styles.add_style("DoubleItalicToggle", WD_STYLE_TYPE.CHARACTER)
    double_toggle.base_style = document.styles["Emphasis"]
    double_toggle.font.italic = True
    paragraph = document.add_paragraph()
    paragraph.add_run("Emphasis ").style = "Emphasis"
    paragraph.add_run("Strong ").style = "Strong"
    paragraph.add_run("Inherited ").style = inherited
    override = paragraph.add_run("Override ")
    override.style = inherited
    override.italic = False
    paragraph.add_run("Toggle").style = double_toggle
    # A preserved table paragraph needs its own paragraph-style chain too.
    table_style = document.styles.add_style("CellMeaning", WD_STYLE_TYPE.PARAGRAPH)
    table_style.font.italic = True
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Cell text"
    cell.paragraphs[0].style = table_style
    return document


def test_word_outputs_keep_referenced_styles_and_their_dependencies(tmp_path):
    document = _emphasis_document()
    expected = {name: etree.tostring(document.styles[name]._element, method="c14n", exclusive=True)
                for name in ("Emphasis", "Strong", "InheritedEmphasis",
                             "DoubleItalicToggle", "CellMeaning")}
    result = _finish(_prepared(document, tmp_path), tmp_path)
    for path in result.documents.values():
        styles = DocxPackage(path).tree("word/styles.xml")
        by_id = {s.get(qn("w:styleId")): s for s in styles.findall(qn("w:style"))}
        for name, original in expected.items():
            assert etree.tostring(by_id[name], method="c14n", exclusive=True) == original


def test_idml_resolves_character_style_inheritance_and_direct_overrides(tmp_path):
    inputs = _prepared(_emphasis_document(), tmp_path)
    result = _finish(inputs, tmp_path, ("indesign",))
    path = result.documents["indesign"]
    story_id = discover_body_story(path, body_style_names(inputs[1].sheet))
    with zipfile.ZipFile(path) as package:
        story = etree.fromstring(package.read(f"Stories/Story_{story_id}.xml"))
    applied = {}
    for run in story.iter("CharacterStyleRange"):
        for text in run.findall("Content"):
            for word in (text.text or "").split():
                applied[word] = run.get("FontStyle", "Regular")
    assert {word: applied[word] for word in ("Emphasis", "Strong", "Inherited", "Override", "Toggle")} == {
        "Emphasis": "Italic", "Strong": "Bold", "Inherited": "Bold Italic",
        "Override": "Bold", "Toggle": "Regular"}
