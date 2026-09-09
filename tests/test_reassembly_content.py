"""Text surgery preserves non-text content when changes are accepted."""
import base64

import docx
import pytest
from lxml import etree

from docproof.config import Config
from docproof.ingest import accept_all_revisions, build_document_model, preflight
from docproof.models import Finding
from docproof.reassembler import apply_tracked_changes, apply_untracked, paragraph_view_text
from docproof.utils.xml_helpers import DocxPackage, paragraph_text, qn, walk_package
from docproof.validator import validate_findings


@pytest.mark.parametrize("mode", ["tracked", "untracked"])
@pytest.mark.parametrize("side", ["before", "after"])
@pytest.mark.parametrize("content", ["drawing", "footnoteReference"])
def test_text_correction_preserves_other_run_content(tmp_path, mode, side, content):
    picture = tmp_path / "pixel.png"
    picture.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8"
        "/x8AAusB9Y9ZQmcAAAAASUVORK5CYII="))
    document = docx.Document()
    p = document.add_paragraph("They ")
    r = p.add_run()
    r.bold = True

    def add_content():
        if content == "drawing":
            r.add_picture(str(picture), width=docx.shared.Inches(1))
        else:
            etree.SubElement(r._r, qn("w:footnoteReference"), {qn("w:id"): "1"})

    if side == "before":
        add_content()
    r.add_text("is")
    if side == "after":
        add_content()
    p.add_run(" ready.")
    source = tmp_path / "source.docx"
    document.save(source)
    cfg = Config(comments=False)
    pkg = preflight(source, "abort")
    model = build_document_model(pkg, cfg)
    p = next(walk_package(pkg)).element
    original_content = etree.tostring(next(p.iter(qn(f"w:{content}"))))
    original_text = paragraph_text(p)
    findings = validate_findings([Finding(
        finding_id="grammar", chunk_id="c", para_id="body-0000",
        error_type="subject_verb_agreement", original_text=original_text,
        corrected_text="They are ready.", occurrence=1,
        explanation="Plural agreement.", confidence="high")], model, "medium")
    assert len(findings) == 1 and findings[0].status == "validated"

    if mode == "tracked":
        stats = apply_tracked_changes(pkg, model, findings, cfg)
        assert stats.applied == ("grammar",) and not stats.skipped
        assert paragraph_view_text(p, "reject") == original_text
        accept_all_revisions(pkg)
    else:
        a = findings[0].anchor
        apply_untracked(pkg, model, {"body-0000": [(a.start, a.end, a.insert_text)]})
    output = tmp_path / "corrected.docx"
    pkg.save(output)
    reloaded = DocxPackage(output)
    p = next(walk_package(reloaded)).element
    assert paragraph_text(p) == "They are ready."
    remaining = list(p.iter(qn(f"w:{content}")))
    assert len(remaining) == 1
    assert etree.tostring(remaining[0]) == original_content
    assert remaining[0].getparent().find(qn("w:rPr")).find(qn("w:b")) is not None
    if content == "drawing":
        assert len(docx.Document(output).inline_shapes) == 1
