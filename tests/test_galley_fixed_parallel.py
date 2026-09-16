"""Offline concurrency tests use barriers, never live model calls."""
from concurrent.futures import ThreadPoolExecutor
from contextvars import ContextVar
from functools import partial
import threading
import time

import pytest

from galley.fixed_parallel import ReadScheduler
from galley.fixed_policy import configuration
from galley.fixed_workflow import SONNET, LUNA, SOL


def test_provider_pools_overlap_preserve_order_and_carry_context():
    cfg = configuration()
    cfg.api.subagent_concurrency = 2
    cfg.api.concurrency_by_provider["openai"] = 3
    scheduler = ReadScheduler(cfg)
    # Luna reads on the ChatGPT subscription (2026-09-16), so it shares the
    # Codex pool with Sol and Astra rather than the OpenAI API pool.
    assert scheduler.lane(LUNA) == scheduler.lane(SOL) == "codex"
    assert scheduler.lane(SONNET) == "claude"
    barrier = threading.Barrier(4)
    marker = ContextVar("test_resource_context", default=None)
    marker.set("book-123")
    def read(i):
        assert marker.get() == "book-123"
        barrier.wait(timeout=3)
        return i
    # A single FIFO pool would be starved by the queued Claude work.
    jobs = [(SONNET, partial(read, i)) for i in range(2)] + [(LUNA, partial(read, i)) for i in range(2, 4)]
    assert scheduler.map(jobs) == list(range(4))


def test_nested_batches_share_limits_without_deadlock_or_oversubscription():
    cfg = configuration()
    cfg.api.subagent_concurrency = 2
    scheduler = ReadScheduler(cfg)
    lock = threading.Lock()
    active = peak = 0
    def read():
        nonlocal active, peak
        with lock:
            active += 1
            peak = max(peak, active)
        time.sleep(.02)
        with lock:
            active -= 1
        return True
    def batch():
        return scheduler.map([(SONNET, partial(scheduler.run, SONNET, read))] * 5)
    with ThreadPoolExecutor(max_workers=2) as pool:
        tasks = [pool.submit(batch) for _ in range(2)]
        assert all(all(t.result(timeout=5)) for t in tasks)
    assert peak == 2


def test_explicit_serial_setting_is_global_across_transports():
    cfg = configuration()
    cfg.api.concurrency = 1
    scheduler = ReadScheduler(cfg)
    lock = threading.Lock()
    active = peak = 0
    def read():
        nonlocal active, peak
        with lock:
            active += 1
            peak = max(peak, active)
        time.sleep(.01)
        with lock:
            active -= 1
    scheduler.map([(model, read) for model in (SONNET, LUNA, SOL) for _ in range(3)])
    assert peak == 1


def test_correction_chains_overlap_but_only_commit_checked_passages(tmp_path):
    from docx import Document
    from galley.fixed_workflow import FixedWorkflow
    from tests.test_galley_fixed_workflow import Readers, ruling
    source = tmp_path / "book.docx"
    doc = Document()
    doc.add_paragraph("Placeholder.")
    doc.save(source)
    barrier = threading.Barrier(2)
    def answer(stage, model, payload, kwargs):
        if stage == "checks_meaning":
            barrier.wait(timeout=3)
            return {"decisions": [{"id": x["id"], "verdict": "reject", "reason": "Needs review."}
                                  for x in payload["changes"]]}
        if stage == "checks_meaning_disputes":
            return {"decisions": [ruling(x, replacement=x["after"].replace("wrong", "right")) for x in payload["sites"]]}
        if stage == "checks_correction":
            assert all("right" in x["after"] and "wrong" not in x["after"] for x in payload["changes"])
    flow = FixedWorkflow(source, tmp_path / "fixed", calls=Readers(handler=answer))
    before = {"p1": "a" * 6000, "p2": "b" * 6000}
    flow.original = dict(before)
    flow.current = {pid: text + " wrong" for pid, text in before.items()}
    flow._checks("checks", before)
    assert flow.current == {pid: text + " right" for pid, text in before.items()}
    assert [h["stage"] for h in flow.history] == ["checks_meaning", "checks_meaning_sonnet", "checks_meaning_disputes"] * 2 + ["checks_correction"] * 2


@pytest.mark.parametrize("stage", ["numbers", "disputes", "read", "comments", "poetry_sections"])
def test_independent_workflow_windows_really_overlap(tmp_path, stage):
    from docx import Document
    from galley.fixed_workflow import FixedWorkflow, _candidate, OPUS
    from tests.test_galley_fixed_workflow import Readers, finding, ruling
    source = tmp_path / "book.docx"
    doc = Document()
    doc.add_paragraph("Placeholder.")
    doc.save(source)
    barrier = threading.Barrier(2)
    expected = {"numbers": "numbers", "disputes": "typed_disputes", "read": "ensemble_sweep_opus",
                "comments": "astra_comment_review", "poetry_sections": "poetry_sections"}[stage]
    seen = []
    lock = threading.Lock()
    def answer(name, model, payload, kwargs):
        if name == "typed_screen":
            return {"decisions": [ruling(site, "apply" if model == SONNET else "drop", site["proposals"][0]["replacement"])
                                  for site in payload["sites"]]}
        if stage == "poetry_sections" and name == "poetry":
            return {"classification": "mixed", "reason": "Mixed document."}
        if name == expected:
            with lock:
                seen.append((name, model))
            # First two arrive concurrently; later complete pairs do too.
            barrier.wait(timeout=3)
            if stage == "poetry_sections":
                return {"paragraphs": [{"id": x["id"], "poetry": False} for x in payload]}
    flow = FixedWorkflow(source, tmp_path / "fixed", calls=Readers(handler=answer))
    flow.original = {"p1": "word " * 5000 + " 123", "p2": "word " * 5000 + " 456"}
    flow.current = dict(flow.original)
    if stage == "numbers":
        flow._numbers()
    elif stage == "disputes":
        rows = [_candidate(finding(pid, text, text + "."), flow.current, OPUS) for pid, text in flow.current.items()]
        flow._adjudicate("typed", rows, force=True)
    elif stage == "read":
        flow._read(expected, OPUS)
    elif stage == "comments":
        flow.questions = [{"id": pid, "para_id": pid, "quote": text, "question": "Which person?",
                           "missing_knowledge": "identity", "reason": "Missing identity."} for pid, text in flow.current.items()]
        flow._comments([], "astra", before=dict(flow.current), model="gpt-6-astra")
    else:
        flow._classify()
    assert len(seen) >= 2


def test_parallel_check_questions_keep_meaning_then_correction_order(tmp_path):
    from docx import Document
    from galley.fixed_workflow import FixedWorkflow
    from tests.test_galley_fixed_workflow import Readers, ruling
    source = tmp_path / "book.docx"
    doc = Document()
    doc.add_paragraph("Placeholder.")
    doc.save(source)
    def answer(stage, model, payload, kwargs):
        if stage.endswith("_sonnet"):
            return {"decisions": [{"id": row["id"], "verdict": "approve", "reason": "Independent approval."}
                                  for row in payload["changes"]]}
        if "changes" in payload:
            return {"decisions": [{"id": row["id"], "verdict": "reject" if
                    stage.endswith("correction") or row["id"] in {"p1", "p3"} else "approve", "reason": "Check intent."}
                    for row in payload["changes"]]}
        if "sites" in payload:
            return {"decisions": [{**ruling(row, "query"), "question": "Which person is intended here?",
                                   "missing_knowledge": "The intended person's identity."} for row in payload["sites"]]}
    flow = FixedWorkflow(source, tmp_path / "fixed", calls=Readers(handler=answer))
    flow.original = {"p" + str(i): str(i) * 2500 for i in range(1, 5)}
    flow.current = {pid: text + "." for pid, text in flow.original.items()}
    flow._checks("checks", flow.original)
    assert [q["para_id"] for q in flow.questions] == ["p1", "p3", "p2", "p4"]
    assert flow.current == flow.original
