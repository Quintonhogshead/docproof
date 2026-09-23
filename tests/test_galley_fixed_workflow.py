"""Real Word intake and mocked readers exercise the fixed editorial sequence."""
from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from docx import Document

from docproof.providers import ProviderResult
from galley.fixed_workflow import (ASTRA, LUNA, OPUS, SOL, SONNET,
                                   FixedWorkflow, FixedWorkflowError, _candidate)


def shared_context(request):
    """The JSON block every window of a read shares through its system prompt."""
    from galley.fixed_workflow import SHARED_CONTEXT_MARKER
    return json.loads(request["system"].split(SHARED_CONTEXT_MARKER, 1)[1])


def finding(pid, quote, replacement, category="grammar", *, action="edit", missing=""):
    return {"para_id": pid, "quote": quote, "occurrence": 1,
            "replacement": replacement, "category": category, "action": action,
            "reason": "A clear local proofreading issue.", "missing_knowledge": missing}


def ruling(site, action="apply", replacement=None):
    return {"id": site["id"], "action": action,
            "replacement": site.get("after", site.get("before", "")) if replacement is None else replacement,
            "reason": "Checked the context.", "missing_knowledge": "",
            "question": ""}


def comment_decision(q, action="retain", quote=None):
    return {"id": q["id"], "action": action, "quote": quote or q["quote"],
            "question": q["question"], "reason": "The author must supply the missing identity.",
            "missing_knowledge": q["missing_knowledge"]}


class Readers:
    """No model transport; schemas/coverage still pass through real analyzers."""
    def __init__(self, *, poetry=False, handler=None, typed=None):
        self.poetry, self.handler, self.typed = poetry, handler, typed
        self.events = []

    def ask(self, stage, **kwargs):
        payload = json.loads(kwargs["user"])
        self.events.append({"stage": stage, "payload": payload, **kwargs})
        if self.handler:
            result = self.handler(stage, kwargs["model"], payload, kwargs)
            if result is not None:
                if "reviewed_check_ids" in kwargs["schema"]["properties"]:
                    result.setdefault("reviewed_check_ids", [s["id"] for s in payload["focused_sites"]])
                if "publication_blockers" in kwargs["schema"]["properties"]:
                    result.setdefault("publication_blockers", [])
                return result
        if stage == "poetry":
            return {"classification": "poetry" if self.poetry else "prose", "reason": "Fixed samples."}
        if stage == "story_sheet":
            return {"narration": "Third person past tense", "characters": [], "notes": []}
        properties = kwargs["schema"]["properties"]
        if "reading_notes" in properties:
            return {"findings": [], "reading_notes": "Fixture continuity read."}
        if "reviewed_ids" in properties:
            owned = payload.get("sites", payload.get("paragraphs", []))
            return {"reviewed_ids": [row["id"] for row in owned], "findings": [],
                    "comment_decisions": [comment_decision(q) for q in payload.get("comments", [])],
                    **({"reviewed_check_ids": [s["id"] for s in payload["focused_sites"]]}
                       if "reviewed_check_ids" in properties else {}),
                    **({"publication_blockers": []} if "publication_blockers" in properties else {}),
                    "editorial_verdict": "ready"}
        if "changes" in payload:
            return {"decisions": [{"id": x["id"], "verdict": "approve", "reason": "Correct."}
                                  for x in payload["changes"]]}
        if "comments" in payload:
            return {"decisions": [comment_decision(q) for q in payload["comments"]]}
        return {"decisions": [ruling(site, "drop") for site in payload["sites"]]}

    def provider(self, stage, cfg):
        parent = self

        class TypedProvider:
            name = "fixed-test"

            def complete_structured(self, **kwargs):
                parent.events.append({"stage": stage, "claude_lane": cfg.api.claude_lane, **kwargs})
                paragraphs = dict(re.findall(r'<paragraph id="([^"]+)">\n(.*?)\n</paragraph>',
                                             kwargs["user"].split("</context>")[-1], re.S))
                types = kwargs["schema"]["$defs"]["RawFinding"]["properties"]["error_type"].get("enum")
                if types is None:
                    types = [kwargs["schema"]["$defs"]["RawFinding"]["properties"]["error_type"]["const"]]
                answer = parent.typed(stage, kwargs["model"], paragraphs, types) if parent.typed else []
                if isinstance(answer, ProviderResult):
                    return answer
                return ProviderResult(parsed={"findings": answer,
                                              "reviewed_paragraph_ids": list(paragraphs)})

        return TypedProvider()

    def assert_complete(self):
        pass

    def usage_summary(self):
        return {"mocked_calls": len(self.events)}


@pytest.fixture
def make_book(tmp_path):
    def create(*paragraphs, name="source.docx"):
        path = tmp_path / name
        doc = Document()
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
        doc.save(path)
        return path
    return create


def _typed_row(pid, source, before, after, key):
    return {"para_id": pid, "error_type": key, "original_text": source,
            "corrected_text": source.replace(before, after), "occurrence": 1,
            "confidence": "high", "explanation": "A clear proofreading error."}


@pytest.fixture(autouse=True)
def local_scans(monkeypatch):
    """Existing editorial cases isolate the local scanner's transport/generators.

    New local integration cases replace these stubs or restore the real adapter;
    the workflow's anchoring, arbitration, checks and sequencing always run.
    """
    from galley import fixed_local
    originals = {"initial": fixed_local.collect_local_candidates,
                 "completion": fixed_local.collect_completion_candidates}
    monkeypatch.setattr(fixed_local, "collect_local_candidates",
                        lambda *args, **kwargs: ([], {"fixture": "local clean"}))
    monkeypatch.setattr(fixed_local, "collect_completion_candidates",
                        lambda *args, **kwargs: ([], {"fixture": "completion clean"}))
    return originals


def _local_row(pid, before, after, *, source="LanguageTool", category="grammar"):
    return {**finding(pid, before, after, category), "source": source}


def test_real_docx_full_fixed_sequence_and_successive_corrected_versions(make_book, tmp_path, monkeypatch):
    book = make_book("We seen teh 20 birds beside a apple. They was bright. It are warm. He walk home.")

    def typed(stage, model, paragraphs, keys):
        result = []
        for pid, text in paragraphs.items():
            if "spelling" in keys:
                result.append(_typed_row(pid, text, "teh", "the", "spelling"))
            if "subject_verb_agreement" in keys:
                result.append(_typed_row(pid, text, "seen", "saw" if model == SONNET else "see", "subject_verb_agreement"))
        return result

    def handler(stage, model, payload, kwargs):
        if stage == "typed_screen":
            return {"decisions": [ruling(x, replacement="aw" if model == SONNET else "e") for x in payload["sites"]]}
        if stage == "typed_disputes":
            return {"decisions": [ruling(x, replacement="aw") for x in payload["sites"]]}
        if stage == "numbers":
            pid, text = next(iter(payload["paragraphs"].items()))
            return {"reviewed_ids": [x["id"] for x in payload["sites"]],
                    "findings": [finding(pid, "20", "twenty", "number_style")],
                    "comment_decisions": [], "editorial_verdict": "ready"}
        edits = {"broken_repair": ("They was", "They were", "broken_sentence"),
                 "ensemble_sweep_opus": ("a apple", "an apple", "grammar"),
                 "ensemble_sweep_sol": ("a apple", "an apple", "grammar"),
                 "opus_read": ("It are", "It is", "grammar"),
                 "astra": ("He walk", "He walks", "grammar")}
        if stage in edits:
            before, after, category = edits[stage]
            row = payload["paragraphs"][0]
            return {"reviewed_ids": [x["id"] for x in payload["paragraphs"]],
                    "findings": [finding(row["id"], before, after, category)],
                    "comment_decisions": [], "editorial_verdict": "ready"}

    readers = Readers(handler=handler, typed=typed)
    # A model unexpectedly launched from prepare/finish is a test failure.
    monkeypatch.setattr("docproof.providers.build_provider", lambda *_: pytest.fail("Unscheduled model call"))
    progress = []
    flow = FixedWorkflow(book, tmp_path / "run", calls=readers,
                         progress=lambda event, **data: progress.append((event, data)))
    result = flow.run()
    assert list(result["accepted"].values()) == [
        "We saw the twenty birds beside an apple. They were bright. It is warm. He walks home."]
    assert result["questions"] == []
    assert [row["stage"] for row in result["stages"]] == [
        "poetry", "story_sheet", "opening_read", "typed", "numbers", "broken_repair", "checks", "ensemble_sweep",
        "continuity", "opus_read", "astra", "final_astra", "astra_gate"]
    assert [data["phase"] for event, data in progress if event == "phase_start"] == [
        "poetry", "story_sheet", "opening_read", "typed", "numbers", "broken_repair", "checks", "ensemble_sweep",
        "continuity", "opus_read", "astra", "final_astra", "astra_gate"]
    assert result["editorial_verdict"] == "ready" and result["final_review"]["core_mechanical_errors"] == 0
    assert all(data["ok"] for event, data in progress if event == "phase_end")
    events = readers.events
    assert len([x for x in events if x["stage"] == "typed"]) == 18
    assert {x["model"] for x in events if x["stage"] == "typed"} == {SONNET, LUNA}
    assert all("number_style" not in x["schema"]["$defs"]["RawFinding"]["properties"]["error_type"].get("enum", [])
               for x in events if x["stage"] == "typed")
    assert all(x["claude_lane"] == "subagent" for x in events if x["stage"] == "typed")
    by_stage = {x["stage"]: x for x in events}
    assert by_stage["story_sheet"]["model"] == OPUS
    assert by_stage["opening_read"]["model"] == OPUS
    assert "type 2 diabetes" in by_stage["opening_read"]["system"]
    assert "type 2 diabetes" not in by_stage["story_sheet"]["system"]
    assert "type 2 diabetes" not in by_stage["poetry"]["system"]
    assert "type 2 diabetes" in by_stage["opus_read"]["system"]
    assert by_stage["typed_disputes"]["model"] == OPUS
    assert by_stage["ensemble_sweep_opus"]["payload"]["paragraphs"] == by_stage["ensemble_sweep_sol"]["payload"]["paragraphs"]
    assert by_stage["ensemble_sweep_sol"]["model"] == SOL
    assert "an apple" in by_stage["opus_read"]["payload"]["paragraphs"][0]["text"]
    assert "It is warm" in by_stage["astra"]["payload"]["paragraphs"][0]["text"]
    assert by_stage["astra"]["model"] == ASTRA
    assert any(x["stage"] == "astra_checks_correction" for x in events)


def test_frontier_structure_context_uses_each_readers_current_book(make_book, tmp_path):
    book = make_book("Chapter One: The Harbr", "The boat arrived.")
    doc = Document(book)
    doc.paragraphs[0].style = "Heading 1"
    doc.save(book)

    def handler(stage, model, payload, kwargs):
        if stage == "opus_read":
            row = payload["paragraphs"][0]
            return {"reviewed_ids": [x["id"] for x in payload["paragraphs"]],
                    "findings": [finding(row["id"], "Harbr", "Harbor", "spelling")],
                    "comment_decisions": [], "editorial_verdict": "ready"}

    readers = Readers(handler=handler)
    FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    fable = next(x for x in readers.events if x["stage"] == "opus_read")
    astra = next(x for x in readers.events if x["stage"] == "astra")
    # The excerpt is shared by every window of a read, so it travels in the
    # system prompt (one cached prefix), never in the per-window payload.
    assert "structure_context" not in fable["payload"]
    assert "The Harbr" in fable["system"]
    assert "The Harbor" in astra["system"] and "The Harbr" not in astra["system"]
    assert '"structure_context":{' in astra["system"] and '"complete_inventory":false' in astra["system"]
    assert "Do not infer missing entries" in astra["system"]


def test_press_prompts_and_focused_evidence_reach_both_final_readers(make_book, tmp_path):
    from galley.press_prompt import EDITORIAL_RULES
    book = make_book('“Wait.” He said. Red, white and blue.',
                     'He walks and waits and watches.')
    readers = Readers()
    result = FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    final = [r for r in readers.events if r["stage"] in {"opus_read", "astra"}]
    assert len(final) == 2 and final[0]["system"] == final[1]["system"]
    for request in final:
        assert all(text in request["system"] for text in EDITORIAL_RULES.values())
        assert {s["check"] for s in request["payload"]["focused_sites"]} >= {
            "dialogue_matrix", "serial_comma", "narrative_tense"}
        # Constant per-check guidance travels once, in the shared legend; a
        # plain roman body paragraph has no metadata entry, by the stated rule.
        assert all("detail" not in s for s in request["payload"]["focused_sites"] if s["check"] == "serial_comma")
        assert request["payload"]["paragraph_metadata"] == {}
        assert "entirely known roman" in shared_context(request)["notes"]["paragraph_metadata"]
        assert shared_context(request)["focused_check_legend"]["serial_comma"]
        assert "story_sheet" not in request["payload"] and "story_sheet" in shared_context(request)
        assert "reviewed_check_ids" in request["schema"]["required"]
    story = next(r for r in readers.events if r["stage"] == "story_sheet")
    assert "first/last\nparagraph IDs" in story["system"]
    assert "Canadian" in story["system"]
    stage = json.loads(Path(result["stages"][-1]["path"]).read_text())
    assert stage["evidence"]["press_audit"]["accepted_sha256"]
    assert stage["evidence"]["press_audit"]["raw_signal_counts"]["sweep_dialogue_tag"] == 1


@pytest.mark.parametrize("stage", ["opus_read", "astra"])
def test_final_reader_cannot_skip_focused_sites(make_book, tmp_path, stage):
    def handler(name, model, payload, kwargs):
        if name == stage:
            return {"reviewed_ids": [p["id"] for p in payload["paragraphs"]],
                    "reviewed_check_ids": [], "findings": [], "comment_decisions": [],
                    "editorial_verdict": "ready"}
    with pytest.raises(FixedWorkflowError, match="focused-check coverage"):
        FixedWorkflow(make_book('Red, white and blue.'), tmp_path / "run",
                      calls=Readers(handler=handler)).run()
    assert not (tmp_path / "run/result.json").exists()


def test_final_reader_can_propose_tracked_title_italics_and_astra_sees_it(make_book, tmp_path):
    from galley.fixed_documents import write_manuscripts, paragraph_views
    from docproof.utils.xml_helpers import DocxPackage, qn
    book = make_book("She read The Great Gatsby yesterday.")
    def handler(stage, model, payload, kwargs):
        if stage == "opus_read":
            p = payload["paragraphs"][0]
            return {"reviewed_ids": [p["id"]],
                    "findings": [finding(p["id"], "The Great Gatsby", "The Great Gatsby", "format")],
                    "comment_decisions": [], "editorial_verdict": "ready"}
    readers = Readers(handler=handler)
    result = FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    assert len(result["formats"]) == 1
    astra = next(r for r in readers.events if r["stage"] == "astra")
    meta = next(iter(astra["payload"]["paragraph_metadata"].values()))
    assert any(r["italic"] is True for r in meta["formatting"])
    check = next(r for r in readers.events if r["stage"] == "opus_read_checks_correction")
    assert check["payload"]["changes"][0]["format_proposals"]
    assert "Long-work titles" in check["system"]
    tracked, clean, _ = write_manuscripts(book, tmp_path / "out", result["accepted"], formats=result["formats"])
    assert paragraph_views(tracked, "reject") == result["original"]
    assert paragraph_views(clean) == result["accepted"]
    assert DocxPackage(tracked).tree("word/document.xml").find('.//' + qn('w:rPrChange')) is not None


def test_astra_profiles_fables_checked_text_not_the_source(make_book, tmp_path):
    book = make_book("He walks and waits and watches.")
    def handler(stage, model, payload, kwargs):
        if stage == "opus_read":
            p = payload["paragraphs"][0]
            return {"reviewed_ids": [p["id"]], "findings": [finding(p["id"], p["text"],
                    "He walked and waited and watched.")], "comment_decisions": [], "editorial_verdict": "ready"}
    readers = Readers(handler=handler)
    FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    values = {r["stage"]: r["payload"]["narrative_profile"]["paragraphs"][0]
              for r in readers.events if r["stage"] in {"opus_read", "astra"}}
    assert values["opus_read"]["present"] == 3 and values["astra"]["present"] == 0
    assert values["astra"]["past"] == 3


def test_real_footnote_and_endnote_parts_reach_final_readers_and_reject_audit(make_book, tmp_path):
    from lxml import etree
    from docproof.utils.xml_helpers import DocxPackage, qn
    from galley.fixed_documents import write_manuscripts, paragraph_views, _report
    book = make_book("Body text with notes.")
    pkg = DocxPackage(book)
    body = pkg.tree("word/document.xml").find(qn("w:body")).find(qn("w:p"))
    for kind in ("footnote", "endnote"):
        part = f"word/{kind}s.xml"
        notes = etree.Element(qn(f"w:{kind}s"), nsmap={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
        note = etree.SubElement(notes, qn(f"w:{kind}"), {qn("w:id"): "1"})
        p = etree.SubElement(note, qn("w:p"))
        marker = etree.SubElement(p, qn("w:r"))
        etree.SubElement(marker, qn(f"w:{kind}Ref"))
        run = etree.SubElement(p, qn("w:r"))
        etree.SubElement(run, qn("w:t")).text = f"The {kind} has teh error."
        pkg.add_part(part, notes)
        reference = etree.SubElement(body, qn("w:r"))
        etree.SubElement(reference, qn(f"w:{kind}Reference"), {qn("w:id"): "1"})
        rels = pkg.tree("word/_rels/document.xml.rels")
        etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                         Id=f"rIdTest{kind}", Type=f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s", Target=f"{kind}s.xml")
        types = pkg.tree("[Content_Types].xml")
        etree.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                         PartName="/" + part, ContentType=f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml")
    for name in ("word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"):
        pkg.mark_modified(name)
    pkg.save(book)
    before = book.read_bytes()
    def handler(stage, model, payload, kwargs):
        if stage == "opus_read":
            return {"reviewed_ids": [p["id"] for p in payload["paragraphs"]],
                    "findings": [finding(p["id"], "teh", "the", "spelling") for p in payload["paragraphs"] if "teh" in p["text"]],
                    "comment_decisions": [], "editorial_verdict": "ready"}
    readers = Readers(handler=handler)
    result = FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    assert book.read_bytes() == before
    for request in (r for r in readers.events if r["stage"] in {"opus_read", "astra"}):
        assert {m["part"] for m in request["payload"]["paragraph_metadata"].values()} >= {
            "word/footnotes.xml", "word/endnotes.xml"}
    tracked, clean, details = write_manuscripts(book, tmp_path / "out", result["accepted"])
    assert paragraph_views(tracked, "reject") == result["original"]
    assert paragraph_views(clean) == result["accepted"]
    assert len([v for v in result["accepted"].values() if "has the error" in v]) == 2
    report = _report(result, details)
    assert "word/footnotes.xml: 1" in report and "word/endnotes.xml: 1" in report


def test_poetry_takes_house_mechanics_from_both_typed_readers_and_the_verse_sweep(make_book, tmp_path):
    """Verse is proofread for house mechanics — spelling, the number stage,
    the deterministic glyph sweeps — by Sonnet and Luna, with the checks, and
    never enters the prose-only stages (story sheet, repair, whole-book reads)."""
    source = "teh Moon\n  waits, 20 times -- quiet\n‘bout now"
    book = make_book(source)

    def typed(stage, model, paragraphs, keys):
        assert stage == "verse"
        assert "subject_verb_agreement" not in keys and "number_style" not in keys
        if "spelling" not in keys:
            return []
        return [_typed_row(pid, text, "teh", "the", "spelling") for pid, text in paragraphs.items()]

    def handler(stage, model, payload, kwargs):
        if stage == "typed_screen":
            # The verse sweep's dash and apostrophe rows reach the pair screen
            # like any other local signal; both readers apply them.
            assert all(p["models"] == ["local:verse"] for site in payload["sites"] for p in site["proposals"])
            return {"decisions": [ruling(site, "apply", site["proposals"][0]["replacement"]) for site in payload["sites"]]}
        if stage == "numbers":
            assert payload["verse_ids"] == [x["para_id"] for x in payload["sites"]][:1]
            pid = payload["sites"][0]["para_id"]
            return {"reviewed_ids": [x["id"] for x in payload["sites"]],
                    "findings": [finding(pid, "20", "twenty", "number_style")],
                    "comment_decisions": [], "editorial_verdict": "ready"}

    readers = Readers(poetry=True, typed=typed, handler=handler)
    result = FixedWorkflow(book, tmp_path / "poetry", calls=readers).run()
    assert list(result["accepted"].values()) == ["the Moon\n  waits, twenty times—quiet\n’bout now"]
    assert result["poetry_only"]
    assert [row["stage"] for row in result["stages"]] == ["poetry", "typed", "numbers", "checks", "poetry_complete"]
    stages = {x["stage"] for x in readers.events}
    assert {"poetry", "verse", "typed_screen", "numbers", "checks_meaning", "checks_correction"} <= stages
    assert not stages & {"story_sheet", "broken_repair", "ensemble_sweep_opus", "ensemble_sweep_sol", "opus_read", "astra"}
    assert {x["model"] for x in readers.events if x["stage"] == "verse"} == {SONNET, LUNA}
    assert {x["model"] for x in readers.events} <= {SONNET, LUNA}
    # The number policy heads the number stage (and the checks that carry its
    # proposal); verse readers and the screen get the editorial brief alone.
    assert all("type 2 diabetes" not in x["system"] for x in readers.events
               if x["stage"] in {"poetry", "verse", "typed_screen"})
    assert all("type 2 diabetes" in x["system"] for x in readers.events if x["stage"] == "numbers")
    verse_policy = [x["system"] for x in readers.events if x["stage"] == "verse"][0]
    assert "never the poem's STRUCTURE" in verse_policy
    saved = json.loads((tmp_path / "poetry/stages/typed.json").read_text())
    verse_local = saved["evidence"]["verse_local"]
    assert verse_local["stage"] == "verse" and verse_local["verse_ids"] == ["body-0000"]
    assert [c["check"] for c in verse_local["checks"]] == ["verse_sweeps"]


def test_embedded_verse_takes_mechanics_from_any_reader_but_never_grammar(make_book, tmp_path):
    verse = "teh Moon\n  waits"

    def typed(stage, model, paragraphs, keys):
        rows = []
        for pid, text in paragraphs.items():
            if stage == "verse" and "spelling" in keys:
                rows.append(_typed_row(pid, text, "teh", "the", "spelling"))
            elif stage == "typed" and "subject_verb_agreement" in keys:
                rows.append(_typed_row(pid, text, "They was", "They were", "subject_verb_agreement"))
        return rows

    def handler(stage, model, payload, kwargs):
        if stage == "poetry":
            return {"classification": "mixed", "reason": "Verse and prose."}
        if stage == "poetry_sections":
            return {"paragraphs": [{"id": row["id"], "poetry": "\n" in row["text"]} for row in payload]}
        if stage == "typed_disputes":
            pytest.fail("Agreed readers must not trigger an Opus dispute")
        if stage in {"ensemble_sweep_opus", "ensemble_sweep_sol", "opus_read", "astra"}:
            poem = payload["paragraphs"][0]
            assert payload["poetry_ids"] == [poem["id"]]
            # A grammar "repair" of the poem is a sentence-level judgment and
            # is dropped; a house-mechanics fix from the same reader stands.
            rows = [finding(poem["id"], "waits", "wait", "grammar")]
            if stage == "opus_read":
                rows.append(finding(poem["id"], "Moon", "Moon.", "punctuation"))
            if stage == "astra":
                rows.append(finding(poem["id"], " waits", " waits", "spelling"))
            return {"reviewed_ids": [x["id"] for x in payload["paragraphs"]], "findings": rows,
                    "comment_decisions": [], "editorial_verdict": "ready"}

    readers = Readers(typed=typed, handler=handler)
    result = FixedWorkflow(make_book(verse, "They was here."), tmp_path / "mixed", calls=readers).run()
    assert list(result["accepted"].values()) == [verse.replace("teh", "the"), "They were here."]
    assert not result["poetry_only"]
    assert {row["model"] for row in readers.events if row["stage"] == "verse"} == {SONNET, LUNA}
    dropped = [h for h in result["history"] if h.get("dropped")]
    reasons = {h["reason"] for h in dropped}
    assert "Verse takes house mechanics only, never a change to its structure" in reasons
    checks = [row for row in readers.events if row["stage"] == "checks_meaning"]
    assert checks and checks[0]["payload"]["verse_ids"] == ["body-0000"]
    assert "Moon." not in result["accepted"]["body-0000"]


def _flow(make_book, tmp_path, readers=None, text="He waited for someone."):
    from types import SimpleNamespace
    from docproof.models import DocumentModel
    flow = FixedWorkflow(make_book(text), tmp_path / "unit", calls=readers or Readers())
    flow.original = {"p": text}
    flow.current = dict(flow.original)
    flow.prose_prepared = SimpleNamespace(doc=DocumentModel(str(flow.source), ()))
    return flow


def test_story_sheet_preserves_interleaved_manuscript_order(make_book, tmp_path):
    flow = _flow(make_book, tmp_path)
    flow.original = {"body-0000": "Opening text.", "table-0-r0-c0-p0": "Table in the middle.",
                     "body-0001": "Following text."}
    flow.current = dict(flow.original)
    flow._story()
    assert flow.calls.events[0]["payload"]["manuscript"] == [
        {"id": "body-0000", "text": "Opening text."},
        {"id": "table-0-r0-c0-p0", "text": "Table in the middle."},
        {"id": "body-0001", "text": "Following text."}]


def test_verse_edits_are_screened_like_prose_but_never_touch_structure(make_book, tmp_path):
    flow = _flow(make_book, tmp_path, text="teh Moon\n  waits")
    flow.poetry_ids = {"p"}
    # Conflicting spelling proposals go to the ordinary pair screen (which
    # drops them here) and never straight to Opus or into a comment.
    candidates = [_candidate(finding("p", "teh", fix, "spelling"), flow.current, SONNET)
                  for fix in ("the", "ten")]
    assert flow._adjudicate("typed", candidates, (SONNET, LUNA)) == []
    assert {x["stage"] for x in flow.calls.events} == {"typed_screen"}
    assert {x["model"] for x in flow.calls.events} == {SONNET, LUNA} and flow.questions == []
    # A sentence-level category never reaches the screen for a poem.
    grammar = _candidate(finding("p", "waits", "wait", "grammar"), flow.current, SONNET)
    assert flow._adjudicate("typed", [grammar], (SONNET, LUNA)) == []
    assert flow.history[-1]["reason"] == "Verse takes house mechanics only, never a change to its structure"
    # Structure is the poet's: a recased line head, an added terminal mark and
    # a lost line break are all dropped at application, whatever the category.
    for before, after in (("waits", "Waits"), ("waits", "waits."), ("\n  ", " ")):
        case = _candidate(finding("p", before, after, "spelling"), flow.current, SONNET)
        flow._apply("typed", [case])
        assert flow.current["p"] == "teh Moon\n  waits", (before, after)
    # A mid-line recase and a real misspelling are mechanics and apply.
    for before, after in (("Moon", "moon"), ("teh", "the")):
        flow._apply("typed", [_candidate(finding("p", before, after, "spelling"), flow.current, SONNET)])
    assert flow.current["p"] == "the moon\n  waits"


@pytest.mark.parametrize("answer", [ProviderResult(stop_reason="refusal"),
    ProviderResult(parsed={"findings": []})])
def test_incomplete_typed_read_is_never_a_clean_book(make_book, tmp_path, answer):
    readers = Readers(typed=lambda *args: answer)
    with pytest.raises(FixedWorkflowError, match="complete|coverage"):
        FixedWorkflow(make_book("A quiet paragraph."), tmp_path / "broken", calls=readers).run()
    assert not any(x["stage"] == "opus_read" for x in readers.events)
    assert not (tmp_path / "broken/result.json").exists()


def test_number_read_requires_every_extracted_site(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        if stage == "numbers":
            return {"reviewed_ids": [], "findings": [], "comment_decisions": [], "editorial_verdict": "ready"}
    readers = Readers(handler=handler)
    with pytest.raises(FixedWorkflowError, match="Number coverage"):
        FixedWorkflow(make_book("There were 20 birds."), tmp_path / "broken", calls=readers).run()


def test_read_rejects_missing_paragraphs_and_broken_repair_scope(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        return {"reviewed_ids": ["p"], "findings": [finding("p", "waited", "paused", "spelling")],
                "comment_decisions": [], "editorial_verdict": "ready"}
    flow = _flow(make_book, tmp_path, Readers(handler=handler))
    proposals, comments, coverage = flow._read("broken_repair", OPUS, ids=["p"])
    assert proposals == comments == []
    assert coverage[0]["paragraph_ids"] == ["p"]
    assert "scope" in flow.history[-1]["rejected_proposal"]["reason"]
    assert "ONLY genuinely broken sentences" in flow.calls.events[0]["system"]
    flow.calls.handler = lambda *args: {"reviewed_ids": [], "findings": [], "comment_decisions": [], "editorial_verdict": "ready"}
    with pytest.raises(FixedWorkflowError, match="paragraph coverage"):
        flow._read("astra", ASTRA)


def test_luna_rejection_is_settled_by_opus_instead_of_becoming_comment(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        if stage == "check_meaning":
            return {"decisions": [{"id": "p", "verdict": "reject", "reason": "Unsure."}]}
        if stage == "check_meaning_disputes":
            return {"decisions": [ruling(payload["sites"][0], replacement="He waited for somebody.")]}
    flow = _flow(make_book, tmp_path, Readers(handler=handler))
    before = dict(flow.current)
    flow.current["p"] = "He waited for somebody."
    flow._checks("check", before)
    assert flow.current["p"] == "He waited for somebody."
    assert flow.questions == []
    assert [x["model"] for x in flow.calls.events] == [LUNA, SONNET, OPUS, LUNA]


def test_formatting_enters_correction_check_and_rejected_format_is_removed(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        if stage == "check_correction":
            assert payload["changes"][0]["format_proposals"][0]["format"] == "italic"
            return {"decisions": [{"id": "p", "verdict": "reject", "reason": "No house basis."}]}
    flow = _flow(make_book, tmp_path, Readers(handler=handler))
    row = _candidate(finding("p", "someone", "someone", "format"), flow.current, OPUS,
                     format_types={"format": "italic"})
    before = flow._apply("opus_read", [row])
    flow._checks("check", before)
    assert flow.formats == []
    assert [x["stage"] for x in flow.calls.events] == ["check_correction", "check_correction_sonnet", "check_correction_disputes"]


def test_comment_resolved_by_rejected_edit_is_reviewed_on_restored_text(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        if stage == "opus_read_checks_meaning":
            return {"decisions": [{"id": "p", "verdict": "reject", "reason": "Invented identity."}]}
        if stage == "opus_read_comment_review":
            assert payload["paragraphs"]["p"] == "He waited for someone."
            return {"decisions": [comment_decision(q) for q in payload["comments"]]}
    flow = _flow(make_book, tmp_path, Readers(handler=handler))
    flow._question("p", flow.current["p"], "Who was he waiting for?", "The intended identity", "An unresolved reference.", "typed")
    decisions = [comment_decision(flow.questions[0], "drop")]
    before = flow._apply("opus_read", [_candidate(finding("p", "someone", "Mary"), flow.current, OPUS)])
    flow._checks("opus_read_checks", before)
    flow._comments(decisions, "opus_read", before=before, model=OPUS)
    assert flow.current["p"] == before["p"]
    assert len(flow.questions) == 1
    assert flow.calls.events[-1]["stage"] == "opus_read_comment_review"


def test_late_correction_elsewhere_refreshes_retained_comment(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        assert stage == "astra_comment_review"
        assert payload["changed_passages"] == [{"para_id": "other", "before": "She was the visitor.", "after": "Jo was the visitor."}]
        return {"decisions": [comment_decision(q, "drop") for q in payload["comments"]]}
    flow = _flow(make_book, tmp_path, Readers(handler=handler), text="The visitor waited.")
    flow.original["other"] = flow.current["other"] = "She was the visitor."
    flow._question("p", flow.current["p"], "Who is the visitor?", "The visitor's identity", "An unresolved identity.", "opus_read")
    before = dict(flow.current)
    flow.current["other"] = "Jo was the visitor."
    flow._comments([comment_decision(flow.questions[0])], "astra", before=before, model=ASTRA)
    assert flow.questions == []
    assert len(flow.calls.events) == 1


def test_final_comment_cannot_silently_anchor_to_wrong_repeated_word(make_book, tmp_path):
    flow = _flow(make_book, tmp_path, text="He saw her, and her visitor left.")
    flow._question("p", flow.current["p"], "Who does the latter her mean?", "The intended identity", "An unclear reference.", "opus_read")
    flow._comments([comment_decision(flow.questions[0], quote="her")], "astra", model=ASTRA)
    assert flow.questions == []
    assert "unambiguous contextual quote" in flow.history[-1]["rejected_proposal"]["reason"]


def test_failed_parallel_batch_cancels_queued_reads(monkeypatch):
    from concurrent.futures import Future
    from galley.fixed_parallel import ReadScheduler
    from galley.fixed_policy import configuration
    scheduled = []
    class ControlledExecutor:
        def __init__(self, **kwargs): pass
        def __enter__(self): return self
        def __exit__(self, *args): return False
        def submit(self, callback, *args):
            future = Future()
            if not scheduled:
                future.set_exception(FixedWorkflowError("failed read"))
            scheduled.append(future)
            return future
    monkeypatch.setattr("galley.fixed_parallel.ThreadPoolExecutor", ControlledExecutor)
    with pytest.raises(FixedWorkflowError, match="failed read"):
        ReadScheduler(configuration()).map([(SONNET, lambda: None)] * 4)
    assert all(future.cancelled() for future in scheduled[1:])


def test_new_astra_question_gets_explicit_final_astra_comment_review(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        if stage == "astra":
            row = payload["paragraphs"][0]
            return {"reviewed_ids": [row["id"]],
                    "findings": [finding(row["id"], "someone", "", "author_question", action="query", missing="The intended identity")],
                    "comment_decisions": [], "editorial_verdict": "ready"}
        if stage == "astra_screen":
            return {"decisions": [{**ruling(x, "query"), "missing_knowledge": "The intended identity", "question": "Who was he waiting for?"}
                                  for x in payload["sites"]]}
    readers = Readers(handler=handler)
    result = FixedWorkflow(make_book("He waited for someone."), tmp_path / "run", calls=readers).run()
    assert len(result["questions"]) == 1
    stages = [r["stage"] for r in readers.events]
    review = stages.index("astra_comment_review")
    assert stages.index("astra") < review < stages.index("final_astra")
    assert readers.events[review]["model"] == ASTRA
    # The surviving question is then judged again by the second Astra reading.
    final = readers.events[stages.index("final_astra")]
    assert [c["id"] for c in final["payload"]["comments"]] == [result["questions"][0]["id"]]


def test_stale_source_fails_before_any_model_call(make_book, tmp_path):
    source = make_book("A quiet paragraph.")
    readers = Readers()
    flow = FixedWorkflow(source, tmp_path / "run", calls=readers)
    doc = Document(source)
    doc.add_paragraph("Changed after intake.")
    doc.save(source)
    with pytest.raises(FixedWorkflowError, match="source changed"):
        flow.run()
    assert readers.events == []
    with pytest.raises(FixedWorkflowError, match="source or fixed recipe changed"):
        FixedWorkflow(source, tmp_path / "run", calls=readers)
    assert readers.events == []


@pytest.mark.parametrize("actual", [["p1", "p2"], ["c1", "p2", "p1"], ["p1", "c2", "p2", "c1"]])
def test_typed_context_coverage_is_normalized_without_changing_response(actual):
    from docproof.models import Chunk, ParagraphRef
    from galley.fixed_workflow import _typed_response
    def p(pid):
        return ParagraphRef(pid, "word/document.xml", "body", "A quiet room.", "Normal")
    chunk = Chunk("chunk-2", (p("p1"), p("p2")), 20, (p("c1"), p("c2")))
    parsed = {"findings": [], "reviewed_paragraph_ids": actual[:]}
    response = ProviderResult(parsed=parsed)
    normalized = _typed_response(response, chunk)
    assert set(normalized.parsed["reviewed_paragraph_ids"]) == {"p1", "p2"}
    assert response.parsed == parsed and response.parsed["reviewed_paragraph_ids"] == actual
    assert normalized is not response and normalized.parsed is not response.parsed


@pytest.mark.parametrize("actual", [["p1", "c1"], ["p1", "p2", "unknown"],
                                   ["p1", "p2", "p2"], ["p1", "p2", "c1", "c1"], [], None])
def test_typed_context_tolerance_never_invents_or_duplicates_owned_coverage(actual):
    from docproof.models import Chunk, ParagraphRef
    from galley.fixed_workflow import _typed_response
    def p(pid):
        return ParagraphRef(pid, "word/document.xml", "body", "A quiet room.", "Normal")
    chunk = Chunk("chunk-2", (p("p1"), p("p2")), 20, (p("c1"),))
    with pytest.raises(FixedWorkflowError, match="Typed paragraph coverage"):
        _typed_response(ProviderResult(parsed={"findings": [], "reviewed_paragraph_ids": actual}), chunk)


def test_replay_rejects_tampered_stage_evidence(make_book, tmp_path):
    source = make_book("A quiet paragraph.")
    directory = tmp_path / "run"
    FixedWorkflow(source, directory, calls=Readers()).run()
    path = directory / "stages/poetry.json"
    data = json.loads(path.read_text())
    data["accepted_sha256"] = "wrong"
    path.write_text(json.dumps(data))
    with pytest.raises(FixedWorkflowError, match="Saved poetry evidence changed"):
        FixedWorkflow(source, directory, calls=Readers()).run()


@pytest.mark.parametrize("decision", ["apply", "drop"])
def test_local_grammar_missed_by_both_readers_requires_pair_screen_then_luna(
        make_book, tmp_path, monkeypatch, local_scans, decision):
    from types import SimpleNamespace
    scanned = []

    class LocalTool:
        picky = False

        def check(self, text):
            scanned.append(text)
            return [SimpleNamespace(offset=text.index("is"), error_length=2,
                replacements=["are"], rule_id="TEST_SUBJECT_AGREEMENT",
                rule_issue_type="grammar", message="The plural subject takes are.")]

        def close(self):
            pass

    def handler(stage, model, payload, kwargs):
        if stage == "typed_screen":
            assert model in {SONNET, LUNA}
            assert len(payload["sites"]) == 1
            proposal = payload["sites"][0]["proposals"][0]
            assert proposal["models"] == ["local:languagetool"]
            return {"decisions": [ruling(site, decision, "are") for site in payload["sites"]]}

    # All local generators and the scanner's real filtering/coverage adapter
    # run; only the external Java transport is replaced.
    monkeypatch.setattr("galley.fixed_local.collect_local_candidates", local_scans["initial"])
    monkeypatch.setattr("galley.fixed_local.default_lt_factory", lambda dictionary: LocalTool())
    readers = Readers(handler=handler)
    result = FixedWorkflow(make_book("They is here."), tmp_path / "local", calls=readers).run()

    assert list(result["accepted"].values()) == ["They are here." if decision == "apply" else "They is here."]
    assert result["questions"] == []
    assert scanned == ["They is here."]
    events = readers.events
    typed = [row for row in events if row["stage"] == "typed"]
    assert {row["model"] for row in typed} == {SONNET, LUNA}
    assert {row["model"] for row in events if row["stage"] == "typed_screen"} == {SONNET, LUNA}
    assert not any(row["stage"] == "typed_disputes" for row in events)
    checks = [row for row in events if row["stage"] in {"checks_meaning", "checks_correction"}]
    assert len(checks) == (2 if decision == "apply" else 0)
    assert all(row["model"] == LUNA for row in checks)
    if checks:
        assert all(row["payload"]["changes"][0]["after"] == "They are here." for row in checks)
    saved = json.loads((tmp_path / "local/stages/typed.json").read_text())
    local = saved["evidence"]["local"]
    assert local["proposal_count"] == 1
    assert next(check for check in local["checks"] if check["check"] == "languagetool")["proposal_count"] == 1


def test_local_completion_runs_after_ensemble_fable_and_astra(make_book, tmp_path, monkeypatch):
    completion_calls = []

    def complete(prepared, original, texts, *args, **kwargs):
        completion_calls.append((kwargs["stage"], dict(texts)))
        pid = next(iter(texts))
        rows = [_local_row(pid, "teh", "the", source="recurrence", category="spelling")] if "teh" in texts[pid] else []
        return rows, {"recurrence_candidates": len(rows)}

    def handler(stage, model, payload, kwargs):
        if stage == "local_completion_screen":
            assert model in {SONNET, LUNA}
            return {"decisions": [ruling(site, replacement=site["proposals"][0]["replacement"])
                                  for site in payload["sites"]]}

    monkeypatch.setattr("galley.fixed_local.collect_completion_candidates", complete)
    readers = Readers(handler=handler)
    result = FixedWorkflow(make_book("She found teh letter."), tmp_path / "completion", calls=readers).run()

    assert list(result["accepted"].values()) == ["She found the letter."]
    assert [stage for stage, _ in completion_calls] == ["completion", "completion_opus_read", "completion_astra", "completion_final_astra"]
    assert all(texts == result["accepted"] for stage, texts in completion_calls[1:])
    events = readers.events
    stages = [row["stage"] for row in events]
    assert max(stages.index("ensemble_sweep_opus"), stages.index("ensemble_sweep_sol")) < stages.index("local_completion_screen")
    assert stages.index("local_completion_screen") < stages.index("local_completion_checks_meaning")
    assert stages.index("local_completion_checks_meaning") < stages.index("local_completion_checks_correction") < stages.index("opus_read")
    assert [row for row in events if row["stage"] == "opus_read"][0]["payload"]["paragraphs"][0]["text"] == "She found the letter."
    assert result["questions"] == []
    for stage in ("opus_read", "astra", "final_astra"):
        saved = json.loads((tmp_path / f"completion/stages/{stage}.json").read_text())
        assert saved["evidence"]["local"] == {"recurrence_candidates": 0}


def test_fable_edit_is_propagated_by_the_post_fable_completion_pass(make_book, tmp_path, monkeypatch):
    def complete(prepared, original, texts, *args, **kwargs):
        if kwargs["stage"] == "completion_opus_read" and texts["body-0000"] != original["body-0000"]:
            return [_local_row("body-0001", "Beckham", "Brooks", source="recurrence", category="spelling")], {"seeded": 1}
        return [], {"seeded": 0}

    def handler(stage, model, payload, kwargs):
        if stage == "opus_read":
            row = payload["paragraphs"][0]
            return {"reviewed_ids": [x["id"] for x in payload["paragraphs"]],
                    "findings": [finding(row["id"], "Beckham", "Brooks", "spelling")] if row["id"] == "body-0000" else [],
                    "comment_decisions": [], "editorial_verdict": "ready"}
        if stage == "local_completion_opus_read_screen":
            return {"decisions": [ruling(site, replacement=site["proposals"][0]["replacement"])
                                  for site in payload["sites"]]}

    monkeypatch.setattr("galley.fixed_local.collect_completion_candidates", complete)
    readers = Readers(handler=handler)
    result = FixedWorkflow(make_book("Beckham smiled.", "Then Beckham left."), tmp_path / "propagate", calls=readers).run()
    assert list(result["accepted"].values()) == ["Brooks smiled.", "Then Brooks left."]
    astra = [row for row in readers.events if row["stage"] == "astra"]
    assert astra and [p["text"] for p in astra[0]["payload"]["paragraphs"]] == ["Brooks smiled.", "Then Brooks left."]
    assert json.loads((tmp_path / "propagate/stages/opus_read.json").read_text())["evidence"]["local"] == {"seeded": 1}


def test_poetry_runs_the_verse_sweep_and_no_prose_collector(make_book, tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("Poetry must not enter a prose proofreading collector")

    monkeypatch.setattr("galley.fixed_local.collect_local_candidates", forbidden)
    monkeypatch.setattr("galley.fixed_local.collect_completion_candidates", forbidden)
    readers = Readers(poetry=True)
    result = FixedWorkflow(make_book("The Moon\n  waits, 20 times\n—quiet"),
                           tmp_path / "poetry-local", calls=readers).run()
    assert result["accepted"] == result["original"]
    assert {row["stage"] for row in readers.events} == {"poetry", "verse", "numbers"}
    saved = json.loads((tmp_path / "poetry-local/stages/typed.json").read_text())
    assert saved["evidence"]["verse_local"]["proposal_count"] == 0


def test_local_collector_cannot_emit_a_candidate_in_embedded_poetry(
        make_book, tmp_path, monkeypatch):
    def handler(stage, model, payload, kwargs):
        if stage == "poetry":
            return {"classification": "mixed", "reason": "Verse and prose."}
        if stage == "poetry_sections":
            return {"paragraphs": [{"id": row["id"], "poetry": "\n" in row["text"]} for row in payload]}

    def leak(prepared, texts, *args, **kwargs):
        pid = next(iter(kwargs["poetry_ids"]))
        return [_local_row(pid, "Moon", "moon", category="spelling")], {}

    monkeypatch.setattr("galley.fixed_local.collect_local_candidates", leak)
    readers = Readers(handler=handler)
    with pytest.raises(FixedWorkflowError, match="protected poetry"):
        FixedWorkflow(make_book("The Moon\n  waits", "She waited by the door."),
                      tmp_path / "mixed-local", calls=readers).run()
    assert not any(row["stage"] in {"typed_disputes", "opus_read", "astra"} for row in readers.events)


def test_real_local_generators_exclude_embedded_poetry_at_both_checkpoints(
        make_book, tmp_path, monkeypatch, local_scans):
    scanned = []

    class LocalTool:
        def check(self, text):
            scanned.append(text)
            return []

        def close(self):
            pass

    def handler(stage, model, payload, kwargs):
        if stage == "poetry":
            return {"classification": "mixed", "reason": "Verse and prose."}
        if stage == "poetry_sections":
            return {"paragraphs": [{"id": row["id"], "poetry": "\n" in row["text"]} for row in payload]}

    monkeypatch.setattr("galley.fixed_local.collect_local_candidates", local_scans["initial"])
    monkeypatch.setattr("galley.fixed_local.collect_completion_candidates", local_scans["completion"])
    monkeypatch.setattr("galley.fixed_local.default_lt_factory", lambda dictionary: LocalTool())
    source = make_book("teh Moon\n  waits, 20 times\n—quiet", "A quiet paragraph.")
    readers = Readers(handler=handler)
    result = FixedWorkflow(source, tmp_path / "real-mixed-local", calls=readers).run()

    assert scanned == ["A quiet paragraph."]
    assert result["accepted"] == result["original"]
    assert result["questions"] == []
    for stage, expected in (("typed", "initial"), ("ensemble_sweep", "completion"),
                            ("opus_read", "completion_opus_read"), ("astra", "completion_astra")):
        saved = json.loads((tmp_path / f"real-mixed-local/stages/{stage}.json").read_text())
        local = saved["evidence"]["local"]
        assert local["stage"] == expected
        assert local["paragraph_ids"] == ["body-0001"]
        assert local["excluded_poetry_ids"] == ["body-0000"]
        assert all(check["paragraph_ids"] == ["body-0001"] for check in local["checks"])


def test_missing_languagetool_prevents_later_stages_after_parallel_reads(
        make_book, tmp_path, monkeypatch, local_scans):
    from galley.fixed_local import FixedLocalError

    monkeypatch.setattr("galley.fixed_local.collect_local_candidates", local_scans["initial"])
    monkeypatch.setattr("docproof.languagetool.AVAILABLE", False)
    readers = Readers()
    with pytest.raises((FixedLocalError, FixedWorkflowError), match="LanguageTool"):
        FixedWorkflow(make_book("A quiet paragraph."), tmp_path / "missing-local", calls=readers).run()
    assert not any(row["stage"] in {"opus_read", "astra"} for row in readers.events)
    assert not (tmp_path / "missing-local/result.json").exists()


def test_stylistic_diagnostic_cannot_become_a_local_proofreading_candidate(
        make_book, tmp_path, monkeypatch):
    def inappropriate(prepared, texts, *args, **kwargs):
        pid = next(iter(texts))
        return [_local_row(pid, texts[pid], texts[pid], source="reading_level", category="reading_level")], {}

    monkeypatch.setattr("galley.fixed_local.collect_local_candidates", inappropriate)
    readers = Readers()
    with pytest.raises(FixedWorkflowError, match="stylistic diagnostic"):
        FixedWorkflow(make_book("A quiet paragraph."), tmp_path / "style-local", calls=readers).run()
    assert not any(row["stage"] == "typed_disputes" for row in readers.events)


def test_unchanged_rejected_local_site_is_not_paid_for_again_at_completion(
        make_book, tmp_path, monkeypatch):
    def initial(prepared, texts, *args, **kwargs):
        return [_local_row(next(iter(texts)), "quiet", "calm")], {}

    def completion(prepared, original, texts, *args, **kwargs):
        return [_local_row(next(iter(texts)), "quiet", "calm", source="recurrence")], {}

    monkeypatch.setattr("galley.fixed_local.collect_local_candidates", initial)
    monkeypatch.setattr("galley.fixed_local.collect_completion_candidates", completion)
    readers = Readers()  # Both screeners reject the stylistic suggestion.
    result = FixedWorkflow(make_book("A quiet paragraph."), tmp_path / "dedup-local", calls=readers).run()
    assert result["accepted"] == result["original"] and result["questions"] == []
    assert sum(row["stage"] == "typed_screen" for row in readers.events) == 2
    assert not any(row["stage"] in {"typed_disputes", "local_completion_screen", "local_completion_disputes",
                                    "local_completion_opus_read_screen", "local_completion_astra_screen"}
                   for row in readers.events)


def test_call_coverage_freezes_all_assigned_inventories_and_only_known_read_context():
    from galley.fixed_workflow import _call_coverage, FRONTIER_SCHEMA, READ_SCHEMA, CHECK_SCHEMA, DECISIONS
    payload = {"paragraphs": [{"id": "p1"}], "context": {"c1": "Read-only context."},
               "focused_sites": [{"id": "f1"}], "comments": [{"id": "q1"}]}
    contract = _call_coverage(payload, FRONTIER_SCHEMA)
    assert contract == {
        "reviewed_ids": {"ids": ["p1"], "id_key": None, "context_ids": ["c1"]},
        "reviewed_check_ids": {"ids": ["f1"], "id_key": None, "context_ids": []},
        "comment_decisions": {"ids": ["q1"], "id_key": "id", "context_ids": []}}
    assert _call_coverage({"sites": [{"id": "n1"}], "paragraphs": {"p1": "Text."}}, READ_SCHEMA)["reviewed_ids"]["ids"] == ["n1"]
    assert _call_coverage({"changes": [{"id": "p1"}]}, CHECK_SCHEMA)["decisions"]["ids"] == ["p1"]
    assert _call_coverage({"sites": [{"id": "d1"}]}, DECISIONS)["decisions"]["ids"] == ["d1"]


@pytest.mark.parametrize("defect", ["missing_quote", "unknown_paragraph", "bad_occurrence", "empty_quote"])
def test_unanchored_reader_proposal_is_audited_without_edit_or_comment(make_book, tmp_path, defect):
    from galley.fixed_workflow import _hash
    flow = FixedWorkflow(make_book("A quiet room."), tmp_path / "run", calls=Readers())
    texts = {"p1": "A quiet room."}
    row = finding("p1", "quiet", "silent")
    if defect == "missing_quote":
        row["quote"] = "A quote longer than the entire paragraph that the reader invented."
    elif defect == "unknown_paragraph":
        row["para_id"] = "another-paragraph"
    elif defect == "bad_occurrence":
        row["occurrence"] = 2
    else:
        row["quote"] = ""
    before = json.loads(json.dumps(row))
    assert flow._reader_candidate("opus_read", row, texts, OPUS) is None
    rejected = flow.history[0]["rejected_proposal"]
    assert rejected["status"] == "rejected_no_anchor"
    assert rejected["finding"] == row == before
    assert rejected["reviewed_sha256"] == _hash(texts)
    assert flow.questions == [] and texts == {"p1": "A quiet room."}
    valid = flow._reader_candidate("opus_read", finding("p1", "quiet", "silent"), texts, OPUS)
    assert valid is not None
    text = texts["p1"]
    assert text[:valid["start"]] + valid["replacement"] + text[valid["end"]:] == "A silent room."


def test_rejected_model_proposal_handling_does_not_hide_local_or_internal_failures(make_book, tmp_path, monkeypatch):
    from types import SimpleNamespace
    flow = FixedWorkflow(make_book("A quiet room."), tmp_path / "run", calls=Readers())
    texts = {"p1": "A quiet room."}
    row = {**finding("p1", "Invented quote.", "Replacement."), "source": "local:sweeps"}
    with pytest.raises(FixedWorkflowError, match="quote does not occur"):
        flow._local_candidates([row], texts=texts, prepared=SimpleNamespace(query_types=(), format_types={}))
    assert flow.history == []
    def unexpected(*args, **kwargs):
        raise RuntimeError("Unrelated integrity failure")
    monkeypatch.setattr("galley.fixed_workflow._candidate", unexpected)
    with pytest.raises(RuntimeError, match="Unrelated integrity failure"):
        flow._reader_candidate("typed", row, texts, SONNET)
    assert flow.history == []


def test_broken_sentence_reader_rejects_unanchored_repairs_but_keeps_coverage(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        return {"reviewed_ids": ["p"], "findings": [finding("p", "An invented longer sentence.", "A rewrite.", "broken_sentence")],
                "comment_decisions": [], "editorial_verdict": "ready"}
    flow = _flow(make_book, tmp_path, Readers(handler=handler))
    proposals, comments, coverage = flow._read("broken_repair", OPUS, ids=["p"])
    assert proposals == comments == []
    assert coverage[0]["paragraph_ids"] == ["p"]
    assert flow.history[0]["rejected_proposal"]["status"] == "rejected_no_anchor"


def test_rejected_proposal_stage_evidence_is_stable_across_reader_completion_order(make_book, tmp_path):
    flow = FixedWorkflow(make_book("A quiet room."), tmp_path / "run", calls=Readers())
    texts = {"p1": "A quiet room."}
    for label, model in (("opus", OPUS), ("sol", SOL)):
        flow._reader_candidate("ensemble_sweep_" + label, finding("p1", "Missing text.", "Replacement."), texts, model)
    flow._record("ensemble_sweep", readings=[])
    path = flow.directory / "stages/ensemble_sweep.json"
    before = path.read_bytes()
    assert len(json.loads(before)["evidence"]["rejected_proposals"]) == 2
    flow.history.reverse()
    flow._record("ensemble_sweep", readings=[])
    assert path.read_bytes() == before


@pytest.mark.parametrize("defect", ["scope", "control", "format_text", "format_unknown", "format_italic"])
def test_invalid_model_proposal_drops_only_bad_suggestion(make_book, tmp_path, defect):
    flow = _flow(make_book, tmp_path)
    bad = finding("p", "someone", "someone", "format")
    options = {"format_types": {"format": "italic"},
               "formatting": {"p": [{"start": 0, "end": len(flow.current["p"]), "italic": False}]}}
    if defect == "scope":
        options["allowed_categories"] = {"broken_sentence"}
    elif defect == "control":
        bad["replacement"] = "some\x00one"
    elif defect == "format_text":
        bad["replacement"] = "Mary"
    elif defect == "format_unknown":
        options["formatting"]["p"][0]["italic"] = None
    else:
        options["formatting"]["p"][0]["italic"] = True
    assert flow._reader_candidate("opus_read", bad, flow.current, OPUS, **options) is None
    good = flow._reader_candidate("opus_read", finding("p", "waited", "waits"), flow.current, OPUS)
    flow._apply("opus_read", [good])
    assert flow.current["p"] == "He waits for someone."
    assert flow.formats == flow.questions == []
    assert flow.history[0]["rejected_proposal"]["finding"] == bad


@pytest.mark.parametrize("defect", ["blank_question", "blank_knowledge", "control", "missing_quote"])
def test_invalid_comment_drops_while_valid_comment_survives(make_book, tmp_path, defect):
    flow = _flow(make_book, tmp_path)
    for missing in ("Identity", "Location"):
        flow._question("p", flow.current["p"], "What is the " + missing + "?", missing, "Unresolved.", "typed")
    decisions = [comment_decision(q) for q in flow.questions]
    bad = decisions[0]
    bad[{"blank_question": "question", "blank_knowledge": "missing_knowledge",
         "control": "question", "missing_quote": "quote"}[defect]] = {
             "blank_question": " ", "blank_knowledge": " ", "control": "Who\x00?", "missing_quote": "Invented quote."}[defect]
    flow._comments(decisions, "astra", model=ASTRA)
    assert len(flow.questions) == 1
    assert flow.questions[0]["missing_knowledge"] == "Location"
    assert flow.history[-2]["rejected_proposal"]["model"] == ASTRA


@pytest.mark.parametrize("action", ["apply", "query"])
def test_invalid_opus_disposition_drops_site_without_losing_good_edit(make_book, tmp_path, action):
    def handler(stage, model, payload, kwargs):
        if stage == "typed_screen":
            return {"decisions": [ruling(site, "apply" if model == SONNET else "drop", site["proposals"][0]["replacement"])
                                  for site in payload["sites"]]}
        assert stage == "typed_disputes"
        return {"decisions": [ruling(site, action, "bad\x00text") if site["before"] == "someone"
                              else ruling(site, "apply", site["proposals"][0]["replacement"]) for site in payload["sites"]]}
    flow = _flow(make_book, tmp_path, Readers(handler=handler))
    rows = [_candidate(finding("p", before, after), flow.current, SONNET)
            for before, after in (("waited", "waits"), ("someone", "Mary"))]
    flow._apply("typed", flow._adjudicate("typed", rows, force=True))
    assert flow.current["p"] == "He waits for someone."
    assert flow.questions == []
    assert len([h for h in flow.history if h.get("rejected_proposal")]) == 1


def test_invalid_check_adjudication_restores_text_and_removes_disputed_format(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        if stage == "check_meaning":
            return {"decisions": [{"id": "p", "verdict": "reject", "reason": "Unsafe edit."}]}
        if stage == "check_meaning_disputes":
            return {"decisions": [ruling({"id": "p"}, "apply", "Bad\x00paragraph.")]}
    flow = _flow(make_book, tmp_path, Readers(handler=handler))
    before = flow._apply("opus_read", [_candidate(finding("p", "waited", "waits"), flow.current, OPUS),
                                  _candidate(finding("p", "someone", "someone", "format"), flow.current, OPUS,
                                             format_types={"format": "italic"})])
    flow._checks("check", before)
    assert flow.current == before
    assert flow.formats == flow.questions == []
    assert len([h for h in flow.history if h.get("rejected_proposal")]) == 1


def test_a_guarded_proposal_is_dropped_at_application_with_a_receipt(make_book, tmp_path):
    """A proposal every stage agreed on can still be the wrong edit: the
    application guard drops it into history instead of rewriting the text."""
    text = "The log said the drone lifted at 18:03 UTC, nine minutes late."
    flow = _flow(make_book, tmp_path, text=text)
    row = _candidate(finding("p", "18:03", "6:03 p.m.", "number_style"), flow.current, LUNA)
    before = flow._apply("numbers", [row])
    assert flow.current["p"] == before["p"] == text
    assert flow.history[-1]["dropped"] == row
    assert flow.history[-1]["reason"].startswith("guard: ") and "24-hour" in flow.history[-1]["reason"]
    # The same stage still applies a correction the guards have no view on.
    flow._apply("numbers", [_candidate(finding("p", "nine", "9", "number_style"), flow.current, LUNA)])
    assert flow.current["p"].endswith("9 minutes late.")


def test_number_scope_noise_does_not_prevent_valid_number_correction(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        if stage == "numbers":
            return {"reviewed_ids": [s["id"] for s in payload["sites"]],
                    "findings": [finding("p", "birds", "bees", "grammar"),
                                 finding("p", "20", "twenty", "number_style")],
                    "comment_decisions": [], "editorial_verdict": "ready"}
    flow = _flow(make_book, tmp_path, Readers(handler=handler), text="There were 20 birds.")
    flow._numbers()
    assert flow.current["p"] == "There were twenty birds."
    assert len([h for h in flow.history if h.get("rejected_proposal")]) == 2


def test_comment_resolution_cannot_depend_on_discarded_reader_proposal(make_book, tmp_path):
    flow = _flow(make_book, tmp_path)
    flow._question("p", flow.current["p"], "Who is the visitor?", "Identity", "Unresolved.", "typed")
    before = dict(flow.current)
    assert flow._reader_candidate("opus_read", finding("p", "someone", "Mary\x00"), flow.current, OPUS) is None
    flow._comments([comment_decision(flow.questions[0], "drop")], "opus_read", before=before, model=OPUS)
    assert len(flow.questions) == 1
    assert flow.calls.events[-1]["stage"] == "opus_read_comment_review"
    assert flow.current == before


# --- final walk-through scope, book map, reader guards ------------------------

def test_walkthrough_prompts_reach_only_the_final_readers_and_their_checks(make_book, tmp_path):
    from galley.press_prompt import CONTINUITY_TASK, FINAL_WALKTHROUGH, FINAL_WALKTHROUGH_CHECK

    def handler(stage, model, payload, kwargs):
        if stage == "opus_read":
            row = payload["paragraphs"][0]
            return {"reviewed_ids": [x["id"] for x in payload["paragraphs"]],
                    "findings": [{**finding(row["id"], "brand new", "brand-new", "usage"), "evidence": []}],
                    "comment_decisions": [], "editorial_verdict": "ready"}

    readers = Readers(handler=handler)
    result = FixedWorkflow(make_book("A brand new board waited."), tmp_path / "scope", calls=readers).run()
    assert list(result["accepted"].values()) == ["A brand-new board waited."]
    systems = {}
    for row in readers.events:
        systems.setdefault(row["stage"], row["system"])
    assert all(FINAL_WALKTHROUGH in systems[stage] for stage in ("opening_read", "opus_read", "astra", "final_astra"))
    assert all(FINAL_WALKTHROUGH not in systems[stage] for stage in systems
               if stage not in {"opening_read", "opus_read", "astra", "final_astra"})
    from galley.press_prompt import FINAL_GATE_TASK
    assert FINAL_GATE_TASK in systems["final_astra"]
    assert all(FINAL_GATE_TASK not in systems[stage] for stage in systems if stage != "final_astra")
    assert CONTINUITY_TASK in systems["continuity"]
    assert all(CONTINUITY_TASK not in systems[stage] for stage in systems if stage != "continuity")
    assert FINAL_WALKTHROUGH_CHECK in systems["opus_read_checks_meaning"]
    assert FINAL_WALKTHROUGH_CHECK not in systems.get("typed_screen", "")
    fable = [row for row in readers.events if row["stage"] == "opus_read"]
    assert all(shared_context(row)["book_map"]["complete_inventory"] is True for row in fable)
    assert all("book_map" not in row["payload"] for row in fable)
    assert "usage" in fable[0]["schema"]["properties"]["findings"]["items"]["properties"]["category"]["enum"]
    from galley.fixed_workflow import READ_SCHEMA
    assert "usage" not in READ_SCHEMA["properties"]["findings"]["items"]["properties"]["category"]["enum"]


def test_book_map_lists_headings_and_running_heads_and_header_edits_round_trip(tmp_path):
    from galley.fixed_documents import write_manuscripts
    document = Document()
    document.add_paragraph("CHAPTER 2", style="Heading 1")
    document.add_paragraph("The first body paragraph of the chapter.")
    document.add_paragraph("The second body paragraph of the chapter.")
    document.add_paragraph("TOP TEN")
    document.add_paragraph("Only one item.")
    document.sections[0].header.paragraphs[0].text = "CHAPTER ONE"
    source = tmp_path / "book.docx"
    document.save(source)
    header_id = None

    def handler(stage, model, payload, kwargs):
        nonlocal header_id
        if stage == "opus_read":
            heads = [x for x in payload["paragraphs"] if x["id"].startswith("header")]
            if heads:
                header_id = heads[0]["id"]
                return {"reviewed_ids": [x["id"] for x in payload["paragraphs"]],
                        "findings": [{**finding(header_id, "CHAPTER ONE", "CHAPTER 1", "structure"), "evidence": []}],
                        "comment_decisions": [], "editorial_verdict": "ready"}

    readers = Readers(handler=handler)
    result = FixedWorkflow(source, tmp_path / "map", calls=readers).run()
    assert result["accepted"][header_id] == "CHAPTER 1"
    book_map = shared_context(next(row for row in readers.events if row["stage"] == "opus_read"))["book_map"]
    assert [(h["text"], h["signal"], h["body_paragraphs"]) for h in book_map["headings"]] == [
        ("CHAPTER 2", "style", 2), ("TOP TEN", "caps_line", 1)]
    assert [(h["location"], h["text"]) for h in book_map["headers_footers"]] == [("header", "CHAPTER ONE")]
    tracked, clean, _ = write_manuscripts(source, tmp_path / "final", result["accepted"])
    assert Document(clean).sections[0].header.paragraphs[0].text == "CHAPTER 1"


@pytest.mark.parametrize("replacement, problem", [
    ("*Huckleberry Finn*", "markup"),
    ("Huckleberry\nFinn", "line break"),
])
def test_reader_replacements_that_are_not_manuscript_text_are_rejected(replacement, problem):
    from galley.fixed_workflow import RejectedModelProposal
    texts = {"p1": "She read Huckleberry Finn twice."}
    with pytest.raises(RejectedModelProposal, match=problem) as info:
        _candidate(finding("p1", "Huckleberry Finn", replacement), texts, OPUS)
    assert info.value.status == "rejected_invalid_proposal"
    with pytest.raises(RejectedModelProposal, match="paragraph start"):
        _candidate(finding("p1", "She", " She"), texts, OPUS)
    assert _candidate(finding("p1", "She read", "She had read"), texts, OPUS)["replacement"] == "had "


def test_markup_from_a_final_reader_lands_in_the_rejected_diagnostics(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        if stage == "opus_read":
            row = payload["paragraphs"][0]
            return {"reviewed_ids": [x["id"] for x in payload["paragraphs"]],
                    "findings": [{**finding(row["id"], "Huckleberry Finn", "*Huckleberry Finn*", "typesetting"), "evidence": []}],
                    "comment_decisions": [], "editorial_verdict": "ready"}

    readers = Readers(handler=handler)
    result = FixedWorkflow(make_book("She read Huckleberry Finn twice."), tmp_path / "markup", calls=readers).run()
    assert result["accepted"] == result["original"]
    rejected = [h for h in result["history"] if h.get("rejected_proposal") and h["stage"] == "opus_read"]
    assert rejected and rejected[0]["rejected_proposal"]["status"] == "rejected_invalid_proposal"
    assert "markup" in rejected[0]["rejected_proposal"]["reason"]


# --- continuity lane ----------------------------------------------------------

def continuity_finding(pid, quote, replacement, evidence, *, action="edit", question="", missing=""):
    return {"para_id": pid, "quote": quote, "occurrence": 1, "replacement": replacement,
            "action": action, "category": "continuity", "reason": "The surname is established elsewhere.",
            "question": question, "missing_knowledge": missing, "evidence": evidence}


def test_continuity_reads_the_whole_book_once_and_opus_rules_with_the_cited_evidence(make_book, tmp_path):
    seen = {}

    def handler(stage, model, payload, kwargs):
        if stage == "continuity":
            seen["payload"] = payload
            assert model == OPUS
            return {"findings": [continuity_finding("body-0001", "Beckham", "Brooks",
                                                    [{"para_id": "body-0000", "quote": "Kai Brooks smiled."}])],
                    "reading_notes": "One surname split."}
        if stage == "continuity_adjudication":
            assert model == OPUS
            seen["sites"] = payload["sites"]
            return {"decisions": [ruling(site, replacement=site["proposals"][0]["replacement"]) for site in payload["sites"]]}
        if stage == "continuity_checks_meaning":
            seen["checks"] = payload["changes"]

    readers = Readers(handler=handler)
    result = FixedWorkflow(make_book("Kai Brooks smiled.", "Then Kai Beckham left."), tmp_path / "continuity", calls=readers).run()
    assert list(result["accepted"].values()) == ["Kai Brooks smiled.", "Then Kai Brooks left."]
    assert [row["id"] for row in seen["payload"]["book"]] == ["body-0000", "body-0001"]
    assert seen["payload"]["complete_book"] is True and seen["payload"]["part"] == [1, 1]
    assert all(row["location"] == "body" for row in seen["payload"]["book"])
    [site] = seen["sites"]
    assert site["evidence_paragraphs"] == {"body-0000": "Kai Brooks smiled."}
    assert site["proposals"][0]["evidence"][0]["quote"] == "Kai Brooks smiled."
    assert seen["checks"][0]["evidence"][0]["para_id"] == "body-0000"
    assert not any(row["stage"] == "continuity_screen" for row in readers.events)
    applied = [h for h in result["history"] if h.get("applied") and h["stage"] == "continuity"]
    assert applied[0]["applied"]["models"] == [OPUS]
    assert [s["stage"] for s in result["stages"]].index("continuity") < [s["stage"] for s in result["stages"]].index("opus_read")


@pytest.mark.parametrize("evidence, reason", [
    ([], "lacks cited evidence"),
    # A citation of the finding's own paragraph is ignored, so it is no evidence.
    ([{"para_id": "body-0001", "quote": "Beckham"}], "lacks cited evidence"),
    ([{"para_id": "body-9999", "quote": "Beckham"}], "another paragraph"),
    ([{"para_id": "body-0000", "quote": "Kai Brooks frowned."}], "does not occur verbatim"),
])
def test_continuity_edits_without_verified_evidence_are_rejected(make_book, tmp_path, evidence, reason):
    def handler(stage, model, payload, kwargs):
        if stage == "continuity":
            return {"findings": [continuity_finding("body-0001", "Beckham", "Brooks", evidence)],
                    "reading_notes": ""}

    readers = Readers(handler=handler)
    result = FixedWorkflow(make_book("Kai Brooks smiled.", "Then Kai Beckham left."), tmp_path / "evidence", calls=readers).run()
    assert result["accepted"] == result["original"]
    assert not any(row["stage"] == "continuity_adjudication" for row in readers.events)
    rejected = [h for h in result["history"] if h.get("rejected_proposal") and h["stage"] == "continuity"]
    assert rejected and reason in rejected[0]["rejected_proposal"]["reason"]


@pytest.mark.parametrize("path", ["opus_drop", "opus_query", "reader_query"])
def test_continuity_drop_and_query_paths(make_book, tmp_path, path):
    evidence = [{"para_id": "body-0000", "quote": "the Rusty Hook Tavern"}]

    def handler(stage, model, payload, kwargs):
        if stage == "continuity":
            if path == "reader_query":
                return {"findings": [continuity_finding("body-0001", "the Mad Crabber", "the Mad Crabber", evidence,
                                                        action="query", question="Which name is the restaurant's?",
                                                        missing="The intended restaurant name")], "reading_notes": ""}
            return {"findings": [continuity_finding("body-0001", "the Mad Crabber", "the Rusty Hook Tavern", evidence)],
                    "reading_notes": ""}
        if stage == "continuity_adjudication":
            if path == "opus_drop":
                return {"decisions": [ruling(site, "drop") for site in payload["sites"]]}
            return {"decisions": [{**ruling(site, "query"), "question": "Which name is the restaurant's?",
                                   "missing_knowledge": "The intended restaurant name"} for site in payload["sites"]]}

    readers = Readers(handler=handler)
    result = FixedWorkflow(make_book("They ate at the Rusty Hook Tavern.", "Later they left the Mad Crabber."),
                           tmp_path / path, calls=readers).run()
    assert result["accepted"] == result["original"]
    if path == "opus_drop":
        # Wilder 2026-09-14: Opus dropped "Mad Crabber" -> "Rusty Hook Tavern" as
        # "not settled by the evidence" and nothing reached the author. A dropped
        # continuity edit is now the author's question.
        from galley.fixed_workflow import DEMOTED_MISSING_KNOWLEDGE
        [q] = result["questions"]
        assert q["quote"] == "Mad Crabber" and q["stage"] == "continuity"
        assert q["question"].startswith("“Mad Crabber” here; the reader proposed “Rusty Hook Tavern”, matching "
                                        "“the Rusty Hook Tavern” elsewhere in the book. Is the change intended?")
        assert q["missing_knowledge"] == DEMOTED_MISSING_KNOWLEDGE
        assert [h["question"] for h in result["history"] if h.get("stage") == "continuity_demoted"] == [q["question"]]
    else:
        assert [q["question"] for q in result["questions"]] == ["Which name is the restaurant's?"]
        assert result["questions"][0]["stage"] == "continuity"
    assert any(row["stage"] == "continuity_adjudication" for row in readers.events) == (path != "reader_query")


def test_a_final_readers_self_citation_is_ignored_and_a_dropped_fact_edit_becomes_a_question(make_book, tmp_path):
    """Wilder 2026-09-14: Fable's "eastern horizon" fix cited its own paragraph
    beside the next one and the evidence gate discarded it; had it reached the
    screen, a drop there would have ended it. Now the self-citation is ignored
    and a dropped fact/logic edit goes to Astra's comment review as a question."""
    from galley.fixed_workflow import DEMOTED_MISSING_KNOWLEDGE
    book = make_book("The sun sinks on the eastern horizon.", "We watch the sunset from the Atlantic shore.")
    seen = {}

    def handler(stage, model, payload, kwargs):
        if stage == "opus_read":
            rows = payload["paragraphs"]
            return {"reviewed_ids": [x["id"] for x in rows],
                    "findings": [{**finding(rows[0]["id"], "eastern horizon", "western horizon", "fact_logic"),
                                  "reason": "The sun sets in the west.",
                                  "evidence": [{"para_id": rows[0]["id"], "quote": "The sun sinks"},
                                               {"para_id": rows[1]["id"], "quote": "watch the sunset"}]}],
                    "comment_decisions": [comment_decision(q) for q in payload.get("comments", [])],
                    "editorial_verdict": "ready"}
        if stage in {"opus_read_checks_meaning", "opus_read_checks_meaning_sonnet"}:
            seen.setdefault("checked", []).append(payload["changes"])
            return {"decisions": [{"id": x["id"], "verdict": "reject", "reason": "Poetic license, not an error."}
                                  for x in payload["changes"]]}
        if stage == "opus_read_comment_review":
            seen["reviewed"] = payload["comments"]

    readers = Readers(handler=handler)
    result = FixedWorkflow(book, tmp_path / "demote", calls=readers).run()
    assert result["accepted"] == result["original"]
    assert not any(h.get("rejected_proposal") for h in result["history"] if h["stage"] == "opus_read")
    [applied] = [h["applied"] for h in result["history"] if h["stage"] == "opus_read" and h.get("applied")]
    assert [e["para_id"] for e in applied["evidence"]] == ["body-0001"], "the self-citation is ignored, the other kept"
    assert seen["checked"][0][0]["after"] == "The sun sinks on the western horizon."
    [q] = result["questions"]
    assert q["quote"] == "eastern" and q["stage"] == "opus_read_checks_meaning"
    assert q["question"] == ("“eastern” here; the reader proposed “western”, matching "
                             "“watch the sunset” elsewhere in the book. Is the change intended? The sun sets in the west.")
    assert q["missing_knowledge"] == DEMOTED_MISSING_KNOWLEDGE
    assert [h["site"] for h in result["history"] if h.get("stage") == "opus_read_checks_meaning_demoted"] == [applied["id"]]
    assert [c["id"] for c in seen["reviewed"]] == [q["id"]], "Fable's comment review judged the demoted question"
    site = {"proposals": [applied]}
    text = result["original"]["body-0000"]
    assert FixedWorkflow._frontier_demotion("opus_read", site, text) == (q["question"], q["missing_knowledge"], q["quote"])
    assert FixedWorkflow._frontier_demotion("typed", site) is None
    assert FixedWorkflow._frontier_demotion("opus_read", {"proposals": [{**applied, "category": "grammar"}]}) is None


def test_reinstatement_reads_dropped_frontier_edits_as_questions():
    from galley.fixed_reinstate import dropped_question_candidates
    current = {"body-0000": "The sun sinks on the eastern horizon."}
    edit = {"id": "f-1", "para_id": "body-0000", "start": 21, "end": 36, "before": "eastern horizon",
            "replacement": "western horizon", "category": "fact_logic", "action": "edit",
            "reason": "The sun sets in the west.", "missing_knowledge": "", "models": [OPUS], "evidence": []}
    usage = {**edit, "id": "f-2", "category": "usage"}
    result = {"history": [{"stage": "opus_read_screened", "decision": {"action": "drop", "reason": "Poetic."},
                           "site": {"proposals": [edit, usage]}}]}
    candidates, unanchored = dropped_question_candidates(result, current)
    assert unanchored == [] and [c["id"] for c in candidates] == ["f-1"]
    [row] = candidates
    assert row["action"] == "query" and row["category"] == "fact_logic"
    assert row["reason"].startswith("“eastern horizon” here; the reader proposed “western horizon”. Is the change intended?")
    assert row["missing_knowledge"]


def test_a_skipped_continuity_read_still_delivers(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        if stage == "continuity":
            return {"_skipped_read": {"stage": stage, "model": model, "reason": "Simulated exhausted read"}}

    readers = Readers(handler=handler)
    result = FixedWorkflow(make_book("A quiet page."), tmp_path / "skipped", calls=readers).run()
    assert result["review_complete"] is False and result["skipped_reads"]
    continuity = json.loads((tmp_path / "skipped/stages/continuity.json").read_text())["evidence"]
    assert continuity["coverage"][0]["status"] == "skipped" and continuity["skipped_reads"]
    assert [row["stage"] for row in readers.events if row["stage"] in {"opus_read", "astra"}] == ["opus_read", "astra"]


def test_a_version_2_workspace_requires_a_fresh_run(make_book, tmp_path):
    source = make_book("A quiet page.")
    flow = FixedWorkflow(source, tmp_path / "v2", calls=object())
    marker = json.loads(flow.manifest.read_text())
    marker["identity"]["version"] = "fixed-proofreading-v2"
    flow.manifest.write_text(json.dumps(marker))
    with pytest.raises(FixedWorkflowError, match="fresh workspace"):
        FixedWorkflow(source, tmp_path / "v2", calls=object())


def test_number_policy_travels_only_with_number_work(make_book, tmp_path):
    """The 3,500-token number policy heads the number stage, the whole-book
    readers, and any request that carries a number proposal; screening,
    checks and comment reviews without one get the editorial brief alone."""
    from galley.fixed_policy import NUMBER_POLICY
    from galley.press_prompt import EDITORIAL_RULES
    book = make_book("We seen teh 20 birds. It are warm.")

    def typed(stage, model, paragraphs, keys):
        pid, text = next(iter(paragraphs.items()))
        rows = []
        if "spelling" in keys:
            rows.append(_typed_row(pid, text, "teh", "the", "spelling"))
        if "subject_verb_agreement" in keys and model == SONNET:
            rows.append(_typed_row(pid, text, "seen", "saw", "subject_verb_agreement"))
        return rows

    readers = Readers(typed=typed)
    flow = FixedWorkflow(book, tmp_path / "run", calls=readers)
    flow.run()
    systems = {}
    for row in readers.events:
        if "system" in row and "user" in row:
            systems.setdefault(row["stage"], row["system"])
    assert systems["story_sheet"].startswith(flow.base_policy)
    assert NUMBER_POLICY not in systems["story_sheet"]
    assert systems["numbers"].startswith(NUMBER_POLICY)
    for stage in ("typed_screen", "checks_meaning", "checks_correction"):
        assert stage in systems, stage
        assert NUMBER_POLICY not in systems[stage]
        assert systems[stage].startswith(flow.editorial_policy)
        assert all(text in systems[stage] for text in EDITORIAL_RULES.values())
    for stage in ("ensemble_sweep_opus", "ensemble_sweep_sol", "opening_read", "opus_read", "astra"):
        assert systems[stage].startswith(NUMBER_POLICY)
    # A screening or check request that carries a number proposal gets it back.
    user = json.dumps({"sites": [{"proposals": [{"category": "number_style"}]}]})
    assert flow._policy_for("typed_screen", user) == flow.policy
    assert flow._policy_for("checks_correction", json.dumps({"changes": [{"categories": ["currency_style"]}]})) == flow.policy
    assert flow._policy_for("checks_correction", json.dumps({"changes": [{"categories": ["spelling"]}]})) == flow.editorial_policy
    assert flow._policy_for("continuity", "{}") == flow.base_policy
    assert flow.identity["policy_sha256"] == flow.identity["policy_sha256"]
    # The identity covers every contract, so a change to any of them starts a fresh workspace.
    assert flow.identity["version"] == "fixed-proofreading-v8"


def test_checks_carry_the_categories_of_accepted_corrections(make_book, tmp_path):
    book = make_book("We seen teh birds.")

    def typed(stage, model, paragraphs, keys):
        pid, text = next(iter(paragraphs.items()))
        return [_typed_row(pid, text, "teh", "the", "spelling")] if "spelling" in keys else []

    readers = Readers(typed=typed)
    FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    check = next(r for r in readers.events if r["stage"] == "checks_meaning")
    assert check["payload"]["changes"][0]["categories"] == ["spelling"]
    assert "categories names the proofreading categories" in check["system"]
    later = [r for r in readers.events if r["stage"].endswith("_checks_meaning") and r["stage"] != "checks_meaning"]
    assert not later or all(c["categories"] == [] for r in later for c in r["payload"]["changes"])


def test_screening_requests_are_short_windows_with_room_to_answer(make_book, tmp_path):
    """43-50 sites per window pushed Sonnet past its 12k output ceiling and
    Luna past complete coverage on the first production book."""
    book = make_book("He waited, watching the door. " * 3)

    def typed(stage, model, paragraphs, keys):
        pid, text = next(iter(paragraphs.items()))
        if model != SONNET or "unnecessary_comma" not in keys:
            return []
        # Thirty distinct single-model proposals: every one must be screened.
        return [{"para_id": pid, "error_type": "unnecessary_comma", "original_text": text,
                 "corrected_text": text[:i] + "." + text[i + 1:], "occurrence": 1,
                 "confidence": "high", "explanation": "Site %d" % i}
                for i, c in enumerate(text) if c in ",." ][:30]

    readers = Readers(typed=typed)
    FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    screens = [r for r in readers.events if r["stage"] == "typed_screen"]
    assert screens
    assert all(r["max_tokens"] == 16000 for r in screens)
    assert all(len(r["payload"]["sites"]) <= 25 for r in screens)
    assert all("at most 25 words" in r["system"] for r in screens)


def test_final_readers_see_only_tense_sites_that_read_against_the_baseline(make_book, tmp_path):
    past = ["She walked to the harbour, waited on the pier, and watched the boats."] * 22
    book = make_book(*past, "He walks and waits and watches.")
    readers = Readers()
    result = FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    for stage in ("opus_read", "astra"):
        rows = [r for r in readers.events if r["stage"] == stage]
        tense = [s for r in rows for s in r["payload"]["focused_sites"] if s["check"] == "narrative_tense"]
        assert len(tense) == 1 and tense[0]["quote"].startswith("He walks")
        assert "past signals=0, present=3" in tense[0]["detail"]
        assert all("sample" not in p for r in rows for p in r["payload"]["narrative_profile"]["paragraphs"])
        assert sum(len(r["payload"]["narrative_profile"]["paragraphs"]) for r in rows) == 23
        assert "(baseline past)" in shared_context(rows[0])["notes"]["focused_sites"]
    stage = json.loads(Path(next(s["path"] for s in result["stages"] if s["stage"] == "opus_read")).read_text())
    assert sum(c["tense_sites_omitted"] for c in stage["evidence"]["coverage"]) == 22
    assert sum(c["focused_counts"]["narrative_tense"] for c in stage["evidence"]["coverage"]) == 1


def test_number_sites_are_read_under_short_labels_and_recorded_by_durable_id(make_book, tmp_path):
    book = make_book("She counted 3 boats and twenty gulls at 4 pm.")
    readers = Readers()
    result = FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    reads = [r for r in readers.events if r["stage"] == "numbers"]
    assert reads and all([s["id"] for s in r["payload"]["sites"]] == ["n01", "n02", "n03"] for r in reads)
    assert all(s["text"] for r in reads for s in r["payload"]["sites"])
    stage = json.loads(Path(next(s["path"] for s in result["stages"] if s["stage"] == "numbers")).read_text())
    assert [s["id"][:7] for s in stage["evidence"]["sites"]] == ["number-"] * 3


def test_final_reader_questions_are_screened_in_the_walkthrough_scope(make_book, tmp_path):
    """Wilder 2026-09-14: Fable and Astra raised eight fact, logic, continuity and
    structure questions; the screen, judging under the plain contract, dropped
    seven as out of scope. The screen now carries the walk-through rider."""
    from galley.press_prompt import WALKTHROUGH_QUERY_RIDER
    book = make_book("The sun sets over the Atlantic at Juno Beach.", "We seen the birds.")

    def typed(stage, model, paragraphs, keys):
        pid, text = next(iter(paragraphs.items()))
        return [_typed_row(pid, text, "seen", "saw", "subject_verb_agreement")] if "subject_verb_agreement" in keys and model == SONNET else []

    def handler(stage, model, payload, kwargs):
        if stage == "astra":
            row = payload["paragraphs"][0]
            return {"reviewed_ids": [x["id"] for x in payload["paragraphs"]],
                    "findings": [{**finding(row["id"], "sun sets over the Atlantic", "", "fact_logic", action="query",
                                            missing="Which coast the sunset is seen from"), "evidence": []}],
                    "comment_decisions": [comment_decision(q) for q in payload.get("comments", [])],
                    "editorial_verdict": "ready"}
        if stage == "astra_screen":
            return {"decisions": [{"id": s["id"], "action": "query", "replacement": "", "reason": "Geography.",
                                   "question": "Which coast?", "missing_knowledge": "The intended coast"} for s in payload["sites"]]}
    readers = Readers(typed=typed, handler=handler)
    result = FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    stages = [r["stage"] for r in readers.events]
    # The fact/logic question never meets the paragraph-level pair screen; it
    # goes to Astra's comment review, which carries the walk-through rider.
    assert "astra_screen" not in stages and "astra_comment_review" in stages
    review = next(r for r in readers.events if r["stage"] == "astra_comment_review")
    from galley.press_prompt import WALKTHROUGH_COMMENT_RIDER
    assert WALKTHROUGH_COMMENT_RIDER in review["system"]
    assert [q["missing_knowledge"] for q in result["questions"]] == ["Which coast the sunset is seen from"]
    assert any(h.get("stage") == "astra_frontier_question" for h in result["history"])
    assert FixedWorkflow._query_rider("typed") == "" and FixedWorkflow._query_rider("opus_read") == WALKTHROUGH_QUERY_RIDER


def test_completed_run_can_reinstate_dropped_walkthrough_questions(make_book, tmp_path):
    from galley.fixed_reinstate import reinstate_walkthrough_questions, FixedReinstateError, STAGE
    from galley.fixed_documents import write_manuscripts, paragraph_views
    from docproof.utils.xml_helpers import DocxPackage, qn
    book = make_book("The sun sets over the Atlantic at Juno Beach.", "Gray whales pass Florida.")
    workspace = tmp_path / "ws"

    def handler(stage, model, payload, kwargs):
        if stage == "astra":
            rows = payload["paragraphs"]
            return {"reviewed_ids": [x["id"] for x in rows],
                    "findings": [{**finding(rows[0]["id"], "sun sets over the Atlantic", "", "fact_logic", action="query",
                                            missing="The intended coast"), "evidence": []},
                                 {**finding(rows[1]["id"], "Gray whales", "", "fact_logic", action="query",
                                            missing="The intended species"), "evidence": []}],
                    "comment_decisions": [comment_decision(q) for q in payload.get("comments", [])],
                    "editorial_verdict": "needs_human"}
        # The original run's screen drops both (the Readers default).
    def dropping(stage, model, payload, kwargs):
        if stage == "astra_comment_review":      # the old behaviour: everything dropped
            return {"decisions": [comment_decision(q, "drop") for q in payload["comments"]]}
        return handler(stage, model, payload, kwargs)
    readers = Readers(handler=dropping)
    result = FixedWorkflow(book, workspace / "runs" / "fixed", calls=readers).run()
    # A reader's own window verdict no longer decides: only the second Astra
    # reading's counted rule can send a book to a person.
    assert result["questions"] == [] and result["editorial_verdict"] == "ready"
    # Simulate a run made before the frontier-question route existed: the
    # screen's drop rows are what reinstatement reads.
    path = workspace / "runs/fixed/result.json"
    saved = json.loads(path.read_text())
    frontier = [h["candidate"] for h in saved["history"] if h.get("stage") == "astra_frontier_question"]
    assert len(frontier) == 2
    for row in frontier:
        site = {"id": "d-" + row["id"][2:], "para_id": row["para_id"], "start": row["start"], "end": row["end"],
                "before": row["before"], "paragraph": saved["accepted"][row["para_id"]], "source": saved["original"][row["para_id"]],
                "proposals": [row]}
        saved["history"].append({"stage": "astra_screened", "site": site,
                                 "decision": {"id": site["id"], "action": "drop", "reason": "Out of scope."}})
    path.write_text(json.dumps(saved))
    (workspace / "runs" / "driver").mkdir(parents=True)
    (workspace / "runs" / "driver" / "package.json").write_text("{}")
    (workspace / "handoff").mkdir()

    def reinstating(stage, model, payload, kwargs):
        if stage.startswith(STAGE) and stage.endswith("_comment_review"):
            return {"decisions": [comment_decision(q, "retain" if "coast" in q["missing_knowledge"] else "drop")
                                  for q in payload["comments"]]}
    again = Readers(handler=reinstating)
    out = reinstate_walkthrough_questions(book, workspace, calls=again)
    assert out["candidates"] == 2 and out["stage"] == STAGE
    assert [q["missing_knowledge"] for q in out["reinstated"]] == ["The intended coast"]
    stages = {r["stage"] for r in again.events}
    assert STAGE + "_screen" not in stages and STAGE + "_comment_review" in stages
    assert not any(r["stage"] in {"typed", "astra", "final_astra", "opus_read", "numbers"} for r in again.events), "nothing is re-read"
    saved = json.loads((workspace / "runs/fixed/result.json").read_text())
    assert [s["stage"] for s in saved["stages"]][-2:] == ["astra_gate", STAGE]
    assert saved["final_review"] == result["final_review"] and saved["editorial_verdict"] == "ready"
    assert saved["accepted"] == result["accepted"] and len(saved["questions"]) == 1
    manifest = json.loads((workspace / "runs/fixed/workflow.json").read_text())
    assert manifest["result_sha256"] == saved["result_sha256"] and manifest["status"] == "completed"
    assert not (workspace / "runs/driver/package.json").exists() and not (workspace / "handoff").exists()
    tracked, clean, details = write_manuscripts(book, tmp_path / "out", saved["accepted"], saved["questions"])
    comments = DocxPackage(tracked).tree("word/comments.xml")
    texts = ["".join(t.text or "" for t in c.iter(qn("w:t"))) for c in comments if c.tag == qn("w:comment")]
    assert any(saved["questions"][0]["question"] in text for text in texts)
    # A second pass is its own receipted stage and adds nothing already present.
    second = reinstate_walkthrough_questions(book, workspace, calls=Readers(handler=reinstating))
    assert second["stage"] == STAGE + "_2" and second["reinstated"] == []
    final = json.loads(path.read_text())
    assert [s["stage"] for s in final["stages"]][-3:] == ["astra_gate", STAGE, STAGE + "_2"] and len(final["questions"]) == 1


# --- Wilder proofreader follow-ups (2026-09-17) -------------------------------

@pytest.mark.parametrize("before, after", [("at around five.", "at around 5:00."),
                                           ("by 10—11 at the latest.", "by 10:00–11:00 AM at the latest.")])
def test_number_reader_cannot_invent_a_clock_reading(make_book, tmp_path, before, after):
    flow = FixedWorkflow(make_book("We surfed " + before), tmp_path / "run", calls=Readers())
    texts = {"p1": "We surfed " + before}
    row = finding("p1", texts["p1"], "We surfed " + after, "number_style")
    assert flow._reader_candidate("numbers", row, texts, SONNET) is None
    [rejected] = [h["rejected_proposal"] for h in flow.history]
    assert "clock reading" in rejected["reason"] or "digits" in rejected["reason"]
    spelled = finding("p1", "“At 3?”", "“At three?”", "number_style")
    assert flow._reader_candidate("numbers", spelled, {"p1": "“At 3?”"}, SONNET)["replacement"] == "three"


def test_consistency_sites_move_together_once_one_swap_is_accepted():
    from galley.fixed_workflow import _harmonize_consistency, _word_swap
    assert _word_swap("My parents are OK with it.", "My parents are okay with it.") == ("OK", "okay")
    assert _word_swap("on this green Earth blessed", "on this green earth blessed") == ("Earth", "earth")
    assert _word_swap("She waited.", "She waited and left.") is None
    site = lambda sid, before, category="term_consistency": {
        "id": sid, "para_id": sid, "before": before, "proposals": [{"category": category, "replacement": before}]}
    sites = [site("a", "My parents are OK with it, of course."),
             site("b", "“Is everything OK?”"),
             site("c", "The OKLAHOMA sign was fine."),
             site("d", "He said it was fine.", "grammar")]
    agreed = {"a": {"id": "a", "action": "apply", "replacement": "My parents are okay with it, of course.", "reason": "x"},
              "b": {"id": "b", "action": "drop", "replacement": "", "reason": "Correct in dialogue."},
              "c": {"id": "c", "action": "drop", "replacement": "", "reason": "Not the term."},
              "d": {"id": "d", "action": "drop", "replacement": "", "reason": "Fine."}}
    updated, log = _harmonize_consistency(sites, agreed)
    assert updated["b"]["action"] == "apply" and updated["b"]["replacement"] == "“Is everything okay?”"
    assert updated["c"] == agreed["c"] and updated["d"] == agreed["d"]
    assert log == [{"site": "b", "swap": ["OK", "okay"], "origin": "a", "dropped_reason": "Correct in dialogue."}]
    # A query stands, and nothing moves when no site was accepted.
    agreed["b"]["action"] = "query"
    assert _harmonize_consistency(sites, agreed)[0]["b"]["action"] == "query"
    agreed["a"]["action"] = "drop"
    assert _harmonize_consistency(sites, agreed) == (agreed, [])


def test_screened_consistency_drop_is_harmonized_with_the_accepted_swap(make_book, tmp_path):
    def answer(stage, model, payload, kwargs):
        assert stage == "typed_screen"
        # The screen sees the minimal span ("OK") at the first site and the
        # scan's sentence window at the second; it accepts one, drops the other.
        return {"decisions": [ruling(s, "apply" if s["para_id"] == "a" else "drop",
                                     s["proposals"][0]["replacement"] if s["para_id"] == "a" else "")
                              for s in payload["sites"]]}
    flow = _flow(make_book, tmp_path, Readers(handler=answer))
    flow.original = {"a": "My parents are OK with it.", "b": "“Is everything OK?”"}
    flow.current = dict(flow.original)
    pids = list(flow.current)
    rows = [_candidate({**finding(pids[0], "My parents are OK with it.", "My parents are okay with it.",
                                  "term_consistency")}, flow.current, "local:consistency"),
            _candidate({**finding(pids[1], "“Is everything OK?”", "“Is everything OK?”",
                                  "term_consistency", action="query")}, flow.current, "local:consistency")]
    flow._apply("typed", flow._adjudicate("typed", rows))
    assert flow.current[pids[0]] == "My parents are okay with it."
    assert flow.current[pids[1]] == "“Is everything okay?”"
    assert [h["swap"] for h in flow.history if h["stage"] == "typed_harmonized"] == [["OK", "okay"]]


def test_a_screen_query_on_a_chapter_label_is_overruled_into_the_fix(make_book, tmp_path):
    def answer(stage, model, payload, kwargs):
        assert stage == "typed_screen"
        return {"decisions": [{**ruling(s, "query", ""), "question": "Should this be Chapter 1?",
                               "missing_knowledge": "The intended label."} for s in payload["sites"]]}
    flow = _flow(make_book, tmp_path, Readers(handler=answer), text="CHAPTER ONE")
    pid = "p"
    rows = [_candidate({**finding(pid, "CHAPTER ONE", "CHAPTER 1", "chapter_label")}, flow.current, "local:chapter_labels")]
    flow._apply("typed", flow._adjudicate("typed", rows))
    assert flow.current[pid] == "CHAPTER 1" and flow.questions == []
    [overruled] = [h for h in flow.history if h["stage"] == "typed_label_query_overruled"]
    assert overruled["question"] == "Should this be Chapter 1?"


def test_a_final_reader_never_demotes_a_label_edit_into_a_question():
    proposal = {"action": "edit", "category": "structure", "before": "CHAPTER ONE", "replacement": "CHAPTER 1",
                "reason": "Label out of style.", "para_id": "h1", "start": 0, "end": 11, "evidence": []}
    assert FixedWorkflow._frontier_demotion("opus_read", {"proposals": [proposal]}, "CHAPTER ONE") is None
    assert FixedWorkflow._frontier_demotion("opus_read", {"proposals": [proposal]}, "The sun set in the east.") is not None


# --- one question, one id ----------------------------------------------------

def test_question_two_stages_raise_is_held_once(make_book, tmp_path):
    """The Gunn run of 2026-09-17: a later stage demoted a rejected edit into
    the question an earlier stage already held. The child check workflow's own
    guard saw only its empty list, so the parent ended up holding one id twice
    and the next comment review's coverage inventory was refused."""
    def handler(stage, model, payload, kwargs):
        if stage == "opus_read_checks_meaning":
            return {"decisions": [{"id": "p", "verdict": "reject", "reason": "Invented identity."}]}
        if stage == "opus_read_checks_meaning_sonnet":
            return {"decisions": [{"id": "p", "verdict": "approve", "reason": "Reads correctly."}]}
        if stage == "opus_read_checks_meaning_disputes":
            return {"decisions": [{"id": "p", "action": "query", "replacement": "",
                                   "question": "Who was he waiting for?",
                                   "missing_knowledge": "The intended identity",
                                   "reason": "An unresolved reference."}]}
    flow = _flow(make_book, tmp_path, Readers(handler=handler))
    flow._question("p", flow.current["p"], "Who was he waiting for?",
                   "The intended identity", "An unresolved reference.", "typed")
    held = flow.questions[0]["id"]
    before = flow._apply("opus_read", [_candidate(finding("p", "someone", "Mary"), flow.current, OPUS)])
    flow._checks("opus_read_checks", before)
    assert [q["id"] for q in flow.questions] == [held]


def test_replaced_comment_takes_the_id_its_new_wording_hashes_to(make_book, tmp_path):
    """An id is a hash of the question's content, so a review that replaces the
    wording re-keys the row and retires the id it left behind — otherwise a
    later stage regenerating the original wording mints that id a second
    time."""
    from galley.fixed_workflow import _question_id
    flow = _flow(make_book, tmp_path, text="He saw the visitor.")
    flow._question("p", flow.current["p"], "Who is the visitor?",
                   "The visitor's identity", "An unresolved identity.", "opus_read")
    retired = flow.questions[0]["id"]
    decision = comment_decision(flow.questions[0], quote="the visitor")
    decision["missing_knowledge"] = "The visitor's name"
    flow._comments([decision], "astra", model=ASTRA)
    assert [q["id"] for q in flow.questions] == [
        _question_id("p", "the visitor", "The visitor's name")]
    kept = flow.questions[0]["id"]
    assert kept != retired
    # The retired wording, raised again by a later stage, is already asked.
    flow._question("p", "He saw the visitor.", "Who is the visitor?",
                   "The visitor's identity", "An unresolved identity.", "final_astra")
    assert [q["id"] for q in flow.questions] == [kept]


def test_duplicate_question_ids_are_named_where_they_become_an_inventory(make_book, tmp_path):
    """Whatever produced them, a duplicate id is reported as a duplicate id and
    not as the coverage contract two layers down refusing the call."""
    flow = _flow(make_book, tmp_path)
    flow._question("p", flow.current["p"], "Who was he waiting for?",
                   "The intended identity", "An unresolved reference.", "typed")
    flow.questions.append(dict(flow.questions[0]))     # only a defect can do this
    with pytest.raises(FixedWorkflowError, match="duplicate ids"):
        flow._comments([comment_decision(flow.questions[0])], "astra", model=ASTRA)


def test_a_failure_that_replays_identically_is_reported_as_exhausted():
    """A resume replays the call cache to reach the same place, so the agent is
    told not to spend the run again; a lock another worker holds is not that."""
    from galley.fixed_calls import FixedCallContractError
    from galley.fixed_workflow import FixedWorkflowBusy, deterministic_failure
    assert deterministic_failure(FixedCallContractError("Coverage inventory needs unique string IDs"))
    assert deterministic_failure(FixedWorkflowError("Preparation silently changed source text"))
    assert deterministic_failure(RuntimeError("wrapped")) is False
    assert deterministic_failure(FixedWorkflowBusy("Another worker owns this fixed proofread")) is False
    try:
        try:
            raise FixedCallContractError("Coverage inventory needs unique string IDs")
        except FixedCallContractError as exc:
            raise RuntimeError("the stage failed") from exc
    except RuntimeError as exc:
        assert deterministic_failure(exc)


def test_a_refused_contract_reaches_the_agent_as_exhausted_recovery(make_book, tmp_path):
    """The driver's own report, not just the classifier: a block the resume
    would reproduce is handed over as exhausted so the agent holds the book."""
    import types
    from galley import fixed_workflow as fw
    from galley.fixed_calls import FixedCallContractError

    def refuse(self):
        raise FixedCallContractError("Coverage inventory needs unique string IDs")

    workspace = tmp_path / "ws"
    workspace.mkdir()
    driver = types.SimpleNamespace(
        workspace=workspace, book=make_book("He waited for someone."),
        budget_usd=10.0, drive_folder_id="", source_id="", slug="test",
        drive_archive_folder_id="", upload=None, verify_upload=None,
        _write_ledger=lambda result: None, _progress=lambda *a, **kw: None)
    original, fw.FixedWorkflow.run = fw.FixedWorkflow.run, refuse
    try:
        result = fw.run_fixed_driver(driver)
    finally:
        fw.FixedWorkflow.run = original
    assert result.outcome == "blocked"
    assert result.stopped_at == "fixed"
    assert result.recovery_exhausted is True


def test_a_comma_edit_and_a_relative_pronoun_swap_cannot_both_read_one_clause(tmp_path, make_book):
    """Georgis 2026-09-15: deleting the comma before "that" (restrictive) and
    swapping that -> which (nonrestrictive) were both approved and shipped
    "pillar which". The pronoun swap is dropped with a receipt; the comma
    edit lands. A pair that agrees with itself is untouched."""
    from galley.fixed_workflow import contradictory_relative_swaps
    text = "I felt like a marble pillar, that had been gashed by time."
    at = text.index(", that")
    comma_out = {"id": "c", "para_id": "p1", "start": at, "end": at + 1, "before": ",", "replacement": "",
                 "category": "unnecessary_comma", "models": [SONNET, LUNA], "format": "", "reason": "x"}
    that_which = {"id": "w", "para_id": "p1", "start": at + 2, "end": at + 6, "before": "that", "replacement": "which",
                  "category": "that_which", "models": [SONNET, LUNA], "format": "", "reason": "x"}
    dropped = contradictory_relative_swaps([comma_out, that_which], {"p1": text})
    assert [row["id"] for row, _ in dropped] == ["w"] and "comma edit stands" in dropped[0][1]
    comma_in = {**comma_out, "start": at + 1, "end": at + 1, "before": "", "replacement": ","}
    which_that = {**that_which, "before": "which", "replacement": "that"}
    assert contradictory_relative_swaps([comma_in, that_which], {"p1": text}) == []
    assert contradictory_relative_swaps([comma_out, which_that], {"p1": text}) == []
    assert contradictory_relative_swaps([comma_in, which_that], {"p1": text})[0][0]["id"] == "w"
    # A comma elsewhere in the sentence is not this clause's comma.
    far = {**comma_out, "start": 0, "end": 0, "before": "", "replacement": ","}
    assert contradictory_relative_swaps([far, which_that], {"p1": text}) == []
    # Through the apply path: the text keeps the author's pronoun.
    book = make_book(text)
    flow = FixedWorkflow(book, tmp_path / "run", calls=Readers())
    flow.original = flow.current = {"p1": text}
    flow._apply("typed", [comma_out, that_which])
    assert flow.current["p1"] == "I felt like a marble pillar that had been gashed by time."
    receipts = [h for h in flow.history if h["stage"] == "typed" and "dropped" in h]
    assert [h["dropped"]["id"] for h in receipts] == ["w"] and receipts[0]["reason"].startswith("guard:")
    assert [h["applied"]["id"] for h in flow.history if h["stage"] == "typed" and "applied" in h] == ["c"]


def test_opening_read_findings_are_screened_with_the_typed_stage(make_book, tmp_path):
    """v8's opening salvo: the Story Sheet model reads the untouched book with
    the final readers' brief. Its edits join the typed candidates and meet the
    Sonnet + Luna screen; its whole-book questions skip that screen."""
    text = "She recieved the letter. The inn stood on the east shore."
    book = make_book(text)
    def handler(stage, model, payload, kwargs):
        if stage == "opening_read":
            row = payload["paragraphs"][0]
            return {"reviewed_ids": [x["id"] for x in payload["paragraphs"]],
                    "findings": [finding(row["id"], "recieved", "received", "spelling"),
                                 {**finding(row["id"], "east shore", "", "continuity", action="query",
                                            missing="Which shore the inn stands on"), "evidence": []}],
                    "comment_decisions": [], "editorial_verdict": "ready"}
        if stage == "typed_screen":
            return {"decisions": [ruling(x, replacement="ei") for x in payload["sites"]]}
    readers = Readers(handler=handler)
    result = FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    assert list(result["accepted"].values()) == ["She received the letter. The inn stood on the east shore."]
    applied = [h for h in result["history"] if h.get("applied")]
    assert [(h["stage"], h["applied"]["origin"]) for h in applied] == [("typed", "opening_read")]
    assert {r["model"] for r in readers.events if r["stage"] == "typed_screen"} == {SONNET, LUNA}
    assert [q["stage"] for q in result["questions"]] == ["opening_read"]
    assert any(h.get("stage") == "typed_frontier_question" for h in result["history"])
    opening = next(r for r in readers.events if r["stage"] == "opening_read")
    assert opening["model"] == OPUS and "focused_sites" in opening["payload"]
    assert opening["payload"]["paragraphs"][0]["text"] == text
    story = next(r for r in readers.events if r["stage"] == "story_sheet")
    assert story["model"] == OPUS
    stages = [s["stage"] for s in result["stages"]]
    assert stages.index("story_sheet") < stages.index("opening_read") < stages.index("typed")
    # An opening-read broken_sentence finding does not trigger the unscreened
    # broken-sentence repair; only the detectors' density does.
    assert not any(r["stage"] == "broken_repair" for r in readers.events)


def test_astra_gate_returns_rejected_paragraphs_to_their_pre_astra_text(make_book, tmp_path):
    book = make_book("He walk home.", "The sky was blue.")
    def handler(stage, model, payload, kwargs):
        if stage == "astra":
            first, second = payload["paragraphs"]
            return {"reviewed_ids": [first["id"], second["id"]],
                    "findings": [finding(first["id"], "He walk", "He walks"),
                                 finding(second["id"], "blue", "azure", "usage")],
                    "comment_decisions": [comment_decision(q) for q in payload.get("comments", [])],
                    "editorial_verdict": "ready"}
        if stage == "astra_gate_meaning":
            return {"decisions": [{"id": x["id"], "verdict": "reject" if "azure" in x["after"] else "approve",
                                   "reason": "Azure is a rewording, not a correction."} for x in payload["changes"]]}
    readers = Readers(handler=handler)
    result = FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    assert list(result["accepted"].values()) == ["He walks home.", "The sky was blue."]
    gate_calls = [r for r in readers.events if r["stage"].startswith("astra_gate")]
    assert {r["model"] for r in gate_calls} == {OPUS}
    meaning = next(r for r in gate_calls if r["stage"] == "astra_gate_meaning")
    assert sorted(x["after"] for x in meaning["payload"]["changes"]) == ["He walks home.", "The sky was azure."]
    correction = next(r for r in gate_calls if r["stage"] == "astra_gate_correction")
    assert [x["after"] for x in correction["payload"]["changes"]] == ["He walks home."]
    stage = json.loads((tmp_path / "run/stages/astra_gate.json").read_text())
    ids = list(result["original"])
    assert stage["evidence"]["gate"] == {"changed_paragraphs": 2, "rejected": {"meaning": [ids[1]], "correction": []},
                                         "unavailable": []}
    assert stage["evidence"]["press_audit"]["accepted_sha256"] == stage["accepted_sha256"]
    # The verdict was counted before the gate and is not revisited.
    assert result["final_review"]["verdict"] == "ready"


def test_astra_gate_asks_nothing_when_astra_changed_nothing(make_book, tmp_path):
    readers = Readers()
    result = FixedWorkflow(make_book("A quiet room."), tmp_path / "run", calls=readers).run()
    assert not any(r["stage"].startswith("astra_gate") for r in readers.events)
    stage = json.loads((tmp_path / "run/stages/astra_gate.json").read_text())
    assert stage["evidence"]["gate"]["changed_paragraphs"] == 0
    assert result["stages"][-1]["stage"] == "astra_gate"


def test_astra_gate_moves_a_question_whose_quote_left_with_the_restored_text(make_book, tmp_path):
    def handler(stage, model, payload, kwargs):
        if stage == "astra_gate_meaning":
            return {"decisions": [{"id": x["id"], "verdict": "reject", "reason": "Invented identity."}
                                  for x in payload["changes"]]}
    flow = _flow(make_book, tmp_path, Readers(handler=handler))
    before = dict(flow.current)
    row = _candidate({**finding("p", "someone", "Mary", "continuity"), "evidence": []}, flow.current, ASTRA)
    flow._apply("astra", [row])
    flow._question("p", "Mary", "Which Mary is this?", "The intended identity", "Two Marys appear.", "astra")
    gate = flow._astra_gate(before, 0)
    assert flow.current["p"] == "He waited for someone."
    assert gate["rejected"]["meaning"] == ["p"]
    moved = next(q for q in flow.questions if q["question"] == "Which Mary is this?")
    assert moved["quote"] == "He waited for someone."
    # The rejected continuity edit is put to the author as a question.
    assert any(h.get("stage") == "astra_gate_demoted" for h in flow.history)
    assert any(q["stage"] == "astra_gate" for q in flow.questions)


def test_house_respellings_survive_every_model_rejection_and_a_late_reintroduction(make_book, tmp_path):
    """Immanuel (2026-09-22): the screen dropped towards/amongst as "an
    established variant" and the checks restored paragraphs where one landed.
    Code now applies them after the first checks and again after the Astra
    gate, so every screen and check before the final reading can refuse and the book still ships the
    house form; a final reader who writes "towards" back is corrected too."""
    book = make_book("She walked towards the sea, and toward the rocks.",
                     "Amongst the gulls, afterwards, he walk home.",
                     "Nothing else happened.")

    def handler(stage, model, payload, kwargs):
        if "changes" in payload and not stage.startswith(("astra_gate", "final_astra")):
            return {"decisions": [{"id": x["id"], "verdict": "reject", "reason": "An established variant."}
                                  for x in payload["changes"]]}
        if stage == "final_astra":
            row = next(p for p in payload["paragraphs"] if "walk home" in p["text"])
            return {"reviewed_ids": [x["id"] for x in payload["paragraphs"]],
                    "findings": [finding(row["id"], "he walk home", "he walks towards home")],
                    "comment_decisions": [], "editorial_verdict": "ready"}

    readers = Readers(handler=handler)
    result = FixedWorkflow(book, tmp_path / "run", calls=readers).run()
    assert list(result["accepted"].values()) == [
        "She walked toward the sea, and toward the rocks.",
        "Among the gulls, afterward, he walks toward home.",
        "Nothing else happened."]
    early = [h["applied"] for h in result["history"] if h["stage"] == "house_respell" and h.get("applied")]
    assert sorted((r["before"], r["replacement"]) for r in early) == [
        ("Amongst", "Among"), ("afterwards", "afterward"), ("towards", "toward")]
    stages = {s["stage"]: json.loads(Path(s["path"]).read_text()) for s in result["stages"]}
    assert stages["checks"]["evidence"]["house_respell"]["applied"] == 3
    assert [(x["before"], x["replacement"]) for x in stages["astra_gate"]["evidence"]["house_respell"]["sites"]] == [
        ("towards", "toward")]
    assert stages["astra_gate"]["evidence"]["press_audit"]["raw_signal_counts"].get("variant_spelling", 0) == 0
    # No screen was ever asked about the respellings.
    assert not any("towards" in json.dumps(x.get("payload", {}).get("sites", "")) for x in readers.events)


def test_final_house_respell_corrects_a_reader_who_writes_the_british_form_back(make_book, tmp_path):
    from types import SimpleNamespace
    from docproof.spellscan import SpellScan
    from docproof.variants import load_variant
    flow = _flow(make_book, tmp_path, text="He walks home.")
    flow.prose_prepared = SimpleNamespace(doc=flow.prose_prepared.doc, variant=load_variant("us"), spell=SpellScan())
    flow.current = {"p": "He walks towards home, backwards and forwards."}
    flow.pending_categories = {"p": {"grammar"}}
    evidence = flow._house_respell("house_respell_final")
    assert flow.current["p"] == "He walks toward home, backward and forward."
    assert evidence["applied"] == 3 and [s["replacement"] for s in evidence["sites"]] == ["forward", "backward", "toward"]
    # A reader's pending categories are left for the next check untouched.
    assert flow.pending_categories == {"p": {"grammar"}}
    assert flow._house_respell("house_respell_final")["proposed"] == 0
