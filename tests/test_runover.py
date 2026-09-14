"""Rejoining paragraphs a typeset export split across page boundaries."""
from __future__ import annotations

from zipfile import ZipFile

from docx import Document
from lxml import etree
import pytest

from docproof import runover
from docproof.ingest import preflight
from docproof.utils.xml_helpers import DocxPackage, paragraph_text, qn, walk_package
from galley.fixed_policy import configuration

INDENT = (248, 300)      # a multi-line true paragraph: first line one indent in
ONE_LINE = (548, 0)      # a one-line true paragraph: the export folds the indent into left
MARGIN = (248, 0)        # a continuation: first line flush at the body margin


def P(text, ind=INDENT, style=None, sect=False, bookmark=False, comment=False):
    return {"text": text, "ind": ind, "style": style, "sect": sect,
            "bookmark": bookmark, "comment": comment}


def typeset_book(path, *specs, header=None, table=None):
    document = Document()
    for spec in specs:
        paragraph = document.add_paragraph(spec["text"], style=spec["style"]) if spec["style"] else \
            document.add_paragraph(spec["text"])
        ppr = paragraph._p.get_or_add_pPr()
        if spec["ind"] is not None:
            ind = etree.SubElement(ppr, qn("w:ind"))
            ind.set(qn("w:left"), str(spec["ind"][0]))
            ind.set(qn("w:firstLine"), str(spec["ind"][1]))
        if spec["sect"]:
            etree.SubElement(ppr, qn("w:sectPr"))
        if spec["bookmark"]:
            start = etree.Element(qn("w:bookmarkStart"), {qn("w:id"): "3", qn("w:name"): "mark"})
            end = etree.Element(qn("w:bookmarkEnd"), {qn("w:id"): "3"})
            paragraph._p.insert(1, start)
            paragraph._p.append(end)
            etree.SubElement(paragraph._p, qn("w:proofErr"), {qn("w:type"): "spellStart"})
        if spec["comment"]:
            document.add_comment(paragraph.runs[0], text="Author note.", author="Author")
    if header is not None:
        document.sections[0].header.paragraphs[0].text = header
    if table is not None:
        document.add_table(rows=1, cols=1).cell(0, 0).text = table
    document.save(path)
    return path


def convention_book(path, *extra, count=12, **kwargs):
    """Enough indented body paragraphs to establish the convention."""
    filler = [P(f"Body paragraph number {i} runs on for a while.") for i in range(count)]
    return typeset_book(path, *filler, *extra, **kwargs)


def body_texts(pkg):
    return [paragraph_text(wp.element) for wp in walk_package(pkg)
            if wp.location == "body" and paragraph_text(wp.element).strip()]


@pytest.fixture
def cfg():
    return configuration()


def test_indent_convention_gate_refuses_block_paragraph_books(tmp_path, cfg):
    flat = typeset_book(tmp_path / "flat.docx", *[P(f"Line {i} of text here.", ind=MARGIN) for i in range(12)],
                        P("continued here.", ind=MARGIN))
    joins, diag = runover.find_runover_joins(DocxPackage(flat), cfg)
    assert joins == [] and diag["gate"] == "no_indent_convention"
    absent = typeset_book(tmp_path / "absent.docx", *[P(f"Line {i} of text here.", ind=None) for i in range(12)])
    joins, diag = runover.find_runover_joins(DocxPackage(absent), cfg)
    assert joins == [] and diag["convention"]["applies"] is False


def test_joins_runover_with_space_seam_and_preserved_space(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx",
                           P("“A nice stroll? There’s nothing nice about a stroll,” he says. “I will"),
                           P("never—ever!—understand how some people enjoy walking.”", ind=MARGIN))
    pkg = DocxPackage(path)
    joins, diag = runover.find_runover_joins(pkg, cfg)
    assert len(joins) == 1 and diag["refusals"] == {}
    head = joins[0]
    assert head.absorbed == ["body-0013"] and head.seams[0].separator == " "
    assert head.seams[0].offset == len("“A nice stroll? There’s nothing nice about a stroll,” he says. “I will") + 1
    rows = runover.apply_runover_joins(pkg, joins)
    joined = body_texts(pkg)[-1]
    assert joined == ("“A nice stroll? There’s nothing nice about a stroll,” he says. “I will "
                      "never—ever!—understand how some people enjoy walking.”")
    assert rows[0]["para_id"] == "body-0012" and rows[0]["baseline_para_id"] == "body-0012"
    element = next(wp.element for wp in walk_package(pkg) if wp.para_id == "body-0012")
    seam = [t for t in element.iter(qn("w:t")) if t.text == " "]
    assert seam and seam[0].get(qn("xml:space")) == "preserve"
    assert element.find(qn("w:pPr")).find(qn("w:ind")).get(qn("w:firstLine")) == "300"


def test_chain_of_three_folds_into_head(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx",
                           P("The first page ends with the sentence"),
                           P("continuing on the second page and", ind=MARGIN),
                           P("finishing on the third.", ind=MARGIN))
    pkg = DocxPackage(path)
    joins, _ = runover.find_runover_joins(pkg, cfg)
    assert [j.absorbed for j in joins] == [["body-0013", "body-0014"]]
    assert [s.offset for s in joins[0].seams] == [38, 72]
    runover.apply_runover_joins(pkg, joins)
    assert body_texts(pkg)[-1] == "The first page ends with the sentence continuing on the second page and finishing on the third."
    assert len(body_texts(pkg)) == 13


@pytest.mark.parametrize("left, right, expected, status", [
    ("She drifted off into an un-", "usually peaceful sleep.", "unusually", "dropped"),
    ("It was a well-", "known fact about her.", "well-known", "kept"),
])
def test_hyphen_seam_consults_the_dictionary(tmp_path, cfg, left, right, expected, status):
    pytest.importorskip("spylls")
    path = convention_book(tmp_path / "book.docx", P(left), P(right, ind=MARGIN))
    pkg = DocxPackage(path)
    joins, _ = runover.find_runover_joins(pkg, cfg)
    assert joins[0].seams[0].hyphen == status
    runover.apply_runover_joins(pkg, joins)
    assert expected in body_texts(pkg)[-1]


def test_hyphen_seam_is_undecided_without_a_dictionary(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx", P("into an un-"), P("usually calm sea.", ind=MARGIN))
    joins, _ = runover.find_runover_joins(DocxPackage(path), cfg, dictionary=None)
    assert joins[0].seams[0].hyphen == "undecided" and joins[0].text == "into an un-usually calm sea."


def test_em_dash_and_ellipsis_seams_join_without_a_space(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx",
                           P("She paused—"), P("then went on.", ind=MARGIN),
                           P("He waited …"), P("and nothing came.", ind=MARGIN))
    joins, _ = runover.find_runover_joins(DocxPackage(path), cfg)
    assert [j.text for j in joins] == ["She paused—then went on.", "He waited …and nothing came."]


def test_section_break_on_the_head_is_kept_and_on_the_continuation_is_transferred(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx",
                           P("Page one ends with a section", sect=True), P("break carried by the head.", ind=MARGIN),
                           P("Page two ends with the"), P("continuation ending the section.", ind=MARGIN, sect=True))
    pkg = DocxPackage(path)
    joins, _ = runover.find_runover_joins(pkg, cfg)
    assert [s.section_break for j in joins for s in j.seams] == ["kept", "transferred"]
    rows = runover.apply_runover_joins(pkg, joins)
    elements = [wp.element for wp in walk_package(pkg) if wp.para_id in {r["baseline_para_id"] for r in rows}]
    for element in elements:
        ppr = element.find(qn("w:pPr"))
        assert ppr[-1].tag == qn("w:sectPr")
    assert len(body_texts(pkg)) == 14


def test_both_section_breaks_refuse_the_join_but_later_pairs_still_join(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx",
                           P("Both halves end", sect=True), P("a section.", ind=MARGIN, sect=True),
                           P("A later pair"), P("still joins.", ind=MARGIN))
    joins, diag = runover.find_runover_joins(DocxPackage(path), cfg)
    assert diag["refusals"] == {"both_section_breaks": 1}
    assert [j.text for j in joins] == ["A later pair still joins."]


def test_heading_style_mismatch_and_gaps_are_refused(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx",
                           P("A body paragraph"), P("CHAPTER TWO", ind=MARGIN, style="Heading 1"),
                           P("Another body paragraph"), P("in a different style.", ind=MARGIN, style="Quote"),
                           P("Then a blank"), P("", ind=MARGIN), P("follows the blank.", ind=MARGIN),
                           P("A colophon line", ind=MARGIN), P("and another colophon line.", ind=MARGIN))
    joins, diag = runover.find_runover_joins(DocxPackage(path), cfg)
    assert joins == []
    assert diag["refusals"] == {"heading": 1, "style_mismatch": 1, "empty_between": 1, "head_not_indented": 2}


def test_one_line_paragraphs_and_absent_geometry_are_not_continuations(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx",
                           P("He looked up from the sand and smiled at her."),
                           P("“Well, what?” Kai asks.", ind=ONE_LINE),
                           P("She said nothing at all."),
                           P("Inherits its style's indent.", ind=None))
    joins, diag = runover.find_runover_joins(DocxPackage(path), cfg)
    assert joins == [] and diag["refusals"] == {}


def test_text_shape_is_not_consulted(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx",
                           P("My mother forces a small, reassuring smile. “It’s all right, Anahita."),
                           P("“You can come down now. It’s okay.”", ind=MARGIN))
    joins, _ = runover.find_runover_joins(DocxPackage(path), cfg)
    assert [j.text for j in joins] == [
        "My mother forces a small, reassuring smile. “It’s all right, Anahita. “You can come down now. It’s okay.”"]


def test_table_header_and_textbox_paragraphs_are_untouched(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx", P("Body head"), P("body tail.", ind=MARGIN),
                           header="Running head", table="Cell text")
    pkg = DocxPackage(path)
    before = {wp.para_id: paragraph_text(wp.element) for wp in walk_package(pkg) if wp.location != "body"}
    joins, _ = runover.find_runover_joins(pkg, cfg)
    runover.apply_runover_joins(pkg, joins)
    after = {wp.para_id: paragraph_text(wp.element) for wp in walk_package(pkg) if wp.location != "body"}
    assert before == after and "Running head" in after.values() and "Cell text" in after.values()


def test_bookmarks_comment_ranges_and_proof_errors_move_in_order(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx", P("The head paragraph"),
                           P("continues here.", ind=MARGIN, bookmark=True, comment=True))
    pkg = DocxPackage(path)
    comments_before = pkg.raw("word/comments.xml")
    joins, _ = runover.find_runover_joins(pkg, cfg)
    rows = runover.apply_runover_joins(pkg, joins)
    element = next(wp.element for wp in walk_package(pkg) if wp.para_id == rows[0]["baseline_para_id"])
    tags = [child.tag.split("}")[1] for child in element]
    assert tags.index("bookmarkStart") < tags.index("commentRangeStart") < tags.index("commentRangeEnd") < tags.index("bookmarkEnd")
    assert "proofErr" in tags and element.find(qn("w:bookmarkStart")).get(qn("w:id")) == "3"
    assert paragraph_text(element) == "The head paragraph continues here."
    out = tmp_path / "joined.docx"
    pkg.save(out)
    with ZipFile(out) as z:
        assert z.read("word/comments.xml") == comments_before


def test_join_is_a_fixed_point_after_save_and_reopen(tmp_path, cfg):
    path = convention_book(tmp_path / "book.docx", P("Head"), P("tail.", ind=MARGIN))
    pkg = DocxPackage(path)
    rows, diag = runover.join_runover_paragraphs(pkg, cfg)
    assert len(rows) == 1 and diag["joins"] == 1
    assert runover.find_runover_joins(pkg, cfg)[0] == []
    out = tmp_path / "joined.docx"
    pkg.save(out)
    assert runover.find_runover_joins(preflight(out, "abort"), cfg)[0] == []
    assert runover.join_runover_paragraphs(preflight(out, "abort"), cfg) == ([], runover.find_runover_joins(preflight(out, "abort"), cfg)[1])
