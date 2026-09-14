"""Real Word revision intake: originals, comments, crash recovery and tampering."""
import json
from zipfile import ZipFile

from docx import Document
from lxml import etree
import pytest

from docproof.ingest import IngestError
from docproof.utils.xml_helpers import qn
from galley import fixed_intake as intake
from galley.fixed_documents import paragraph_views, write_manuscripts


def revised_book(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("She ")
    run = paragraph.add_run("recieved")
    run.bold = True
    document.add_comment(run, text="Preserve this author note.", author="Author")
    paragraph.add_run(" two letters.")
    old = etree.Element(qn("w:del"), {qn("w:id"): "7", qn("w:author"): "Editor"})
    deleted = etree.SubElement(old, qn("w:r"))
    etree.SubElement(deleted, qn("w:delText")).text = "received"
    run._r.addprevious(old)
    new = etree.Element(qn("w:ins"), {qn("w:id"): "8", qn("w:author"): "Editor"})
    run._r.addprevious(new)
    new.append(run._r)
    header = document.sections[0].header.paragraphs[0].add_run("Running head")
    wrapper = etree.Element(qn("w:ins"), {qn("w:id"): "9"})
    header._r.addprevious(wrapper)
    wrapper.append(header._r)
    path = tmp_path / "Writer - Book 1.docx"
    document.save(path)
    return path


def test_revision_baseline_preserves_original_comments_and_other_parts(tmp_path, monkeypatch):
    source = revised_book(tmp_path)
    before = source.read_bytes()
    baseline, evidence = intake.prepare_source(source, tmp_path / "run")
    assert source.read_bytes() == before
    assert (tmp_path / "run/intake/original" / source.name).read_bytes() == before
    assert paragraph_views(baseline) == paragraph_views(source)
    with ZipFile(source) as original, ZipFile(baseline) as accepted:
        changed = {n for n in original.namelist() if original.read(n) != accepted.read(n)}
        assert changed == {"word/document.xml", "word/header1.xml"}
        assert original.read("word/comments.xml") == accepted.read("word/comments.xml")
    assert Document(baseline).paragraphs[0].runs[1].bold
    corrected = {pid: text.replace("recieved", "received") for pid, text in paragraph_views(baseline).items()}
    tracked, clean, _ = write_manuscripts(baseline, tmp_path / "final", corrected)
    assert paragraph_views(tracked, "reject") == paragraph_views(baseline)
    assert paragraph_views(tracked) == paragraph_views(clean) == corrected
    monkeypatch.setattr(intake, "accept_all_revisions", lambda *a: pytest.fail("Baseline rebuilt on resume"))
    assert intake.prepare_source(source, tmp_path / "run") == (baseline, evidence)


@pytest.mark.parametrize("member", ["original", "accepted", "receipt"])
def test_tampered_intake_cannot_be_certified(tmp_path, member):
    source = revised_book(tmp_path)
    baseline, evidence = intake.prepare_source(source, tmp_path / "run")
    target = (tmp_path / "run/intake/receipt.json" if member == "receipt" else
              tmp_path / "run/intake" / member / source.name)
    target.write_bytes(target.read_bytes() + b"changed")
    with pytest.raises(intake.FixedIntakeError, match="changed"):
        intake.validate_intake(tmp_path / "run", evidence, baseline)


def test_interruption_before_baseline_publication_can_retry(tmp_path, monkeypatch):
    source = revised_book(tmp_path)
    before = source.read_bytes()
    publish = intake.os.replace
    def interrupt(src, dst):
        if str(dst).endswith("/intake"):
            raise OSError("Interrupted intake publication")
        return publish(src, dst)
    monkeypatch.setattr(intake.os, "replace", interrupt)
    with pytest.raises(OSError, match="Interrupted"):
        intake.prepare_source(source, tmp_path / "run")
    assert source.read_bytes() == before
    assert not (tmp_path / "run/intake").exists()
    monkeypatch.setattr(intake.os, "replace", publish)
    baseline, evidence = intake.prepare_source(source, tmp_path / "run")
    assert intake.validate_intake(tmp_path / "run", evidence) == baseline


def test_unsupported_revision_stops_before_workflow_or_model_calls(tmp_path):
    source = revised_book(tmp_path)
    document = Document(source)
    props = document.add_table(rows=1, cols=1).cell(0, 0)._tc.get_or_add_tcPr()
    props.append(etree.Element(qn("w:cellIns"), {qn("w:id"): "10"}))
    document.save(source)
    before = source.read_bytes()
    from galley.fixed_workflow import FixedWorkflow
    with pytest.raises(IngestError, match="Can't resolve"):
        FixedWorkflow(source, tmp_path / "run", calls=object())
    assert source.read_bytes() == before
    assert not (tmp_path / "run/workflow.json").exists()
    assert not (tmp_path / "run/intake").exists()


def test_clean_manuscript_keeps_its_source_and_identity(tmp_path):
    source = tmp_path / "clean.docx"
    document = Document()
    document.add_paragraph("A quiet room.")
    document.save(source)
    assert intake.prepare_source(source, tmp_path / "run") == (source, None)
    assert not (tmp_path / "run/intake").exists()


def test_source_change_refuses_to_rebuild_existing_baseline(tmp_path):
    source = revised_book(tmp_path)
    baseline, evidence = intake.prepare_source(source, tmp_path / "run")
    source.write_bytes(source.read_bytes() + b"changed")
    with pytest.raises(intake.FixedIntakeError, match="incoming manuscript changed"):
        intake.prepare_source(source, tmp_path / "run")
    assert intake.validate_intake(tmp_path / "run", evidence) == baseline


def test_deleted_paragraph_mark_uses_words_accepted_join_and_style(tmp_path):
    from test_ingest import _tracked_doc
    source = _tracked_doc(tmp_path)
    original = source.read_bytes()
    baseline, evidence = intake.prepare_source(source, tmp_path / "run")
    joined = next(p for p in Document(baseline).paragraphs if p.text == "The first half and the second half.")
    assert joined.style.name == "Quote"
    assert source.read_bytes() == original
    assert intake.validate_intake(tmp_path / "run", evidence) == baseline


def runover_book(tmp_path, *, revisions=False):
    """A typeset export: indented paragraphs, one split across a page boundary."""
    from test_runover import INDENT, MARGIN, P, typeset_book
    filler = [P(f"Body paragraph number {i} runs on for a while.", ind=INDENT) for i in range(12)]
    path = typeset_book(tmp_path / "Writer - Galley.docx", *filler,
                        P("“Fine,” he says. “I will", ind=INDENT),
                        P("never understand the rules of this house.”", ind=MARGIN))
    if revisions:
        document = Document(path)
        run = document.paragraphs[-1].runs[0]
        run.text = "never understand the "
        added = document.paragraphs[-1].add_run("rules of this house.”")
        wrapper = etree.Element(qn("w:ins"), {qn("w:id"): "12", qn("w:author"): "Editor"})
        added._r.addprevious(wrapper)
        wrapper.append(added._r)
        document.save(path)
    return path


def test_runover_book_without_revisions_gets_a_joined_baseline(tmp_path, monkeypatch):
    source = runover_book(tmp_path)
    before = source.read_bytes()
    baseline, evidence = intake.prepare_source(source, tmp_path / "run")
    assert baseline != source and source.read_bytes() == before
    receipt = json.loads((tmp_path / "run/intake/receipt.json").read_text())
    assert receipt["version"] == "fixed-intake-v2" == evidence["version"]
    assert receipt["changed_parts"] == ["word/document.xml"]
    assert receipt["resolved_revision_elements"] == {}
    assert receipt["paragraph_id_space"] == "accepted-before-join"
    assert receipt["indent_convention"]["applies"] is True
    [join] = receipt["runover_joins"]
    assert join["para_id"] == "body-0012" and join["absorbed"] == ["body-0013"]
    assert join["seam_offsets"] == [len("“Fine,” he says. “I will") + 1]
    views = paragraph_views(baseline)
    assert views[join["baseline_para_id"]] == "“Fine,” he says. “I will never understand the rules of this house.”"
    assert receipt["paragraphs"] == len(views)
    assert intake.validate_intake(tmp_path / "run", evidence, baseline) == baseline
    monkeypatch.setattr(intake, "join_runover_paragraphs", lambda *a, **k: pytest.fail("Baseline rebuilt on resume"))
    assert intake.prepare_source(source, tmp_path / "run") == (baseline, evidence)


def test_runover_receipt_ids_round_trip_between_source_and_baseline(tmp_path):
    source = runover_book(tmp_path)
    baseline, _ = intake.prepare_source(source, tmp_path / "run")
    receipt = json.loads((tmp_path / "run/intake/receipt.json").read_text())
    [join] = receipt["runover_joins"]
    incoming, accepted = paragraph_views(source), paragraph_views(baseline)
    head, [tail] = incoming[join["para_id"]], [incoming[pid] for pid in join["absorbed"]]
    assert accepted[join["baseline_para_id"]] == head + join["separators"][0] + tail
    assert len(accepted) == len(incoming) - 1


def test_revisions_are_accepted_before_runovers_are_joined(tmp_path):
    source = runover_book(tmp_path, revisions=True)
    baseline, evidence = intake.prepare_source(source, tmp_path / "run")
    receipt = json.loads((tmp_path / "run/intake/receipt.json").read_text())
    assert receipt["changed_parts"] == ["word/document.xml"]
    assert receipt["resolved_revision_elements"] == {"word/document.xml": 1}
    assert len(receipt["runover_joins"]) == 1
    assert "“Fine,” he says. “I will never understand the rules of this house.”" in paragraph_views(baseline).values()
    assert intake.validate_intake(tmp_path / "run", evidence, baseline) == baseline


def test_v1_intake_receipt_requires_a_fresh_workspace(tmp_path):
    source = runover_book(tmp_path)
    intake.prepare_source(source, tmp_path / "run")
    path = tmp_path / "run/intake/receipt.json"
    receipt = json.loads(path.read_text())
    receipt["version"] = "fixed-intake-v1"
    path.write_text(json.dumps(receipt))
    with pytest.raises(intake.FixedIntakeError, match="fresh workspace"):
        intake.prepare_source(source, tmp_path / "run")


def test_unjoined_baseline_cannot_be_published(tmp_path, monkeypatch):
    from docproof import runover
    source = runover_book(tmp_path)
    monkeypatch.setattr(runover, "apply_runover_joins", lambda pkg, joins: [])
    with pytest.raises(runover.RunoverError, match="fixed point"):
        intake.prepare_source(source, tmp_path / "run")
    assert not (tmp_path / "run/intake").exists()
