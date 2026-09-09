"""The 2026-09-07 delivery review: a promised plan line vanished, two stale
comments shipped, the letter overclaimed, the counts disagreed, the headings
named the wrong book, and a settlement cited the wrong edit."""
from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from galley import plan_ledger as pl
from galley.premises import quotes_unbalanced, stale_queries

PLAN = """
```
0.  chapter sweep (Luna 6 windows + Sonnet 6 windows $0)   $0.35   est. 120-200 findings
2.  mechanical ladder (Luna+Sonnet ensemble, Luna verifier) $0.91
4c. $0 subagent lanes: whole-book continuity (Opus),
    chapter continuity (Opus), type-and-compare (Sonnet)   $0      query-only, imported
7.  verify (rotated reread, 2 passes) + settle + certify   $0
TOTAL $2.00  ·  stop: $2/finding marginal, ONE wave
```
3. **The book arguably wants copy-editing.** Recorded here and for the letter,
"""


# --- the plan ledger -----------------------------------------------------------

def test_plan_items_are_the_priced_numbered_lines_only():
    labels = [it.label for it in pl.plan_items(PLAN)]
    # 4c is a promise — "$0" on its own line is a price. The caveat "3." has
    # none and is not; TOTAL is not an item.
    assert labels == ["0", "2", "4c", "7"]


def test_the_ledger_records_and_audits(tmp_path):
    ledger = tmp_path / "runs" / pl.LEDGER_NAME
    pl.record(ledger, "0", "ran", evidence="runs/sweep_windows/")
    pl.record(ledger, "2", "ran", evidence="runs/ladder/findings.json")
    pl.record(ledger, "4c", "ran", evidence="runs/continuity/")
    with pytest.raises(ValueError):
        pl.record(ledger, "7", "maybe")
    status, detail = pl.check(PLAN, pl.load_ledger(ledger), ledger_exists=True)
    assert status == "fail" and "line 7" in detail
    pl.record(ledger, "7", "ran", evidence="runs/final2/")
    status, detail = pl.check(PLAN, pl.load_ledger(ledger), ledger_exists=True)
    assert status == "pass", detail


def test_a_promise_that_vanished_fails_and_is_named_for_the_letter(tmp_path):
    """Plan line 4c, 2026-09-07: deferred-to and never run."""
    plan = "1. ladder  $1.00\n4c. whole-book continuity (Opus)  $0\nTOTAL $1\n"
    ledger = tmp_path / pl.LEDGER_NAME
    pl.record(ledger, "1", "ran", evidence="runs/ladder/")
    status, detail = pl.check(plan, pl.load_ledger(ledger), ledger_exists=True)
    assert status == "fail" and "4c" in detail and "no ledger entry" in detail
    pl.record(ledger, "4c", "skipped", reason="")
    status, detail = pl.check(plan, pl.load_ledger(ledger), ledger_exists=True)
    assert status == "fail" and "without a reason" in detail
    pl.record(ledger, "4c", "skipped", reason="session ran out of turns")
    assert pl.check(plan, pl.load_ledger(ledger), ledger_exists=True)[0] == "pass"
    rows = pl.not_done(plan, pl.load_ledger(ledger))
    assert [(it.label, st) for it, st, _ in rows] == [("4c", "skipped")]
    assert rows[0][2] == "session ran out of turns"


def test_no_ledger_at_all_fails_when_the_plan_made_promises():
    assert pl.check("1. ladder $1\n", None, ledger_exists=False)[0] == "fail"
    assert pl.check("no items here", None, ledger_exists=False)[0] == "skip"


def test_certify_reads_the_ledger_beside_the_run(tmp_path):
    from galley.manifest import _certify_plan_ledger
    ws = tmp_path / "ws"
    (ws / "runs" / "final").mkdir(parents=True)
    run = ws / "runs" / "final"
    assert _certify_plan_ledger(run).status == "skip"          # no PLAN.md
    (ws / "PLAN.md").write_text("1. ladder  $1.00\nTOTAL $1\n", "utf-8")
    c = _certify_plan_ledger(run)
    assert c.status == "fail" and "plan_ledger.json" in c.detail
    pl.record(ws / "runs" / pl.LEDGER_NAME, "1", "ran", evidence="runs/final")
    assert _certify_plan_ledger(run).status == "pass"


# --- stale comments ------------------------------------------------------------

def test_quote_balance_is_the_premise():
    assert quotes_unbalanced("“Not at night.")
    assert not quotes_unbalanced("“Not at night.”")
    assert quotes_unbalanced('he said "no')
    assert not quotes_unbalanced("no quotes here")


def test_a_comment_whose_premise_settle_fixed_is_stale():
    rows = [
        {"finding_id": "f-0044", "para_id": "body-0037", "status": "query",
         "queried": True, "error_type": "unclosed_quote"},
        {"finding_id": "f-0045", "para_id": "body-0038", "status": "query",
         "queried": True, "error_type": "unclosed_quote"},
        {"finding_id": "f-0090", "para_id": "body-0038", "status": "query",
         "queried": True, "error_type": "spelling"},          # no predicate
    ]
    delivered = {"body-0037": "“Not at night.”",              # fixed → stale
                 "body-0038": "“Marley—” “You said the rules were.”"}  # balanced
    stale = stale_queries(rows, delivered)
    assert [r["finding_id"] for r, _ in stale] == ["f-0044", "f-0045"]
    # One still unclosed, one now closed; a paragraph missing from the
    # delivered view is skipped, never judged.
    still = stale_queries(rows, {"body-0037": "“Not at night.",
                                 "body-0038": "“Marley—” “You said the rules were.”"})
    assert [r["finding_id"] for r, _ in still] == ["f-0045"]
    assert stale_queries(rows, {}) == []


# --- the letter ----------------------------------------------------------------

def _src(tmp_path, rows, plan="", workspace=None):
    from galley.journal import JournalSources
    return JournalSources(run_dir=tmp_path, workspace=workspace,
                          envelope={"findings": rows}, plan=plan)


APPLIED = [
    {"finding_id": "f-1", "para_id": "p1", "applied": True, "status": "validated",
     "original_text": "went", "corrected_text": "gone",
     "anchor": {"delete_text": "went", "insert_text": "gone"}},
    {"finding_id": "f-2", "para_id": "p2", "applied": True, "status": "validated",
     "original_text": "all of the sudden", "corrected_text": "all of a sudden",
     "anchor": {"delete_text": "of the sudden", "insert_text": "of a sudden"}},
    {"finding_id": "f-3", "para_id": "p3", "applied": True, "status": "validated",
     "format": True, "error_type": "speaker_split", "original_text": "", "corrected_text": ""},
]
COMMENTS = [
    {"finding_id": "q-1", "para_id": "p4", "status": "query", "queried": True,
     "error_type": "galley_read", "explanation": "Is Marley the fourth or fifth Quinn?",
     "original_text": "the fourth generation"},
    {"finding_id": "q-2", "para_id": "p5", "status": "query", "queried": True,
     "error_type": "speaker_split", "explanation": "I've separated the dialogue.",
     "original_text": "“It’s been alive"},
    {"finding_id": "q-3", "para_id": "p6", "status": "query", "queried": True,
     "error_type": "sweep_dash", "explanation": "House style sets an em dash unspaced.",
     "original_text": "edition — contains"},
]


def test_edit_shapes_are_counted_not_asserted():
    from galley.letter import edit_shapes
    assert edit_shapes(APPLIED) == {"single_word": 1, "multiword": 1,
                                    "paragraph_marks": 1, "formatting": 0}


def test_the_summary_separates_corrections_from_paragraph_breaks(tmp_path):
    from galley.casefile import CaseFile
    from galley.letter import _summary_line
    line = _summary_line(_src(tmp_path, APPLIED + COMMENTS), CaseFile())
    assert "**2 tracked correction(s)** and **1 paragraph break(s)** (3 tracked changes in all)" in line
    assert "**3 margin comment(s)**" in line


def test_questions_and_notes_are_told_apart(tmp_path):
    from galley.casefile import CaseFile
    from galley.letter import _section_decisions, split_comments
    questions, notes = split_comments(COMMENTS)
    assert [r["finding_id"] for r in questions] == ["q-1"]
    assert [r["finding_id"] for r in notes] == ["q-2", "q-3"]
    text = "\n".join(_section_decisions(_src(tmp_path, COMMENTS), CaseFile()))
    assert "1 author question(s)" in text
    assert "2 note(s)" in text
    assert "### Questions for you (1)" in text
    assert "### Notes on what was done (2)" in text
    assert "Each asks something only you can answer" not in text


def test_the_letter_names_plan_lines_that_did_not_run(tmp_path):
    from galley.letter import _section_plan_ledger
    ws = tmp_path / "ws"; (ws / "runs").mkdir(parents=True)
    plan = "1. ladder $1\n4c. whole-book continuity (Opus)  $0\nTOTAL $1\n"
    pl.record(ws / "runs" / pl.LEDGER_NAME, "1", "ran", evidence="runs/final")
    pl.record(ws / "runs" / pl.LEDGER_NAME, "4c", "deferred",
              reason="the session ended before the continuity read")
    out = "\n".join(_section_plan_ledger(_src(tmp_path, [], plan=plan, workspace=ws)))
    assert "## Plan lines that did not run" in out
    assert "**4c**" in out and "deferred: the session ended" in out
    assert "**1**" not in out


def test_the_headings_take_the_delivered_name(tmp_path):
    from galley.casefile import CaseFile
    from galley.letter import (render_letter, render_style_sheet,
                               render_verification_report)
    cf = CaseFile(book="Test - Book One.docx")
    src = _src(tmp_path, APPLIED)
    letter = render_letter(cf, tmp_path / "out", evidence=src,
                           title="Test - Book Two").read_text("utf-8")
    assert letter.startswith("# Proofreading letter — Test - Book Two")
    sheet = render_style_sheet(cf, tmp_path / "out", evidence=src,
                               title="Test - Book Two").read_text("utf-8")
    assert sheet.startswith("# Style sheet — Test - Book Two")
    report = render_verification_report(src, tmp_path / "out", cf=cf,
                                        title="Test - Book Two").read_text("utf-8")
    assert report.startswith("# Verification report — Test - Book Two")
    assert "Book One" not in letter.splitlines()[0]


# --- the settlement owner -------------------------------------------------------

def test_settlement_records_carry_a_build_stable_owner_key(monkeypatch):
    from galley import settle
    rec = settle.SettlementRecord("r-1", 1, "absorb", "f-0210", "a", "b", "", "x")
    monkeypatch.setattr(settle, "_apply_decision",
                        lambda *a, **k: (rec, [], {}))
    working = {"f-0210": {"para_id": "body-0007", "original_text": "slept on the windowsill hardly moving",
                          "occurrence": 1, "corrected_text": "slept on the windowsill, hardly moving",
                          "error_type": "galley_settle"}}
    out, _, _ = settle.apply_decision(SimpleNamespace(), SimpleNamespace(owner_key="f-0210"),
                                      working, {}, 1, verified_by="x")
    assert out.owner_row_key == ["body-0007", "slept on the windowsill hardly moving", 1,
                                 "slept on the windowsill, hardly moving", "galley_settle"]
    again = settle.SettlementRecord.from_json(out.to_json())
    assert again.owner_row_key == out.owner_row_key


def test_the_journal_resolves_the_owner_against_the_build_it_renders():
    """f-0210 was body-0011 in one build and body-0019 in the next; the
    windowsill fix is f-0216 in the delivered build."""
    from galley.journal import _rows_by_key, owner_label
    rows = [{"finding_id": "f-0210", "para_id": "body-0019", "original_text": "Wasn’t you scared?",
             "occurrence": 1, "corrected_text": "Weren’t you scared?", "error_type": "galley_read"},
            {"finding_id": "f-0216", "para_id": "body-0007", "original_text": "slept on the windowsill hardly moving",
             "occurrence": 1, "corrected_text": "slept on the windowsill, hardly moving", "error_type": "galley_settle"}]
    idx = _rows_by_key(rows)
    keyed = {"owner_finding_id": "f-0210",
             "owner_row_key": ["body-0007", "slept on the windowsill hardly moving", 1,
                               "slept on the windowsill, hardly moving", "galley_settle"]}
    assert owner_label(keyed, idx) == "f-0216"
    old = {"owner_finding_id": "f-0210"}
    assert owner_label(old, idx) == "f-0210"
    orphan = {"owner_finding_id": "f-0210", "owner_row_key": ["gone", "", 1, "", ""]}
    assert owner_label(orphan, idx) == "f-0210 (id from an earlier build)"


# --- the outcome's evidence ----------------------------------------------------

def test_outcome_evidence_says_which_unresolved_is_which(tmp_path):
    from galley.outcome import evidence_of
    (tmp_path / "findings.json").write_text(json.dumps({"findings": APPLIED + COMMENTS}), "utf-8")
    ev = evidence_of(tmp_path, source_paras={"p1": "went home"})
    assert ev["open_author_queries"] == 3
    assert ev["unresolved_internal"] == 0
    assert ev["unresolved_queries"] == ev["unresolved_internal"]


# --- the author letter ---------------------------------------------------------

def _docx_text(path):
    import docx
    return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)


def test_the_author_letter_says_what_the_author_needs_and_nothing_else(tmp_path):
    from galley.casefile import CaseFile
    from galley.letter import AUTHOR_LETTER_NAME, render_author_letter
    ws = tmp_path / "ws"; (ws / "runs").mkdir(parents=True)
    plan = "1. ladder $1\n4c. whole-book continuity (Opus)  $0\nTOTAL $1\n"
    pl.record(ws / "runs" / pl.LEDGER_NAME, "1", "ran", evidence="runs/final")
    pl.record(ws / "runs" / pl.LEDGER_NAME, "4c", "skipped",
              reason="the session ended before it ran")
    src = _src(tmp_path, APPLIED + COMMENTS, plan=plan, workspace=ws)
    src.approval = {"mechanical_only": True, "max_spend_usd": 10.0}
    path = render_author_letter(CaseFile(book="Test - Book One.docx"),
                                tmp_path / "out", evidence=src,
                                title="Test - Book Two.docx")
    assert path.name == AUTHOR_LETTER_NAME
    text = _docx_text(path)
    assert "A note from your proofreader" in text
    assert "Test - Book Two" in text and "Book One" not in text
    assert "2 tracked correction(s) and 1 paragraph break(s)" in text
    assert "We did not line-edit" in text
    assert "Questions for you (1)" in text
    assert "Is Marley the fourth or fifth Quinn?" in text
    assert "Things we did that you should know about (2)" in text
    assert "I've separated the dialogue." in text
    assert "Not done in this pass" in text and "whole-book continuity" in text
    # Nothing internal reaches the author.
    assert "$" not in text and "f-1" not in text and "q-1" not in text
    assert "body-" not in text and "p4" not in text


def test_the_handoff_carries_the_author_letter(tmp_path):
    import docx as _docx
    from galley.driver import build_handoff, DriverError
    ws = tmp_path / "ws"; dv = ws / "deliverable"; dv.mkdir(parents=True)
    d = _docx.Document(); d.add_paragraph("x"); d.save(str(dv / "Test - Book One - Atmosphere Press Proofreader.docx"))
    for n in ("letter.md", "style-sheet.md", "decision-log.md", "verification.md"):
        (dv / n).write_text("# x\n", "utf-8")
    (dv / "outcome.json").write_text("{}", "utf-8")
    with pytest.raises(DriverError, match="author letter"):
        build_handoff(ws, "Test - Book One.docx", tmp_path / "h",
                      outcome_sources=[dv / "outcome.json"])
    d2 = _docx.Document(); d2.add_paragraph("letter"); d2.save(str(dv / "author-letter.docx"))
    written = build_handoff(ws, "Test - Book One.docx", tmp_path / "h",
                            outcome_sources=[dv / "outcome.json"])
    assert any(p.name == "Test - Book Two - Author Letter.docx" for p in written)
    assert any(p.name == "Test - Book Two - clean.docx" for p in written)
    assert len(written) == 8


# --- who overruled ---------------------------------------------------------------

def test_every_phase_session_carries_its_phase_in_the_environment(tmp_path):
    from galley import driver as gd
    drv = gd.Driver(book=tmp_path / "b.docx", slug="s",
                    workspace_root=tmp_path / "ws")
    spec = drv._spec("deliver", {"PATH": "/usr/bin",
                                 "CLAUDE_CODE_OAUTH_TOKEN": "tok"})
    assert spec.env[gd.BRAIN_PHASE_ENV] == "deliver"
    assert spec.env["PATH"] == "/usr/bin"          # the rest is untouched


def _outcome_args(run, **kw):
    base = dict(run=str(run), source=None, config="config/default.yaml",
                set="done", reason="seeded test book; density is expected",
                by="", done_value=None, needs_human_value=None,
                rewrite_share=None, edit_density=None, json=False)
    base.update(kw)
    return SimpleNamespace(**base)


def _seed_run(tmp_path):
    run = tmp_path / "run"; run.mkdir()
    (run / "findings.json").write_text(json.dumps({"findings": []}), "utf-8")
    return run


def test_an_overrule_is_attributed_to_the_brain_that_made_it(tmp_path, monkeypatch, capsys):
    """The brain MAY overrule needs_human (owner's policy, 2026-09-07). The
    record must say the brain did — the first delivery's outcome.json, log
    and HubSpot value all said "human", and nobody had."""
    from docproof.__main__ import _galley_outcome
    from galley.driver import BRAIN_PHASE_ENV
    run = _seed_run(tmp_path)
    monkeypatch.setenv(BRAIN_PHASE_ENV, "deliver")
    assert _galley_outcome(_outcome_args(run)) == 0
    oc = json.loads((run / "outcome.json").read_text("utf-8"))
    assert oc["outcome"] == "done"
    assert oc["set_by"] == "galley brain (deliver phase)"


def test_a_person_at_a_terminal_is_still_a_person(tmp_path, monkeypatch):
    from docproof.__main__ import _galley_outcome
    from galley.driver import BRAIN_PHASE_ENV
    run = _seed_run(tmp_path)
    monkeypatch.delenv(BRAIN_PHASE_ENV, raising=False)
    _galley_outcome(_outcome_args(run))
    assert json.loads((run / "outcome.json").read_text("utf-8"))["set_by"] == "human"
    _galley_outcome(_outcome_args(run, by="Quinton"))
    assert json.loads((run / "outcome.json").read_text("utf-8"))["set_by"] == "Quinton"
