"""Anchor-only retries preserve findings, validated windows, and accounting."""
import copy
import json
import uuid

import docx
import pytest

from docproof import resource_ledger as ledger
from docproof.models import Usage
from docproof.providers.base import NormalizedUsage, ProviderResult
from galley import verify
from .test_verification_checkpoints import MODEL, POLICY, isolated_coverage, run


class Provider:
    name = "anchor-test"
    effort = "high"

    def __init__(self, *bodies):
        self.bodies, self.calls = list(bodies), []

    def complete_structured(self, **request):
        operation = uuid.uuid4().hex
        fields = dict(receipt_id=operation, operation_id=operation, model=MODEL,
                      transport="test", max_output_tokens=request["max_tokens"])
        ledger.append_usage(**fields, status="started")
        self.calls.append(copy.deepcopy(request))
        assert self.bodies, "Unexpected extra model call"
        result = ProviderResult(parsed=copy.deepcopy(self.bodies.pop(0)),
                                usage=NormalizedUsage(10, 4, billed=False))
        ledger.append_usage(**fields, usage=result.usage)
        return result


def finding(pid, quote):
    return dict(para_id=pid, quote=quote, problem="A mechanical typo.",
                suggestion="cat", severity="high")


def invoke(run, provider):
    with verify.verification_invocation(run, provider, MODEL, **POLICY) as invocation:
        options = {**POLICY, "command_id": invocation.command_id}
        uc, uw = Usage(), Usage()
        changes = verify.verify_run(run, provider, MODEL, uc, run_walk=False, **options)
        walk = verify.verify_run(run, provider, MODEL, uw, run_changes=False, **options)
        verify.write_artifacts(run, changes, walk, model=MODEL, engine="provider",
            usage_changes=uc, usage_walk=uw, applied=len(verify.applied_edits(run)),
            paragraphs=len(verify.accepted_text(run)))
        return invocation.mark_complete(changes, walk), invocation.command_id, changes, walk, uc, uw


def diagnostics(run):
    return [json.loads(path.read_text()) for path in
            (run / verify._CHECKPOINT_DIR).glob("*/*/rejected/*/*.json")]


def test_exact_quote_and_id_repair_keeps_all_findings_and_original_identity(run):
    document = docx.Document()
    document.add_paragraph("The “dog” sleeps.")
    document.save(run / "book.docx")
    pid = next(iter(verify.accepted_text(run)))
    original = {"findings": [finding(pid, "dog"), finding(pid, '"dog"'),
                              finding("unknown", "sleeps")]}
    corrections = {"anchors": [{"index": 1, "para_id": pid, "quote": "“dog”"},
                                {"index": 2, "para_id": pid, "quote": "sleeps"}]}
    provider = Provider({"problems": []}, original, corrections)
    complete, _, _, walk, _, usage = invoke(run, provider)
    assert complete and usage.api_calls == 2 and len(provider.calls) == 3
    expected = copy.deepcopy(original)
    expected["findings"][1]["quote"] = "“dog”"
    expected["findings"][2]["para_id"] = pid
    actual = [{key: getattr(row, key) for key in original["findings"][0]} for row in walk.residuals]
    assert actual == expected["findings"]
    primary, retry = provider.calls[1:]
    assert primary["schema"] == verify._walk_schema()[0]
    assert primary["system"] == verify._WALK_SYSTEM + verify._context_block("", "proofreader")
    assert primary["user"] == verify._walk_user(list(verify.accepted_text(run).items()))
    assert retry["system"] != primary["system"] and "anchors schema" in retry["system"]
    assert retry["schema_name"] == "anchors" and "anchors" in retry["schema"]["properties"]
    assert retry["model"] == primary["model"] and retry["max_tokens"] == primary["max_tokens"]
    assert verify.validate_complete_pass(run, run, provider, MODEL, **POLICY)
    records = {row["stage"]: row for row in diagnostics(run)}
    assert records["initial_rejected"]["parsed"] == original
    assert records["anchor_repair"]["parsed"] == corrections
    assert records["anchor_repair"]["repair_accepted"] is True
    assert records["anchor_repair"]["reconstructed_sha256"] == verify._digest(expected)
    assert records["anchor_repair"]["request_sha256"] == records["initial_rejected"]["request_sha256"]
    assert records["anchor_repair"]["actual_request_sha256"] != records["initial_rejected"]["actual_request_sha256"]


@pytest.mark.parametrize("damage", ["delete", "add", "duplicate", "judgment", "whole_response",
                                    "wrong_quote", "wrong_id", "valid_row", "boolean_index"])
def test_bad_repair_cannot_lose_findings_or_complete_coverage(run, damage):
    pid = next(iter(verify.accepted_text(run)))
    original = {"findings": [finding(pid, "sleeps"), finding(pid, "DOG")]}
    anchor = {"index": 1, "para_id": pid, "quote": "dog"}
    reply = {"anchors": [anchor]}
    if damage == "delete":
        reply["anchors"] = []
    elif damage == "add":
        reply["anchors"].append({**anchor, "index": 2})
    elif damage == "duplicate":
        reply["anchors"].append(dict(anchor))
    elif damage == "judgment":
        anchor["problem"] = "Actually clean."
    elif damage == "whole_response":
        reply = {"findings": []}
    elif damage == "wrong_quote":
        anchor["quote"] = "DOG"
    elif damage == "wrong_id":
        anchor["para_id"] = "outside-the-read"
    elif damage == "valid_row":
        anchor["index"] = 0
    else:
        anchor["index"] = True
    provider, usage = Provider(original, reply), Usage()
    result = verify.verify_run(run, provider, MODEL, usage, run_changes=False, **POLICY)
    assert len(provider.calls) == usage.api_calls == 2
    assert result.residuals == [] and not result.ran_walk
    assert not result.verification_provenance["walk"]["complete"]
    assert result.verification_provenance["walk"]["windows"] == {}
    assert verify.UNREAD == list(verify.accepted_text(run))
    records = {row["stage"]: row for row in diagnostics(run)}
    assert records["initial_rejected"]["parsed"] == original
    assert records["anchor_repair"]["repair_accepted"] is False


@pytest.mark.parametrize("kind", ["schema", "severity", "capacity"])
def test_non_anchor_failures_keep_one_plain_retry(run, kind):
    pid = next(iter(verify.accepted_text(run)))
    row = finding(pid, "dog")
    body = {"findings": [row]}
    if kind == "schema":
        del row["problem"]
    elif kind == "severity":
        row["severity"] = "urgent"
    else:
        body["findings"] *= verify.MAX_RESIDUALS_PER_READ + 1
    provider, usage = Provider(body, body), Usage()
    result = verify.verify_run(run, provider, MODEL, usage, run_changes=False, **POLICY)
    assert len(provider.calls) == usage.api_calls == 2
    assert provider.calls[0] == provider.calls[1]
    assert not result.verification_provenance["walk"]["complete"]
    assert {row["stage"] for row in diagnostics(run)} == {"initial_rejected", "retry_rejected"}


def test_resume_preserves_completed_reads_and_saved_budget(run):
    document = docx.Document()
    for n in range(2):
        document.add_paragraph(f"Paragraph {n} dog. " + "Quiet words. " * 320)
    document.save(run / "book.docx")
    ids = list(verify.accepted_text(run))
    bad = {"findings": [finding(ids[1], "DOG")]}
    first = Provider({"problems": []}, {"findings": []}, bad, {"anchors": []})
    path = run / "resources.jsonl"
    env = {**ledger.context_env(path, "source", "config"), ledger.GROUP_ENV: "review",
           ledger.MAX_CALLS_ENV: "6", ledger.MAX_OUTPUT_ENV: "100000"}
    with ledger.use_context(env):
        complete, command_id, changes, walk, _, _ = invoke(run, first)
    assert not complete and len(first.calls) == 4
    previous = {}
    for proof in (changes.verification_provenance["changes"], walk.verification_provenance["walk"]):
        for key in proof["windows"]:
            file = run / verify._CHECKPOINT_DIR / proof["scope"] / proof["invocation_id"] / (key + ".json")
            previous[file] = file.read_bytes()
    assert len(previous) == 2
    second = Provider(bad, {"anchors": [{"index": 0, "para_id": ids[1], "quote": "dog"}]})
    with ledger.use_context(env):
        complete, resumed_id, changes, walk, uc, uw = invoke(run, second)
    assert complete and resumed_id == command_id
    assert len(second.calls) == uw.api_calls == 2 and uc.api_calls == 0
    assert all(file.read_bytes() == body for file, body in previous.items())
    assert changes.recovered_usage.api_calls == 1 and walk.recovered_usage.api_calls == 3
    assert verify.validate_complete_pass(run, run, second, MODEL, **POLICY)
    totals = ledger.summarize(path)
    assert totals["new_attempts"] == 6 and totals["output_tokens"] == 24
    assert totals["groups"]["review"]["remaining_calls"] == 0
    with ledger.use_context(env), pytest.raises(ledger.ResourceBudgetExceeded):
        Provider({"findings": []}).complete_structured(**second.calls[0])
