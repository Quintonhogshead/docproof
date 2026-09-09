"""Local reconciliation preserves the reviewed source and never calls a model."""
import json
import zipfile

import docx
import pytest
from lxml import etree

from docproof.utils.xml_helpers import DocxPackage, qn
from galley import astra_review as ar


def _revision(p, kind, word_id, text):
    node = etree.SubElement(p._p, qn("w:" + kind), {
        qn("w:id"): str(word_id), qn("w:author"): "Prior editor",
        qn("w:date"): "2026-09-01T12:00:00Z",
    })
    run = etree.SubElement(node, qn("w:r"))
    etree.SubElement(run, qn("w:delText" if kind == "del" else "w:t")).text = text


def _anchor(p, comment_id):
    p._p.insert(0, etree.Element(qn("w:commentRangeStart"), {qn("w:id"): comment_id}))
    etree.SubElement(p._p, qn("w:commentRangeEnd"), {qn("w:id"): comment_id})
    run = etree.SubElement(p._p, qn("w:r"))
    etree.SubElement(run, qn("w:commentReference"), {qn("w:id"): comment_id})


@pytest.fixture
def run(tmp_path):
    path = tmp_path / "run"
    path.mkdir()
    document = docx.Document()
    p = document.add_paragraph("The ")
    _revision(p, "del", 11, "teh")
    _revision(p, "ins", 12, "cat")
    p.add_run(" reads ")
    title = p.add_run("Atlas")
    title.italic = True
    change = etree.SubElement(title._r.get_or_add_rPr(), qn("w:rPrChange"), {
        qn("w:id"): "13", qn("w:author"): "Prior editor",
    })
    etree.SubElement(change, qn("w:rPr"))
    p.add_run(".")
    _anchor(p, "7")
    p = document.add_paragraph()
    _revision(p, "del", 21, "Sadie")
    _revision(p, "ins", 22, "Saide")
    p.add_run(" came home.")
    _anchor(p, "8")
    p = document.add_paragraph("A plain ")
    p.add_run("eror").italic = True
    p.add_run(" survives here.")
    document.add_paragraph("An untouched paragraph remains exactly as written.")
    document.save(path / "book.docx")
    package = DocxPackage(path / "book.docx")
    comments = etree.Element(qn("w:comments"))
    for cid, question in (("7", "Is Atlas a title?"), ("8", "Is this return intentional?")):
        comment = etree.SubElement(comments, qn("w:comment"), {
            qn("w:id"): cid, qn("w:author"): "Prior editor",
        })
        p = etree.SubElement(comment, qn("w:p"))
        r = etree.SubElement(p, qn("w:r"))
        etree.SubElement(r, qn("w:t")).text = question
    package.add_part("word/comments.xml", comments)
    package.save(path / "book.docx")
    for filename in ar.REQUIRED_ARTIFACTS:
        payload = {"findings": [{"finding_id": "old-1", "status": "applied"}]} if filename == "findings.json" else {}
        (path / filename).write_text(json.dumps(payload))
    (path / "VOICE.md").write_text("Preserve intentional repetition and established character names.")
    return path


def _review(packet):
    return {
        "schema_version": 1, "packet_sha256": packet["packet_sha256"],
        "editorial_verdict": "ready", "verdict_reason": "The exact repairs below resolve the editorial defects.",
        "coverage": {
            **{key: packet[key] for key in (
                "source_sha256", "accepted_sha256", "revision_sha256", "comment_sha256", "findings_sha256", "issue_sha256")},
            **packet["counts"], "full_manuscript_read": True,
        },
        "revision_review": {"all_reviewed": True, "default_action": "keep", "exceptions": []},
        "finding_review": {"all_reviewed": True, "default_action": "prior_disposition_stands", "exceptions": []},
        "issue_decisions": [{"issue_id": issue["id"], "action": "drop",
                             "reason": "This fixture records no unresolved editorial issue."}
                            for issue in packet["issue_index"]],
        "comment_decisions": [
            {"comment_id": comment["id"], "action": "retain_author_question",
             "reason": "The context does not settle the substantive question."}
            for comment in packet["comments"]
        ],
        "actions": [],
    }


def _action(kind, para_id, *, quote="", replacement="", revision_ids=(), comment_ids=()):
    return {
        "id": "repair-1", "kind": kind, "para_id": para_id, "quote": quote,
        "replacement": replacement, "revision_ids": list(revision_ids),
        "comment_ids": list(comment_ids), "finding_ids": [], "issue_ids": [],
        "reason": "The full book supplies a clear and exact correction.",
    }


def _freeze(run, configure, *, context_paths=()):
    packet = ar.build_packet(run, context_paths=context_paths)
    review = _review(packet)
    configure(review, packet)
    ar.validate_review(review, packet)
    receipt = ar._receipt_result(
        review, packet, response_id="resp-offline-fixture",
        inputs={"docx_path": str(run / "book.docx"), "context_paths": [str(p) for p in context_paths]},
    )
    ar._atomic(run / ar.PACKET_FILE, packet)
    ar._atomic(run / ar.RECEIPT_FILE, receipt)
    return packet, receipt


def _parts(path):
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        return {name: archive.read(name) for name in archive.namelist()}


def _paragraphs(path, *, reject=False):
    """Independent small-fixture OOXML reader, including source italic properties."""
    root = etree.fromstring(_parts(path)["word/document.xml"])
    result = []
    for p in root.find(qn("w:body")).findall(qn("w:p")):
        chars = []
        for t in p.iter(qn("w:t"), qn("w:delText")):
            kinds = {etree.QName(a).localname for a in t.iterancestors()}
            if kinds & ({"ins", "moveTo"} if reject else {"del", "moveFrom"}):
                continue
            run = next(a for a in t.iterancestors() if a.tag == qn("w:r"))
            props = run.find(qn("w:rPr"))
            old = None if props is None else props.find(qn("w:rPrChange"))
            if reject and old is not None:
                props = old.find(qn("w:rPr"))
            italic = props is not None and props.find(qn("w:i")) is not None
            chars.extend((ch, italic) for ch in (t.text or ""))
        result.append(chars)
    return result


def _texts(path, *, reject=False):
    return ["".join(char for char, _ in paragraph) for paragraph in _paragraphs(path, reject=reject)]


def _drop(review, packet):
    review["comment_decisions"][0]["action"] = "drop"
    review["actions"] = [_action("remove_comment", "body-0000", comment_ids=["7"])]


def _edit(review, packet):
    review["actions"] = [_action("edit_text", "body-0002", quote="eror", replacement="error")]


def test_drop_exact_comment_preserves_other_comment_revisions_and_all_text(run):
    from galley.astra_reconcile import reconcile_run

    frozen, _ = _freeze(run, _drop)
    before = _parts(run / "book.docx")
    original_text = _paragraphs(run / "book.docx")
    original_source = _paragraphs(run / "book.docx", reject=True)
    result = reconcile_run(run)
    after = ar.build_packet(run)
    assert result["delivery_ready"] and not result["repair_required"]
    assert [comment["id"] for comment in after["comments"]] == ["8"]
    assert after["comments"][0] == frozen["comments"][1]
    assert after["revisions"] == frozen["revisions"]
    assert _paragraphs(run / "book.docx") == original_text
    assert _paragraphs(run / "book.docx", reject=True) == original_source
    root = etree.fromstring(_parts(run / "book.docx")["word/document.xml"])
    for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
        assert [e.get(qn("w:id")) for e in root.iter(qn("w:" + tag))] == ["8"]
    after_parts = _parts(run / "book.docx")
    assert set(after_parts) == set(before)
    assert {k: v for k, v in after_parts.items() if k not in {"word/document.xml", "word/comments.xml"}} == {
        k: v for k, v in before.items() if k not in {"word/document.xml", "word/comments.xml"}
    }


def test_text_repair_is_tracked_preserving_source_formatting_and_prior_revisions(run):
    from galley.astra_reconcile import reconcile_run

    frozen, _ = _freeze(run, _edit)
    before_parts = _parts(run / "book.docx")
    source = _paragraphs(run / "book.docx", reject=True)
    result = reconcile_run(run)
    current = ar.build_packet(run)
    assert result["delivery_ready"]
    assert _texts(run / "book.docx") == [
        "The cat reads Atlas.", "Saide came home.", "A plain error survives here.",
        "An untouched paragraph remains exactly as written.",
    ]
    assert _paragraphs(run / "book.docx", reject=True) == source
    assert current["source_sha256"] == frozen["source_sha256"]
    assert current["comments"] == frozen["comments"]
    # XML paths and positional review IDs can shift; preserved revision payloads must not.
    prior = {r["word_id"]: r["data"] for r in frozen["revisions"]}
    final = {r["word_id"]: r["data"] for r in current["revisions"]}
    assert all(final[key] == value for key, value in prior.items())
    new = [r for r in current["revisions"] if r["word_id"] not in prior]
    assert {r["kind"] for r in new} == {"del", "ins"}
    assert all(italic for _, italic in _paragraphs(run / "book.docx")[2][8:13])
    assert _parts(run / "book.docx")["word/comments.xml"] == before_parts["word/comments.xml"]


def test_reverting_harmful_name_revision_restores_source_without_rejecting_other_edits(run):
    from galley.astra_reconcile import reconcile_run

    def revert(review, packet):
        ids = [r["id"] for r in packet["revisions"] if r["word_id"] in {"21", "22"}]
        review["revision_review"]["exceptions"] = [
            {"revision_id": rid, "action": "revert", "reason": "Sadie is the established name."} for rid in ids
        ]
        review["actions"] = [_action("revert_revision", "body-0001", revision_ids=ids)]

    frozen, _ = _freeze(run, revert)
    source = _paragraphs(run / "book.docx", reject=True)
    result = reconcile_run(run)
    current = ar.build_packet(run)
    assert result["delivery_ready"]
    assert _texts(run / "book.docx")[:2] == ["The cat reads Atlas.", "Sadie came home."]
    assert _paragraphs(run / "book.docx", reject=True) == source
    assert [r["data"] for r in current["revisions"]] == [
        r["data"] for r in frozen["revisions"] if r["word_id"] not in {"21", "22"}
    ]


def test_reconcile_resume_is_idempotent_and_keeps_review_evidence_immutable(run, monkeypatch):
    from galley.astra_reconcile import SOURCE_FILE, reconcile_run, validate_reconciliation

    frozen, receipt = _freeze(run, _edit)
    evidence = {name: (run / name).read_bytes() for name in (ar.PACKET_FILE, ar.RECEIPT_FILE)}
    original_bytes = (run / "book.docx").read_bytes()
    def forbidden(*args, **kwargs):
        raise AssertionError("Reconciliation must not call the reviewer")
    monkeypatch.setattr(ar, "review_run", forbidden)
    first = reconcile_run(run)
    assert (run / SOURCE_FILE).read_bytes() == original_bytes
    corrected_bytes = (run / "book.docx").read_bytes()
    current = ar.build_packet(run)
    assert reconcile_run(run) == first
    assert validate_reconciliation(run, receipt, frozen, current) == first
    assert ar.validate_receipt(run) == first
    assert (run / "book.docx").read_bytes() == corrected_bytes
    assert {name: (run / name).read_bytes() for name in evidence} == evidence
    assert (run / SOURCE_FILE).read_bytes() == original_bytes


@pytest.mark.parametrize("target", ["document", "artifact", "context"])
def test_unreviewed_changes_after_reconciliation_block_handoff(run, target):
    from galley.astra_reconcile import reconcile_run, validate_reconciliation

    frozen, receipt = _freeze(run, _drop)
    reconcile_run(run)
    if target == "document":
        package = DocxPackage(run / "book.docx")
        root = package.tree("word/document.xml")
        next(root.iter(qn("w:t"))).text = "An unauthorized alteration "
        package.mark_modified("word/document.xml")
        package.save(run / "book.docx")
    elif target == "artifact":
        (run / "findings.json").write_text('{"findings": []}')
    else:
        (run / "VOICE.md").write_text("Reverse the previously reviewed style decision.")
    with pytest.raises(ar.AstraReviewError):
        validate_reconciliation(run, receipt, frozen, ar.build_packet(run))
    with pytest.raises(ar.AstraReviewError):
        ar.validate_receipt(run)


def test_explicit_external_context_remains_bound_after_reconciliation(run, tmp_path):
    from galley.astra_reconcile import reconcile_run

    context = tmp_path / "author-ruling.md"
    context.write_text("Atlas is a published title.")
    _freeze(run, _drop, context_paths=[context])
    reconcile_run(run)
    assert ar.validate_receipt(run)["delivery_ready"]
    context.write_text("Atlas is a character, not a title.")
    with pytest.raises(ar.AstraReviewError):
        ar.validate_receipt(run)


def test_unsafe_text_edit_intersecting_prior_revision_is_atomic(run):
    from galley.astra_reconcile import reconcile_run

    def unsafe(review, packet):
        safe = _action("edit_text", "body-0002", quote="eror", replacement="error")
        overlapping = _action("edit_text", "body-0000", quote="cat reads", replacement="dog reads")
        overlapping["id"] = "repair-2"
        review["actions"] = [safe, overlapping]
    _freeze(run, unsafe)
    before = {p.name: p.read_bytes() for p in run.iterdir() if p.is_file()}
    with pytest.raises(ar.AstraReviewError):
        reconcile_run(run)
    assert {name: (run / name).read_bytes() for name in before} == before
    assert not ar.validate_receipt(run)["delivery_ready"]


@pytest.mark.parametrize("target", ["review", "packet", "source"])
def test_mutated_immutable_evidence_cannot_authorize_a_reconciled_document(run, target):
    from galley.astra_reconcile import SOURCE_FILE, reconcile_run

    _freeze(run, _edit)
    reconcile_run(run)
    if target == "review":
        receipt = json.loads((run / ar.RECEIPT_FILE).read_text())
        receipt["review"]["actions"][0]["replacement"] = "a different unreviewed word"
        ar._atomic(run / ar.RECEIPT_FILE, receipt)
    elif target == "packet":
        packet = json.loads((run / ar.PACKET_FILE).read_text())
        packet["accepted_paragraphs"][0]["text"] = "Unreviewed evidence"
        ar._atomic(run / ar.PACKET_FILE, packet)
    else:
        (run / SOURCE_FILE).write_bytes((run / "book.docx").read_bytes())
    with pytest.raises(ar.AstraReviewError):
        ar.validate_receipt(run)


@pytest.mark.parametrize("kind", ["replace_comment", "add_author_query"])
def test_author_question_changes_have_exact_text_and_valid_surviving_anchors(run, kind):
    from galley.astra_reconcile import reconcile_run

    question = "Does the return refer to the earlier visit or to a second arrival?"
    def configure(review, packet):
        if kind == "replace_comment":
            review["comment_decisions"][1]["action"] = "replace_question"
            review["actions"] = [_action(kind, "body-0001", replacement=question, comment_ids=["8"])]
        else:
            review["actions"] = [_action(kind, "body-0003", quote="untouched paragraph", replacement=question)]
    frozen, _ = _freeze(run, configure)
    before_source = _paragraphs(run / "book.docx", reject=True)
    result = reconcile_run(run)
    current = ar.build_packet(run)
    assert result["delivery_ready"]
    assert current["accepted_paragraphs"] == frozen["accepted_paragraphs"]
    assert current["revisions"] == frozen["revisions"]
    assert _paragraphs(run / "book.docx", reject=True) == before_source
    assert current["comments"][0] == frozen["comments"][0]
    if kind == "replace_comment":
        comment = current["comments"][1]
        assert len(current["comments"]) == 2
        assert comment["anchors"] == frozen["comments"][1]["anchors"]
    else:
        assert current["comments"][1] == frozen["comments"][1]
        assert len(current["comments"]) == 3
        comment = current["comments"][2]
        assert {anchor["para_id"] for anchor in comment["anchors"]} == {"body-0003"}
        assert {anchor["kind"] for anchor in comment["anchors"]} == {
            "commentRangeStart", "commentRangeEnd", "commentReference",
        }
    assert comment["text"] == question
    assert not current["coverage_issues"]


@pytest.mark.parametrize("crash_after_replace", [False, True])
def test_crash_on_either_side_of_atomic_document_replace_can_resume(run, monkeypatch, crash_after_replace):
    from galley import astra_reconcile as rec

    _freeze(run, _edit)
    original_replace = rec.os.replace
    original_atomic = rec._atomic
    if crash_after_replace:
        def interrupted_atomic(path, data):
            if path.name == rec.RECONCILIATION_FILE and data.get("status") == "completed":
                raise RuntimeError("Simulated crash after manuscript replacement")
            original_atomic(path, data)
        monkeypatch.setattr(rec, "_atomic", interrupted_atomic)
    else:
        def interrupted_replace(source, destination):
            if str(destination).endswith("book.docx"):
                raise RuntimeError("Simulated crash before manuscript replacement")
            return original_replace(source, destination)
        monkeypatch.setattr(rec.os, "replace", interrupted_replace)
    with pytest.raises(RuntimeError, match="Simulated crash"):
        rec.reconcile_run(run)
    assert json.loads((run / rec.RECONCILIATION_FILE).read_text())["status"] == "prepared"
    assert _texts(run / "book.docx")[2] == (
        "A plain error survives here." if crash_after_replace else "A plain eror survives here.")
    monkeypatch.setattr(rec.os, "replace", original_replace)
    monkeypatch.setattr(rec, "_atomic", original_atomic)
    result = rec.reconcile_run(run)
    assert result["delivery_ready"] and not result["repair_required"]
    assert _texts(run / "book.docx")[2] == "A plain error survives here."
    assert ar.validate_receipt(run) == result
