"""Subscription reviews prove actual bounded coverage, with no live model calls."""
import copy
import json

import docx
import pytest
from lxml import etree

from docproof.utils.xml_helpers import DocxPackage, qn
from galley import astra_review as ar
from galley import astra_subscription as sub


@pytest.fixture
def run(tmp_path):
    path = tmp_path / "run"
    path.mkdir()
    book = docx.Document()
    for i in range(24):
        book.add_paragraph(f"Chapter {i}. Sadie visits the current. " + "The river shimmers. " * 100)
    book.save(path / "book.docx")
    for name in ar.REQUIRED_ARTIFACTS:
        value = {"findings": [{"para_id": "body-0001", "state": "dropped", "original_text": "Sadie", "reason": "The name is correct"}]} if name == "findings.json" else {}
        (path / name).write_text(json.dumps(value))
    (path / "DECISIONS.md").write_text("Sadie is a different character from Saide. The magic path is called Straight.")
    return path


@pytest.fixture
def cross_book_run(tmp_path):
    """Match the live context-review failure: two revisions and comment 0."""
    path = tmp_path / "cross-book"
    path.mkdir()
    book = docx.Document()
    book.add_paragraph("Sadie and Saide are different people.")
    paragraph = book.add_paragraph("The ")
    for kind, text in (("del", "curent"), ("ins", "current")):
        change = etree.SubElement(paragraph._p, qn("w:" + kind), {qn("w:id"): "1" if kind == "del" else "2"})
        run_node = etree.SubElement(change, qn("w:r"))
        etree.SubElement(run_node, qn("w:delText" if kind == "del" else "w:t")).text = text
    paragraph.add_run(" flows north.")
    book.add_paragraph("The current current, if you will.")
    anchor = book.add_paragraph("The records distinguish Sadie from Saide.")
    anchor._p.insert(0, etree.Element(qn("w:commentRangeStart"), {qn("w:id"): "0"}))
    etree.SubElement(anchor._p, qn("w:commentRangeEnd"), {qn("w:id"): "0"})
    etree.SubElement(etree.SubElement(anchor._p, qn("w:r")), qn("w:commentReference"), {qn("w:id"): "0"})
    book.add_paragraph("Sadie returns to the river.")
    book.save(path / "book.docx")
    package = DocxPackage(path / "book.docx")
    # This citation fixture needs no unused factory style definitions.
    styles = package.tree("word/styles.xml")
    for child in list(styles):
        styles.remove(child)
    package.mark_modified("word/styles.xml")
    comments = etree.Element(qn("w:comments"))
    comment = etree.SubElement(comments, qn("w:comment"), {qn("w:id"): "0"})
    text = etree.SubElement(etree.SubElement(etree.SubElement(comment, qn("w:p")), qn("w:r")), qn("w:t"))
    text.text = "Is Sadie the same person as Saide?"
    package.add_part("word/comments.xml", comments)
    package.save(path / "book.docx")
    for name in ar.REQUIRED_ARTIFACTS:
        data = {"findings": [{"para_id": "body-0004", "state": "dropped", "original_text": "Sadie"}]} if name == "findings.json" else {}
        (path / name).write_text(json.dumps(data))
    return path


def cross_book_context(plan):
    chunk = next(c for c in plan["chunks"] if c["phase"] == "context")
    review = chunk_good(chunk)
    assert chunk["owned_ids"]["supplemental"] == ["evidence-000001"]
    review["investigations"] = [{
        "question": "Check the two tracked revisions and whether the name comment is answered elsewhere.",
        "evidence_ids": ["evidence-000001", "revision-000001", "revision-000002", "0"],
        "para_ids": ["body-0000", "body-0001", "body-0003"]}]
    return chunk, review


def chunk_good(chunk):
    return {"chunk_id": chunk["chunk_id"], "chunk_sha256": chunk["chunk_sha256"],
            "all_evidence_reviewed": True, "reviewed_ids": copy.deepcopy(chunk["owned_ids"]),
            "revision_review": {"all_reviewed": True, "default_action": "keep", "exceptions": []},
            "finding_review": {"all_reviewed": True, "default_action": "prior_disposition_stands", "exceptions": []},
            "comment_decisions": [{"comment_id": cid, "action": "retain_author_question", "reason": "No unique intended answer is supported."} for cid in chunk["owned_ids"]["comments"]],
            "issue_decisions": [{"issue_id": iid, "action": "drop", "reason": "The flagged form is already corrected."} for iid in chunk["owned_ids"]["issues"]],
            "actions": [], "guide_notes": [], "investigations": []}


def final_good(packet):
    return {"schema_version": 1, "packet_sha256": packet["packet_sha256"],
            "editorial_verdict": "ready", "verdict_reason": "All recorded review findings reconcile with the source book.",
            "coverage": {**{k: packet[k] for k in ("source_sha256", "accepted_sha256", "revision_sha256", "comment_sha256", "findings_sha256", "issue_sha256")}, **packet["counts"], "full_manuscript_read": True},
            "revision_review": {"all_reviewed": True, "default_action": "keep", "exceptions": []},
            "finding_review": {"all_reviewed": True, "default_action": "prior_disposition_stands", "exceptions": []},
            "comment_decisions": [], "issue_decisions": [], "actions": []}


class Runner:
    def __init__(self, packet, fail_at=None, mutate=None):
        self.packet = packet
        self.calls = []
        self.fail_at = fail_at
        self.mutate = mutate

    def __call__(self, prompt, schema, **kwargs):
        self.calls.append((prompt, schema, kwargs))
        if self.fail_at == len(self.calls):
            raise RuntimeError("subscription allowance unavailable")
        evidence = json.loads(prompt.split("\nEVIDENCE\n", 1)[1])
        if "chunk" in evidence:
            result = chunk_good(evidence["chunk"])
        else:
            if "complete_final_input_path" in evidence:
                from pathlib import Path
                evidence = ar._load(Path(evidence["complete_final_input_path"]))
            result = {"review": final_good(self.packet),
                      "reviewed_chunks": [{"chunk_id": c["chunk_id"], "review_sha256": c["review_sha256"]}
                                          for c in evidence["coverage_manifest"]["chunks"]],
                      "investigation_resolutions": [{"investigation_id": i["investigation_id"],
                                                     "resolution": "The full source resolves the issue.", "para_ids": i["para_ids"]}
                                                    for i in evidence["investigations"]]}
        if self.mutate:
            self.mutate(result)
        return result


def test_plan_complete_deterministic_and_bounded(run):
    packet = ar.build_packet(run)
    plan = sub.plan_review(packet, max_chunk_bytes=40_000)
    assert plan == sub.plan_review(packet, max_chunk_bytes=40_000)
    assert len(plan["chunks"]) > 2
    assert all(sub._bytes(c) <= 24_000 for c in plan["chunks"])
    ids = [pid for c in plan["chunks"] for pid in c["owned_ids"]["paragraphs"]]
    assert len(ids) == len(set(ids)) == 24
    assert set(ids) == {p["id"] for p in packet["accepted_paragraphs"]}
    assert [fid for c in plan["chunks"] for fid in c["owned_ids"]["findings"]] == packet["finding_ids"]
    assert all(c["phase"] == "context" for c in plan["chunks"][:next(i for i,c in enumerate(plan["chunks"]) if c["phase"] == "manuscript")])
    assert any(c["supporting_paragraphs"] for c in plan["chunks"] if c["phase"] == "manuscript")
    assert "Sadie is a different character" in ar._json(plan)


def test_source_fragments_preserve_long_strings_and_metadata():
    text = "🦋" * 20_000
    fragments = list(sub._fragment({"notes": text, "empty": [], "metadata": {"x": 1}}, [], 8000))
    text_fragments = [r for r in fragments if r["path"] == ["notes"]]
    assert "".join(r["value"] for r in text_fragments) == text
    assert text_fragments[0]["text_start"] == 0
    assert text_fragments[-1]["text_end"] == len(text)
    assert all(sub._bytes(r) <= 8050 for r in fragments)
    assert any(r["value"] == [] for r in fragments)


def test_completed_review_has_verified_source_and_chunk_coverage_and_caches(run):
    packet = ar.build_packet(run)
    runner = Runner(packet)
    result = sub.review_run(run, runner=runner)
    assert result["transport"] == "codex" and result["delivery_ready"]
    assert result["actual_cost_usd"] is None
    assert result["coverage_manifest"]["counts"]["paragraphs"] == 24
    assert (run / ar.SOURCE_FILE).read_bytes() == (run / "book.docx").read_bytes()
    assert sub.validate_coverage_receipt(run, result, packet) == result
    calls = len(runner.calls)
    assert sub.review_run(run, runner=runner) == result
    assert len(runner.calls) == calls
    assert calls == result["request_count"]
    assert all("UNTRUSTED EVIDENCE" in prompt for prompt, _, _ in runner.calls)
    assert runner.calls[-1][1] == sub.FINAL_SCHEMA
    assert all(schema == sub.CHUNK_SCHEMA for _,schema,_ in runner.calls[:-1])


def test_failure_resumes_only_missing_chunks_and_is_not_editorial(run):
    packet = ar.build_packet(run)
    runner = Runner(packet, fail_at=2)
    with pytest.raises(ar.AstraReviewError, match="operational"):
        sub.review_run(run, runner=runner)
    failure = ar._load(run / ar.RECEIPT_FILE)
    assert failure["status"] == "operational_failure"
    assert "review" not in failure and "editorial_verdict" not in failure
    first_id = runner.calls[0][2]["request_id"]
    runner.fail_at = None
    sub.review_run(run, runner=runner)
    assert sum(call[2]["request_id"] == first_id for call in runner.calls) == 1


@pytest.mark.parametrize("kind", ["paragraphs", "supplemental", "findings"])
def test_missing_read_ids_fail_operationally(run, kind):
    packet = ar.build_packet(run)
    def mutate(review):
        if "reviewed_ids" in review and review["reviewed_ids"][kind]:
            review["reviewed_ids"][kind].pop()
    runner = Runner(packet, mutate=mutate)
    with pytest.raises(ar.AstraReviewError, match="coverage"):
        sub.review_run(run, runner=runner)
    assert ar._load(run / ar.RECEIPT_FILE)["status"] == "operational_failure"


def test_modified_cached_review_invalidates_complete_receipt(run):
    packet = ar.build_packet(run)
    result = sub.review_run(run, runner=Runner(packet))
    plan = ar._load(run / sub.DIRECTORY / "plan.json")
    path = run / sub.DIRECTORY / (plan["chunks"][0]["chunk_id"] + "-review.json")
    edited = ar._load(path)
    edited["all_evidence_reviewed"] = False
    ar._atomic(path, edited)
    with pytest.raises(ar.AstraReviewError, match="coverage"):
        sub.validate_coverage_receipt(run, result, packet)
    with pytest.raises(ar.AstraReviewError, match="coverage"):
        ar.validate_receipt(run)


def test_existing_api_receipt_never_falls_back_to_subscription(run):
    ar._atomic(run / ar.RECEIPT_FILE, {"status": "pending", "response_id": "resp-test"})
    runner = Runner(ar.build_packet(run))
    with pytest.raises(ar.AstraReviewError, match="different transport"):
        sub.review_run(run, runner=runner)
    assert not runner.calls


def test_pending_evidence_or_chunk_changes_rejected(run):
    packet = ar.build_packet(run)
    with pytest.raises(ar.AstraReviewError):
        sub.review_run(run, runner=Runner(packet, fail_at=1))
    with pytest.raises(ar.AstraReviewError, match="chunk planning"):
        sub.review_run(run, runner=Runner(packet), max_chunk_bytes=40_000)
    (run / "DECISIONS.md").write_text("Changed after review started")
    with pytest.raises(ar.AstraReviewError, match="evidence changed"):
        sub.review_run(run, runner=Runner(packet))


def test_missing_production_artifacts_rejected_before_runner(run):
    (run / "settlement.json").unlink()
    runner = Runner(None)
    with pytest.raises(ar.AstraReviewError, match="Missing required"):
        sub.review_run(run, runner=runner)
    assert not runner.calls


def test_large_final_aggregation_uses_complete_file_manifest(run, monkeypatch):
    packet = ar.build_packet(run)
    runner = Runner(packet)
    original = sub._final_input
    def huge(*args, **kwargs):
        data = original(*args, **kwargs)
        data["extra"] = "a" * 200_000
        return data
    monkeypatch.setattr(sub, "_final_input", huge)
    result = sub.review_run(run, runner=runner)
    assert result["delivery_ready"]
    evidence = json.loads(runner.calls[-1][0].split("\nEVIDENCE\n", 1)[1])
    assert "complete_final_input_path" in evidence
    assert len(ar._load(run / sub.DIRECTORY / "final-input.json")["extra"]) == 200_000
    assert len(runner.calls[-1][0].encode("utf-8")) < sub.DEFAULT_MAX_CHUNK_BYTES
    assert "chunk_review_files" not in evidence
    assert "single authoritative representation" in evidence["file_read_requirement"]
    final_input = ar._load(run / sub.DIRECTORY / "final-input.json")
    assert len(final_input["chunk_reviews"]) == len(runner.calls) - 1
    for review, bound in zip(final_input["chunk_reviews"], final_input["coverage_manifest"]["chunks"]):
        assert ar._hash(review) == bound["review_sha256"]


def test_final_must_acknowledge_every_chunk(run):
    def mutate(review):
        if "reviewed_chunks" in review:
            review["reviewed_chunks"].pop()
    with pytest.raises(ar.AstraReviewError, match="acknowledge every"):
        sub.review_run(run, runner=Runner(ar.build_packet(run), mutate=mutate))


def test_final_must_resolve_each_cross_book_investigation(run):
    def mutate(review):
        if "reviewed_ids" in review and review["reviewed_ids"]["paragraphs"]:
            pid = review["reviewed_ids"]["paragraphs"][0]
            review["investigations"] = [{"question": "Is Sadie a separate character?", "evidence_ids": [pid], "para_ids": [pid]}]
        elif "investigation_resolutions" in review:
            review["investigation_resolutions"] = []
    with pytest.raises(ar.AstraReviewError, match="every cross-book investigation"):
        sub.review_run(run, runner=Runner(ar.build_packet(run), mutate=mutate))


def test_context_investigation_accepts_real_cross_book_revision_and_comment_citations(cross_book_run):
    packet = ar.build_packet(cross_book_run)
    plan = sub.plan_review(packet)
    context, review = cross_book_context(plan)
    original_hash = ar._hash(review)
    assert sub._validate_chunk(review, context, packet, known_citation_ids=sub._citation_ids(plan)) == review
    assert ar._hash(review) == original_hash
    reviews = [review if c == context else chunk_good(c) for c in plan["chunks"]]
    proof = sub._manifest(packet, plan, reviews)
    assert proof["counts"] == plan["counts"]
    assert not review["reviewed_ids"]["revisions"] and not review["reviewed_ids"]["comments"]


def test_manuscript_guide_can_cite_prior_shared_context(cross_book_run):
    packet = ar.build_packet(cross_book_run)
    plan = sub.plan_review(packet)
    chunk = next(c for c in plan["chunks"] if c["phase"] == "manuscript")
    review = chunk_good(chunk)
    review["guide_notes"] = [{"claim": "The shared context supports keeping both distinct names.",
                              "evidence_ids": ["evidence-000001"], "para_ids": ["body-0000"]}]
    assert not chunk["owned_ids"]["supplemental"]
    sub._validate_chunk(review, chunk, packet, known_citation_ids=sub._citation_ids(plan))


@pytest.mark.parametrize("field,unknown", [("evidence_ids", "revision-999999"),
    ("evidence_ids", "evidence-999999"), ("evidence_ids", "missing-comment"),
    ("evidence_ids", "body-9999"), ("para_ids", "body-9999")])
def test_cross_book_citations_still_reject_nonexistent_ids(cross_book_run, field, unknown):
    packet = ar.build_packet(cross_book_run)
    plan = sub.plan_review(packet)
    chunk, review = cross_book_context(plan)
    review["investigations"][0][field].append(unknown)
    with pytest.raises(ar.AstraReviewError, match="unknown evidence"):
        sub._validate_chunk(review, chunk, packet, known_citation_ids=sub._citation_ids(plan))


def test_cross_book_citation_does_not_expand_review_ownership(cross_book_run):
    packet = ar.build_packet(cross_book_run)
    plan = sub.plan_review(packet)
    chunk, review = cross_book_context(plan)
    review["reviewed_ids"]["revisions"] = ["revision-000001"]
    with pytest.raises(ar.AstraReviewError, match="revisions coverage"):
        sub._validate_chunk(review, chunk, packet, known_citation_ids=sub._citation_ids(plan))


@pytest.mark.parametrize("action", [
    {"kind": "remove_comment", "para_id": "body-0003", "quote": "", "replacement": "",
     "comment_ids": ["0"]},
    {"kind": "edit_text", "para_id": "body-0002", "quote": "current current", "replacement": "current"},
])
def test_cross_book_citation_does_not_authorize_unowned_actions(cross_book_run, action):
    packet = ar.build_packet(cross_book_run)
    plan = sub.plan_review(packet)
    chunk, review = cross_book_context(plan)
    review["actions"] = [{"id": chunk["chunk_id"] + "-1", "reason": "A proposed change needs its owning review.",
                          **{k: [] for k in ("revision_ids", "comment_ids", "finding_ids", "issue_ids")}, **action}]
    with pytest.raises(ar.AstraReviewError, match="unknown evidence IDs|owning chunk"):
        sub._validate_chunk(review, chunk, packet, known_citation_ids=sub._citation_ids(plan))


def test_previously_completed_cross_book_response_resumes_without_regeneration(cross_book_run, monkeypatch):
    packet = ar.build_packet(cross_book_run)
    plan = sub.plan_review(packet)
    context, context_review = cross_book_context(plan)
    def mutate(result):
        if result.get("chunk_id") == context["chunk_id"]:
            result.update(copy.deepcopy(context_review))
        elif "review" in result:
            result["review"]["comment_decisions"] = [{"comment_id": "0", "action": "retain_author_question",
                                                      "reason": "Final test disposition."}]
    generated = Runner(packet, mutate=mutate)
    cache = {}
    def cached_runner(prompt, schema, **kwargs):
        key = kwargs["request_id"]
        if key not in cache:
            cache[key] = (prompt, copy.deepcopy(schema), generated(prompt, schema, **kwargs))
        old_prompt, old_schema, result = cache[key]
        assert prompt == old_prompt and schema == old_schema
        return copy.deepcopy(result)
    def previous_validator(*args, **kwargs):
        raise ar.AstraReviewError("Subscription guide/investigation cites unreviewed evidence")
    with monkeypatch.context() as previous:
        previous.setattr(sub, "_validate_chunk", previous_validator)
        with pytest.raises(ar.AstraReviewError, match="unreviewed evidence"):
            sub.review_run(cross_book_run, runner=cached_runner)
    assert len(generated.calls) == 1
    result = sub.review_run(cross_book_run, runner=cached_runner)
    assert result["delivery_ready"]
    first_request = generated.calls[0][2]["request_id"]
    assert sum(call[2]["request_id"] == first_request for call in generated.calls) == 1
    assert len(generated.calls) == len(plan["chunks"]) + 1
    assert sub.validate_coverage_receipt(cross_book_run, result, packet) == result
