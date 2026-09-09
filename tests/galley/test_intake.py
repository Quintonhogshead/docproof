"""Agent intake uses real Word formatting and correction writers, fake readers."""
from __future__ import annotations

import base64
import json
import re
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Pt

from docproof.audit import run_audit
from docproof.cleancopy import has_markup
from docproof.ingest import build_document_model, preflight
from docproof.models import Anchor, Finding
from docproof.providers import NormalizedUsage, ProviderResult
from docproof.reassembler import apply_tracked_changes, paragraph_view_text
from docproof.utils.xml_helpers import DocxPackage, qn, walk_package
from galley import agent, driver, intake
from galley.manifest import sha256_file
from galley.state_machine import RunStateMachine
from galley.verify import build_fingerprints


class StyleReader:
    def __init__(self):
        self.calls = []
        self.fail_at = None

    def complete_structured(self, **kwargs):
        self.calls.append(kwargs)
        if self.fail_at == len(self.calls):
            raise RuntimeError("reader interrupted")
        rows = []
        for match in re.finditer(r'<p id="([^"]+)"[^>]*>(.*?)</p>|<blank id="([^"]+)"/>',
                                 kwargs["user"], re.S):
            pid, text, blank_id = match.groups()
            role = ("spacing" if blank_id else
                    "chapter # / title" if text.startswith("Chapter ") else
                    "scene break" if text == "***" else "body")
            rows.append({"para_id": pid or blank_id, "role": role, "flag": ""})
        return ProviderResult(parsed={"paragraphs": rows},
                              usage=NormalizedUsage(input_tokens=100,
                                                    output_tokens=20, billed=False))


@pytest.fixture
def reader(monkeypatch):
    fake = StyleReader()
    monkeypatch.setattr(intake, "build_provider", lambda cfg: fake)
    return fake


@pytest.fixture
def book(tmp_path):
    doc = Document()
    doc.add_paragraph("Chapter One", "Heading 1")
    doc.add_paragraph()  # Removing this changes positional ids before proofing.
    p = doc.add_paragraph()
    r = p.add_run("Teh door was ")
    r.font.name, r.font.size = "Arial", Pt(20)
    p.add_run("open").italic = True
    p.add_run(".")
    doc.add_paragraph("She waited outside.")
    doc.add_paragraph("***")
    doc.add_paragraph("Nobody answered.")
    doc.add_picture(BytesIO(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGOoqKgA"
        "AALUAWneSzH7AAAAAElFTkSuQmCC")))
    path = tmp_path / "Smith - Book 1.docx"
    doc.save(path)
    return path


def test_format_preserves_original_content_emphasis_and_plain_style(book, tmp_path, reader):
    before = book.read_bytes()
    ws = tmp_path / "ws"
    formatted = intake.format_for_proof(book, ws)
    assert formatted.name == book.name
    assert book.read_bytes() == before
    assert (formatted.parents[1] / "original" / book.name).read_bytes() == before
    doc = Document(formatted)
    body = next(p for p in doc.paragraphs if p.text.startswith("Teh"))
    assert body.style.font.name == "Times New Roman"
    assert body.style.font.size.pt == 12
    assert body.runs[0].font.name is None and body.runs[0].font.size is None
    assert next(r for r in body.runs if r.text == "open").italic
    assert doc.paragraphs[0].style.font.size.pt == 14
    assert not any(p.text == "" for p in doc.paragraphs if not p._p.findall(".//" + qn("w:drawing")))
    assert len(doc.inline_shapes) == 1
    receipt = json.loads((formatted.parents[1] / intake.RECEIPT_NAME).read_text())
    assert receipt["verified"] is True
    assert receipt["original_sha256"] == sha256_file(book)
    assert receipt["formatted_sha256"] == sha256_file(formatted)
    assert all(c["ok"] for c in receipt["verifications"])
    assert receipt["lane"] == "subagent"
    assert reader.calls


def test_complete_intake_reuses_exact_bytes_without_reading_again(book, tmp_path, reader, monkeypatch):
    ws = tmp_path / "ws"
    formatted = intake.format_for_proof(book, ws)
    saved = formatted.read_bytes()
    monkeypatch.setattr(intake, "configuration", lambda: pytest.fail("reformatted completed intake"))
    assert intake.format_for_proof(book, ws) == formatted
    assert formatted.read_bytes() == saved


@pytest.mark.parametrize("target", ["formatted", "original", "receipt", "missing"])
def test_damaged_completed_intake_cannot_start_a_different_proofread(book, tmp_path, reader, target):
    ws = tmp_path / "ws"
    formatted = intake.format_for_proof(book, ws)
    directory = formatted.parents[1]
    if target == "formatted":
        formatted.write_bytes(b"changed")
    elif target == "original":
        (directory / "original" / book.name).write_bytes(b"changed")
    elif target == "receipt":
        (directory / intake.RECEIPT_NAME).write_text("{}")
    else:
        formatted.unlink()
    with pytest.raises(intake.IntakeError):
        intake.format_for_proof(book, ws)
    assert len(reader.calls) == 1


def test_interrupted_formatting_resumes_completed_windows(book, tmp_path, reader, monkeypatch):
    cfg = intake.configuration()
    cfg.prep.max_paragraphs_per_window = 2
    monkeypatch.setattr(intake, "configuration", lambda: cfg)
    reader.fail_at = 2
    ws = tmp_path / "ws"
    with pytest.raises(RuntimeError, match="interrupted"):
        intake.format_for_proof(book, ws)
    assert not (ws / "state.json").exists()
    assert not list(ws.rglob(intake.RECEIPT_NAME))
    first_window = reader.calls[0]["user"]
    reader.fail_at = None
    assert intake.format_for_proof(book, ws).is_file()
    assert sum(c["user"] == first_window for c in reader.calls) == 1


def test_unclassified_paragraphs_stop_intake_then_can_retry(book, tmp_path, reader, monkeypatch):
    monkeypatch.setattr(intake, "build_provider", lambda cfg: type("Silent", (), {
        "complete_structured": lambda self, **kw: ProviderResult(parsed={"paragraphs": []})})())
    ws = tmp_path / "ws"
    with pytest.raises(intake.IntakeError, match="unclassified"):
        intake.format_for_proof(book, ws)
    assert not list(ws.rglob(intake.RECEIPT_NAME))
    monkeypatch.setattr(intake, "build_provider", lambda cfg: reader)
    assert intake.format_for_proof(book, ws).is_file()


def test_lost_content_fails_before_proofing_or_delivery(book, tmp_path, reader, monkeypatch):
    from docproof.prep.writers import clean
    real_write = clean.write_clean

    def lose_paragraph(pkg, *args, **kwargs):
        stats = real_write(pkg, *args, **kwargs)
        body = pkg.tree("word/document.xml").find(qn("w:body"))
        paragraph = next(p for p in body.findall(qn("w:p"))
                         if "Teh" in "".join(p.itertext()))
        body.remove(paragraph)
        pkg.mark_modified("word/document.xml")
        return stats

    monkeypatch.setattr(clean, "write_clean", lose_paragraph)
    monkeypatch.setattr(driver.Driver, "run", lambda self: pytest.fail("proofreading started"))
    with pytest.raises(intake.prep.VerificationFailed):
        agent._run_driver(book=book, slug="smith", workspace_root=tmp_path / "ws",
                          astra_review=False, log=lambda _: None)
    assert not list((tmp_path / "ws").rglob("state.json"))
    assert not list((tmp_path / "ws").rglob(intake.RECEIPT_NAME))


def test_legacy_in_progress_source_is_not_reformatted(book, tmp_path, monkeypatch):
    ws = driver.seed_workspace(book, "smith", workspace_root=tmp_path / "ws")
    machine = RunStateMachine.load(ws / "state.json")
    machine.advance("audited", by="test")
    machine.save(ws / "state.json")
    monkeypatch.setattr(intake, "build_provider", lambda cfg: pytest.fail("reformatted legacy run"))
    assert intake.format_for_proof(book, ws) == ws / "source" / book.name
    assert RunStateMachine.load(ws / "state.json").current == "audited"


def test_changed_book_starts_profile_in_a_new_revision(book, tmp_path, reader, monkeypatch):
    root = tmp_path / "ws"
    formatted = intake.format_for_proof(book, root / "smith")
    ws = driver.seed_workspace(formatted, "smith", workspace_root=root)
    machine = RunStateMachine.load(ws / "state.json")
    machine.advance("audited", by="test")
    machine.save(ws / "state.json")
    (ws / "runs" / "old.txt").write_text("old proofread")
    doc = Document(book)
    doc.add_paragraph("An added final sentence.")
    doc.save(book)

    def begin(self):
        assert self.start_phase is None
        driver.seed_workspace(self.book, self.slug, workspace_root=self.workspace_root,
                              on_source_change=self.on_source_change)
        state = RunStateMachine.load(self.workspace / "state.json")
        assert state.revision == 2 and not state.current
        assert state.source_sha256 == sha256_file(self.book)
        assert "An added final sentence." in [p.text for p in Document(self.book).paragraphs]
        assert (self.workspace / "runs.rev1" / "old.txt").is_file()

    monkeypatch.setattr(driver.Driver, "run", begin)
    agent._run_driver(book=book, slug="smith", workspace_root=root,
                      start_phase="verify", astra_review=False, log=lambda _: None)
    assert len(list(ws.rglob(intake.RECEIPT_NAME))) == 2


def test_same_formatted_source_preserves_resume_phase(book, tmp_path, reader, monkeypatch):
    root = tmp_path / "ws"
    formatted = intake.format_for_proof(book, root / "smith")
    ws = driver.seed_workspace(formatted, "smith", workspace_root=root)
    before = (ws / "state.json").read_bytes()

    def resume(self):
        assert self.start_phase == "verify"
        assert self.book == formatted
        assert (ws / "state.json").read_bytes() == before

    monkeypatch.setattr(driver.Driver, "run", resume)
    agent._run_driver(book=book, slug="smith", workspace_root=root,
                      start_phase="verify", astra_review=False, log=lambda _: None)
    assert len(reader.calls) == 1


def test_agent_formats_then_corrects_and_exports_both_book_2_files(book, tmp_path, reader, monkeypatch):
    events = []

    def proofread(self):
        # This is the actual agent's handoff to the driver: ids and source
        # fingerprints are first recorded AFTER blank paragraphs were removed.
        ws = driver.seed_workspace(self.book, self.slug, workspace_root=self.workspace_root)
        cfg = intake.configuration()
        pkg = preflight(self.book, "abort")
        model = build_document_model(pkg, cfg)
        p = next(p for p in model.paragraphs if p.text.startswith("Teh"))
        assert p.para_id == "body-0001"
        state = RunStateMachine.load(ws / "state.json")
        assert state.source_sha256 == sha256_file(self.book) != sha256_file(book)
        baseline = {wp.para_id: paragraph_view_text(wp.element, "reject")
                    for wp in walk_package(pkg)}
        finding = Finding("f-1", "chunk-0", p.para_id, "spelling", p.text, 1,
                          "Correct the spelling.", "test", "high", status="validated",
                          anchor=Anchor(0, 3, "Teh", "The"))
        stats = apply_tracked_changes(pkg, model, [finding], cfg)
        assert stats.applied == ("f-1",) and not stats.skipped
        reviewed = ws / "deliverable" / "reviewed.docx"
        pkg.save(reviewed)
        rejected = {wp.para_id: paragraph_view_text(wp.element, "reject")
                    for wp in walk_package(DocxPackage(reviewed))}
        assert run_audit(baseline, rejected).passed
        fingerprints = build_fingerprints(ws / "deliverable")
        assert fingerprints["build_sha256"] == sha256_file(reviewed)
        outcome = ws / "deliverable" / "outcome.json"
        outcome.write_text('{"outcome":"done"}')
        return driver.build_handoff(ws, self.book.name, ws / "handoff",
                                    outcome_sources=[outcome], partial=True)

    monkeypatch.setattr(driver.Driver, "run", proofread)
    paths = agent._run_driver(book=book, slug="smith", workspace_root=tmp_path / "ws",
                              astra_review=False, log=lambda _: None, progress=events.append)
    tracked = next(p for p in paths if p.name == "Smith - Book 2.docx")
    clean = next(p for p in paths if p.name == "Smith - Book 2 - clean.docx")
    assert has_markup(tracked) and not has_markup(clean)
    assert "The door was open." in [p.text for p in Document(clean).paragraphs]
    for path in (tracked, clean):
        paragraph = Document(path).paragraphs[1]
        assert paragraph.style.font.name == "Times New Roman"
        assert paragraph.style.font.size.pt == 12
        assert len(Document(path).inline_shapes) == 1
    assert [(e["event"], e["phase"]) for e in events] == [
        ("phase_start", "formatting"), ("phase_end", "formatting")]
