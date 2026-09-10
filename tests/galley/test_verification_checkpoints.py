"""Local fake-model tests for exact verification reuse and crash recovery."""
import copy
import json

import docx
import pytest

from docproof.contract import build_envelope
from docproof.models import Usage
from docproof.providers.base import NormalizedUsage, ProviderResult
from galley import verify
from galley.settle import Settler, SettleOptions


MODEL = "gpt-5.6-luna"
POLICY = dict(engine="provider", pass_id="primary", required_pass_ids=("primary",),
              policy_id=verify.VERIFICATION_POLICY, config_sha256="a" * 64)


class Provider:
    name = "test-provider"
    effort = "high"

    def __init__(self, interrupt_at=None, body=None, billed=True):
        self.calls = []
        self.interrupt_at = interrupt_at
        self.body = body
        self.billed = billed

    def complete_structured(self, **kwargs):
        self.calls.append(kwargs)
        if len(self.calls) == self.interrupt_at:
            raise RuntimeError("interrupted model transport")
        body = self.body if self.body is not None else {kwargs["schema_name"]: []}
        return ProviderResult(parsed=copy.deepcopy(body),
                              usage=NormalizedUsage(1000, 200, billed=self.billed))


@pytest.fixture(autouse=True)
def isolated_coverage(monkeypatch):
    monkeypatch.setattr(verify, "UNREAD", [])
    monkeypatch.setattr(verify, "UNREAD_BATCHES", [])
    monkeypatch.setattr(verify, "_LOSSES", [])


@pytest.fixture
def run(tmp_path):
    path = tmp_path / "run"
    path.mkdir()
    document = docx.Document()
    document.add_paragraph("The dog sleeps.")
    document.save(path / "book.docx")
    pid = next(iter(verify.accepted_text(path)))
    (path / "findings.json").write_text(json.dumps({"findings": [{
        "para_id": pid, "original_text": "teh", "corrected_text": "the",
        "status": "validated", "error_type": "spelling"}]}))
    return path


def complete(run, provider=None, **overrides):
    provider = provider or Provider()
    options = {**POLICY, **overrides}
    model = options.pop("model", MODEL)
    uc, uw = Usage(), Usage()
    changes = verify.verify_run(run, provider, model, uc, run_walk=False, **options)
    walk = verify.verify_run(run, provider, model, uw, run_changes=False, **options)
    verify.write_artifacts(run, changes, walk, model=model, engine=options["engine"],
                           usage_changes=uc, usage_walk=uw,
                           applied=len(verify.applied_edits(run)),
                           paragraphs=len(verify.accepted_text(run)))
    return provider, changes, walk


def reusable(run, provider=None, **overrides):
    options = {**POLICY, **overrides}
    model = options.pop("model", MODEL)
    return verify.reusable_clean_verification(run, provider or Provider(), model, **options)


def settle_clean_entry(run, monkeypatch, provider, **overrides):
    options = {**POLICY, **overrides}
    model = options.pop("model", MODEL)
    settler = Settler(run, cfg=None, manuscript=run / "book.docx", error_dir=run,
        provider=provider, options=SettleOptions(until_clean=True, engine=options["engine"],
            model=model, context=options.get("context", ""),
            verification_pass=options["pass_id"], verification_policy=options["policy_id"],
            required_verification_passes=options["required_pass_ids"],
            verification_config_sha256=options["config_sha256"]))
    monkeypatch.setattr(settler, "_intake_corrections", lambda: [])
    monkeypatch.setattr(settler, "_reconcile_comments", lambda: None)
    monkeypatch.setattr(settler, "_finalize", lambda: None)
    settler.run()
    return settler


def test_clean_entry_reuses_complete_exact_policy_without_model_calls(run, monkeypatch):
    complete(run)
    assert reusable(run)
    provider = Provider()
    result = settle_clean_entry(run, monkeypatch, provider)
    assert provider.calls == []
    assert result.usage.api_calls == 0
    assert any("Reused complete clean verification" in n for n in result.settlement.notes)


@pytest.mark.parametrize("changed", ["document", "source", "edit", "context", "model", "engine",
                                      "effort", "schema", "policy", "config", "pass_policy"])
def test_changed_semantics_force_fresh_model_calls(run, monkeypatch, changed):
    complete(run)
    options, provider = {}, Provider()
    if changed == "document":
        document = docx.Document(run / "book.docx")
        document.paragraphs[0].text = "The dog wakes."
        document.save(run / "book.docx")
    elif changed == "source":
        original, accepted = verify.paragraph_views(run)
        original[next(iter(original))] = "A different original manuscript sentence."
        monkeypatch.setattr(verify, "paragraph_views", lambda _run: (original, accepted))
    elif changed == "edit":
        payload = json.loads((run / "findings.json").read_text())
        payload["findings"][0]["original_text"] = "tge"
        (run / "findings.json").write_text(json.dumps(payload))
    elif changed == "context":
        options["context"] = "This book follows an updated name convention."
    elif changed == "model":
        options["model"] = "gpt-5.6-sol"
    elif changed == "engine":
        options["engine"] = "subagent"
    elif changed == "effort":
        provider.effort = "low"
    elif changed == "schema":
        schema, name = verify._walk_schema()
        schema["description"] = "Changed verification response contract."
        monkeypatch.setattr(verify, "_walk_schema", lambda: (schema, name))
    elif changed == "policy":
        options["policy_id"] = "mechanical-verification-v2"
    elif changed == "config":
        options["config_sha256"] = "b" * 64
    else:
        options["required_pass_ids"] = ("primary", "independent-second")
    assert not reusable(run, provider, **options)
    settle_clean_entry(run, monkeypatch, provider, **options)
    assert len(provider.calls) == 2


@pytest.mark.parametrize("damage", ["legacy", "incomplete", "missing_window", "tampered_receipt",
                                     "unread", "dirty", "partial", "failed", "missing_findings"])
def test_unproven_or_incomplete_artifacts_never_reuse(run, monkeypatch, damage):
    complete(run)
    path = run / "finished_walk.json"
    payload = json.loads(path.read_text())
    if damage == "legacy":
        payload.pop("verification_provenance")
    elif damage == "incomplete":
        payload["verification_provenance"]["complete"] = False
    elif damage == "missing_window":
        proof = payload["verification_provenance"]
        key = next(iter(proof["windows"]))
        (run / verify._CHECKPOINT_DIR / proof["scope"] / proof["invocation_id"] / (key + ".json")).unlink()
    elif damage == "tampered_receipt":
        proof = payload["verification_provenance"]
        key = next(iter(proof["windows"]))
        receipt = run / verify._CHECKPOINT_DIR / proof["scope"] / proof["invocation_id"] / (key + ".json")
        body = json.loads(receipt.read_text())
        body["parsed"] = {}
        receipt.write_text(json.dumps(body))
    elif damage == "unread":
        payload["unread_paragraphs"] = [next(iter(verify.accepted_text(run)))]
    elif damage == "dirty":
        payload["unverified_paragraphs"] = [next(iter(verify.accepted_text(run)))]
    elif damage == "partial":
        payload["paragraphs_verified"] = list(verify.accepted_text(run))
    elif damage == "failed":
        payload["ran"] = False
    else:
        (run / "findings.json").unlink()
    path.write_text(json.dumps(payload))
    assert not reusable(run)
    provider = Provider()
    if damage == "missing_findings":
        with pytest.raises(RuntimeError, match="incomplete"):
            settle_clean_entry(run, monkeypatch, provider)
    else:
        settle_clean_entry(run, monkeypatch, provider)
    assert len(provider.calls) == (1 if damage == "missing_findings" else 2)


def test_completed_invocation_is_a_fresh_independent_read_on_repeat(run):
    provider, _, first = complete(run)
    _, _, second = complete(run, provider)
    assert len(provider.calls) == 4
    assert first.verification_provenance["walk"]["invocation_id"] != second.verification_provenance["walk"]["invocation_id"]
    assert second.recovered_usage.api_calls == 0


@pytest.mark.parametrize("boundary", ["pass", "purpose", "context"])
def test_distinct_read_does_not_resume_another_pass_or_purpose(run, boundary):
    document = docx.Document()
    for i in range(3):
        document.add_paragraph(f"Paragraph {i}. " + "The dog sleeps. " * 260)
    document.save(run / "book.docx")
    provider = Provider(interrupt_at=2)
    with pytest.raises(RuntimeError):
        verify.verify_run(run, provider, MODEL, Usage(), run_changes=False, **POLICY)
    other = dict(POLICY)
    if boundary == "pass":
        other.update(pass_id="second", required_pass_ids=("primary", "second"))
    elif boundary == "purpose":
        other["purpose"] = "deliberate-independent-review"
    else:
        other["context"] = "A changed voice convention."
    provider = Provider()
    result = verify.verify_run(run, provider, MODEL, Usage(), run_changes=False, **other)
    assert len(provider.calls) == 3
    assert result.recovered_usage.api_calls == 0


def test_interrupted_walk_resumes_only_complete_windows_and_restores_cost(run):
    document = docx.Document()
    for i in range(3):
        document.add_paragraph(f"Paragraph {i}. " + "The dog sleeps. " * 260)
    document.save(run / "book.docx")
    provider, original_usage = Provider(interrupt_at=2), Usage()
    with pytest.raises(RuntimeError):
        verify.verify_run(run, provider, MODEL, original_usage, run_changes=False, **POLICY)
    provider, new_usage = Provider(), Usage()
    result = verify.verify_run(run, provider, MODEL, new_usage, run_changes=False, **POLICY)
    assert len(provider.calls) == 2
    assert new_usage.api_calls == 2 and new_usage.input_tokens == 2000
    assert result.recovered_usage.api_calls == 1 and result.recovered_usage.input_tokens == 1000
    assert result.verification_provenance["walk"]["complete"]
    assert len(result.verification_provenance["walk"]["windows"]) == 3
    skipped = verify.VerifyRunResult([], [], False, False)
    verify.write_artifacts(run, skipped, result, model=MODEL, engine="provider",
                           usage_changes=Usage(), usage_walk=new_usage, applied=1, paragraphs=3)
    expected = Usage()
    for _ in range(3):
        expected.add(NormalizedUsage(1000, 200), model=MODEL)
    artifact = json.loads((run / "finished_walk.json").read_text())
    assert artifact["cost"] == build_envelope(findings=(), usage=expected, fallback_model=MODEL)["cost"]


def test_recovered_subscription_usage_retains_zero_billing(run):
    document = docx.Document()
    for i in range(2):
        document.add_paragraph(f"Paragraph {i}. " + "The dog sleeps. " * 260)
    document.save(run / "book.docx")
    with pytest.raises(RuntimeError):
        verify.verify_run(run, Provider(interrupt_at=2, billed=False), MODEL,
                          Usage(), run_changes=False, **POLICY)
    usage = Usage()
    result = verify.verify_run(run, Provider(billed=False), MODEL, usage,
                               run_changes=False, **POLICY)
    assert result.recovered_usage.by_model[MODEL]["billed"] is False
    skipped = verify.VerifyRunResult([], [], False, False)
    verify.write_artifacts(run, skipped, result, model=MODEL, engine="provider",
                           usage_changes=Usage(), usage_walk=usage, applied=1, paragraphs=2)
    assert json.loads((run / "finished_walk.json").read_text())["cost"]["total_usd"] == 0


def test_identical_change_batches_are_separate_required_windows(run, monkeypatch):
    monkeypatch.setattr(verify, "DEFAULT_CHANGE_BATCH", 1)
    # The callable's default is captured at import, so also request batch size 1
    # through a thin adapter while exercising the real verification loop.
    original_verify = verify.verify_changes
    monkeypatch.setattr(verify, "verify_changes", lambda *a, **kw:
                        original_verify(*a, batch_size=1, **kw))
    payload = json.loads((run / "findings.json").read_text())
    payload["findings"] *= 2
    (run / "findings.json").write_text(json.dumps(payload))
    provider = Provider()
    result = verify.verify_run(run, provider, MODEL, Usage(), run_walk=False, **POLICY)
    assert len(provider.calls) == 2
    proof = result.verification_provenance["changes"]
    assert proof["complete"] and len(proof["windows"]) == 2


@pytest.mark.parametrize("body", [{}, {"findings": [{"para_id": "missing", "quote": "bad",
    "problem": "typo", "suggestion": "good", "severity": "high"}]}])
def test_invalid_success_body_is_never_a_cached_complete_window(run, body):
    provider = Provider(body=body)
    first = verify.verify_run(run, provider, MODEL, Usage(), run_changes=False, **POLICY)
    assert first.verification_provenance["walk"]["complete"] is False
    assert first.ran_walk is False and verify.UNREAD
    skipped = verify.VerifyRunResult([], [], False, False)
    verify.write_artifacts(run, skipped, first, model=MODEL, engine="provider",
                           usage_changes=Usage(), usage_walk=Usage(), applied=1, paragraphs=1)
    from galley.manifest import _certify_finished_walk
    assert _certify_finished_walk(run).status == "fail"
    provider.body = None
    second = verify.verify_run(run, provider, MODEL, Usage(), run_changes=False, **POLICY)
    assert len(provider.calls) == 3
    assert second.verification_provenance["walk"]["complete"] is True


def test_partial_write_invalidates_full_proof(run):
    complete(run)
    delta = verify.verify_delta(run, list(verify.accepted_text(run)), Provider(), MODEL, Usage())
    verify.write_artifacts(run, delta, delta, model=MODEL, engine="provider",
                           usage_changes=Usage(), usage_walk=Usage(), applied=1, paragraphs=1,
                           para_ids=list(verify.accepted_text(run)), merge=True)
    assert "verification_provenance" not in json.loads((run / "finished_walk.json").read_text())
    assert not reusable(run)


def test_legacy_direct_caller_still_runs_fresh_without_policy_proof(run):
    provider = Provider()
    first = verify.verify_run(run, provider, MODEL, Usage())
    second = verify.verify_run(run, provider, MODEL, Usage())
    assert len(provider.calls) == 4
    assert first.verification_provenance == second.verification_provenance == {}


def command(run, provider, *, output=None, crash_after_write=False):
    output = output or run
    with verify.verification_invocation(run, provider, MODEL, output_dir=output, **POLICY) as invocation:
        uc, uw = Usage(), Usage()
        options = {**POLICY, "command_id": invocation.command_id}
        changes = verify.verify_run(run, provider, MODEL, uc, run_walk=False, **options)
        walk = verify.verify_run(run, provider, MODEL, uw, run_changes=False, **options)
        verify.write_artifacts(output, changes, walk, model=MODEL, engine="provider",
            usage_changes=uc, usage_walk=uw, applied=1,
            paragraphs=len(verify.accepted_text(run)), source_run_dir=run)
        if crash_after_write:
            raise RuntimeError("crash before command completion")
        assert invocation.mark_complete(changes, walk)
        return invocation.command_id, changes, walk, uc, uw


def test_command_recovers_completed_changes_after_interruption_in_walk(run):
    with pytest.raises(RuntimeError):
        command(run, Provider(interrupt_at=2))
    provider = Provider()
    _, changes, walk, uc, uw = command(run, provider)
    assert len(provider.calls) == 1 and provider.calls[0]["schema_name"] == "findings"
    assert uc.api_calls == 0 and changes.recovered_usage.api_calls == 1
    assert uw.api_calls == 1 and walk.recovered_usage.api_calls == 0
    expected = Usage()
    expected.add(NormalizedUsage(1000, 200), model=MODEL)
    assert json.loads((run / "change_verify.json").read_text())["cost"] == build_envelope(
        findings=(), usage=expected, fallback_model=MODEL)["cost"]


def test_crash_after_artifact_write_recovers_then_next_command_is_fresh(run):
    with pytest.raises(RuntimeError, match="before command completion"):
        command(run, Provider(), crash_after_write=True)
    provider = Provider()
    recovered_id, changes, walk, uc, uw = command(run, provider)
    assert provider.calls == [] and uc.api_calls == uw.api_calls == 0
    assert changes.recovered_usage.api_calls == walk.recovered_usage.api_calls == 1
    next_id, changes, walk, uc, uw = command(run, provider)
    assert next_id != recovered_id and len(provider.calls) == 2
    assert changes.recovered_usage.api_calls == walk.recovered_usage.api_calls == 0


def test_side_output_command_binds_proof_to_source_manuscript(run, tmp_path):
    output = tmp_path / "reports"
    command(run, Provider(), output=output)
    payload = json.loads((output / "finished_walk.json").read_text())
    assert not list(output.glob("*.docx"))
    assert payload["build_sha256"] == verify.build_fingerprints(run)["build_sha256"]
    assert payload["verification_provenance"]["complete"] is True


@pytest.mark.parametrize("gate", ["changes", "walk"])
def test_saturated_explicit_read_is_unread_and_cannot_certify(run, monkeypatch, gate):
    if gate == "changes":
        monkeypatch.setattr(verify, "MAX_PROBLEMS", 1)
        provider = Provider(body={"problems": [{"index": 1, "verdict": "wrong_rule",
                                               "detail": "Problem", "fix": "the"}]})
    else:
        monkeypatch.setattr(verify, "MAX_RESIDUALS", 1)
        pid = next(iter(verify.accepted_text(run)))
        provider = Provider(body={"findings": [{"para_id": pid, "quote": "dog",
            "problem": "Problem", "suggestion": "cat", "severity": "high"}]})
    result = verify.verify_run(run, provider, MODEL, Usage(),
        run_changes=gate == "changes", run_walk=gate == "walk", **POLICY)
    assert result.verification_provenance[gate]["complete"] is False
    assert verify.UNREAD_BATCHES if gate == "changes" else verify.UNREAD
    skipped = verify.VerifyRunResult([], [], False, False)
    verify.write_artifacts(run, result if gate == "changes" else skipped,
        result if gate == "walk" else skipped, model=MODEL, engine="provider",
        usage_changes=Usage(), usage_walk=Usage(), applied=1, paragraphs=1)
    from galley.manifest import _certify_change_verify, _certify_finished_walk
    assert (_certify_change_verify(run) if gate == "changes" else _certify_finished_walk(run)).status == "fail"
