"""The interior-corrections stage, end to end, with a fake Drive, a fake HubSpot
and no model at all.

What it is held to: the file it corrects is the highest integer
"<surname> - Book N.idml" in the author's "Interior Design" folder and nothing
else; what comes back is "<surname> - Book N.5.idml" with the two-sheet
spreadsheet beside it, in the same folder; every correction the author sent
lands on exactly one sheet; the CRM moves on exactly once; a second pass
neither re-applies nor re-uploads; and nothing is ever guessed — a rival Book
N.5, a missing folder, an empty form each stop the book and tell a person.

The submission is a tracked-changes Word file, which the intake reads without a
model, and the model passes are off, so the whole round trip is deterministic.
"""
from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from app.jobs import JobRunner, JobStore
from app.settings import Paths
from app.watch import corrections as corrlib
from app.watch import naming
from app.watch import tick as ticklib
from app.watch.drive import DriveFile, FOLDER_MIME
from app.watch.hubspot import HubSpotError
from app.watch.settings import WatchSettings
from app.watch.stages import (CORRECTIONS_DONE, CORRECTIONS_FAILED,
                              CORRECTIONS_PROP, OUTPUT_PROP, SOURCE_PROP)
from app.watch.state import WatchState
from docproof.providers import ProviderResult

from .conftest import FIXTURES
from .fakes import FakeProvider, drive_entry, fake_drive
from .test_corrections_extract import make_tracked_docx

FOLDER = "1AbCdEfGhIjKlMnOp"
AUTHOR = "sf-johnson"
INTERIOR = "sf-johnson-interior"
IDML = (FIXTURES / "layout.idml").read_bytes()
FILE_URL = "https://api.hubapi.com/files/sub-1?filename=Johnson%20corrections.docx"

SOURCE = "Johnson - Book 3.idml"
HAND_OFF = {"Johnson - Book 3.5.idml", "Johnson - Book 3.5 - corrections.xlsx",
            "Johnson - Book 3.5 - notes.md"}


def ws(**over) -> WatchSettings:
    fields = dict(folder_id=FOLDER, model="claude-haiku-4-5",
                  client_id="client-1", client_secret="secret-1",
                  hubspot_enabled=True, hubspot_object="0-970",
                  hubspot_key_property="author_last_name",
                  hubspot_status_property="docproof",
                  hubspot_format_ready_value="Ready for Formatting",
                  hubspot_format_done_value="Formatting Complete",
                  subfolders_enabled=True, require_source_label=True,
                  hubspot_first_property="firstname",
                  hubspot_last_property="lastname",
                  corrections_enabled=True,
                  # Every test below this line was written against the old
                  # HubSpot-status gate, so it keeps that mode explicitly —
                  # form intake (the stage's new default) has its own fixture,
                  # `form_ws`, and its own tests further down.
                  corrections_intake="hubspot",
                  hubspot_corrections_file_property="corr_file",
                  hubspot_corrections_text_property="corr_text",
                  corrections_model_passes=False,
                  # A zero quiet period keeps every existing test's round trip
                  # immediate, exactly as it ran before the hold existed; the
                  # hold itself is exercised on its own below.
                  corrections_quiet_seconds=0)
    fields.update(over)
    return WatchSettings(**fields)


def form_ws(**over) -> WatchSettings:
    """`ws()`, switched to form intake: the corrections form's own submissions
    are the trigger instead of a HubSpot status. The two CRM properties
    (`hubspot_corrections_file_property` / `_text_property`) are kept set even
    though form mode never reads them — the stage's preflight is shared with
    hubspot mode and still checks for them."""
    fields = dict(corrections_intake="form",
                 corrections_form_first_property="firstname",
                 corrections_form_last_property="lastname",
                 corrections_form_book_property="book_title",
                 corrections_form_file_property="files",
                 corrections_form_notes_property="notes",
                 corrections_quiet_seconds=0)
    fields.update(over)
    return ws(**fields)


def ready(**extra) -> dict:
    return {"docproof": "Ready for Corrections", "firstname": "Quinton",
            "lastname": "Johnson", "corr_file": FILE_URL, **extra}


def folder_entry(name, parent) -> dict:
    entry = drive_entry(name, mime=FOLDER_MIME)
    entry["parents"] = [parent]
    return entry


def in_folder(name, parent=INTERIOR, **kw) -> dict:
    entry = drive_entry(name, **kw)
    entry["parents"] = [parent]
    return entry


def tree(*extra_interior: tuple[str, dict]) -> dict:
    files = {AUTHOR: folder_entry("Quinton Johnson", FOLDER),
             INTERIOR: folder_entry("Interior Design", AUTHOR),
             "idml-3": in_folder(SOURCE)}
    for fid, entry in extra_interior:
        files[fid] = entry
    return files


def submission(tmp_path) -> bytes:
    """A tracked-changes list: one change the file carries, one it does not."""
    path = make_tracked_docx(tmp_path / "corrections.docx", [
        [("", "She opened the door, the room was "), ("del", "empty"),
         ("ins", "bare"), ("", ".")],
        [("", "The moon was "), ("del", "blue"), ("ins", "red"), ("", ".")],
    ])
    return path.read_bytes()


def make_opener(tmp_path, files=None, hubspot=None):
    opener = fake_drive(files if files is not None else tree(),
                        hubspot=hubspot if hubspot is not None
                        else {"Johnson": ready()})
    opener.content["idml-3"] = IDML
    opener.content["sub-1"] = submission(tmp_path)
    return opener


def run(home, settings, opener, **kw):
    return ticklib.tick(home, settings, opener=opener,
                        get_key=lambda name: "refresh-1", **kw)


def uploads_in(opener) -> dict:
    return {e["name"]: e for e in opener.files.values()
            if e.get("appProperties", {}).get(OUTPUT_PROP)}


def patches(opener) -> list:
    return [c for c in opener.calls if "api.hubapi.com" in c.full_url
            and c.get_method() == "PATCH"]


# --- names (pure) --------------------------------------------------------------

@pytest.mark.parametrize("name,version", [
    ("Johnson - Book 3.idml", 3.0), ("Johnson - Book3.idml", 3.0),
    ("Johnson — Book 10.idml", 10.0), ("Johnson - Book 3.5.idml", 3.5),
])
def test_idml_version_reads_the_designers_series(name, version):
    assert naming.idml_version(name) == ("johnson", version)


@pytest.mark.parametrize("name", [
    "Johnson - Book 3.docx", "Johnson - Book Original.idml", "notes.idml",
])
def test_idml_version_ignores_what_is_not_an_export(name):
    assert naming.idml_version(name) is None


def test_source_is_an_integer_export_with_the_records_surname():
    assert naming.is_idml_source_name("Johnson - Book 3.idml", "JOHNSON")
    assert not naming.is_idml_source_name("Johnson - Book 3.5.idml", "Johnson")
    assert not naming.is_idml_source_name("Smith - Book 3.idml", "Johnson")


def test_hand_off_names_step_by_a_half_and_mirror_the_spacing():
    names = naming.corrections_hand_off_names("Johnson - Book 3.idml")
    assert names["idml"] == "Johnson - Book 3.5.idml"
    assert names["sheet"] == "Johnson - Book 3.5 - corrections.xlsx"
    assert naming.corrections_base("Johnson - Book3.idml") == "Johnson - Book3.5"
    with pytest.raises(ValueError):
        naming.corrections_base("Johnson - Book 3.5.idml")


def test_the_hand_off_is_recognised_as_output():
    assert naming.is_idml_output_name("Johnson - Book 3.5.idml")
    assert naming.is_idml_output_name("Johnson - Book 3.5 - corrections.xlsx")
    assert not naming.is_idml_output_name("Johnson - Book 3.idml")


@pytest.mark.parametrize("name,version", [
    ("Smith - Book 4.indd", 4.0), ("Smith - Book4.indd", 4.0),
    ("Smith — Book 10.indd", 10.0),
])
def test_indd_version_reads_the_designers_indd_series(name, version):
    assert naming.indd_version(name) == ("smith", version)


@pytest.mark.parametrize("name", [
    "Smith - Book 4.idml", "Smith - Book Original.indd", "notes.indd",
])
def test_indd_version_ignores_what_is_not_an_indd_export(name):
    assert naming.indd_version(name) is None


def test_indd_source_is_an_integer_export_with_the_records_surname():
    assert naming.is_indd_source_name("Smith - Book 4.indd", "SMITH")
    assert not naming.is_indd_source_name("Smith - Book 4.5.indd", "Smith")
    assert not naming.is_indd_source_name("Jones - Book 4.indd", "Smith")


def _df(name, fid, props=None, mime="application/octet-stream"):
    return DriveFile(id=fid, name=name, mime_type=mime,
                     app_properties=props or {})


def test_pick_source_takes_the_highest_integer_export():
    listing = [_df("Johnson - Book 3.idml", "a"), _df("Johnson - Book 4.idml", "b"),
               _df("Johnson - Book 3.5.idml", "c"), _df("Johnson - Book 4.pdf", "d")]
    chosen, why = corrlib.pick_source(listing, "Johnson")
    assert chosen.id == "b" and why == ""


def test_pick_source_refuses_a_tie_and_reports_a_finished_latest():
    tie = [_df("Johnson - Book 4.idml", "a"), _df("Johnson - Book 4.idml", "b")]
    assert corrlib.pick_source(tie, "Johnson") == (None, "tie")
    done = [_df("Johnson - Book 4.idml", "a", {CORRECTIONS_PROP: CORRECTIONS_DONE})]
    assert corrlib.pick_source(done, "Johnson") == (None, "done")
    failed = [_df("Johnson - Book 4.idml", "a", {CORRECTIONS_PROP: CORRECTIONS_FAILED})]
    assert corrlib.pick_source(failed, "Johnson") == (None, "failed")
    assert corrlib.pick_source([_df("Johnson - Book 3.5.idml", "c")],
                               "Johnson") == (None, "none")


def test_pick_source_spots_an_indd_with_no_idml_export():
    listing = [_df("Johnson - Book 4.indd", "a")]
    assert corrlib.pick_source(listing, "Johnson") == (None, "indd-only")


def test_pick_source_spots_an_idml_behind_the_latest_indd():
    listing = [_df("Johnson - Book 3.idml", "a"), _df("Johnson - Book 4.indd", "b")]
    assert corrlib.pick_source(listing, "Johnson") == (None, "idml-stale")


def test_pick_source_is_unbothered_by_an_indd_alongside_its_own_idml():
    # The .idml is at the same number as the .indd (the ordinary case, the
    # designer's export sitting right beside their working file) or ahead of
    # it (a later export than the last .indd DocWatch happened to see) —
    # either way this is not "stale" and the export is picked normally.
    same = [_df("Johnson - Book 4.idml", "a"), _df("Johnson - Book 4.indd", "b")]
    chosen, why = corrlib.pick_source(same, "Johnson")
    assert chosen.id == "a" and why == ""
    ahead = [_df("Johnson - Book 5.idml", "a"), _df("Johnson - Book 4.indd", "b")]
    chosen, why = corrlib.pick_source(ahead, "Johnson")
    assert chosen.id == "a" and why == ""


def test_proof_pdf_is_the_one_under_the_same_stem():
    idml = _df("Johnson - Book 3.idml", "a")
    listing = [idml, _df("Johnson - Book 3.pdf", "p"), _df("Johnson - Book 2.pdf", "q")]
    assert corrlib.proof_pdf_for(listing, idml).id == "p"
    assert corrlib.proof_pdf_for([idml], idml) is None


# --- the round trip --------------------------------------------------------------

def test_a_ready_form_is_applied_and_the_book_3_5_goes_back_beside_the_export(
        tmp_path):
    opener = make_opener(tmp_path)
    report = run(tmp_path, ws(), opener)

    assert not report.failed, report.failed
    assert report.corrected and report.corrected[0].startswith(SOURCE)
    placed = uploads_in(opener)
    assert HAND_OFF <= set(placed)
    for name in HAND_OFF:
        # In the Interior Design folder, pointing at the export it came from.
        assert placed[name]["parents"] == [INTERIOR]
        assert placed[name]["appProperties"][SOURCE_PROP] == "idml-3"
    # Nothing was written into the author folder or the parent.
    assert all(e["parents"] == [INTERIOR] for e in placed.values())

    # The CRM moved on, once, to the designer's cue.
    assert opener.hubspot["hs-Johnson"]["properties"]["docproof"] == \
        "Corrections Applied"
    assert len(patches(opener)) == 1
    # The export carries the stage's marker, written last.
    assert opener.files["idml-3"]["appProperties"][CORRECTIONS_PROP] == \
        CORRECTIONS_DONE


def test_the_spreadsheet_accounts_for_every_correction_on_one_sheet_each(
        tmp_path):
    opener = make_opener(tmp_path)
    run(tmp_path, ws(), opener)
    sheet_id = uploads_in(opener)["Johnson - Book 3.5 - corrections.xlsx"]["id"]
    path = tmp_path / "sheet.xlsx"
    path.write_bytes(opener.content[sheet_id])
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Applied", "Not applied"]

    def rows(sheet):
        cells = list(sheet.iter_rows(values_only=True))
        header_at = next(i for i, r in enumerate(cells) if r and r[0] == "#")
        header = list(cells[header_at])
        return [dict(zip(header, r)) for r in cells[header_at + 1:]
                if r and r[0] is not None]

    applied = rows(wb["Applied"])
    not_applied = rows(wb["Not applied"])
    # Two tracked changes went in: one the file carries, one it does not.
    assert len(applied) == 1 and len(not_applied) == 1
    assert "empty" in applied[0]["Find"] and "bare" in applied[0]["Replace"]
    assert "blue" in not_applied[0]["Find"]
    assert "not found" in not_applied[0]["Status"].lower()
    # The corrected IDML really carries the change.
    from docproof.corrections.idml import read_stories
    corrected = tmp_path / "book.idml"
    corrected.write_bytes(
        opener.content[uploads_in(opener)["Johnson - Book 3.5.idml"]["id"]])
    text = "\n".join(p.text for s in read_stories(corrected) for p in s.paragraphs)
    assert "the room was bare." in text and "the room was empty." not in text


# --- delivery verification --------------------------------------------------------

def test_uploads_are_read_back_from_drive_before_the_hand_off_counts(tmp_path):
    opener = make_opener(tmp_path)
    before = len(opener.calls)
    report = run(tmp_path, ws(), opener)

    assert not report.failed, report.failed
    placed = uploads_in(opener)
    # Every landed artifact's Drive id was read back with a GET — not just
    # trusted because the upload's own response carried an id.
    gets = [c for c in opener.calls[before:]
           if c.get_method() == "GET"
           and any(f"/drive/v3/files/{e['id']}" in c.full_url
                   for e in placed.values())]
    assert len(gets) >= len(HAND_OFF)
    assert opener.hubspot["hs-Johnson"]["properties"]["docproof"] == \
        "Corrections Applied"


def test_a_bad_first_readback_is_reuploaded_and_then_succeeds(tmp_path, monkeypatch):
    """One mismatch — Drive says the xlsx landed a different size than the file
    on disk — earns one reupload, not an immediate failure. A dropped
    connection or a stale id is often gone on a retry."""
    opener = make_opener(tmp_path)
    real_get_file = corrlib.drive.get_file
    spoiled = {"done": False}

    def flaky_get_file(token, file_id, *, opener, with_parents=False):
        found = real_get_file(token, file_id, opener=opener,
                              with_parents=with_parents)
        if not spoiled["done"] and found.name.endswith("corrections.xlsx"):
            spoiled["done"] = True
            return replace(found, size=found.size + 1)
        return found

    monkeypatch.setattr(corrlib.drive, "get_file", flaky_get_file)
    report = run(tmp_path, ws(), opener)

    assert not report.failed, report.failed
    assert spoiled["done"]                    # the bad reading really happened
    placed = uploads_in(opener)
    assert HAND_OFF <= set(placed)
    assert opener.hubspot["hs-Johnson"]["properties"]["docproof"] == \
        "Corrections Applied"


def test_a_persistently_bad_readback_fails_before_hubspot_or_the_marker(
        tmp_path, monkeypatch):
    opener = make_opener(tmp_path)
    real_get_file = corrlib.drive.get_file

    def bad_get_file(token, file_id, *, opener, with_parents=False):
        found = real_get_file(token, file_id, opener=opener,
                              with_parents=with_parents)
        if found.name.endswith("3.5.idml"):
            return replace(found, md5_checksum="0" * 32)
        return found

    monkeypatch.setattr(corrlib.drive, "get_file", bad_get_file)
    report = run(tmp_path, ws(), opener)

    assert report.failed and "did not land intact" in report.failed[0][1]
    # HubSpot never moved, and the source was never marked done.
    assert opener.hubspot["hs-Johnson"]["properties"]["docproof"] == \
        "Ready for Corrections"
    assert CORRECTIONS_PROP not in opener.files["idml-3"]["appProperties"]
    rec = WatchState.load(tmp_path / "state.json").get("idml-3")
    assert rec.corrections_attempts == 1
    assert "Johnson - Book 3.5.idml" not in rec.corrections_uploaded


def test_a_second_pass_neither_re_applies_nor_re_uploads(tmp_path):
    opener = make_opener(tmp_path)
    run(tmp_path, ws(), opener)
    before = dict(uploads_in(opener))
    # The status has moved on, so HubSpot no longer lists the record as ready.
    report = run(tmp_path, ws(), opener)
    assert report.corrected == []
    assert uploads_in(opener) == before
    assert len(patches(opener)) == 1
    store = JobStore(Paths(tmp_path))
    assert len([j for j in store.all() if j.kind == "corrections"]) == 1


def test_the_stage_reads_only_the_highest_export_and_never_a_half_step(tmp_path):
    files = tree(("idml-2", in_folder("Johnson - Book 2.idml")),
                 ("idml-25", in_folder("Johnson - Book 2.5.idml",
                                       props={OUTPUT_PROP: "1"})))
    opener = make_opener(tmp_path, files=files)
    opener.content["idml-2"] = IDML
    run(tmp_path, ws(), opener)
    placed = uploads_in(opener)
    assert "Johnson - Book 3.5.idml" in placed
    assert "Johnson - Book 2.5.idml" not in {
        n for n, e in placed.items() if e["id"].startswith("up-")}
    assert CORRECTIONS_PROP not in opener.files["idml-2"]["appProperties"]


def test_a_rival_book_3_5_is_never_overwritten(tmp_path):
    files = tree(("rival", in_folder("Johnson - Book 3.5.idml")))
    opener = make_opener(tmp_path, files=files)
    report = run(tmp_path, ws(), opener)
    assert report.corrected == []
    assert uploads_in(opener) == {}
    assert any("will not overwrite" in reason
               for _, reason in report.needs_human)
    assert opener.hubspot["hs-Johnson"]["properties"]["docproof"] == \
        "Ready for Corrections"


def test_no_interior_design_folder_is_a_missing_source(tmp_path):
    files = {AUTHOR: folder_entry("Quinton Johnson", FOLDER)}
    opener = make_opener(tmp_path, files=files)
    report = run(tmp_path, ws(), opener)
    assert report.missing_source and "Interior Design" in report.missing_source[0][1]
    assert uploads_in(opener) == {}


def test_an_indd_with_no_idml_export_explains_what_to_export(tmp_path):
    files = {AUTHOR: folder_entry("Quinton Johnson", FOLDER),
             INTERIOR: folder_entry("Interior Design", AUTHOR),
             "indd-4": in_folder("Johnson - Book 4.indd")}
    opener = make_opener(tmp_path, files=files)
    report = run(tmp_path, ws(), opener)
    assert report.missing_source
    reason = report.missing_source[0][1]
    assert "File → Export" in reason and "InDesign Markup (IDML)" in reason
    assert "Johnson - Book 4.idml" in reason
    assert uploads_in(opener) == {}
    assert report.corrected == []


# --- multi-book authors ------------------------------------------------------

RED_FOLDER, RED_INTERIOR, RED_IDML = "red-folder", "red-interior", "red-idml"
BLUE_FOLDER, BLUE_INTERIOR, BLUE_IDML = "blue-folder", "blue-interior", "blue-idml"
BOOK_HAND_OFF = {"Johnson - Book 1.5.idml", "Johnson - Book 1.5 - corrections.xlsx",
                 "Johnson - Book 1.5 - notes.md"}


def two_book_files() -> dict:
    return {
        AUTHOR: folder_entry("Quinton Johnson", FOLDER),
        RED_FOLDER: folder_entry("The Red Book", AUTHOR),
        RED_INTERIOR: folder_entry("Interior Design", RED_FOLDER),
        RED_IDML: in_folder("Johnson - Book 3.idml", parent=RED_INTERIOR),
        BLUE_FOLDER: folder_entry("Blue Tide", AUTHOR),
        BLUE_INTERIOR: folder_entry("Interior Design", BLUE_FOLDER),
        BLUE_IDML: in_folder("Johnson - Book 1.idml", parent=BLUE_INTERIOR),
    }


def test_a_multi_book_author_is_routed_by_the_records_book_title(tmp_path):
    record = ready(book_title="Blue Tide")
    opener = make_opener(tmp_path, files=two_book_files(),
                         hubspot={"Johnson": record})
    opener.content[BLUE_IDML] = IDML
    report = run(tmp_path, ws(), opener)

    assert not report.failed, report.failed
    assert report.corrected and report.corrected[0].startswith(
        "Johnson - Book 1.idml")
    placed = uploads_in(opener)
    assert BOOK_HAND_OFF <= set(placed)
    for name in BOOK_HAND_OFF:
        # In Blue Tide's own "Interior Design" folder, not the Red Book's.
        assert placed[name]["parents"] == [BLUE_INTERIOR]


def test_a_multi_book_author_with_no_title_needs_a_person(tmp_path):
    opener = make_opener(tmp_path, files=two_book_files())    # no book_title
    report = run(tmp_path, ws(), opener)
    assert any("book_title" in reason for _, reason in report.needs_human)
    assert uploads_in(opener) == {}


def test_a_form_with_nothing_attached_needs_a_person(tmp_path):
    record = ready()
    record["corr_file"] = ""
    opener = make_opener(tmp_path, hubspot={"Johnson": record})
    report = run(tmp_path, ws(), opener)
    assert any("no corrections" in reason for _, reason in report.needs_human)
    assert uploads_in(opener) == {}
    assert len(patches(opener)) == 0


def test_typed_text_alone_needs_a_model_and_says_so(tmp_path):
    """No key for the reader here, so a typed list cannot be read: the export is
    marked failed with the reason and a person is told, rather than three silent
    retries."""
    record = ready()
    record["corr_file"] = ""
    record["corr_text"] = "Page 3: change 'empty' to 'bare'."
    opener = make_opener(tmp_path, hubspot={"Johnson": record})
    report = run(tmp_path, ws(), opener)
    assert any("could not be read" in reason for _, reason in report.needs_human)
    assert opener.files["idml-3"]["appProperties"][CORRECTIONS_PROP] == \
        CORRECTIONS_FAILED
    assert uploads_in(opener) == {}


def test_the_stage_stands_aside_when_off_or_half_configured(tmp_path):
    opener = make_opener(tmp_path)
    report = run(tmp_path, ws(corrections_enabled=False), opener)
    assert report.corrected == [] and uploads_in(opener) == {}
    with pytest.raises(ticklib.NotConfigured):
        run(tmp_path, ws(subfolders_enabled=False), opener)
    with pytest.raises(ticklib.NotConfigured):
        run(tmp_path, ws(hubspot_corrections_file_property="",
                         hubspot_corrections_text_property=""), opener)


def test_read_only_hubspot_still_delivers_but_writes_nothing(tmp_path):
    opener = make_opener(tmp_path)
    run(tmp_path, ws(hubspot_write_back=False), opener)
    assert "Johnson - Book 3.5.idml" in uploads_in(opener)
    assert len(patches(opener)) == 0
    assert opener.files["idml-3"]["appProperties"][CORRECTIONS_PROP] == \
        CORRECTIONS_DONE


def test_state_records_the_job_before_the_run_and_the_input_kind(tmp_path):
    opener = make_opener(tmp_path)
    run(tmp_path, ws(), opener)
    rec = WatchState.load(tmp_path / "state.json").get("idml-3")
    assert rec.corrections_job_id and rec.corrections_hubspot_done
    assert rec.corrections_marked == CORRECTIONS_DONE
    assert rec.corrections_input_kind == "docx"
    assert rec.corrections_input_name == "Johnson corrections.docx"
    assert rec.subfolder_id == INTERIOR


# --- the quiet period ------------------------------------------------------------

def test_a_ready_record_is_held_and_then_released_past_its_quiet_period(
        tmp_path):
    opener = make_opener(tmp_path)
    settings = ws(corrections_quiet_seconds=10800)          # three hours

    held = run(tmp_path, settings, opener)
    assert held.corrected == []
    assert held.waiting >= 1
    assert uploads_in(opener) == {}
    state = WatchState.load(tmp_path / "state.json")
    assert len(state.corrections_pending) == 1
    entry = next(iter(state.corrections_pending.values()))
    assert entry.author == "Quinton Johnson"
    assert len(entry.submissions) == 1

    # Back-date the hold past its own quiet period, the way waiting three
    # hours would — a test cannot wait three hours, so it moves the clock the
    # record's own state remembers instead.
    for pending in state.corrections_pending.values():
        pending.first_seen = pending.last_submission_at = \
            "2020-01-01T00:00:00+00:00"
    state.save()

    released = run(tmp_path, settings, opener)
    assert released.corrected and released.corrected[0].startswith(SOURCE)
    assert "Johnson - Book 3.5.idml" in uploads_in(opener)


def test_a_new_submission_during_the_hold_resets_the_quiet_period(tmp_path):
    opener = make_opener(tmp_path)
    settings = ws(corrections_quiet_seconds=10800)

    run(tmp_path, settings, opener)
    state = WatchState.load(tmp_path / "state.json")
    record_id = next(iter(state.corrections_pending))
    assert len(state.corrections_pending[record_id].submissions) == 1
    # Back-dated so the two runs' clocks cannot land in the same second and
    # make the "it moved" assertion below a coin flip.
    state.corrections_pending[record_id].first_seen = "2020-01-01T00:00:00+00:00"
    state.corrections_pending[record_id].last_submission_at = \
        "2020-01-01T00:00:00+00:00"
    state.save()
    first_wait = corrlib.ready_at(state.corrections_pending[record_id], settings)

    # A second submission — a different file — arrives before the hold ends.
    opener.hubspot["hs-Johnson"]["properties"]["corr_file"] = FILE_URL + "&v=2"
    run(tmp_path, settings, opener)
    state2 = WatchState.load(tmp_path / "state.json")
    assert len(state2.corrections_pending[record_id].submissions) == 2
    second_wait = corrlib.ready_at(state2.corrections_pending[record_id], settings)
    assert second_wait > first_wait


def test_pending_summary_reports_the_authors_ready_at(tmp_path):
    opener = make_opener(tmp_path)
    settings = ws(corrections_quiet_seconds=10800)
    report = run(tmp_path, settings, opener)
    assert report.waiting

    state = WatchState.load(tmp_path / "state.json")
    rows = corrlib.pending_summary(state, settings)
    assert len(rows) == 1
    row = rows[0]
    assert row["author"] == "Quinton Johnson"
    assert row["submissions"] == 1
    assert row["ready"] is False
    assert row["ready_at"]                    # a real ISO timestamp, not blank


def test_a_dry_run_reports_the_plan_and_applies_nothing(tmp_path):
    opener = make_opener(tmp_path)
    settings = ws()                            # corrections_quiet_seconds=0
    state = WatchState(tmp_path / "state.json")
    store = JobStore(Paths(tmp_path).ensure())
    runner = JobRunner(store, settings.app_settings(tmp_path),
                       config_path=ticklib.config_path(), notify_home=tmp_path)
    report = ticklib.TickReport()

    corrlib.run_stage("drive-token", tmp_path, settings, state, runner, store,
                      mock=False, opener=opener, hs_token="hubspot-token",
                      report=report, dry_run=True)

    assert report.corrected and "would apply" in report.corrected[0]
    assert report.corrected[0].startswith(SOURCE)
    assert uploads_in(opener) == {}
    assert store.all() == []
    rec = WatchState.load(tmp_path / "state.json").get("idml-3")
    assert not rec.corrections_job_id


# --- form-poll mode ---------------------------------------------------------------

def _form_row(when: str, **values) -> dict:
    return {"submittedAt": when, "recordId": "hs-Johnson",
            "values": [{"name": k, "value": v} for k, v in values.items()]}


def test_form_poll_folds_every_submission_into_one_job(tmp_path, monkeypatch):
    opener = make_opener(tmp_path)
    opener.content["form-sub-1"] = submission(tmp_path)
    rows = [
        _form_row("2026-09-01T00:00:00Z",
                 files=("https://api.hubapi.com/files/form-sub-1"
                        "?filename=Round%201.docx")),
        _form_row("2026-09-02T00:00:00Z", notes="Change 'gone' to 'here'."),
    ]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows)
    monkeypatch.setattr(corrlib, "extract_provider", lambda: (
        FakeProvider(results=[ProviderResult(parsed={"edits": [
            {"find": "gone", "replace": "here", "instruction": "swap"}]})]),
        "fake-model"))
    settings = ws(corrections_form_poll=True,
                 corrections_form_file_property="files",
                 corrections_form_notes_property="notes")

    report = run(tmp_path, settings, opener)

    assert not report.failed, report.failed
    assert report.corrected
    rec = WatchState.load(tmp_path / "state.json").get("idml-3")
    assert len(rec.corrections_submissions) == 2

    store = JobStore(Paths(tmp_path))
    job = next(j for j in store.all() if j.kind == "corrections")
    payload = json.loads(
        (Path(job.results_dir) / "corrections.json").read_text("utf-8"))
    sources = {o.get("source") for key in ("applied_items", "flagged", "no_op")
              for o in payload["apply"][key] if o.get("source")}
    # Every edit that survived to the report carries which of the two folded
    # submissions it came from.
    assert sources == set(rec.corrections_submissions)


def test_form_poll_403_falls_back_to_the_records_own_properties(
        tmp_path, monkeypatch):
    opener = make_opener(tmp_path)

    def boom(*a, **k):
        raise HubSpotError("no forms scope on this token")

    monkeypatch.setattr(corrlib.hubspot, "form_submissions", boom)
    settings = ws(corrections_form_poll=True)

    report = run(tmp_path, settings, opener)

    assert not report.failed, report.failed
    assert report.corrected and report.corrected[0].startswith(SOURCE)
    assert "Johnson - Book 3.5.idml" in uploads_in(opener)


def test_form_poll_leaves_rounds_before_the_start_date_alone(
        tmp_path, monkeypatch):
    """A submission older than `corrections_form_start_after` is a round the
    press handled by hand before form-poll mode went live: it is never folded,
    so the next job carries only what came in since."""
    opener = make_opener(tmp_path)
    opener.content["form-sub-1"] = submission(tmp_path)
    rows = [
        _form_row("2026-08-01T00:00:00Z",
                 files=("https://api.hubapi.com/files/form-sub-1"
                        "?filename=Old%20round.docx")),
        _form_row("2026-09-02T00:00:00Z", notes="Change 'gone' to 'here'."),
    ]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows)
    monkeypatch.setattr(corrlib, "extract_provider", lambda: (
        FakeProvider(results=[ProviderResult(parsed={"edits": [
            {"find": "gone", "replace": "here", "instruction": "swap"}]})]),
        "fake-model"))
    settings = ws(corrections_form_poll=True,
                 corrections_form_file_property="files",
                 corrections_form_notes_property="notes",
                 corrections_form_start_after="2026-08-15")

    report = run(tmp_path, settings, opener)

    assert not report.failed, report.failed
    rec = WatchState.load(tmp_path / "state.json").get("idml-3")
    assert len(rec.corrections_submissions) == 1
    assert "Old round" not in rec.corrections_input_name


# --- form intake (`corrections_intake == "form"`) ---------------------------------
#
# Here the corrections form's own submissions are the trigger: nobody in
# HubSpot flips anything, so the fake HubSpot below never carries a record at
# "Ready for Corrections" — it carries, at most, a Projects record or two
# named for the author, purely for the write-back `_match_hubspot_record`
# looks up once a book is delivered.

def test_form_intake_runs_from_the_forms_own_submissions_and_moves_hubspot(
        tmp_path, monkeypatch):
    opener = make_opener(tmp_path, hubspot={"Johnson": {"firstname": "Quinton",
                                                        "lastname": "Johnson"}})
    opener.content["form-sub-1"] = submission(tmp_path)
    rows = [
        _form_row("2026-09-01T00:00:00Z", firstname="Quinton", lastname="Johnson",
                 files=("https://api.hubapi.com/files/form-sub-1"
                        "?filename=Round%201.docx")),
        _form_row("2026-09-02T00:00:00Z", firstname="Quinton", lastname="Johnson",
                 notes="Change 'gone' to 'here'."),
    ]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows)
    monkeypatch.setattr(corrlib, "extract_provider", lambda: (
        FakeProvider(results=[ProviderResult(parsed={"edits": [
            {"find": "gone", "replace": "here", "instruction": "swap"}]})]),
        "fake-model"))
    settings = form_ws()

    report = run(tmp_path, settings, opener)

    assert not report.failed, report.failed
    assert report.corrected and report.corrected[0].startswith(SOURCE)
    placed = uploads_in(opener)
    assert "Johnson - Book 3.5.idml" in placed
    assert "Johnson - Book 3.5 - corrections.xlsx" in placed
    rec = WatchState.load(tmp_path / "state.json").get("idml-3")
    assert len(rec.corrections_submissions) == 2
    assert rec.corrections_pending_key == "form:quinton|johnson"
    # Exactly one Johnson record sits in the fake HubSpot, so its status moved
    # — even though nothing there was ever flagged "Ready for Corrections".
    assert opener.hubspot["hs-Johnson"]["properties"]["docproof"] == \
        "Corrections Applied"
    assert len(patches(opener)) == 1
    assert report.needs_human == []


def test_form_intake_with_two_matching_hubspot_records_delivers_but_leaves_hubspot_alone(
        tmp_path, monkeypatch):
    opener = make_opener(tmp_path, hubspot={
        "Johnson1": {"firstname": "Quinton", "lastname": "Johnson"},
        "Johnson2": {"firstname": "Quinton", "lastname": "Johnson"}})
    rows = [_form_row("2026-09-01T00:00:00Z", firstname="Quinton",
                      lastname="Johnson", files=FILE_URL)]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows)
    settings = form_ws()

    report = run(tmp_path, settings, opener)

    assert not report.failed, report.failed
    assert "Johnson - Book 3.5.idml" in uploads_in(opener)
    assert patches(opener) == []
    assert any("no single Projects record" in reason
              for _, reason in report.needs_human)


def test_form_intake_row_with_no_last_name_needs_a_person(tmp_path, monkeypatch):
    opener = make_opener(tmp_path, hubspot={})
    rows = [_form_row("2026-09-01T00:00:00Z", firstname="Quinton", lastname="",
                      files=FILE_URL)]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows)
    settings = form_ws()

    report = run(tmp_path, settings, opener)

    assert any("no first or last name" in reason
              for _, reason in report.needs_human)
    assert uploads_in(opener) == {}
    assert report.corrected == []


def test_form_intake_is_held_and_then_released_past_its_quiet_period(
        tmp_path, monkeypatch):
    opener = make_opener(tmp_path, hubspot={"Johnson": {"firstname": "Quinton",
                                                        "lastname": "Johnson"}})
    rows = [_form_row("2026-09-01T00:00:00Z", firstname="Quinton",
                      lastname="Johnson", files=FILE_URL)]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows)
    settings = form_ws(corrections_quiet_seconds=10800)

    held = run(tmp_path, settings, opener)
    assert held.corrected == []
    assert held.waiting >= 1
    assert uploads_in(opener) == {}
    state = WatchState.load(tmp_path / "state.json")
    assert len(state.corrections_pending) == 1
    entry = next(iter(state.corrections_pending.values()))
    assert entry.record_id == "form:quinton|johnson"
    assert entry.author == "Quinton Johnson"
    assert entry.first == "Quinton" and entry.last == "Johnson"
    assert len(entry.submissions) == 1

    # Back-date the hold past its own quiet period, the way waiting three
    # hours would — a test cannot wait three hours, so it moves the clock the
    # record's own state remembers instead.
    for pending in state.corrections_pending.values():
        pending.first_seen = pending.last_submission_at = \
            "2020-01-01T00:00:00+00:00"
    state.save()

    released = run(tmp_path, settings, opener)
    assert released.corrected and released.corrected[0].startswith(SOURCE)
    assert "Johnson - Book 3.5.idml" in uploads_in(opener)


def test_form_intake_a_second_row_resets_the_quiet_period(tmp_path, monkeypatch):
    opener = make_opener(tmp_path, hubspot={"Johnson": {"firstname": "Quinton",
                                                        "lastname": "Johnson"}})
    rows = [_form_row("2026-09-01T00:00:00Z", firstname="Quinton",
                      lastname="Johnson", files=FILE_URL)]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows)
    settings = form_ws(corrections_quiet_seconds=10800)

    run(tmp_path, settings, opener)
    state = WatchState.load(tmp_path / "state.json")
    key = next(iter(state.corrections_pending))
    assert len(state.corrections_pending[key].submissions) == 1
    # Back-dated so the two runs' clocks cannot land in the same second and
    # make the "it moved" assertion below a coin flip.
    state.corrections_pending[key].first_seen = "2020-01-01T00:00:00+00:00"
    state.corrections_pending[key].last_submission_at = \
        "2020-01-01T00:00:00+00:00"
    state.save()
    first_wait = corrlib.ready_at(state.corrections_pending[key], settings)

    # A second submission arrives before the hold ends.
    rows2 = rows + [_form_row("2026-09-02T00:00:00Z", firstname="Quinton",
                              lastname="Johnson", notes="one more thing")]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows2)
    run(tmp_path, settings, opener)
    state2 = WatchState.load(tmp_path / "state.json")
    assert len(state2.corrections_pending[key].submissions) == 2
    second_wait = corrlib.ready_at(state2.corrections_pending[key], settings)
    assert second_wait > first_wait


def test_form_intake_leaves_rounds_before_the_start_date_alone(
        tmp_path, monkeypatch):
    opener = make_opener(tmp_path, hubspot={"Johnson": {"firstname": "Quinton",
                                                        "lastname": "Johnson"}})
    opener.content["form-sub-1"] = submission(tmp_path)
    rows = [
        _form_row("2026-08-01T00:00:00Z", firstname="Quinton", lastname="Johnson",
                 files=("https://api.hubapi.com/files/form-sub-1"
                        "?filename=Old%20round.docx")),
        _form_row("2026-09-02T00:00:00Z", firstname="Quinton", lastname="Johnson",
                 notes="Change 'gone' to 'here'."),
    ]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows)
    monkeypatch.setattr(corrlib, "extract_provider", lambda: (
        FakeProvider(results=[ProviderResult(parsed={"edits": [
            {"find": "gone", "replace": "here", "instruction": "swap"}]})]),
        "fake-model"))
    settings = form_ws(corrections_form_start_after="2026-08-15")

    report = run(tmp_path, settings, opener)

    assert not report.failed, report.failed
    rec = WatchState.load(tmp_path / "state.json").get("idml-3")
    assert len(rec.corrections_submissions) == 1
    assert "Old round" not in rec.corrections_input_name


def test_form_intake_read_failure_needs_a_person_and_runs_nothing(
        tmp_path, monkeypatch):
    opener = make_opener(tmp_path, hubspot={"Johnson": {"firstname": "Quinton",
                                                        "lastname": "Johnson"}})

    def boom(*a, **k):
        raise HubSpotError("no forms scope on this token")

    monkeypatch.setattr(corrlib.hubspot, "form_submissions", boom)
    settings = form_ws()

    report = run(tmp_path, settings, opener)

    assert not report.failed, report.failed
    assert len(report.needs_human) == 1
    assert "could not read the form" in report.needs_human[0][1]
    assert uploads_in(opener) == {}


def test_form_intake_only_record_selects_the_group_by_name(tmp_path, monkeypatch):
    opener = make_opener(tmp_path, hubspot={"Johnson": {"firstname": "Quinton",
                                                        "lastname": "Johnson"}})
    rows = [_form_row("2026-09-01T00:00:00Z", firstname="Quinton",
                      lastname="Johnson", files=FILE_URL)]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows)
    settings = form_ws()
    state = WatchState(tmp_path / "state.json")
    report = ticklib.TickReport()

    works = corrlib.discover("hubspot-token", "drive-token", settings, state,
                             opener=opener, report=report,
                             only_record="Quinton Johnson")

    assert len(works) == 1
    assert works[0].first == "Quinton" and works[0].last == "Johnson"

    # A key or a HubSpot id name the same group just as well.
    state2 = WatchState(tmp_path / "state2.json")
    works_by_key = corrlib.discover(
        "hubspot-token", "drive-token", settings, state2, opener=opener,
        report=ticklib.TickReport(), only_record="form:quinton|johnson")
    assert len(works_by_key) == 1


def test_form_intake_pending_summary_carries_first_last_and_mode(
        tmp_path, monkeypatch):
    opener = make_opener(tmp_path, hubspot={"Johnson": {"firstname": "Quinton",
                                                        "lastname": "Johnson"}})
    rows = [_form_row("2026-09-01T00:00:00Z", firstname="Quinton",
                      lastname="Johnson", files=FILE_URL)]
    monkeypatch.setattr(corrlib.hubspot, "form_submissions",
                        lambda *a, **k: rows)
    settings = form_ws(corrections_quiet_seconds=10800)

    report = run(tmp_path, settings, opener)
    assert report.waiting

    state = WatchState.load(tmp_path / "state.json")
    rows2 = corrlib.pending_summary(state, settings)
    assert len(rows2) == 1
    row = rows2[0]
    assert row["first"] == "Quinton" and row["last"] == "Johnson"
    assert row["mode"] == "form"
