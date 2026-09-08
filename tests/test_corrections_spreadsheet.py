"""The two-sheet corrections spreadsheet: every correction the run received lands
on exactly one of "Applied" / "Not applied", with its page, and the summary
counts reconcile — driven end to end through `apply_corrections`, and as a pure
function over a hand-built payload."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from docproof.corrections.run import apply_corrections, verify_corrections
from docproof.corrections.spreadsheet import (HEADER, HEADER_ROW, SHEET_APPLIED,
                                              SHEET_NOT_APPLIED,
                                              write_spreadsheet)

FIXTURES = Path(__file__).parent / "fixtures"
LAYOUT = FIXTURES / "layout.idml"


def _sheets(path: Path):
    wb = load_workbook(path)
    return wb, wb.sheetnames


def _rows(ws) -> list[dict]:
    header = [c.value for c in ws[HEADER_ROW]]
    assert header == list(HEADER)
    out = []
    for r in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if r[0] is None:
            continue
        out.append(dict(zip(header, r)))
    return out


def _summary(ws) -> dict:
    return {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value
            for i in range(1, HEADER_ROW - 1)}


def test_a_typed_list_lands_every_edit_on_exactly_one_sheet(tmp_path):
    edits = [
        {"find": "was empty", "replace": "was vacant", "instruction": "vacant"},
        {"find": "not in the book at all", "replace": "x",
         "instruction": "fix this"},
    ]
    got = apply_corrections(LAYOUT, edits, tmp_path)
    assert got.spreadsheet is not None
    assert got.spreadsheet.name == "layout_corrections.xlsx"
    assert got.spreadsheet.exists()

    wb, names = _sheets(got.spreadsheet)
    assert names == [SHEET_APPLIED, SHEET_NOT_APPLIED]
    applied, not_applied = _rows(wb[SHEET_APPLIED]), _rows(wb[SHEET_NOT_APPLIED])
    assert [r["Edit id(s)"] for r in applied] == ["c1"]
    assert [r["Edit id(s)"] for r in not_applied] == ["c2"]
    assert applied[0]["Status"] == "Applied exactly"
    assert applied[0]["Find"] == "was empty"
    assert applied[0]["Replace"] == "was vacant"
    assert applied[0]["#"] == 1 and not_applied[0]["#"] == 1
    # The not-found row reads as a sentence, not a status code.
    assert not_applied[0]["Status"] == "The text to change was not found"
    assert "not_found" not in (not_applied[0]["Reason / detail"] or "")
    # A typed list carries no page, so none is invented.
    for r in applied + not_applied:
        assert r["Page (as marked)"] in (None, "")
        assert r["Page confirmed"] == "—"
    # The summary block reconciles on both sheets.
    for ws in (wb[SHEET_APPLIED], wb[SHEET_NOT_APPLIED]):
        s = _summary(ws)
        assert s["Total corrections received"] == 2
        assert s["Applied"] == 1 and s["Not applied"] == 1
        assert s["Applied"] + s["Not applied"] == s["Total corrections received"]
        assert s["Source file"] == "layout.idml"
        assert s["Corrected file"] == "layout_corrected.idml"
        assert ws.freeze_panes == f"A{HEADER_ROW + 1}"
        assert ws.auto_filter.ref.startswith(f"A{HEADER_ROW}:")


def test_a_marked_proof_lands_every_comment_on_exactly_one_sheet(tmp_path):
    comments = [
        {"id": "p1-1", "page": 1, "kind": "highlight",
         "instruction": "vacant", "anchor": "empty"},
        {"id": "p1-2", "page": 1, "kind": "note",
         "instruction": "fix this", "anchor": "somewhere"},
        {"id": "p1-3", "page": 2, "kind": "note",
         "instruction": "please check the running head", "anchor": ""},
    ]
    edits = [
        {"find": "was empty", "replace": "was vacant",
         "instruction": "vacant", "source": "p1-1", "page": 1},
        {"find": "not in the book at all", "replace": "x",
         "instruction": "fix this", "source": "p1-2", "page": 1},
    ]
    got = apply_corrections(LAYOUT, edits, tmp_path, comments=comments)
    wb, names = _sheets(got.spreadsheet)
    assert names == [SHEET_APPLIED, SHEET_NOT_APPLIED]
    applied, not_applied = _rows(wb[SHEET_APPLIED]), _rows(wb[SHEET_NOT_APPLIED])
    ids = sorted(r["Comment id"] for r in applied + not_applied)
    assert ids == ["p1-1", "p1-2", "p1-3"]          # each comment exactly once
    assert [r["Comment id"] for r in applied] == ["p1-1"]
    by = {r["Comment id"]: r for r in not_applied}
    assert by["p1-2"]["Status"] == "The text to change was not found"
    assert by["p1-3"]["Status"] == "Could not be turned into an edit"
    assert applied[0]["Marked text"] == "empty"
    assert applied[0]["Correction (as written)"] == "vacant"
    assert applied[0]["Edit id(s)"] == "c1"
    # A page was marked, so it is carried — and only called confirmed when the
    # engine aligned it to an InDesign folio.
    for r in applied + not_applied:
        assert r["Page (as marked)"] in (1, 2)
        assert r["Page confirmed"] in ("Yes", "No")
        assert r["Page in IDML"] not in (None, "")
    payload = json.loads(got.report_json.read_text(encoding="utf-8"))
    s = _summary(wb[SHEET_APPLIED])
    assert s["Total corrections received"] == payload["comments"]["total"] == 3
    assert s["Applied"] == 1 and s["Not applied"] == 2
    # The JSON gained the itemised applied outcomes the sheet reads, additively.
    assert [o["id"] for o in payload["apply"]["applied_items"]] == ["c1"]
    assert payload["apply"]["applied_items"][0]["page"] == 1
    assert payload["apply"]["applied"] == 1


def test_verify_mode_writes_the_sheet_too(tmp_path):
    edits = [{"find": "was empty", "replace": "was vacant"}]
    first = apply_corrections(LAYOUT, edits, tmp_path / "apply")
    got = verify_corrections(LAYOUT, first.corrected_idml, edits,
                             tmp_path / "verify")
    assert got.spreadsheet is not None and got.spreadsheet.exists()
    wb, _ = _sheets(got.spreadsheet)
    applied = _rows(wb[SHEET_APPLIED])
    assert [r["Edit id(s)"] for r in applied] == ["c1"]
    assert applied[0]["Status"] == "Applied exactly"
    assert _rows(wb[SHEET_NOT_APPLIED]) == []


def _payload(**over) -> dict:
    base = {
        "generated_at": "2026-09-08T12:00:00+00:00",
        "source": "/x/book.idml", "source_name": "book.idml",
        "after": "/x/book_corrected.idml", "mode": "apply",
        "parse": {"edits": 3, "issues": []},
        "apply": {
            "applied": 1,
            "applied_items": [
                {"id": "c1", "status": "applied", "find": "teh", "replace": "the",
                 "instruction": "typo", "detail": "", "page": 7,
                 "page_label": "vii", "source": "p7-1"}],
            "flagged": [
                {"id": "c2", "status": "not_found", "find": "gone",
                 "replace": "here", "instruction": "swap", "detail": "",
                 "page": 9, "source": "p9-1"}],
            "no_op": [
                {"id": "c3", "status": "no_op", "find": "same",
                 "replace": "same", "instruction": "stet", "detail": "",
                 "page": 0, "source": ""}],
        },
        "verify": {"reconciliations": [
            {"id": "c1", "status": "applied_exactly", "detail": ""}],
            "discrepancies": []},
        "comments": {"total": 0, "unresolved": 0, "items": []},
        "pages": {"placed": 3, "total": 4, "labeled": 1, "cited": 2},
    }
    base.update(over)
    return base


def test_pure_write_over_a_payload_places_pages_and_labels(tmp_path):
    out = write_spreadsheet(_payload(), tmp_path / "book_corrections.xlsx",
                            corrected_name="book_corrected.idml")
    wb, names = _sheets(out)
    assert names == [SHEET_APPLIED, SHEET_NOT_APPLIED]
    applied, not_applied = _rows(wb[SHEET_APPLIED]), _rows(wb[SHEET_NOT_APPLIED])
    assert len(applied) == 1 and len(not_applied) == 2
    a = applied[0]
    # A labelled page: the folio is what InDesign shows, and it is confirmed.
    assert (a["Page (as marked)"], a["Page in IDML"], a["Page confirmed"]) == \
        (7, "vii", "Yes")
    assert a["Status"] == "Applied exactly"
    assert a["Comment id"] == "p7-1"
    by = {r["Edit id(s)"]: r for r in not_applied}
    # A page with no label stands as the physical page, unconfirmed.
    assert (by["c2"]["Page (as marked)"], by["c2"]["Page in IDML"],
            by["c2"]["Page confirmed"]) == (9, 9, "No")
    assert by["c2"]["Status"] == "The text to change was not found"
    # No page at all: nothing invented.
    assert by["c3"]["Page (as marked)"] in (None, "")
    assert by["c3"]["Page in IDML"] in (None, "")
    assert by["c3"]["Page confirmed"] == "—"
    assert by["c3"]["Status"] == "No change needed"
    s = _summary(wb[SHEET_NOT_APPLIED])
    assert s["Total corrections received"] == 3
    assert s["Applied"] == 1 and s["Not applied"] == 2
    assert s["Pages placed / total"] == "3 / 4"
    assert s["Corrected file"] == "book_corrected.idml"
    assert s["Generated"] == "2026-09-08 12:00:00 UTC"


def test_pure_write_over_comments_uses_the_comment_as_the_unit(tmp_path):
    items = [
        {"id": "p7-1", "page": 7, "page_label": "vii", "kind": "highlight",
         "instruction": "typo", "anchor": "teh", "disposition": "applied",
         "edit_ids": ["c1"], "detail": ""},
        {"id": "p9-1", "page": 9, "kind": "highlight", "instruction": "swap",
         "anchor": "gone", "disposition": "flagged", "edit_ids": ["c2"],
         "detail": "not found on page 9"},
        {"id": "p9-2", "page": 9, "kind": "note", "instruction": "stet",
         "anchor": "same", "disposition": "no_op", "edit_ids": ["c3"],
         "detail": "correct as set"},
        {"id": "p10-1", "page": 10, "kind": "note", "instruction": "??",
         "anchor": "", "disposition": "not_extracted", "edit_ids": [],
         "detail": ""},
    ]
    payload = _payload(comments={"total": 4, "unresolved": 2, "items": items})
    payload["verify"]["reconciliations"][0]["status"] = "deviates"
    payload["verify"]["reconciliations"][0]["detail"] = "landed as 'The'"
    out = write_spreadsheet(payload, tmp_path / "s.xlsx")
    wb, _ = _sheets(out)
    applied, not_applied = _rows(wb[SHEET_APPLIED]), _rows(wb[SHEET_NOT_APPLIED])
    assert [r["Comment id"] for r in applied] == ["p7-1"]
    assert applied[0]["Status"] == "Applied (differs from what was asked)"
    assert "landed as" in applied[0]["Reason / detail"]
    assert applied[0]["Find"] == "teh" and applied[0]["Replace"] == "the"
    assert applied[0]["Page in IDML"] == "vii"
    by = {r["Comment id"]: r for r in not_applied}
    assert set(by) == {"p9-1", "p9-2", "p10-1"}
    assert by["p9-1"]["Status"] == "The text to change was not found"
    assert by["p9-1"]["Reason / detail"] == "not found on page 9"
    assert by["p9-2"]["Status"] == "No change needed"
    assert by["p9-2"]["Reason / detail"] == "correct as set"
    assert by["p10-1"]["Status"] == "Could not be turned into an edit"
    assert by["p10-1"]["Edit id(s)"] in (None, "")
    assert by["p10-1"]["Page confirmed"] == "No"
    s = _summary(wb[SHEET_APPLIED])
    assert s["Total corrections received"] == 4 == payload["comments"]["total"]
    assert s["Applied"] + s["Not applied"] == 4
