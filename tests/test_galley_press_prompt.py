"""Source completeness, concrete reading aids, and protected source formatting."""
import hashlib
import json

import pytest
from docx import Document

from docproof.models import ParagraphRef
from docproof.utils.xml_helpers import DocxPackage
from galley import press_prompt
from galley.press_checks import focused_checks, source_formatting, current_formatting, citation_context


def para(pid, text, *, part="word/document.xml", location="body", style="Normal"):
    return ParagraphRef(pid, part, location, text, style)


def test_every_source_rule_paragraph_and_dialogue_table_cell_is_accounted_for():
    source = json.loads(press_prompt.SOURCE.read_text())
    coverage = json.loads(press_prompt.COVERAGE.read_text())
    units = press_prompt.source_units(source)
    assert len(source["rules"]) == source["rule_count"] == 120
    assert len(units) == coverage["source_unit_count"] == 149
    assert coverage["source_sha256"] == hashlib.sha256(press_prompt.SOURCE.read_bytes()).hexdigest()
    for key in ("usage", "note", "default_trigger"):
        assert coverage["wrapper_metadata"][key]["source"] == source[key]
        assert coverage["wrapper_metadata"][key]["handling"]
    rows = {r["source_id"]: r for r in coverage["coverage"]}
    assert len(rows) == len(coverage["coverage"])
    assert set(rows) == set(units)
    for rule in source["rules"]:
        assert units[rule["id"]] == rule["text"]
    for key, text in units.items():
        row = rows[key]
        assert row["source_text_sha256"] == hashlib.sha256(text.encode()).hexdigest()
        assert row["owners"] and row["handling"] and row["disposition"]
        assert set(row["policy_sections"]) <= set(press_prompt.EDITORIAL_RULES)
    assert len([k for k in rows if "/table-1/row-" in k]) == 5


def test_extraction_preserves_exceptions_and_resolves_conflicts_explicitly():
    text = press_prompt.editorial_policy()
    source = json.loads(press_prompt.SOURCE.read_text())
    assert len(text + press_prompt.FRONTIER_TASK + press_prompt.STORY_TASK) < len(source["prompt_text"]) / 2
    for required in ("Oxford", "Macquarie", "Canadian", "3:00 p.m.", "3:00 pm", "no invented",
                     "NBSP", "interrobangs", "verbatim", "BOTH comma and lowercase",
                     "Never lowercase I", "She continued typing", "multi-paragraph",
                     "scientific truths", "internally consistent scene", "No majority"):
        if required == "No majority":
            assert "Frequency is not authorial intent" in text
        else:
            assert required.casefold() in text.casefold()
    assert "silent exceptions\ndo not apply" in text
    assert "never the poem's STRUCTURE" in text and "receives house mechanics only" in text
    assert "never mechanically convert" in text
    assert "reviewed_check_ids" in press_prompt.FRONTIER_TASK


@pytest.mark.parametrize("mark", [",", ".", "?", "!", "…"])
@pytest.mark.parametrize("subject", ["he", "He", "she", "She", "they", "They", "we", "We", "it", "It", "you", "You"])
@pytest.mark.parametrize("reverse", [False, True])
def test_every_dialogue_matrix_combination_is_scripted(mark, subject, reverse):
    seam = '”' + mark if reverse else mark + '”'
    evidence = focused_checks([para("p", '“Wait' + seam + ' ' + subject + ' whispered.')])
    sites = [s for s in evidence["sites"] if s["check"] == "dialogue_matrix"]
    assert len(sites) == 1
    assert sum(evidence["dialogue_matrix"].values()) == 1
    assert ("quote_then_mark" if reverse else "mark_then_quote") in sites[0]["detail"]


def test_dialogue_keeps_action_beats_as_judgment_and_excludes_I():
    evidence = focused_checks([para("p", '“Wait.” She continued typing. “Yes,” I said.')])
    assert len([s for s in evidence["sites"] if s["check"] == "dialogue_matrix"]) == 1
    assert "action beat" in evidence["sites"][0]["detail"]


def test_serial_inventory_includes_valid_lists_nor_and_nonlists_without_editing():
    texts = ["Red, white and blue.", "Red, white, and blue.", "She did not sing, dance nor play.",
             "Son, take the gun and shoot.", "She ran and jumped."]
    evidence = focused_checks([para(str(i), t) for i, t in enumerate(texts)])
    serial = [s for s in evidence["sites"] if s["check"] == "serial_comma"]
    assert [s["para_id"] for s in serial] == ["0", "1", "2", "3"]
    assert all("replacement" not in s for s in serial)


def test_profile_includes_contiguous_scenes_and_strips_dialogue_in_notes():
    paragraphs = [para("past", "She walked and waited. He stood and watched. They knew and heard.")]
    paragraphs += [para(str(i), "She walks and looks. He waits and watches. She turns and smiles.") for i in range(2)]
    paragraphs += [para("note", '“She walks, runs, smiles and waits,” he said.',
                       part="word/footnotes.xml", location="footnote")]
    evidence = focused_checks(paragraphs)
    note = next(r for r in evidence["tense_profile"]["paragraphs"] if r["para_id"] == "note")
    assert note["present"] == 0
    assert len(evidence["tense_profile"]["paragraphs"]) == 4
    assert all(s["para_id"] in {"past", "0", "1", "note"} for s in evidence["sites"])


def test_reference_context_keeps_both_sides_and_declares_pattern_limitations():
    paragraphs = [para("cite", "Smith (2019) found the effect. See chapter 2."),
                  para("head", "References", style="Heading1"),
                  para("ref", "Smith, A. (2019). The Book."),
                  para("chapter", "Chapter 2", style="Heading1")]
    value = citation_context(paragraphs)
    assert {p["id"] for p in value["paragraphs"]} == {p.para_id for p in paragraphs}
    assert next(p for p in value["paragraphs"] if p["id"] == "ref")["reference_section"]
    assert not value["complete_reference_inventory"]


def test_formatting_does_not_guess_away_inherited_italics(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("The Book").italic = True
    p.add_run(" is here.")
    p2 = doc.add_paragraph("Another title")
    doc.styles["Normal"].font.italic = True
    p2.runs[0].italic = False
    source = tmp_path / "format.docx"
    doc.save(source)
    marks = list(source_formatting(DocxPackage(source)).values())
    assert marks[0][:8] == [True] * 8
    assert all(m is None for m in marks[0][8:])
    assert all(m is False for m in marks[1])


def test_current_formatting_tracks_approved_title_and_marks_changed_text_unknown():
    original = {"p": "The Book was here."}
    current = {"p": "The Book is here."}
    approved = [{"para_id": "p", "snapshot": original["p"], "format": "italic", "start": 0, "end": 8}]
    ranges = current_formatting(original, current, {"p": [False] * len(original["p"])}, approved)["p"]
    assert ranges[0] == {"start": 0, "end": 8, "italic": True}
    assert any(r["italic"] is None for r in ranges)


def test_changed_extracted_policy_cannot_resume_old_recipe(tmp_path, monkeypatch):
    from galley.fixed_workflow import FixedWorkflow, FixedWorkflowError
    book = Document()
    book.add_paragraph("A quiet page.")
    source = tmp_path / "source.docx"
    book.save(source)
    FixedWorkflow(source, tmp_path / "run", calls=object())
    monkeypatch.setattr(press_prompt, "FRONTIER_TASK", press_prompt.FRONTIER_TASK + "\nNew rule.")
    with pytest.raises(FixedWorkflowError, match="recipe changed"):
        FixedWorkflow(source, tmp_path / "run", calls=object())


def test_seam_hyphen_sites_need_a_dictionary_and_an_unknown_part():
    known = {"the", "road", "county", "well", "known", "fact"}
    knows = lambda w: w.lower() in known
    paragraphs = [para("a", "The Cala-veras road."), para("b", "Calaveras County."),
                  para("c", "A well-known fact."), para("d", "The Tusca-loosa road.")]
    focused = focused_checks(paragraphs, knows=knows)
    seams = [s for s in focused["sites"] if s["check"] == "seam_hyphen"]
    assert [(s["para_id"], s["quote"]) for s in seams] == [("a", "Cala-veras"), ("d", "Tusca-loosa")]
    assert "occurs 1 time(s) unhyphenated" in seams[0]["detail"]
    assert "neither 'Tusca' nor 'loosa'" in seams[1]["detail"] and focused["counts"]["seam_hyphen"] == 2
    without = focused_checks(paragraphs)
    assert without["counts"]["seam_hyphen"] == 0 and not any(s["check"] == "seam_hyphen" for s in without["sites"])


def test_book_map_is_a_complete_inventory():
    from galley.press_checks import book_map
    paragraphs = [para("h1", "CHAPTER ONE", style="Heading1"), para("p1", "Body one."), para("p2", "Body two."),
                  para("h2", "ACKNOWLEDGEMENTS"), para("p3", "Thanks."),
                  para("r1", "12 | ANA AND ATLAS", part="word/header1.xml", location="header"),
                  para("r2", "", part="word/header2.xml", location="header")]
    result = book_map(paragraphs, lambda style: style.startswith("Heading"))
    assert result["complete_inventory"] is True
    assert [(h["id"], h["signal"], h["body_paragraphs"]) for h in result["headings"]] == [
        ("h1", "style", 2), ("h2", "caps_line", 1)]
    assert [h["id"] for h in result["headers_footers"]] == ["r1"]
    assert result["total_body_paragraphs"] == 3


def test_walkthrough_and_continuity_prompts_are_bound_into_the_recipe_identity(monkeypatch):
    before = press_prompt.policy_identity()
    for name in ("FINAL_WALKTHROUGH", "FINAL_WALKTHROUGH_CHECK", "CONTINUITY_TASK"):
        monkeypatch.setattr(press_prompt, name, getattr(press_prompt, name) + "\nNew rule.")
        assert press_prompt.policy_identity() != before
        monkeypatch.undo()
        assert press_prompt.policy_identity() == before


# --- where the chapters start and stop ----------------------------------------

def test_matter_regions_splits_front_body_and_back():
    """The gate waives a placeholder outside the chapters, so it has to know
    where they are: a copyright-page credit and an author biography are matter,
    a hole in chapter two is not."""
    from galley.press_checks import matter_regions
    paragraphs = [para("t1", "The Spies From Camp X"), para("c1", "Cover design by XXX"),
                  para("h1", "CHAPTER ONE", style="Heading1"), para("p1", "Body one."),
                  para("h2", "CHAPTER TWO", style="Heading1"), para("p2", "Body two."),
                  para("h3", "ACKNOWLEDGEMENTS", style="Heading1"), para("p3", "Thanks."),
                  para("r1", "12 | CAMP X", part="word/header1.xml", location="header")]
    regions = matter_regions(paragraphs, lambda style: style.startswith("Heading"))
    assert regions["front"] == {"t1", "c1", "r1"}
    assert regions["body"] == {"h1", "p1", "h2", "p2"}
    assert regions["back"] == {"h3", "p3"}


def test_matter_regions_keeps_an_authors_note_that_opens_the_book_in_front():
    """Position decides: the same heading is front matter before the chapters
    and back matter after them."""
    from galley.press_checks import matter_regions
    front_first = [para("h0", "AUTHOR'S NOTE", style="Heading1"), para("p0", "A word first."),
                   para("h1", "CHAPTER ONE", style="Heading1"), para("p1", "Body.")]
    regions = matter_regions(front_first, lambda style: style.startswith("Heading"))
    assert regions["front"] == {"h0", "p0"} and regions["back"] == set()
    assert regions["body"] == {"h1", "p1"}


def test_a_book_with_no_findable_headings_is_all_body():
    """The stricter reading when the structure is unknown: nothing is waived."""
    from galley.press_checks import matter_regions
    paragraphs = [para("p1", "Just prose."), para("p2", "More prose.")]
    regions = matter_regions(paragraphs, lambda style: False)
    assert regions["body"] == {"p1", "p2"}
    assert regions["front"] == set() and regions["back"] == set()
