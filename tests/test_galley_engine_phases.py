"""Code-owned phases preserve coverage, residuals, and restart identity."""
from __future__ import annotations

import json
from types import SimpleNamespace

import pytest

from galley.engine_phases import EnginePhaseError, EnginePhases, review_action


def _driver(tmp_path):
    import docx
    import shutil

    (tmp_path / "runs").mkdir(exist_ok=True)
    (tmp_path / "runs" / "mech.yaml").write_text("api: {model: claude-opus-5}\n")
    book = tmp_path / "source.docx"
    document = docx.Document()
    document.add_paragraph("The book was quiet.")
    document.save(book)
    run = tmp_path / "runs" / "final"
    run.mkdir(exist_ok=True)
    shutil.copyfile(book, run / "book - Atmosphere Press Proofreader.docx")
    (run / "findings.json").write_text(json.dumps({"source": str(book), "findings": []}))
    return SimpleNamespace(
        workspace=tmp_path, book=book, budget_usd=10, review_rounds=2,
        review_calls=40, _engine_env={}, _final_run=lambda: run,
        _resource_env=lambda phase, env: dict(env), timeout_for=lambda phase: 60)


def _complete_verification(phases, monkeypatch):
    from docproof.__main__ import main
    from galley import verify
    from galley.engine_phases import PASS_IDS, POLICY_ID, READER_MODELS
    from tests.galley.test_verification_checkpoints import Provider

    monkeypatch.setattr(verify, "UNREAD", [])
    monkeypatch.setattr(verify, "UNREAD_BATCHES", [])
    monkeypatch.setattr(verify, "_LOSSES", [])
    # Validate the same provider identity that issued the fake structured reads.
    monkeypatch.setattr("docproof.providers.subagent.SubagentProvider",
                        lambda *args, **kwargs: Provider())
    run = phases.driver._final_run()
    for pass_id, model in READER_MODELS.items():
        output = run if pass_id == "primary" else run / "verification" / pass_id
        assert main([
            "galley", "verify", str(run), "--out", str(output),
            "--config", str(phases._config()), "--engine", "subagent", "--model", model,
            "--verification-pass", pass_id, "--verification-policy", POLICY_ID,
            "--required-verification-passes", ",".join(PASS_IDS),
            "--context", str(phases._context(pass_id)),
        ]) == 0
    return run


@pytest.mark.parametrize("damage", ["missing_pass", "incomplete", "policy", "model", "config"])
def test_full_coverage_requires_both_complete_reads_under_the_current_policy(
        tmp_path, monkeypatch, damage):
    phases = EnginePhases(_driver(tmp_path), lambda spec: None)
    run = _complete_verification(phases, monkeypatch)
    assert phases._coverage(run, require_full_passes=True)
    path = run / "verification" / "type-compare" / "finished_walk.json"
    if damage == "missing_pass":
        path.unlink()
    else:
        payload = json.loads(path.read_text())
        if damage == "incomplete":
            payload["verification_provenance"]["complete"] = False
        elif damage == "policy":
            payload["verification_provenance"]["identity"]["policy"]["pass_id"] = "primary"
        elif damage == "model":
            payload["model"] = "some-other-reviewer"
            payload["verification_provenance"]["identity"]["model"] = "some-other-reviewer"
        elif damage == "config":
            phases._config().write_text("api: {model: changed-model}\n")
        path.write_text(json.dumps(payload))
    assert not phases._coverage(run, require_full_passes=True)


def test_delta_coverage_counts_skipped_display_paragraphs_in_the_accepted_book(tmp_path, monkeypatch):
    import docx
    import shutil
    from docproof.config import load_config
    from galley.settle import Settler, _prepare_source
    from galley.verify import VerifyRunResult

    driver = _driver(tmp_path)
    document = docx.Document()
    document.add_paragraph("The Book Title", style="Title")
    document.add_paragraph("The book was quiet.")
    document.save(driver.book)
    shutil.copyfile(driver.book, driver._final_run() / "book - Atmosphere Press Proofreader.docx")
    phases = EnginePhases(driver, lambda spec: None)
    run = _complete_verification(phases, monkeypatch)
    cfg = load_config(str(phases._config()))
    settler = Settler(run, cfg=cfg, manuscript=driver.book, error_dir="config/error_types")
    settler._source = _prepare_source(cfg, driver.book, "config/error_types")[0]
    assert list(settler._source.values()) == ["The book was quiet."]
    # A real delta merge over the prose must preserve the whole accepted
    # manuscript's coverage count, including the unedited Title paragraph.
    settler._merge_coverage(VerifyRunResult([], [], True, True), list(settler._source))
    assert json.loads((run / "finished_walk.json").read_text())["paragraphs"] == 2
    assert phases._coverage(run)


@pytest.mark.parametrize("coverage,ids,prior,rounds,affected,expected", [
    (False, set(), set(), 0, 0, "blocked"),
    (False, {"r-1"}, set(), 1, 1, "blocked"),
    (True, set(), set(), 0, 0, "final_review"),
    (True, {"r-1"}, set(), 0, 1, "local_repair"),
    (True, {"r-1"}, set(), 0, 20, "broad_repair"),
    (True, {"r-1"}, {"r-1"}, 1, 1, "final_review_with_residuals"),
    (True, {"r-2"}, {"r-1"}, 2, 1, "final_review_with_residuals"),
])
def test_review_action_never_calls_small_or_exhausted_residuals_clean(
        coverage, ids, prior, rounds, affected, expected):
    assert review_action(
        coverage_complete=coverage, open_ids=ids, prior_ids=prior,
        rounds=rounds, max_rounds=2, affected_paragraphs=affected,
        reviewed_paragraphs=100) == expected


def test_verification_preserves_both_complete_named_reads(tmp_path, monkeypatch):
    phases = EnginePhases(_driver(tmp_path), lambda spec: None)
    calls = []

    def command(*args, **kwargs):
        calls.append((args, kwargs))
        for output in args[3]:
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text('{"fixture": true}')

    monkeypatch.setattr(phases, "_command", command)
    monkeypatch.setattr(phases, "_coverage", lambda *args, **kwargs: True)
    phases.verify()
    assert len(calls) == 2
    for (phase, name, args, outputs), kwargs in calls:
        assert phase == "verify"
        assert args[args.index("--required-verification-passes") + 1] == "primary,type-compare"
        assert args[args.index("--verification-policy") + 1] == "mechanical-verification-v1"
        assert "--paragraphs" not in args
        assert len(outputs) == 2
        assert kwargs["acceptable"] == (0, 1)
    assert calls[0][0][2][calls[0][0][2].index("--verification-pass") + 1] == "primary"
    assert calls[1][0][2][calls[1][0][2].index("--verification-pass") + 1] == "type-compare"
    assert "Independent slow comparison" in (
        phases.directory / "type-compare-context.txt").read_text(encoding="utf-8")
    run = phases.driver._final_run()
    assert phases._initial_coverage_saved(run)
    # Current primary artifacts may change after repairs; the archived full
    # reads must remain intact and bound to the original receipt.
    (run / "finished_walk.json").write_text('{"delta": true}')
    assert phases._initial_coverage_saved(run)
    receipt = json.loads((phases.directory / "initial-coverage.json").read_text())
    archived = phases.directory / next(iter(receipt["artifacts"]))
    archived.write_text('{"changed": true}')
    assert not phases._initial_coverage_saved(run)


def test_incomplete_read_receipt_can_resume_and_complete_the_same_operation(tmp_path):
    driver = _driver(tmp_path)
    output = driver._final_run() / "finished_walk.json"
    calls = []

    def execute(spec):
        calls.append(spec)
        output.write_text(json.dumps({"complete": len(calls) == 2}))
        return SimpleNamespace(returncode=1 if len(calls) == 1 else 0, limit="", tail="")

    phases = EnginePhases(driver, execute)
    kwargs = {"acceptable": (0, 1),
              "validate": lambda: json.loads(output.read_text())["complete"]}
    with pytest.raises(EnginePhaseError):
        phases._command("verify", "recover-read", ["verify"], [output], **kwargs)
    receipt_path = phases.directory / "recover-read.json"
    assert json.loads(receipt_path.read_text())["status"] == "incomplete"
    resumed = EnginePhases(driver, execute)
    resumed._command("verify", "recover-read", ["verify"], [output], **kwargs)
    assert json.loads(receipt_path.read_text())["status"] == "completed"
    assert len(calls) == 2
    resumed._command("verify", "recover-read", ["verify"], [output], **kwargs)
    assert len(calls) == 2


def test_completed_command_is_reused_only_while_outputs_match(tmp_path):
    output = tmp_path / "runs" / "evidence.json"
    calls = []

    def execute(spec):
        calls.append(spec)
        output.write_text('{"ran": true}')
        return SimpleNamespace(returncode=0, limit="", tail="")

    phases = EnginePhases(_driver(tmp_path), execute)
    phases._command("audit", "audit-check", ["audit"], [output])
    phases._command("audit", "audit-check", ["audit"], [output])
    assert len(calls) == 1
    output.write_text('{"ran": false}')
    phases._command("audit", "audit-check", ["audit"], [output])
    assert json.loads(output.read_text()) == {"ran": True}
    assert len(calls) == 1                   # exact saved command output restored


def test_command_failure_or_missing_output_cannot_create_completion_receipt(tmp_path):
    output = tmp_path / "runs" / "evidence.json"
    result = SimpleNamespace(returncode=0, limit="timeout", tail="stopped")
    phases = EnginePhases(_driver(tmp_path), lambda spec: result)
    with pytest.raises(EnginePhaseError):
        phases._command("audit", "failed-audit", ["audit"], [output])
    receipt = json.loads((phases.directory / "failed-audit.json").read_text())
    assert receipt["status"] != "completed"
    result.limit = ""
    with pytest.raises(EnginePhaseError):
        phases._command("audit", "missing-audit", ["audit"], [output])
    receipt = json.loads((phases.directory / "missing-audit.json").read_text())
    assert receipt["status"] != "completed"


@pytest.mark.parametrize("changed_input", ["context", "findings"])
def test_completed_verification_cannot_be_reused_after_its_evidence_changes(
        tmp_path, changed_input):
    driver = _driver(tmp_path)
    run = driver._final_run()
    context = tmp_path / "voice-notes.md"
    context.write_text("The speaker uses deliberate fragments.")
    output = run / "finished_walk.json"

    calls = []
    def execute(spec):
        calls.append(spec)
        output.write_text('{"ran": true}')
        return SimpleNamespace(returncode=0, limit="", tail="")

    phases = EnginePhases(driver, execute)
    args = ["verify", str(run), "--context", str(context)]
    phases._command("verify", "verify-evidence", args, [output])
    if changed_input == "context":
        context.write_text("The speaker uses deliberate fragments and invented names.")
    else:
        (run / "findings.json").write_text(json.dumps({
            "source": str(driver.book), "findings": [{"finding_id": "new-evidence"}]}))
    phases._command("verify", "verify-evidence", args, [output])
    assert len(calls) == 2                  # new evidence triggers a fresh check
    assert len(list((phases.directory / "attempts" / "verify-evidence").glob("*.json"))) == 1


def test_missing_secondary_read_cannot_advance_settlement(tmp_path, monkeypatch):
    driver = _driver(tmp_path)
    run = driver._final_run()
    # Old, incomplete primary files must not stand in for the required two reads.
    for name in ("change_verify.json", "finished_walk.json", "settlement.json"):
        (run / name).write_text('{"ran": false}')
    attempts = []
    def incomplete(spec):
        attempts.append(spec)
        return SimpleNamespace(returncode=1, limit="", tail="read interrupted")
    phases = EnginePhases(driver, incomplete)
    advanced = []
    monkeypatch.setattr(phases, "_advance", advanced.append)
    monkeypatch.setattr("galley.settle.open_items", lambda run: [])
    monkeypatch.setattr("galley.verify.accepted_text", lambda run: {"body-0000": "Text."})
    monkeypatch.setattr("galley.verify.build_fingerprints", lambda run: {"build": "current"})
    with pytest.raises(EnginePhaseError) as error:
        phases.settle()
    assert error.value.retryable and error.value.kind == "incomplete"
    assert len(attempts) == 1               # recover missing reading before advancing
    assert advanced == []


def test_exhausted_round_cap_cannot_buy_an_extra_round_for_missing_settlement(tmp_path, monkeypatch):
    from galley.manifest import sha256_file

    driver = _driver(tmp_path)
    run = driver._final_run()
    phases = EnginePhases(driver, lambda spec: pytest.fail("round cap must prevent execution"))
    identity = {"source": sha256_file(driver.book), "run": str(run.resolve()),
                "config": sha256_file(phases._config())}
    (phases.directory / "review-loop.json").write_text(json.dumps({
        "schema_version": 1, "identity": identity, "rounds": 2,
        "prior_ids": ["r-old"], "history": []}))
    monkeypatch.setattr(phases, "_coverage", lambda *args, **kwargs: True)
    monkeypatch.setattr(phases, "_initial_coverage_saved", lambda run: True)
    monkeypatch.setattr(phases, "_advance", lambda state: pytest.fail("missing evidence cannot advance"))
    monkeypatch.setattr("galley.settle.open_items", lambda run: [
        SimpleNamespace(id="r-new", para_id="body-0000")])
    with pytest.raises(EnginePhaseError):
        phases.settle()


def test_unchanged_residuals_stop_repair_and_remain_visible_for_final_review(tmp_path, monkeypatch):
    driver = _driver(tmp_path)
    run = driver._final_run()
    phases = EnginePhases(driver, lambda spec: None)
    commands, advanced = [], []

    def command(phase, name, arguments, outputs, **kwargs):
        commands.append(name)
        for output in outputs:
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text('{"fixture": true}')

    monkeypatch.setattr(phases, "_command", command)
    monkeypatch.setattr(phases, "_coverage", lambda *args, **kwargs: True)
    monkeypatch.setattr(phases, "_advance", advanced.append)
    monkeypatch.setattr("galley.settle.open_items", lambda run: [
        SimpleNamespace(id="r-persistent", para_id="body-0000")])
    monkeypatch.setattr("galley.verify.accepted_text", lambda run: {
        f"body-{i:04d}": "Paragraph." for i in range(100)})
    phases.verify()
    commands.clear()
    phases.settle()
    state = json.loads((phases.directory / "review-loop.json").read_text())
    assert commands == ["settle-1"]
    assert state["rounds"] == 1
    assert state["history"][-1]["action"] == "final_review_with_residuals"
    assert state["history"][-1]["open_ids"] == ["r-persistent"]
    assert advanced == ["settled"]


def test_interrupted_broad_reread_finishes_before_resumed_settlement_can_advance(tmp_path, monkeypatch):
    driver = _driver(tmp_path)
    commands, advanced, events = [], [], []
    interrupted = True
    state_path = driver.workspace / "runs" / "driver" / "engine" / "review-loop.json"

    def command(phase, name, arguments, outputs, **kwargs):
        commands.append(name)
        for output in outputs:
            output.write_text('{"fixture": true}')

    def reread(*, cycle=0):
        nonlocal interrupted
        events.append(("verify", cycle))
        assert json.loads(state_path.read_text())["pending_verify_cycle"] == cycle
        if interrupted:
            interrupted = False
            raise EnginePhaseError("Primary completed; independent second read interrupted")

    def open_items(run):
        events.append(("open_items", None))
        return [SimpleNamespace(id="r-persistent", para_id="body-0000")]

    def coordinator():
        phases = EnginePhases(driver, lambda spec: None)
        monkeypatch.setattr(phases, "_command", command)
        monkeypatch.setattr(phases, "_coverage", lambda *args, **kwargs: True)
        monkeypatch.setattr(phases, "_initial_coverage_saved", lambda run: True)
        monkeypatch.setattr(phases, "_advance", advanced.append)
        monkeypatch.setattr(phases, "verify", reread)
        return phases

    monkeypatch.setattr("galley.settle.open_items", open_items)
    with pytest.raises(EnginePhaseError, match="second read interrupted"):
        coordinator().settle()
    state = json.loads(state_path.read_text())
    assert state["rounds"] == 1 and state["pending_verify_cycle"] == 1
    assert advanced == []
    events.clear()
    coordinator().settle()
    assert events[0] == ("verify", 1)
    assert commands == ["settle-1"]
    assert advanced == ["settled"]
    assert not json.loads(state_path.read_text()).get("pending_verify_cycle")


@pytest.mark.parametrize("damage,expected_calls", [("missing_artifact", 0), ("missing_window", 1)])
def test_saved_reader_windows_recover_without_rebuying_complete_work(
        tmp_path, monkeypatch, damage, expected_calls):
    import docx
    import shutil
    from docproof.__main__ import main
    from galley import verify
    from tests.galley.test_verification_checkpoints import Provider

    driver = _driver(tmp_path)
    document = docx.Document()
    for n in range(3):
        document.add_paragraph(f"Section {n}. " + "The book was quiet. " * 220)
    document.save(driver.book)
    shutil.copyfile(driver.book, driver._final_run() / "book - Atmosphere Press Proofreader.docx")
    commands, providers = [], []
    def execute(spec):
        commands.append(spec.argv[spec.argv.index("--verification-pass") + 1])
        return SimpleNamespace(returncode=main(spec.argv[1:]), limit="", tail="")
    phases = EnginePhases(driver, execute)
    run = _complete_verification(phases, monkeypatch)
    path = run / "verification" / "type-compare" / "finished_walk.json"
    proof = json.loads(path.read_text())["verification_provenance"]
    assert len(proof["windows"]) > 1
    if damage == "missing_artifact":
        path.unlink()
    else:
        key = next(iter(proof["windows"]))
        (run / verify._CHECKPOINT_DIR / proof["scope"] / proof["invocation_id"] / (key + ".json")).unlink()
    def factory(*args, **kwargs):
        provider = Provider()
        providers.append(provider)
        return provider
    monkeypatch.setattr("docproof.providers.subagent.SubagentProvider", factory)
    phases.verify()
    assert phases._coverage(run, require_full_passes=True)
    assert commands == ["type-compare"]
    assert sum(len(p.calls) for p in providers) == expected_calls
    restored = json.loads(path.read_text())["verification_provenance"]
    assert restored["invocation_id"] == proof["invocation_id"]
    assert phases._initial_coverage_saved(run)
    phases.verify(cycle=1)
    assert commands == ["type-compare"]            # same build stays proven
    assert sum(len(p.calls) for p in providers) == expected_calls


def test_missing_coverage_archive_is_rebuilt_from_proven_reads_without_model_work(tmp_path, monkeypatch):
    phases = EnginePhases(_driver(tmp_path), lambda spec: pytest.fail("proven reads cannot repeat"))
    run = _complete_verification(phases, monkeypatch)
    phases.verify()
    receipt = json.loads((phases.directory / "initial-coverage.json").read_text())
    target = phases.directory / next(iter(receipt["artifacts"]))
    target.write_text('{"damaged": true}')
    assert not phases._initial_coverage_saved(run)
    phases.verify()
    assert phases._initial_coverage_saved(run)
    assert list(target.parent.glob(target.name + ".damaged-*"))


def test_running_engine_receipt_adopts_independently_proven_completed_output(tmp_path):
    output = tmp_path / "runs" / "evidence.json"
    driver = _driver(tmp_path)
    phases = EnginePhases(driver, lambda spec: pytest.fail("proof must avoid another command"))
    output.write_text('{"complete": true}')
    path = phases.directory / "recovered-read.json"
    path.write_text('{"status": "running"}')
    # A malformed/incomplete engine receipt has no authority. The explicit
    # semantic validator proves the output, and its durable receipt is rebuilt.
    phases._command("verify", "recovered-read", ["verify"], [output],
                    validate=lambda: json.loads(output.read_text())["complete"])
    assert json.loads(path.read_text())["status"] == "completed"
    assert json.loads(path.read_text())["recovered"] is True


def test_running_receipt_resumes_through_budget_authority_and_keeps_attempt(tmp_path):
    output = tmp_path / "runs" / "evidence.json"
    calls = []
    driver = _driver(tmp_path)
    def execute(spec):
        calls.append(spec)
        output.write_text('{"complete": true}')
        return SimpleNamespace(returncode=0, limit="", tail="")
    phases = EnginePhases(driver, execute)
    phases._command("audit", "recovered", ["audit"], [output])
    path = phases.directory / "recovered.json"
    saved = json.loads(path.read_text())
    saved["status"] = "running"
    path.write_text(json.dumps(saved))
    output.unlink()
    phases._command("audit", "recovered", ["audit"], [output])
    assert len(calls) == 2
    assert len(list((phases.directory / "attempts" / "recovered").glob("*.json"))) == 1
    assert json.loads(path.read_text())["status"] == "completed"


def test_active_budget_owner_does_not_lose_its_operation_receipt(tmp_path):
    from galley.execution_budget import ExecutionBudgetError
    driver = _driver(tmp_path)
    class Budget:
        def assert_available(self, phase):
            assert phase == "code-verify"
            raise ExecutionBudgetError("A live reader owns this phase")
    driver._execution_budget = Budget
    phases = EnginePhases(driver, lambda spec: pytest.fail("must not overlap the active reader"))
    receipt = phases.directory / "active-read.json"
    receipt.write_text('{"status": "running", "owner": "existing"}')
    before = receipt.read_bytes()
    with pytest.raises(ExecutionBudgetError, match="live reader"):
        phases._command("verify", "active-read", ["verify"], [])
    assert receipt.read_bytes() == before


@pytest.mark.parametrize("code,limit,retryable", [(1, "", True), (0, "", True),
    (2, "", False), (5, "", False), (130, "", False), (143, "", False),
    (1, "timeout", False)])
def test_command_recovery_does_not_reset_hard_limits(tmp_path, code, limit, retryable):
    phases = EnginePhases(_driver(tmp_path), lambda spec:
                          SimpleNamespace(returncode=code, limit=limit, tail="failed"))
    with pytest.raises(EnginePhaseError) as error:
        phases._command("audit", "failure", ["audit"], [tmp_path / "missing.json"])
    assert error.value.retryable is retryable


@pytest.mark.parametrize("limit,tail", [("credentials", "Sign in required"),
                                      ("", "Invalid API key")])
def test_command_credentials_use_the_existing_sign_in_recovery(tmp_path, limit, tail):
    from galley.driver import CredentialsError
    phases = EnginePhases(_driver(tmp_path), lambda spec:
                          SimpleNamespace(returncode=1, limit=limit, tail=tail))
    with pytest.raises(CredentialsError):
        phases._command("audit", "failure", ["audit"], [tmp_path / "missing.json"])
    receipt = json.loads((phases.directory / "failure.json").read_text())
    assert receipt["status"] == "failed"
    assert receipt["reason"] == tail


@pytest.mark.parametrize("field,value", [("output_copies", ["corrupt"]),
    ("outputs", ["corrupt"]), ("output_copies", {"target": 123}),
    ("identity", ["corrupt"])])
def test_malformed_receipt_metadata_recovers_through_current_proof(tmp_path, field, value):
    output = tmp_path / "artifact.json"
    commands = []
    def execute(spec):
        commands.append(spec)
        output.write_text('{"complete": true}')
        return SimpleNamespace(returncode=0, limit="", tail="")
    phases = EnginePhases(_driver(tmp_path), execute)
    phases._command("audit", "shape", ["audit"], [output])
    path = phases.directory / "shape.json"
    saved = json.loads(path.read_text())
    saved[field] = value
    if field == "identity":
        saved["fingerprint"] = "damaged"
    path.write_text(json.dumps(saved))
    phases._command("audit", "shape", ["audit"], [output],
                    validate=lambda: json.loads(output.read_text()) == {"complete": True})
    assert len(commands) == 1
    assert json.loads(path.read_text())["status"] == "completed"


def test_missing_settlement_output_is_restored_without_an_extra_repair_round(tmp_path, monkeypatch):
    from galley.manifest import sha256_file
    driver = _driver(tmp_path)
    run = driver._final_run()
    output = run / "settlement.json"
    calls = []
    def execute(spec):
        calls.append(spec)
        output.write_text('{"rounds": 2, "records": [], "open": ["r-persistent"]}')
        return SimpleNamespace(returncode=0, limit="", tail="")
    phases = EnginePhases(driver, execute)
    phases._command("settle", "settle-2", ["settle"], [output])
    output.unlink()
    identity = {"source": sha256_file(driver.book), "run": str(run.resolve()),
                "config": sha256_file(phases._config())}
    (phases.directory / "review-loop.json").write_text(json.dumps({
        "schema_version": 1, "identity": identity, "rounds": 2,
        "prior_ids": ["r-persistent"], "history": []}))
    monkeypatch.setattr(phases, "_coverage", lambda *a, **k: True)
    monkeypatch.setattr(phases, "_initial_coverage_saved", lambda run: True)
    monkeypatch.setattr("galley.settle.open_items", lambda run: [SimpleNamespace(id="r-persistent", para_id="body-0000")])
    advanced = []
    monkeypatch.setattr(phases, "_advance", advanced.append)
    phases.settle()
    assert len(calls) == 1 and output.is_file()
    assert advanced == ["settled"]
    assert json.loads((phases.directory / "review-loop.json").read_text())["rounds"] == 2


@pytest.mark.parametrize("changed", ["source", "config"])
def test_receipt_recovery_never_reuses_a_different_approval(tmp_path, changed):
    driver = _driver(tmp_path)
    output = tmp_path / "runs" / "audit.json"
    def execute(spec):
        output.write_text('{"ran": true}')
        return SimpleNamespace(returncode=0, limit="", tail="")
    phases = EnginePhases(driver, execute)
    phases._command("audit", "audit", ["audit"], [output])
    if changed == "config":
        phases._config().write_text("api: {model: changed-model}\n")
    else:
        driver.book.write_bytes(driver.book.read_bytes() + b"changed source")
    with pytest.raises(EnginePhaseError, match="approved inputs") as error:
        phases._command("audit", "audit", ["audit"], [output])
    assert not error.value.retryable


def test_settlement_finishes_missing_secondary_read_before_advancing(tmp_path, monkeypatch):
    from docproof.__main__ import main
    from tests.galley.test_verification_checkpoints import Provider
    driver = _driver(tmp_path)
    commands, providers = [], []
    def execute(spec):
        commands.append(spec.argv[spec.argv.index("--verification-pass") + 1])
        return SimpleNamespace(returncode=main(spec.argv[1:]), limit="", tail="")
    phases = EnginePhases(driver, execute)
    run = _complete_verification(phases, monkeypatch)
    (run / "verification" / "type-compare" / "finished_walk.json").unlink()
    def factory(*args, **kwargs):
        p = Provider()
        providers.append(p)
        return p
    monkeypatch.setattr("docproof.providers.subagent.SubagentProvider", factory)
    (run / "settlement.json").write_text('{"rounds": 0, "records": [], "open": []}')
    advanced = []
    monkeypatch.setattr(phases, "_advance", advanced.append)
    phases.settle()
    assert commands == ["type-compare"]
    assert sum(len(p.calls) for p in providers) == 0
    assert advanced == ["settled"] and phases._initial_coverage_saved(run)


def test_recovery_does_not_mutate_an_active_reader_transaction(tmp_path, monkeypatch):
    import fcntl
    from galley.verify import verification_invocation
    phases = EnginePhases(_driver(tmp_path), lambda spec: pytest.fail("active read cannot repeat"))
    run = _complete_verification(phases, monkeypatch)
    provider, model, options = phases._reader("primary")
    invocation = verification_invocation(run, provider, model, output_dir=run, **options)
    state = invocation.root / "state.json"
    before = state.read_bytes()
    with (invocation.root / "lock").open("a") as lock:
        fcntl.flock(lock.fileno(), fcntl.LOCK_EX)
        with pytest.raises(EnginePhaseError) as error:
            phases._recover_read(run, "primary", run)
        assert error.value.kind == "active" and not error.value.retryable
    assert state.read_bytes() == before
