"""Author teasers v4: DeepSeek V4 Pro writes from the whole book, Opus 5.5 adjudicates."""
from __future__ import annotations

from dataclasses import replace
import hashlib
import io
import json
from pathlib import Path
import time
from types import SimpleNamespace

from docx import Document
from fastapi.testclient import TestClient
import pytest

from app import teasers
from app.jobs import Job, JobRunner, JobStore
from app.main import create_app
from app.settings import Paths, Settings
from app.teasers import Queue, TeaserError, accept_adjudication, accept_draft, approval_issues
from app.teaser_delivery import deliver
from docproof.providers.base import NormalizedUsage, ProviderResult
from docproof.teasers import (ADJUDICATOR_EFFORT, ADJUDICATOR_MODEL, AUTHOR_WARNING, VERSION,
                              WRITER_EFFORT, WRITER_MODEL, adjudicator, writer)
from docproof.teasers.document import write_document
from docproof.teasers.models import (Adjudication, Correction, Draft, Manuscript, OptionRuling, Teaser,
                                     apply_adjudication, bibliographic, digest, draft_issues, word_count)

BOOK = ("The Ferry Ledger\nby Ada Smith\n"
        "Mara returns to the island to repair her father's ferry, the Lark.\n"
        "Her brother Finn wants to sell the boat to a mainland company.\n"
        "In the last chapter Finn admits he forged their father's will.\n")

SENTENCE = ("Mara returns to the island to repair her father's ferry and finds her brother "
            "waiting at the landing with an offer from the mainland. ")


def teaser(number, words=170, angle=None, lead=""):
    """Three paragraphs, `words` words in total, a distinct opening per option."""
    body = (lead + f"Option{number} " + SENTENCE * 20).split()[:words]
    third = len(body) // 3
    return Teaser(number=number, angle=angle or f"Angle {number}",
                  paragraphs=[" ".join(body[:third]), " ".join(body[third:2 * third]),
                              " ".join(body[2 * third:])])


def package(**changes):
    teasers_ = [teaser(n) for n in range(1, 6)]
    for number, value in changes.items():
        teasers_[int(number[1:]) - 1] = value
    return Draft(title="The Ferry Ledger", author="Ada Smith", teasers=teasers_)


def ruling_for(draft, *, corrections=(), rulings=None):
    rulings = rulings or {}
    return Adjudication(draft_sha256=digest(draft), corrections=list(corrections), options=[
        OptionRuling(number=n, ruling=rulings.get(n, "corrected" if any(c.option == n for c in corrections)
                                                  else "accurate"),
                     note="Finn's forgery is a last-chapter reveal." if rulings.get(n) == "rewrite" else "")
        for n in range(1, 6)])


def fix(option=2, paragraph=1, before=None, after="his sister", evidence="Her brother Finn wants to sell the boat"):
    before = before or "her brother"
    return Correction(option=option, paragraph=paragraph, before=before, after=after, kind="factual_error",
                      evidence=evidence, explanation="Wrong relationship.")


@pytest.fixture
def queued(tmp_path):
    path = tmp_path / "manuscript.docx"
    doc = Document()
    for line in BOOK.strip().splitlines():
        doc.add_paragraph(line)
    doc.save(path)
    job = Job(id="format-1", filename="Smith - Book Original.docx", source_path=str(path),
              model="test", mode="now", kind="prep", state="done", owner_id="owner")
    queue = Queue(tmp_path / "watch")
    queue.configure(enabled=True)
    queue.add(job)
    return queue, job, queue.claim("worker")


def drafted(queue, task, draft=None):
    return accept_draft(queue, task, {"draft": (draft or package()).model_dump(),
                                      "receipts": [{"at": time.time(), "model": WRITER_MODEL}]})


# ---- the package contract --------------------------------------------------

@pytest.mark.parametrize("words,paragraphs,valid", [(150, 3, True), (200, 3, True), (149, 3, False),
                                                    (201, 3, False), (170, 2, False), (170, 4, False)])
def test_three_paragraphs_and_150_to_200_words(words, paragraphs, valid):
    t = teaser(1, words)
    if paragraphs != 3:
        text = " ".join(t.paragraphs).split()
        size = len(text) // paragraphs + 1
        t = t.model_copy(update={"paragraphs": [" ".join(text[i:i + size]) for i in range(0, len(text), size)]})
        assert len(t.paragraphs) == paragraphs
    assert (not draft_issues(package(t1=t))) is valid


def test_five_distinct_numbered_options():
    assert draft_issues(package()) == []
    assert draft_issues(Draft(title="", author="", teasers=package().teasers[:4]))
    assert any("duplicates" in i for i in draft_issues(package(t2=teaser(1).model_copy(update={"number": 2}))))
    assert any("angle" in i for i in draft_issues(package(t2=teaser(2, angle="Angle 1"))))


def test_title_and_author_only_when_the_manuscript_prints_them():
    book = Manuscript(BOOK)
    kept = bibliographic(package(), book)
    assert (kept.title, kept.author) == ("The Ferry Ledger", "Ada Smith")
    guessed = bibliographic(package().model_copy(update={"title": "Ferry Tales", "author": "A. Smith Jr."}), book)
    assert (guessed.title, guessed.author) == ("", "")


def test_evidence_matching_ignores_case_punctuation_and_quote_style():
    book = Manuscript("“Don’t,” she said — and the LARK sailed.")
    assert book.quotes("\"don't\" she said, and the lark sailed")
    assert not book.quotes("she said the lark sank")
    assert not book.quotes("  ")


# ---- the adjudicator's corrections ----------------------------------------

def test_minimal_correction_applies_exactly_and_nothing_else_changes():
    draft = package(t2=teaser(2, lead="Her brother waits. "))
    ruling = ruling_for(draft, corrections=[fix(before="Her brother waits.", after="Her brother Finn waits.")])
    corrected = apply_adjudication(draft, ruling, Manuscript(BOOK))
    assert corrected.teasers[1].paragraphs[0].startswith("Her brother Finn waits.")
    assert corrected.teasers[0] == draft.teasers[0] and corrected.teasers[1].paragraphs[1:] == draft.teasers[1].paragraphs[1:]


@pytest.mark.parametrize("change,message", [
    (dict(evidence="Finn sold the Lark to the mainland company."), "not a verbatim"),
    (dict(before="nowhere in this paragraph"), "exactly once"),
    (dict(before="her brother", after="her brother"), "different"),
    (dict(after=""), "nonempty"),
    (dict(paragraph=4), "does not exist"),
    (dict(evidence="Her brother"), "evidence is one quoted passage"),
    (dict(before=" ".join(["word"] * 36)), "at most 35 words"),
])
def test_bad_corrections_refuse_the_whole_ruling(change, message):
    draft = package(t2=teaser(2, lead="Her brother waits. "))
    base = dict(before="Her brother waits.", after="Finn waits.")
    base.update(change)
    with pytest.raises(ValueError, match=message):
        apply_adjudication(draft, ruling_for(draft, corrections=[fix(**base)]), Manuscript(BOOK))


def test_rulings_must_agree_with_corrections():
    draft = package(t2=teaser(2, lead="Her brother waits. "))
    correction = fix(before="Her brother waits.", after="Finn waits.")
    book = Manuscript(BOOK)
    with pytest.raises(ValueError, match="only an option ruled corrected"):
        apply_adjudication(draft, ruling_for(draft, corrections=[correction], rulings={2: "accurate"}), book)
    with pytest.raises(ValueError, match="no corrections"):
        apply_adjudication(draft, ruling_for(draft, rulings={3: "corrected"}), book)
    with pytest.raises(ValueError, match="different draft"):
        apply_adjudication(draft, ruling_for(package()), book)


def test_a_correction_that_breaks_the_length_is_refused():
    draft = package(t2=teaser(2, 151, lead="Her brother waits at the landing. "))
    ruling = ruling_for(draft, corrections=[fix(before="Her brother waits at the landing.", after="Finn.")])
    with pytest.raises(ValueError, match="After the corrections"):
        apply_adjudication(draft, ruling, Manuscript(BOOK))


def test_more_than_sixty_replaced_words_is_a_rewrite_not_a_correction():
    draft = package()
    corrections = []
    for paragraph in (1, 2):
        words = draft.teasers[1].paragraphs[paragraph - 1].split()
        corrections.append(fix(paragraph=paragraph, before=" ".join(words[:31]),
                               after="Finn " + " ".join(words[1:31])))
    with pytest.raises(ValueError, match="is a rewrite"):
        apply_adjudication(draft, ruling_for(draft, corrections=corrections), Manuscript(BOOK))


# ---- the queue --------------------------------------------------------------

def test_formatting_queues_the_whole_manuscript_once(queued):
    queue, job, task = queued
    assert task["version"] == VERSION and task["state"] == "queued"
    assert "Finn admits he forged" in task["manuscript"]
    Path(job.source_path).unlink()
    assert queue.add(job) == task["id"] and len(queue.list()) == 1
    assert queue.claim("other") is None
    with pytest.raises(TeaserError):
        queue.owned(task["id"], "other")


def test_disabled_failed_and_nonformatting_jobs_do_not_queue(queued):
    queue, job, _ = queued
    queue.configure(enabled=False)
    assert queue.add(replace(job, id="other")) is None
    queue.configure(enabled=True)
    assert queue.add(replace(job, id="other", state="failed")) is None
    assert queue.add(replace(job, id="other", kind="promo")) is None


def test_accurate_package_is_approved_as_written(queued):
    queue, _, task = queued
    task = drafted(queue, task)
    assert task["state"] == "drafted" and task["drafts"][-1]["model"] == WRITER_MODEL
    draft = Draft.model_validate(task["drafts"][-1]["content"])
    task = accept_adjudication(queue, task, {"ruling": ruling_for(draft).model_dump()})
    assert task["state"] == "approved" and approval_issues(task) == []
    assert len(task["drafts"]) == 1


def test_corrections_are_saved_as_their_own_bound_draft(queued):
    queue, _, task = queued
    draft = package(t2=teaser(2, lead="Her brother waits. "))
    task = drafted(queue, task, draft)
    ruling = ruling_for(draft, corrections=[fix(before="Her brother waits.", after="Her brother Finn waits.")])
    task = accept_adjudication(queue, task, {"ruling": ruling.model_dump()})
    assert task["state"] == "approved" and approval_issues(task) == []
    last = task["drafts"][-1]
    assert last["operation"] == "correction" and last["model"] == ADJUDICATOR_MODEL
    assert last["base_sha256"] == digest(draft) and "Finn waits" in last["content"]["teasers"][1]["paragraphs"][0]
    # Any other package in last place is refused at delivery.
    task["drafts"].append({**task["drafts"][0]})
    assert approval_issues(task)


def test_refused_ruling_changes_nothing(queued):
    queue, _, task = queued
    task = drafted(queue, task)
    draft = Draft.model_validate(task["drafts"][-1]["content"])
    bad = ruling_for(draft, corrections=[fix(option=1, before="Option1", after="X",
                                             evidence="invented passage about a ferry race")])
    with pytest.raises(TeaserError, match="refused"):
        accept_adjudication(queue, task, {"ruling": bad.model_dump()})
    saved = queue.get(task["id"])
    assert saved["state"] == "drafted" and saved["adjudications"] == [] and len(saved["drafts"]) == 1


def test_rewrite_returns_only_flagged_options_to_the_writer(queued):
    queue, _, task = queued
    task = drafted(queue, task)
    draft = Draft.model_validate(task["drafts"][-1]["content"])
    task = accept_adjudication(queue, task, {"ruling": ruling_for(draft, rulings={4: "rewrite"}).model_dump()})
    assert task["state"] == "revise" and task["rewrite"]["options"] == [4]
    assert "forgery" in task["rewrite"]["notes"]["4"]
    touched = package(t1=teaser(1, 180), t4=teaser(4, 190))
    with pytest.raises(TeaserError, match="only the options"):
        drafted(queue, task, touched)
    task = drafted(queue, task, package(t4=teaser(4, 190)))
    assert task["state"] == "drafted" and task["drafts"][-1]["operation"] == "rewrite"


def test_repeated_rewrites_start_a_fresh_package(queued):
    queue, _, task = queued
    task = drafted(queue, task)
    for round_number in range(teasers.MAX_REWRITE_ROUNDS + 1):
        draft = Draft.model_validate(task["drafts"][-1]["content"])
        task = accept_adjudication(queue, task, {"ruling": ruling_for(draft, rulings={4: "rewrite"}).model_dump()})
        if task["state"] == "revise":
            task = drafted(queue, task, package(t4=teaser(4, 160 + round_number)))
    assert task["state"] == "retry_wait" and task["resume_state"] == "queued" and task["failures"] == 1


def test_resent_results_after_a_lost_reply_are_harmless(queued):
    queue, _, task = queued
    task = drafted(queue, task)
    again = drafted(queue, task)
    assert len(again["drafts"]) == 1
    draft = Draft.model_validate(task["drafts"][-1]["content"])
    payload = {"ruling": ruling_for(draft).model_dump()}
    task = accept_adjudication(queue, task, payload)
    assert accept_adjudication(queue, task, payload)["adjudications"] == task["adjudications"]


def test_unfinished_earlier_workflow_restarts_with_its_audit(queued):
    queue, _, task = queued
    old = {k: v for k, v in task.items() if k not in ("manuscript", "drafts", "adjudications")}
    old.update(version=3, chunks=[{"id": 1, "paragraphs": [{"id": 1, "text": "Mara returns."},
                                                            {"id": 2, "text": "Finn waits."}]}],
               storysheet={"private": "Sol's facts"}, drafts=[{"old": "draft"}], reviews=[], failures=4)
    queue.save(old, "story_ready")
    queue.recover()
    migrated = queue.claim("worker")
    assert migrated["version"] == VERSION and migrated["state"] == "queued"
    assert migrated["manuscript"] == "Mara returns.\nFinn waits." and migrated["drafts"] == []
    assert not migrated.get("failures")
    prior = migrated["prior_workflows"][-1]
    assert prior["storysheet"] == {"private": "Sol's facts"} and "chunks" not in prior


def test_daily_writer_allowance_pauses_without_counting(queued):
    queue, _, task = queued
    task["generation_times"] = [time.time()] * teasers.MAX_WRITER_CALLS_PER_DAY
    queue.save(task)
    queue.save(queue.get(task["id"]), "queued")
    with queue.connect() as conn:
        conn.execute("UPDATE tasks SET lease=0")
    assert queue.claim("worker") is None
    saved = queue.get(task["id"])
    assert saved["state"] == "retry_wait" and not saved.get("failures")


def test_outages_are_transient_and_book_problems_are_counted(queued):
    queue, _, task = queued
    assert teasers.is_transient("DeepInfra call failed: APITimeoutError")
    assert teasers.is_transient("You've hit your session limit · resets 3pm")
    assert not teasers.is_transient("Opus's ruling failed the checks 3 times")
    before = time.time()
    queue.retry(task, "DeepInfra call failed: 503", counted=False)
    saved = queue.get(task["id"])
    assert saved["state"] == "retry_wait" and not saved.get("failures")
    assert saved["retry_at"] - before <= 125


def test_counted_backoff_is_capped_at_half_an_hour(queued):
    queue, _, task = queued
    task["failures"] = 20
    before = time.time()
    queue.retry(task, "Opus's ruling failed the checks.")
    assert queue.get(task["id"])["retry_at"] - before <= teasers.MAX_BACKOFF_SECONDS + 5


def test_duplicate_worker_error_does_not_extend_saved_retry(queued):
    from app.routes.teasers import dispatch, WorkerMessage
    queue, _, task = queued
    queue.retry(task, "Temporary provider failure", delay=30)
    before = queue.get(task["id"])
    app = SimpleNamespace(state=SimpleNamespace(watch=SimpleNamespace(home=queue.root.parent)))
    result = dispatch(app, WorkerMessage(action="error", worker="worker", task_id=task["id"],
                                        payload={"error": "Duplicate transport report"}))["task"]
    assert result["failures"] == before["failures"] and result["retry_at"] == before["retry_at"]


# ---- the writer -------------------------------------------------------------

class FakeCompletions:
    def __init__(self, answers):
        self.answers, self.calls = list(answers), []

    def create(self, **kw):
        self.calls.append(kw)
        answer = self.answers.pop(0)
        if isinstance(answer, Exception):
            raise answer
        body = {"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(answer.model_dump())}}],
                "usage": {"prompt_tokens": 90_000, "completion_tokens": 9_000, "estimated_cost": 0.15,
                          "completion_tokens_details": {"reasoning_tokens": 7_000}}}
        return SimpleNamespace(model_dump=lambda: body)


def fake_writer(*answers):
    completions = FakeCompletions(answers)
    return writer.Writer(client=SimpleNamespace(chat=SimpleNamespace(completions=completions))), completions


def test_writer_reads_the_whole_book_with_reasoning_on_and_fixes_length_in_conversation(tmp_path):
    long = package(t3=teaser(3, 214))
    w, completions = fake_writer(long, package())
    draft, receipts = writer.write(BOOK, tmp_path, writer=w)
    assert draft == package() and len(receipts) == 2
    first, second = completions.calls
    assert first["model"] == WRITER_MODEL == "deepseek-ai/DeepSeek-V4-Pro"
    assert first["reasoning_effort"] == WRITER_EFFORT == "high"
    assert "Finn admits he forged" in first["messages"][1]["content"]  # the whole book, ending included
    assert "EDITORIAL STANDARD" in first["messages"][0]["content"]
    assert second["messages"][2]["role"] == "assistant" and "Option 3 has 214 words" in second["messages"][3]["content"]
    assert receipts[0]["reasoning_tokens"] == 7_000 and receipts[0]["cost_usd"] == 0.15
    # A finished package is saved: a lost hand-off never pays for a second write.
    assert writer.write(BOOK, tmp_path, writer=fake_writer()[0]) == (draft, receipts)


def test_writer_gives_up_after_its_fix_rounds(tmp_path):
    long = package(t3=teaser(3, 214))
    w, _ = fake_writer(*[long] * (writer.FIX_ROUNDS + 1))
    with pytest.raises(writer.WriterError, match="mechanical checks"):
        writer.write(BOOK, tmp_path, writer=w)


def test_writer_outage_is_its_own_error(tmp_path):
    w, _ = fake_writer(TimeoutError("read timed out"))
    with pytest.raises(writer.WriterUnavailable, match="DeepInfra call failed"):
        writer.write(BOOK, tmp_path, writer=w)


def test_rewrite_keeps_every_option_the_adjudicator_passed(tmp_path):
    current = package()
    wandered = package(t1=teaser(1, 199), t4=teaser(4, 180), t5=teaser(5, 160))
    w, completions = fake_writer(wandered)
    draft, _ = writer.rewrite(BOOK, current, {4: "Finn's forgery is a late reveal."}, tmp_path, writer=w)
    assert draft.teasers[3] == teaser(4, 180)
    assert [draft.teasers[i] for i in (0, 1, 2, 4)] == [current.teasers[i] for i in (0, 1, 2, 4)]
    request = completions.calls[0]["messages"][-1]["content"]
    assert "option(s) 4" in request and "late reveal" in request


# ---- the adjudicator ------------------------------------------------------

class FakeLane:
    def __init__(self, *answers):
        self.answers, self.calls = list(answers), []

    def complete_structured(self, **kw):
        self.calls.append(kw)
        return ProviderResult(parsed=self.answers.pop(0), usage=NormalizedUsage(input_tokens=95_000, output_tokens=3_000))


def test_adjudicator_reads_the_whole_book_and_is_told_why_a_ruling_was_refused(tmp_path):
    draft = package(t2=teaser(2, lead="Her brother waits. "))
    bad = ruling_for(draft, corrections=[fix(before="Her brother waits.", after="Finn waits.",
                                             evidence="Finn sank the ferry out of spite")])
    good = ruling_for(draft, corrections=[fix(before="Her brother waits.", after="Finn waits.")])
    lane = FakeLane(bad.model_dump(), good.model_dump())
    ruling, corrected, receipts = adjudicator.adjudicate(BOOK, draft, tmp_path, lane=lane)
    assert ruling == good and corrected.teasers[1].paragraphs[0].startswith("Finn waits.")
    assert [c["model"] for c in lane.calls] == [ADJUDICATOR_MODEL] * 2 and len(receipts) == 2
    assert "Finn admits he forged" in lane.calls[0]["user"] and digest(draft) in lane.calls[0]["user"]
    assert "not a verbatim manuscript passage" in lane.calls[1]["user"]
    # A validated ruling is saved and reused.
    assert adjudicator.adjudicate(BOOK, draft, tmp_path, lane=FakeLane())[0] == good


def test_adjudicator_gives_up_after_its_rounds(tmp_path):
    draft = package()
    bad = ruling_for(draft, rulings={2: "corrected"}).model_dump()
    with pytest.raises(adjudicator.AdjudicatorError, match="failed the checks"):
        adjudicator.adjudicate(BOOK, draft, tmp_path, lane=FakeLane(*[bad] * adjudicator.ROUNDS))


def test_adjudicator_runs_opus_on_the_subscription(monkeypatch):
    monkeypatch.setattr("docproof.agent_lane.require_cli_for", lambda model: None)
    lane = adjudicator.provider()
    assert lane.model == ADJUDICATOR_MODEL == "claude-opus-5-5" and lane.effort == ADJUDICATOR_EFFORT


# ---- worker, end to end -----------------------------------------------------

def test_worker_runs_a_book_from_queue_to_delivery(queued, tmp_path, monkeypatch):
    from app import teaser_delivery as delivery
    from app.routes.teasers import dispatch, WorkerMessage
    from docproof.teasers import worker
    queue, _, task = queued
    app = SimpleNamespace(state=SimpleNamespace(watch=SimpleNamespace(home=queue.root.parent)))

    class Client:
        def call(self, action, task_id="", payload=None):
            return dispatch(app, WorkerMessage(protocol=worker.PROTOCOL, action=action, worker="worker",
                                               task_id=task_id, payload=payload or {}))

    monkeypatch.setattr(delivery, "token_for", lambda home, **kw: "google")
    monkeypatch.setattr(delivery, "ensure_folder", lambda *a, **k: "folder")
    monkeypatch.setattr(delivery.drive, "search_files", lambda *a, **k: [])
    monkeypatch.setattr(delivery, "resume_import", lambda *a, **k: "doc")
    monkeypatch.setattr(delivery, "verify_document", lambda *a, **k: "https://docs.google.com/document/d/doc")
    monkeypatch.setattr(delivery, "deliver_guides", lambda *a, **k: None)
    first = package(t4=teaser(4, lead="Finn forged the will. "))
    w, _ = fake_writer(first, package(t4=teaser(4, 175)))

    class Lane:
        def __init__(self):
            self.count = 0

        def complete_structured(self, **kw):
            self.count += 1
            current = Draft.model_validate(queue.get(task["id"])["drafts"][-1]["content"])
            decision = ruling_for(current, rulings={4: "rewrite"} if self.count == 1 else {})
            return ProviderResult(parsed=decision.model_dump(), usage=NormalizedUsage())

    result = worker.process(task, Client(), tmp_path, write_with=w, adjudicate_with=Lane())
    assert result["state"] == "complete" and result["document_url"].endswith("/doc")
    assert [d["operation"] for d in result["drafts"]] == ["generation", "rewrite"]
    assert len(result["adjudications"]) == 2 and approval_issues(result) == []


def test_old_workers_are_told_to_upgrade(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCPROOF_AGENT_TOKEN", "secret-long-enough-for-the-agent-gate")
    app = create_app(tmp_path, start_runner=False, web=True, session_secret="a-session-secret", https_only=False)
    headers = {"Authorization": "Bearer secret-long-enough-for-the-agent-gate"}
    with TestClient(app) as client:
        assert client.post("/api/teasers/worker", json={}).status_code == 401
        assert client.get("/api/teasers").status_code == 401
        assert client.put("/api/teasers/settings", json={"enabled": True}).status_code == 401
        old = client.post("/api/teasers/worker", headers=headers, json={"protocol": 3, "action": "poll", "worker": "fly"})
        assert old.json() == {"task": None, "upgrade_required": True}
        new = client.post("/api/teasers/worker", headers=headers, json={"protocol": 4, "action": "poll", "worker": "fly"})
        assert new.status_code == 200 and new.json() == {"task": None}


# ---- the document and Google delivery -------------------------------------

def test_document_has_the_warning_five_unranked_options_and_nothing_internal(tmp_path):
    draft = package()
    doc = Document(write_document(tmp_path / "teasers.docx", draft, book_label="Smith - Book Original"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert text.startswith("The Ferry Ledger\nAda Smith") and AUTHOR_WARNING in text
    assert text.index("Option 1") < text.index("Option 2") < text.index("Option 5")
    assert all(p in text for t in draft.teasers for p in t.paragraphs)
    assert "Recommended" not in text and "DeepSeek" not in text and "Opus" not in text
    untitled = Document(write_document(tmp_path / "b.docx", draft.model_copy(update={"title": "", "author": ""}),
                                       book_label="Smith - Book Original"))
    assert untitled.paragraphs[0].text == "Smith - Book Original"


def test_delivery_requires_adjudication_and_is_idempotent(queued, monkeypatch):
    from app import teaser_delivery as delivery
    queue, _, task = queued
    task = drafted(queue, task)
    with pytest.raises(TeaserError, match="Only an approved"):
        deliver(queue, task, queue.root.parent, token="google")
    draft = Draft.model_validate(task["drafts"][-1]["content"])
    task = accept_adjudication(queue, task, {"ruling": ruling_for(draft).model_dump()})
    calls = []
    monkeypatch.setattr(delivery, "ensure_folder", lambda *a, **kw: "folder")
    monkeypatch.setattr(delivery.drive, "search_files", lambda *a, **kw: [])
    monkeypatch.setattr(delivery, "resume_import", lambda *a, **kw: calls.append("upload") or "doc")
    monkeypatch.setattr(delivery, "verify_document", lambda *a, **kw: "https://docs.google.com/document/d/doc/edit")
    monkeypatch.setattr(delivery, "deliver_guides", lambda *a, **kw: calls.append("guides"))
    task = deliver(queue, task, queue.root.parent, token="google")
    assert task["state"] == "complete" and delivery.delivery_key(task).endswith("-v4")
    deliver(queue, task, queue.root.parent, token="google")
    assert calls == ["upload", "guides"]


def test_delivery_is_withheld_from_an_unbound_package(queued, monkeypatch):
    queue, _, task = queued
    task = drafted(queue, task)
    draft = Draft.model_validate(task["drafts"][-1]["content"])
    task = accept_adjudication(queue, task, {"ruling": ruling_for(draft).model_dump()})
    task["drafts"].append({**task["drafts"][-1], "sha256": "tampered"})
    queue.save(task)
    with pytest.raises(TeaserError, match="Upload withheld"):
        deliver(queue, queue.get(task["id"]), queue.root.parent, token="google")


def test_readback_requires_every_paragraph_and_the_warning():
    from app import teaser_delivery as delivery
    draft = package()
    body = "\n".join([AUTHOR_WARNING] + [p for t in draft.teasers for p in t.paragraphs])
    meta = {"mimeType": delivery.drive.GOOGLE_DOC_MIME, "parents": ["folder"], "webViewLink": "https://x"}
    def opener(request):
        return io.BytesIO(body.encode() if "export" in request.full_url else json.dumps(meta).encode())
    assert delivery.verify_document("token", "doc", draft, "folder", opener=opener) == "https://x"
    body = body.replace(AUTHOR_WARNING, "")
    with pytest.raises(TeaserError, match="missing"):
        delivery.verify_document("token", "doc", draft, "folder", opener=opener)


def test_delivery_waits_for_both_guides_and_reuses_saved_ids(queued, monkeypatch):
    from app import teaser_delivery as delivery
    queue, _, task = queued
    uploaded, records = [], {}
    def upload(token, folder_id, path, **kwargs):
        file_id = Path(path).suffix[1:]
        uploaded.append(file_id)
        records[file_id] = {"mimeType": kwargs["mime_type"], "parents": [folder_id],
                            "md5Checksum": hashlib.md5(Path(path).read_bytes()).hexdigest(),
                            "webViewLink": "https://drive.google.com/file/d/" + file_id}
        return file_id
    monkeypatch.setattr(delivery.drive, "search_files", lambda *a, **k: [])
    monkeypatch.setattr(delivery.drive, "upload", upload)
    monkeypatch.setattr(delivery.drive, "_json_call",
                        lambda request, **k: records[request.full_url.split("/files/")[1].split("?")[0]])
    delivery.deliver_guides(queue, task, "token", "folder")
    delivery.deliver_guides(queue, task, "token", "folder")
    assert uploaded == ["xlsx", "pdf"] and task["guide_url"].endswith("/pdf")
    records["pdf"]["md5Checksum"] = "wrong"
    with pytest.raises(ValueError, match="complete two-page guide"):
        delivery.deliver_guides(queue, task, "token", "folder")


def test_folder_search_rejects_case_insensitive_name_match(queued, monkeypatch):
    from app import teaser_delivery as delivery
    queue, _, _ = queued
    old = SimpleNamespace(id="old", name="Author teasers", is_folder=True)
    new = SimpleNamespace(id="new", name="author teasers", is_folder=True)
    queue.configure(folder_id="old")
    monkeypatch.setattr(delivery.drive, "get_file", lambda token, file_id, **kw: old if file_id == "old" else new)
    monkeypatch.setattr(delivery.drive, "search_files", lambda *a, **k: [old])
    created = []
    monkeypatch.setattr(delivery.drive, "create_folder", lambda token, parent, name, **kw: created.append(name) or "new")
    assert delivery.ensure_folder(queue, "token") == "new"
    assert delivery.ensure_folder(queue, "token") == "new"
    assert created == ["author teasers"]


def test_google_upload_recovers_lost_completion_without_second_document(queued, tmp_path):
    from app.teaser_delivery import resume_import
    queue, _, task = queued
    path = tmp_path / "upload.docx"
    path.write_bytes(b"document bytes")
    requests = []
    session = "https://www.googleapis.com/upload/drive/v3/files?upload_id=test"

    def opener(request):
        requests.append(request)
        if request.method == "POST":
            response = io.BytesIO(b"")
            response.headers = {"Location": session}
            return response
        # The session must be durable before sending any document bytes.
        assert queue.get(task["id"])["upload_session"] == session
        if request.data:
            raise TimeoutError("Google accepted the bytes but the reply was lost")
        return io.BytesIO(b'{"id":"one-native-doc"}')

    with pytest.raises(Exception, match="reply was lost"):
        resume_import(queue, task, "google", "folder", path, opener=opener)
    restored = queue.get(task["id"])
    assert resume_import(queue, restored, "google", "folder", path, opener=opener) == "one-native-doc"
    assert [request.method for request in requests] == ["POST", "PUT", "PUT"]
    assert requests[-1].get_header("Content-range") == "bytes */14"


def test_google_upload_resumes_remaining_bytes(queued, tmp_path):
    from email.message import Message
    from urllib.error import HTTPError
    from app.teaser_delivery import resume_import
    queue, _, task = queued
    path = tmp_path / "upload.docx"
    path.write_bytes(b"abcdefghij")
    task["upload_session"] = "https://www.googleapis.com/upload/drive/v3/files?upload_id=test"
    queue.save(task)
    requests = []

    def opener(request):
        requests.append(request)
        if not request.data:
            headers = Message()
            headers["Range"] = "bytes=0-3"
            raise HTTPError(request.full_url, 308, "Resume Incomplete", headers, io.BytesIO())
        assert request.data == b"efghij"
        assert request.get_header("Content-range") == "bytes 4-9/10"
        return io.BytesIO(b'{"id":"one-native-doc"}')

    assert resume_import(queue, task, "google", "folder", path, opener=opener) == "one-native-doc"
    assert len(requests) == 2


def test_completion_hook_queues_before_archiving(queued, tmp_path, monkeypatch):
    queue, job, _ = queued
    store = JobStore(Paths(tmp_path / "app").ensure())
    store.save(job)
    runner = JobRunner(store, Settings(), config_path=Path("config/default.yaml"), notify_home=queue.root.parent)
    events = []
    monkeypatch.setattr("app.teasers.enqueue_completed", lambda home, j: events.append(("teaser", j.id)))
    monkeypatch.setattr(runner, "_archive_done", lambda jid: events.append(("archive", jid)))
    monkeypatch.setattr(runner, "_notify_done", lambda jid: None)
    runner._finish(job.id)
    assert events == [("teaser", job.id), ("archive", job.id)]


def test_fixed_reader_allows_only_the_known_disabled_tool_notice():
    from galley.codex_runner import _safe_events, _fixed_response_complete, _DISABLED_CODE_MODE_NOTICE
    def events(message):
        return _safe_events(io.BytesIO((json.dumps({"type": "item.completed", "item": {
            "type": "error", "message": message}}) + '\n' + json.dumps({"type": "turn.completed"}) + '\n').encode()))
    assert _fixed_response_complete(events(_DISABLED_CODE_MODE_NOTICE))
    assert not _fixed_response_complete(events("A real model error"))
